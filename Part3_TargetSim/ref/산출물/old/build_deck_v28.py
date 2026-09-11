# -*- coding: utf-8 -*-
"""Chapter 3 발표자료(pptx) 를 처음부터 만든다.

    python build_deck.py            # ../Chapter3_target_simulation.pptx 를 새로 쓴다

틀은 Chapter 2 발표자료(네이비판 LIG 틀) 와 같다 : 표지 · 목차 · 간지 · 본문(머리띠 + 흰 패널 + 꼬리말) · 정리.
그림은 ../figures/ 의 PNG (make_figures.py), 노트는 ../Chapter3_발표대본.md (sync_notes.py 의 parse) 에서 채운다.
"""
import io
import os
import sys

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from sync_notes import parse as parse_script          # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.normpath(os.path.join(HERE, ".."))
FIG_DIR = os.path.join(OUT_DIR, "figures")
DECK = os.path.join(OUT_DIR, "Chapter3_target_simulation.pptx")
SCRIPT_MD = os.path.join(OUT_DIR, "Chapter3_발표대본.md")

# ───────────────────────────────────────────── 디자인 토큰 (Chapter 2 네이비판과 같음)
DARK = "13233F"
BG = "E8EDF5"
PANEL = "FFFFFF"
TEXT = "1A2233"
MUTED = "5A6478"
LINE = "B9C0CC"
WHITE = "FFFFFF"
ACCENT = "C2521C"
BLUE = "2B57A6"
GREEN = "2F855A"
ON_DARK = "C9D6EC"
ON_DARK_DIM = "8FA6CC"
RULE_SOFT = "D0D5DE"
CARD = "EEF1F7"
PEACH = "FBEADF"
SOFT = "F4F6FA"
MINT = "E6F3EA"
CIRCLE_DK = "1E3157"
CIRCLE_LN = "34507F"
CIRCLE_TX = "9DB3D6"
SUM_TX = "D6DFF0"

KO = "맑은 고딕"
EN = "Arial"
MONO = "Consolas"

SW, SH = 13.333, 7.5
M = 0.5
INNER_W = SW - 2 * M
BAR_Y, BAR_H = 0.45, 0.85
PANEL_Y, PANEL_H = 1.36, 5.52
BODY_TOP, BODY_BOT = 1.50, 6.82
FOOT_LINE_Y = 6.95
FOOT_TEXT_Y = 7.00
CX0, CX1 = 0.62, 12.72              # 본문 좌우 끝
CW = CX1 - CX0                      # 12.10

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"

CHIPS = ["기동 판단", "자세 갱신", "속도 변환", "위치 적분", "LLA 변환"]


# ───────────────────────────────────────────── 기본 도우미
def rgb(s):
    return RGBColor.from_string(s)


def set_background(slide, color):
    csld = slide._element.find(qn("p:cSld"))
    old = csld.find(qn("p:bg"))
    if old is not None:
        csld.remove(old)
    xml = ('<p:bg xmlns:p="%s" xmlns:a="%s"><p:bgPr><a:solidFill><a:srgbClr val="%s"/></a:solidFill>'
           '<a:effectLst/></p:bgPr></p:bg>' % (P_NS, A_NS, color))
    csld.insert(0, parse_xml(xml))


def txt(slide, text, x, y, w, h, size=12, color=TEXT, bold=False, font=KO, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, wrap=True, spacing=None, space_after=0, name=None):
    """text 는 문자열 또는 (문자열, 옵션dict) 목록. 목록이면 문단마다 서식이 다르다.
    문자열 안의 런 서식: 옵션 'runs' 에 [(text, {bold, color, font, size}), ...] 를 주면 한 문단에 여러 런."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    items = text if isinstance(text, list) else [(text, {})]
    for i, item in enumerate(items):
        s, opt = item if isinstance(item, tuple) else (item, {})
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = opt.get("align", align)
        ls = opt.get("spacing", spacing)
        if ls:
            p.line_spacing = Pt(ls)
        sa = opt.get("space_after", space_after)
        if sa:
            p.space_after = Pt(sa)
        sb = opt.get("space_before", 0)
        if sb:
            p.space_before = Pt(sb)
        runs = opt.get("runs") or [(s, {})]
        for rs, ro in runs:
            r = p.add_run()
            r.text = rs
            f = r.font
            f.name = ro.get("font", opt.get("font", font))
            f.size = Pt(ro.get("size", opt.get("size", size)))
            f.bold = ro.get("bold", opt.get("bold", bold))
            f.color.rgb = rgb(ro.get("color", opt.get("color", color)))
            if ro.get("italic", opt.get("italic")):
                f.italic = True
    return box


def rect(slide, x, y, w, h, fill, line=None, lw=0.75, rounded=False, radius=0.06, name=None):
    shape = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        sp.name = name
    if rounded:
        sp.adjustments[0] = radius
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = rgb(fill)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = rgb(line)
        sp.line.width = Pt(lw)
    sp.shadow.inherit = False
    sp.text_frame.word_wrap = False
    return sp


def ellipse(slide, x, y, d, fill, line=None, lw=1.0):
    sp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    sp.fill.solid()
    sp.fill.fore_color.rgb = rgb(fill)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = rgb(line)
        sp.line.width = Pt(lw)
    sp.shadow.inherit = False
    return sp


def hline(slide, x1, x2, y, color=LINE, lw=0.75):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y), Inches(x2), Inches(y))
    ln.line.color.rgb = rgb(color)
    ln.line.width = Pt(lw)
    return ln


SLIDENUM_XML = (
    '<a:fld xmlns:a="%s" id="{1D0B0B0B-0000-4000-8000-00000000C0DE}" type="slidenum">'
    '<a:rPr lang="ko-KR" altLang="en-US" sz="%d" b="0" dirty="0">'
    '<a:solidFill><a:srgbClr val="%s"/></a:solidFill>'
    '<a:latin typeface="%s"/><a:ea typeface="%s"/></a:rPr>'
    '<a:t>1</a:t></a:fld>')


def slide_number(slide, color=TEXT, size=10):
    box = slide.shapes.add_textbox(Inches(6.2), Inches(FOOT_TEXT_Y), Inches(1.0), Inches(0.3))
    box.name = "SlideNum"
    tf = box.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p._p.append(parse_xml(SLIDENUM_XML % (A_NS, int(size * 100), color, EN, EN)))
    return box


def footer(slide, color=TEXT):
    hline(slide, M, SW - M, FOOT_LINE_Y, LINE, 0.75)
    slide_number(slide, color)


def notes(slide, paras):
    slide.notes_slide.notes_text_frame.text = "\n".join(paras)


def pic(slide, name, x, y, w=None, h=None):
    path = os.path.join(FIG_DIR, name)
    iw, ih = Image.open(path).size
    if w is None and h is None:
        raise ValueError("w 나 h 하나는 줘야 한다")
    if w is None:
        w = h * iw / float(ih)
    if h is None:
        h = w * ih / float(iw)
    p = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    p.name = "Image " + name
    return p, w, h


# ───────────────────────────────────────────── 페이지 틀
def content_chrome(slide, title, subtitle=None):
    set_background(slide, BG)
    rect(slide, M, PANEL_Y, INNER_W, PANEL_H, PANEL, name="Panel")
    rect(slide, M, BAR_Y, INNER_W, BAR_H, DARK, name="Bar")
    paras = [(title, {"size": 19 if subtitle else 20, "bold": True, "color": WHITE, "space_after": 3})]
    if subtitle:
        paras.append((subtitle, {"size": 11, "color": ON_DARK, "spacing": 14}))
    txt(slide, paras, M + 0.20, BAR_Y, 11.6, BAR_H, anchor=MSO_ANCHOR.MIDDLE, name="Title")
    footer(slide, TEXT)


def chips(slide, active, y=BODY_TOP):
    """루프 다섯 단계 칩. active 는 이름 또는 None."""
    cw, gap, h = 1.42, 0.20, 0.36
    total = len(CHIPS) * cw + (len(CHIPS) - 1) * gap
    x = SW / 2 - total / 2
    for i, nm in enumerate(CHIPS):
        on = nm == active
        rect(slide, x, y, cw, h, ACCENT if on else CARD, ACCENT if on else ON_DARK, 1.0, rounded=True, radius=0.25)
        txt(slide, nm, x, y, cw, h, size=11, color=WHITE if on else MUTED, bold=on, align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE)
        x += cw
        if i < len(CHIPS) - 1:
            txt(slide, "▸", x, y, gap, h, size=10, color=ON_DARK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font=KO)
            x += gap


def card(slide, x, y, w, h, title, lines, tone=BLUE, fill=CARD, title_size=13, body_size=11, pad=0.22,
         title_after=4, line_spacing=15):
    """둥근 카드 : 색 제목 + 회색 본문 줄들. lines 는 문자열 목록 (각각 문단)."""
    rect(slide, x, y, w, h, fill, WHITE, 1.0, rounded=True, radius=0.06)
    paras = [(title, {"size": title_size, "bold": True, "color": tone, "space_after": title_after})]
    for s in lines:
        if isinstance(s, tuple):
            paras.append(s)
        else:
            paras.append((s, {"size": body_size, "color": TEXT, "spacing": line_spacing}))
    txt(slide, paras, x + pad, y + pad * 0.7, w - 2 * pad, h - pad * 1.2, anchor=MSO_ANCHOR.TOP)


def callout(slide, x, y, w, h, text, sub=None, size=13, tone=ACCENT, fill=PEACH, align=PP_ALIGN.CENTER):
    rect(slide, x, y, w, h, fill, WHITE, 1.0, rounded=True, radius=0.08)
    paras = [(text, {"size": size, "bold": True, "color": tone, "align": align, "space_after": 3})]
    if sub:
        paras.append((sub, {"size": 10.5, "color": MUTED, "align": align, "spacing": 14}))
    txt(slide, paras, x + 0.24, y, w - 0.48, h, anchor=MSO_ANCHOR.MIDDLE, align=align)


def code_box(slide, x, y, w, h, lines, size=9.5, spacing=12.5, title=None):
    rect(slide, x, y, w, h, CARD, None, rounded=True, radius=0.04)
    top = y + 0.12
    if title:
        txt(slide, title, x + 0.18, top, w - 0.36, 0.26, size=10.5, bold=True, color=TEXT)
        top += 0.30
    paras = [(s, {"size": size, "font": MONO, "color": TEXT, "spacing": spacing}) for s in lines]
    txt(slide, paras, x + 0.18, top, w - 0.36, h - (top - y) - 0.1, font=MONO)


def table(slide, x, y, w, rows, col_w, row_h=0.42, size=11, head_size=11.5, first_col_bold=False,
          align_center_cols=(), hdr_fill=BLUE):
    """rows[0] 이 머리행. col_w 는 인치 목록(합이 w 와 같아야 함)."""
    n, m = len(rows), len(rows[0])
    gf = slide.shapes.add_table(n, m, Inches(x), Inches(y), Inches(w), Inches(row_h * n))
    tb = gf.table
    tb.first_row = True
    tb.horz_banding = False
    tb.vert_banding = False
    for j, cw_ in enumerate(col_w):
        tb.columns[j].width = Inches(cw_)
    for i in range(n):
        tb.rows[i].height = Inches(row_h)
        for j in range(m):
            cell = tb.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.10)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            val = rows[i][j]
            cell.fill.solid()
            if i == 0:
                cell.fill.fore_color.rgb = rgb(hdr_fill)
            else:
                cell.fill.fore_color.rgb = rgb(WHITE if i % 2 else SOFT)
            tf = cell.text_frame
            tf.word_wrap = True
            opt = {}
            if isinstance(val, tuple):
                val, opt = val
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER if (i == 0 or j in align_center_cols) else PP_ALIGN.LEFT
            r = p.add_run()
            r.text = val
            f = r.font
            f.name = opt.get("font", KO)
            f.size = Pt(head_size if i == 0 else opt.get("size", size))
            f.bold = True if i == 0 else opt.get("bold", first_col_bold and j == 0)
            f.color.rgb = rgb(WHITE if i == 0 else opt.get("color", TEXT))
    return gf


def numbered_rows(slide, items, x, y, w, dy=0.64, size=16, dark=True, circle=0.38, accent_from=None):
    """정리 장의 번호 원 + 문장. dark=True 면 어두운 배경용 색."""
    for i, s in enumerate(items):
        yy = y + i * dy
        is_acc = accent_from is not None and i + 1 >= accent_from
        if dark:
            fc, lc, tc, sc = (ACCENT, ACCENT, WHITE, WHITE) if is_acc else (CIRCLE_DK, CIRCLE_LN, CIRCLE_TX, SUM_TX)
        else:
            fc, lc, tc, sc = (ACCENT, ACCENT, WHITE, TEXT) if is_acc else (BLUE, BLUE, WHITE, TEXT)
        ellipse(slide, x, yy + 0.04, circle, fc, lc, 1.0)
        txt(slide, str(i + 1), x, yy + 0.04, circle, circle, size=13, bold=True, color=tc, align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE)
        txt(slide, s, x + circle + 0.20, yy, w - circle - 0.20, 0.46, size=size, color=sc, anchor=MSO_ANCHOR.MIDDLE)


# ───────────────────────────────────────────── 특수 페이지
def build_cover(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(s, DARK)
    txt(s, "Chapter 3", 10.30, 0.55, 2.5, 0.35, size=11, bold=True, color=WHITE, font=EN, align=PP_ALIGN.RIGHT, wrap=False)
    txt(s, "표적 궤적 모의", 4.30, 2.70, 8.50, 1.00, size=40, bold=True, color=WHITE, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, "표적을 정의하고, 0.1초마다 움직이고, 위경도로 그리기까지", 4.30, 3.72, 8.50, 0.42, size=16, color=ON_DARK, align=PP_ALIGN.RIGHT)
    txt(s, "개념편  ·  코드 구현은 다음 발표에서", 4.30, 4.22, 8.50, 0.36, size=12, color=ON_DARK_DIM, align=PP_ALIGN.RIGHT)
    txt(s, "레이다 시스템 소프트웨어 스터디  ·  Part 3", 6.80, 4.80, 6.00, 0.35, size=12, color=ON_DARK, align=PP_ALIGN.RIGHT)
    return s


def build_index(prs, rows, tail):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(s, BG)
    rect(s, 0, 0, 4.90, SH, DARK)
    txt(s, "INDEX", 1.25, 0.55, 2.2, 0.60, size=30, bold=True, color=WHITE, font=EN, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
    hline(s, 2.65, 12.00, 0.86, LINE, 0.75)
    txt(s, "목차", 0.45, 3.05, 4.00, 1.40, size=30, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    slide_number(s, TEXT)
    y0, dy = 1.30, 1.22
    for i, (num, name, desc) in enumerate(rows):
        y = y0 + i * dy
        txt(s, num, 5.55, y, 1.10, 0.62, size=34, bold=True, color=DARK, font=EN, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
        txt(s, name, 6.80, y, 3.55, 0.62, size=19, color=TEXT, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
        txt(s, desc, 10.45, y, 2.40, 0.62, size=11, color=MUTED, anchor=MSO_ANCHOR.MIDDLE, spacing=14)
        if i < len(rows) - 1:
            hline(s, 5.55, 12.80, y + 0.92, RULE_SOFT, 0.5)
    txt(s, tail, 5.55, 6.35, 7.30, 0.40, size=11.5, color=MUTED)
    return s


def build_divider(prs, num, title, items, foot):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(s, BG)
    rect(s, 0, 0, 4.90, SH, DARK)
    txt(s, num, 1.25, 0.55, 2.2, 0.60, size=30, bold=True, color=WHITE, font=EN, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
    hline(s, 2.65, 12.00, 0.86, LINE, 0.75)
    txt(s, title, 0.45, 3.05, 4.00, 1.40, size=30, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    slide_number(s, TEXT)
    n = len(items)
    dy = 0.70 if n <= 6 else 0.62
    y = 1.62 if n <= 6 else 1.45
    for i, it in enumerate(items):
        txt(s, str(i + 1), 5.55, y, 0.60, 0.50, size=16, bold=True, color=BLUE, font=EN, anchor=MSO_ANCHOR.MIDDLE)
        txt(s, it, 6.30, y, 6.50, 0.50, size=16, bold=True, color=TEXT, anchor=MSO_ANCHOR.MIDDLE)
        hline(s, 5.55, 12.80, y + dy - 0.12, RULE_SOFT, 0.5)
        y += dy
    txt(s, foot, 5.55, 6.42, 7.25, 0.40, size=11, color=MUTED)
    return s


def build_summary(prs, items, tail, accent_from=6):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(s, DARK)
    hline(s, 2.00, 12.83, 0.86, LINE, 0.75)
    txt(s, "정리", 0.95, 1.22, 11.5, 0.58, size=30, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    numbered_rows(s, items, 0.90, 2.00, 11.6, dy=0.62, size=16, dark=True, accent_from=accent_from)
    txt(s, tail, 0.95, 6.46, 11.5, 0.42, size=13, bold=True, color=WHITE)
    hline(s, M, SW - M, FOOT_LINE_Y, LINE, 0.75)
    slide_number(s, WHITE)
    return s


# ───────────────────────────────────────────── 본문 페이지들
def s03(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    content_chrome(s, "00.  과제 3이 묻는 것", "표적을 정의하고, 0.1초마다 움직이고, 위경도로 그려라")
    pic(s, "fig01_overview.png", CX0, 1.50, w=7.4)
    x, w, y = 8.20, 4.52, 1.50
    specs = [("① 표적 구조체를 정의한다", ["표적 최대 10개 · 표적마다 기동 최대 30개", "초기 위치는 위경도·고도, 속도는 동체 속력 하나로"], BLUE),
             ("② 위치·속도는 ECEF 에서 갱신한다", ["동체 속력을 ECEF 성분으로 바꾸는 방법도 고민 (추가 항목)"], BLUE),
             ("③ 60초 동안 0.1초 간격으로", ["플랫폼 하나 (정지) · 대함 표적 · 대공 표적"], ACCENT),
             ("④ 위경도로 출력하고 그림으로", ["플랫폼과 표적의 궤적을 지도 위에"], ACCENT)]
    hs = [1.02, 0.86, 0.86, 0.86]
    for (t, ls, tone), h in zip(specs, hs):
        card(s, x, y, w, h, t, ls, tone=tone, title_size=12.5, body_size=10.5, pad=0.20, line_spacing=14)
        y += h + 0.10
    callout(s, CX0, 5.42, CW, 1.32,
            "한 줄로 :  표적을 정의하고,  0.1초마다 움직이고,  위경도로 그려라",
            "표적마다 600개의 점이 찍힌다.  지난 과제의 좌표변환 함수는 그대로 쓰고, 그 위에 시간과 기동을 얹는다.")
    return s


def s04(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    content_chrome(s, "00.  과제 2와 무엇이 다른가", "한 점을 한 번 변환하던 일이, 매 0.1초 600번 반복하는 일이 된다")
    pic(s, "fig02_p2_vs_p3.png", CX0 + 0.30, 1.50, w=11.5)
    y = 4.62
    card(s, CX0, y, 5.96, 1.12, "같은 것 — 좌표변환 함수 넷은 그대로",
         ["LLA ↔ ECEF,  동체 → NED,  NED → ECEF.   제공된 참고 코드에 그대로 있다.",
          "안테나 좌표계만 이번엔 등장하지 않는다 (재는 쪽이 아니라 만드는 쪽)."],
         tone=BLUE, body_size=10.5, line_spacing=14)
    card(s, CX0 + 6.14, y, 5.96, 1.12, "새로운 것 — 속도 · 시간 · 기동",
         ["속도 :  화살표는 위치와 다르게 변환한다 (회전만)",
          "시간 :  새 위치 = 지금 위치 + 속도 × 0.1 s.    기동 :  몇 g 로 어느 축을 도는가"],
         tone=ACCENT, fill=PEACH, body_size=10.5, line_spacing=14)
    callout(s, CX0, 5.90, CW, 0.84, "과제 2의 변환 체인을 재료로 쓰고,  그 위에 시간 루프와 기동 모델을 얹는다", fill=CARD, tone=TEXT, size=12.5)
    return s


def s06(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    content_chrome(s, "01.  이번에 쓰는 좌표계 넷", "입력·출력은 위경도,  계산은 ECEF,  수평면은 NED,  속력은 동체")
    pic(s, "fig03_four_frames.png", CX0, 1.50, w=8.4)
    x, w = 9.28, 3.44
    card(s, x, 1.50, w, 3.95, "이번 과제에서의 역할",
         [("LLA", {"size": 11.5, "bold": True, "color": BLUE, "space_before": 2}),
          ("초기값을 받고, 매 스텝 결과를 적는 형식", {"size": 10.5, "color": TEXT, "space_after": 6}),
          ("ECEF", {"size": 11.5, "bold": True, "color": BLUE}),
          ("위치와 속도를 600번 갱신하는 자리", {"size": 10.5, "color": TEXT, "space_after": 6}),
          ("NED", {"size": 11.5, "bold": True, "color": GREEN}),
          ("표적 자리의 수평면. 자세각의 기준이자 속도 방향을 만드는 자리", {"size": 10.5, "color": TEXT, "space_after": 6}),
          ("동체", {"size": 11.5, "bold": True, "color": ACCENT}),
          ("속력 V 가 적힌 자리. 앞으로 V, 옆·아래는 0", {"size": 10.5, "color": TEXT})],
         tone=TEXT, title_size=13)
    callout(s, CX0, 5.72, CW, 1.02,
            "원점이 표적 자신인 NED 가 하나 더 필요하다 — 자세각과 수평면은 표적이 서 있는 자리 기준이기 때문",
            "지난 과제의 NED 는 플랫폼 자리였다.  이번 NED 는 표적이 움직이면 같이 따라간다.", fill=CARD, tone=TEXT, size=12.5)
    return s


def s07(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    content_chrome(s, "01.  왜 ECEF 에서 움직이나", "위경도는 각도라 더할 수 없고,  ECEF 는 미터라 더하면 끝난다")
    pic(s, "fig04_why_ecef.png", CX0, 1.50, w=7.4)
    x, w, y = 8.20, 4.52, 1.50
    card(s, x, y, w, 1.12, "① 위경도는 각도다", ["위도 1° 는 111 km, 경도 1° 는 94 km (32°N).", "각도에 미터를 더할 수 없다."], tone=ACCENT, fill=PEACH, body_size=10.5, line_spacing=14)
    card(s, x, y + 1.24, w, 1.12, "② ECEF 는 미터다", ["세 축이 전부 미터라 벡터를 그대로 더한다.", "p' = p + v × Δt,  한 줄."], tone=BLUE, body_size=10.5, line_spacing=14)
    card(s, x, y + 2.48, w, 1.26, "③ 플랫폼과 무관한 절대 좌표", ["플랫폼이 움직이거나 흔들려도 표적 궤적은 그대로.", "플랫폼 기준 NED 에서 움직이면 플랫폼 흔들림이 표적에 섞인다."], tone=BLUE, body_size=10.5, line_spacing=14)
    callout(s, CX0, 5.42, CW, 1.32,
            "흐름 :  위경도로 받는다  →  ECEF 로 바꾼다  →  600번 내내 ECEF 에서 더한다  →  기록할 때만 위경도로 되돌린다",
            "과제 요구 (2) \"표적 상태는 ECEF 좌표계에서 갱신\" 의 이유가 이것이다.")
    return s


def s08(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    content_chrome(s, "01.  위치와 속도는 다르게 변환한다", "위치는 원점이 있어야 말이 되고,  속도는 화살표 하나면 된다")
    chips(s, "속도 변환")
    pic(s, "fig05_pos_vs_vel.png", CX0 + 0.30, 1.97, w=11.5)
    y = 5.02
    card(s, CX0, y, 5.96, 0.94, "위치 :  p_new = R · p + t", ["회전하고 원점 차이만큼 옮긴다.  지난 과제에서 계속 하던 그 일."], tone=BLUE, body_size=10.5)
    card(s, CX0 + 6.14, y, 5.96, 0.94, "속도 :  v_new = R · v", ["축이 돌아간 만큼만 돌린다.  t 를 더하면 그 순간 속도가 아니게 된다."], tone=ACCENT, fill=PEACH, body_size=10.5)
    callout(s, CX0, 6.08, CW, 0.66, "과제의 \"추가 고민\" — 동체 속도 V_heading 을 ECEF 성분으로 바꾸는 방법 — 의 답이 오른쪽 한 줄이다", fill=CARD, tone=TEXT, size=12)
    return s


def s09(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    content_chrome(s, "01.  동체 속도를 ECEF 속도로 — 회전 두 번", "자세각으로 한 번, 표적 자신의 위도·경도로 한 번.  평행이동은 없다")
    chips(s, "속도 변환")
    pic(s, "fig06_vel_chain.png", CX0 + 0.30, 1.97, w=11.5)
    y, h = 4.32, 1.42
    card(s, CX0, y, 3.90, h, "① 동체 :  [ V, 0, 0 ]", ["표적은 앞으로만 간다.", "속력 하나만 받아도 벡터가 정해지는 이유."], tone=ACCENT, fill=PEACH, body_size=10.5, line_spacing=14)
    card(s, CX0 + 4.10, y, 3.90, h, "② 자세각으로 회전", ["과제 2의 C_ned←body (roll · pitch · yaw) 그대로.", "yaw 180° → 북쪽 성분이 −200, 즉 남쪽으로 200."], tone=GREEN, fill=MINT, body_size=10.5, line_spacing=14)
    card(s, CX0 + 8.20, y, 3.90, h, "③ 위도·경도로 회전", ["과제 2의 C_ecef←ned 그대로.  단, 위경도는 표적 자신의 것.", "참고 함수는 플랫폼 위치를 더하므로 그 인자에 0 을 넣는다."], tone=BLUE, body_size=10.5, line_spacing=14)
    callout(s, CX0, 5.90, CW, 0.84,
            "검산 :  결과 벡터의 크기는 항상 V 와 같아야 한다  →  200.000 m/s.    플랫폼 위치를 잘못 더하면 6,372 km/s 가 나온다",
            size=12)
    return s


def s10(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    content_chrome(s, "01.  시간 적분 — 새 위치 = 지금 위치 + 속도 × 0.1 초", "한 스텝 안에서는 속도가 일정하다고 본다.  그래서 곱셈 한 번이면 정확하다")
    chips(s, "위치 적분")
    pic(s, "fig07_euler.png", CX0, 1.97, w=7.4)
    x, w, y = 8.20, 4.52, 1.97
    card(s, x, y, w, 1.08, "왜 이렇게 단순해도 되나", ["0.1초 동안은 방향도 크기도 안 바뀐다고 본다.", "곡선은 매 스텝 방향을 다시 정해서 만든다."], tone=BLUE, body_size=10.5, line_spacing=14)
    card(s, x, y + 1.18, w, 1.08, "룬지쿠타 4차는?", ["가속도가 위치의 함수인 궤도 문제용.", "이번엔 속도를 우리가 정해 주므로 필요 없다."], tone=BLUE, body_size=10.5, line_spacing=14)
    card(s, x, y + 2.36, w, 0.86, "숫자 감", ["대공 한 스텝 20 m · 60초 12 km,   대함 한 스텝 3 m · 60초 1.8 km"], tone=ACCENT, fill=PEACH, body_size=10.5, line_spacing=14)
    callout(s, CX0, 5.42, CW, 1.32,
            "p(t + Δt) = p(t) + v(t) · Δt      ECEF 의 x, y, z 세 성분에 각각",
            "Δt = 0.1 s,  600 스텝.   v(t) 는 직전 단계(속도 변환) 에서 매 스텝 새로 만든 값이다.", fill=CARD, tone=TEXT, size=14)
    return s


def s11(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    content_chrome(s, "01.  지구는 둥글다 — 60초 직진하면 11 m 떠오른다", "처음 속도를 그대로 쓰면 접선을 따라 지구에서 멀어진다.  매 스텝 수평면을 다시 잡아야 한다")
    chips(s, "LLA 변환")
    pic(s, "fig08_curvature.png", CX0, 1.97, w=7.4)
    x, w, y = 8.20, 4.52, 1.97
    rows = [["실험 (대공 표적 60초 직진)", "60초 뒤 고도"],
            ["A  처음 ECEF 속도를 60초 내내 그대로", ("311.3 m", {"bold": True, "color": ACCENT})],
            ["B  매 스텝 현재 위경도로 NED 를 다시 잡음", ("300.0 m", {"bold": True, "color": GREEN})]]
    table(s, x, y, w, rows, [3.12, 1.40], row_h=0.50, size=10.5, head_size=10.5, align_center_cols=(1,))
    card(s, x, y + 1.66, w, 1.02, "왜 11 m 인가", ["접선을 따라 d 만큼 가면 땅은 d² / 2R 만큼 내려간다.", "12,000² / (2 × 6,371,000) = 11.3 m.   실험값과 일치."], tone=BLUE, body_size=10.5, line_spacing=14)
    card(s, x, y + 2.78, w, 0.96, "그래서 5단계 LLA 변환은", ["출력용만이 아니라, 다음 스텝의 수평면(NED 축) 을 만드는 기준이다."], tone=ACCENT, fill=PEACH, body_size=10.5, line_spacing=14)
    callout(s, CX0, 5.90, CW, 0.84, "위경도를 구해야 NED 축이 나오고,  NED 축이 있어야 속도 방향이 나온다  —  이 순서를 지키면 수평 비행이 수평으로 유지된다", size=12)
    return s


def s12(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    content_chrome(s, "01.  한 스텝의 순서 — 다섯 단계", "표적 하나가 0.1초 뒤로 가는 일.  이 순서를 600번, 표적 수만큼 반복한다")
    pic(s, "fig09_step_loop.png", CX0 + 0.30, 1.50, w=11.5)
    y, h = 4.62, 1.10
    w5 = (CW - 4 * 0.14) / 5
    specs = [("① 에 필요한 것", "기동 목록,  지금 시각 t", ACCENT, PEACH),
             ("② 에 필요한 것", "각속도 ω = n·g / V,  Δt", ACCENT, PEACH),
             ("③ 에 필요한 것", "자세각 3개,  표적 위경도", GREEN, MINT),
             ("④ 에 필요한 것", "ECEF 위치·속도,  Δt = 0.1 s", BLUE, CARD),
             ("⑤ 에 필요한 것", "ECEF → LLA 변환 함수", BLUE, CARD)]
    for i, (t, l, tone, fill) in enumerate(specs):
        card(s, CX0 + i * (w5 + 0.14), y, w5, h, t, [l], tone=tone, fill=fill, title_size=11.5, body_size=10.5, pad=0.18)
    callout(s, CX0, 5.90, CW, 0.84, "표적 10개 × 600 스텝 = 6,000번.   순서가 중요한 건 ③ ④ ⑤ — 자세가 정해져야 속도, 속도가 있어야 위치, 위치가 있어야 새 수평면", fill=CARD, tone=TEXT, size=12)
    return s


def s14(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    content_chrome(s, "02.  기동이란 — 하중배수 g", "과제의 Gravity Value 는 \"몇 g 로 도는가\".  속력은 그대로, 방향만 바뀐다")
    chips(s, "자세 갱신")
    pic(s, "fig10_gload.png", CX0, 1.97, w=7.4)
    x, w, y = 8.20, 4.52, 1.97
    card(s, x, y, w, 1.10, "하중배수 n", ["안쪽으로 당기는 가속도가 중력가속도 g 의 몇 배인가.", "a = n × g.   과제의 Gravity Value 가 이 n 이다."], tone=ACCENT, fill=PEACH, body_size=10.5, line_spacing=14)
    card(s, x, y + 1.20, w, 1.10, "가속도는 속도에 수직", ["속도 방향으로는 밀지 않고 옆으로만 당긴다.", "그래서 속력은 안 변하고 방향만 변한다."], tone=BLUE, body_size=10.5, line_spacing=14)
    card(s, x, y + 2.40, w, 1.20, "숫자 감", ["200 m/s 표적이 3 g 로 돌면 1초에 8.4°,", "30 m/s 배는 0.1 g 만 돼도 1초에 1.9°."], tone=BLUE, body_size=10.5, line_spacing=14)
    callout(s, CX0, 5.72, CW, 1.02, "g 값 하나로 표적이 얼마나 빨리 도는지가 정해진다  →  다음 장의 등속 원운동 공식", size=12.5)
    return s


def s15(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    content_chrome(s, "02.  원운동 공식 — 각속도와 선회 반경", "고등학교 물리의 등속 원운동 그대로.  같은 g 라도 느린 표적이 훨씬 빨리 돈다")
    chips(s, "자세 갱신")
    pic(s, "fig11_circle.png", CX0, 1.97, w=6.0)
    x, w, y = 6.90, 5.82, 1.97
    rows = [["표적", "하중배수", "각속도", "선회 반경", "비고"],
            ["대공 200 m/s", "1 g", "2.81 °/s", "4,079 m", "180° 에 64 s"],
            ["대공 200 m/s", ("3 g", {"bold": True, "color": ACCENT}), ("8.43 °/s", {"bold": True, "color": ACCENT}), ("1,360 m", {"bold": True, "color": ACCENT}), "180° 에 21 s"],
            ["대공 200 m/s", "5 g", "14.05 °/s", "816 m", "180° 에 13 s"],
            ["대함 30 m/s", "0.1 g", "1.87 °/s", "918 m", "90° 에 48 s"]]
    table(s, x, y, w, rows, [1.50, 0.90, 1.05, 1.05, 1.32], row_h=0.44, size=11, head_size=11, align_center_cols=(1, 2, 3, 4))
    card(s, x, y + 2.42, w, 1.18, "각속도가 속력에 반비례한다", ["ω = n·g / V.   배는 작은 g 로도 잘 돌고, 빠른 항공기는 큰 g 가 있어야 돈다.", "그래서 배 기동에는 0.1 g 같은 값을 넣어야 그럴듯하다."], tone=BLUE, body_size=10.5, line_spacing=14)
    callout(s, CX0, 5.72, CW, 1.02, "프로그램에서는 한 줄 :  ω = n·g / V  를 구해 놓고,  매 스텝 자세각에  ω × 0.1 s  를 더한다   (3 g 면 한 스텝에 0.84°)", size=12.5)
    return s


def s16(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    content_chrome(s, "02.  세 축 — Roll · Yaw · Pitch", "TurnType 은 동체의 어느 축을 돌리는가.  Roll 만 속도 방향을 바꾸지 않는다")
    chips(s, "자세 갱신")
    pic(s, "fig12_axes.png", CX0, 1.97, w=8.4)
    x, w = 9.28, 3.44
    rows = [["TurnType", "축", "효과"],
            ["0", "—", "직진"],
            ["1  Roll", "x (앞뒤)", "자세만 기움"],
            ["2  Yaw", "z (위아래)", "좌우 선회"],
            ["3  Pitch", "y (날개)", "상승 · 강하"]]
    table(s, x, 1.97, w, rows, [1.04, 1.06, 1.34], row_h=0.40, size=10.5, head_size=10.5, align_center_cols=(0, 1, 2))
    card(s, x, 4.12, w, 1.25, "Roll 은 준비 동작", ["롤 90° 뒤에 피치로 당기면 위가 아니라 옆으로 돈다.", "실제 비행기의 선회 : 기울이고, 당긴다."], tone=ACCENT, fill=PEACH, title_size=12, body_size=10, line_spacing=13)
    callout(s, CX0, 5.62, CW, 1.12,
            "축을 돌린 뒤 속도가 어느 쪽인지는 따로 계산하지 않는다  —  자세각만 바꿔 두면 9장의 회전 두 번이 새 속도 방향을 만든다",
            "세 축의 각속도는 모두 ω = n·g / V 로 같게 둔다 (정의의 문제, 발표에서 명시).", fill=CARD, tone=TEXT, size=12)
    return s


def s17(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    content_chrome(s, "02.  기동 시간표 — StartTime · EndTime", "기동 하나 = { g, TurnType, StartTime, EndTime }.  표적마다 최대 30개")
    chips(s, "기동 판단")
    pic(s, "fig13_timeline.png", CX0 + 0.30, 1.97, w=11.5)
    y, h = 4.72, 1.20
    card(s, CX0, y, 3.90, h, "구간 안이면 적용, 밖이면 직진", ["StartTime ≤ t < EndTime.", "없으면 TurnType 0 으로 본다."], tone=BLUE, body_size=10.5, line_spacing=14)
    card(s, CX0 + 4.10, y, 3.90, h, "정해 둘 규칙 셋", ["부호 : n < 0 이면 반대 방향.   겹침 : 먼저 적힌 것.", "끝난 뒤 : 그 순간 자세를 유지하고 직진."], tone=ACCENT, fill=PEACH, body_size=10.5, line_spacing=14)
    card(s, CX0 + 8.20, y, 3.90, h, "최대 30개", ["배열 30칸 + 개수 하나.", "과제 조건에는 기동 목록이 없다 → 예시를 우리가 정한다."], tone=BLUE, body_size=10.5, line_spacing=14)
    callout(s, CX0, 6.06, CW, 0.68, "예시 기동을 넣지 않으면 직진 두 개만 그려진다 — 기동 30개짜리 구조체의 의미가 보이도록 발표용 시나리오를 하나 더 만든다", size=11.5)
    return s


STRUCT_LINES = [
    "typedef struct {",
    "    FLOAT64     dGravity;       // 하중배수 [g], 부호 = 방향",
    "    INT32       nTurnType;      // 0 없음  1 Roll  2 Yaw  3 Pitch",
    "    FLOAT64     dStartTime;     // [s]",
    "    FLOAT64     dEndTime;       // [s]",
    "} ST_Maneuver;",
    "",
    "typedef struct {",
    "    /* 입력 — 사람이 적는 칸 */",
    "    ST_Lla      stInitLla;      // 위도·경도·고도",
    "    FLOAT64     dVheading;      // 동체 속력 [m/s]",
    "    ST_Attitude stInitAtt;      // roll · pitch · yaw",
    "    INT32       nManeuverCnt;   // <= MAX_MANEUVER (30)",
    "    ST_Maneuver astManeuver[MAX_MANEUVER];",
    "    /* 상태 — 프로그램이 갱신 */",
    "    ST_Vec3     stPosEcef;      // ECEF 위치 [m]",
    "    ST_Vec3     stVelEcef;      // ECEF 속도 [m/s]",
    "    ST_Attitude stAtt;          // 현재 자세",
    "    ST_Lla      stLla;          // 현재 위경도 (출력용)",
    "} ST_Target;",
]

SCENARIO_LINES = [
    "typedef struct {",
    "    INT32     nTargetCnt;              // <= MAX_TARGET (10)",
    "    ST_Target astTarget[MAX_TARGET];",
    "    ST_Target stPlatform;              // 속도 0, 기동 0개",
    "    FLOAT64   dSimTime;                // 60 s",
    "    FLOAT64   dDt;                     // 0.1 s",
    "} ST_Scenario;",
]


def s19(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    content_chrome(s, "03.  표적 구조체 — 표적 카드 한 장", "입력 칸은 사람이 적고, 상태 칸은 프로그램이 갱신한다.  구조체 모양에 과제 요구가 보이도록")
    pic(s, "fig14_target_card.png", CX0, 1.50, w=6.4)
    code_box(s, 7.30, 1.50, 5.42, 4.30, STRUCT_LINES, size=9.0, spacing=11.6, title="C 구조체  (ST_ 접두사 · typedef, 코딩 규칙대로)")
    callout(s, CX0, 5.98, CW, 0.76, "위 다섯 줄이 입력, 아래 네 줄이 상태  —  \"초기값은 위경도, 갱신은 ECEF\" 라는 과제 요구가 구조체 모양에 그대로 드러난다", fill=CARD, tone=TEXT, size=12)
    return s


def s20(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    content_chrome(s, "03.  시나리오와 전체 흐름", "카드가 여러 장 모이면 시나리오.  플랫폼도 속도 0 인 카드 한 장으로 둔다")
    pic(s, "fig15_pipeline.png", CX0 + 0.30, 1.50, w=11.5)
    y = 4.42
    code_box(s, CX0, y, 5.90, 2.32, SCENARIO_LINES, size=9.5, spacing=13, title="시나리오 구조체")
    card(s, CX0 + 6.10, y, 6.00, 1.06, "플랫폼도 표적 카드로", ["속도 0, 기동 0개인 표적일 뿐.  갱신 함수가 하나로 통일되고,", "플랫폼이 움직이는 시나리오가 와도 카드 값만 바꾸면 된다."], tone=BLUE, body_size=10.5, line_spacing=14)
    card(s, CX0 + 6.10, y + 1.18, 6.00, 1.14, "CSV 한 줄 = 한 표적의 한 스텝", ["t,  id,  lat,  lon,  alt   (+ 플랫폼 기준 N · E · D,  거리,  고각)", "60초 × 10 Hz × 표적 수.   그림은 파이썬으로 따로 그린다."], tone=GREEN, fill=MINT, body_size=10.5, line_spacing=14)
    return s


def s21(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    content_chrome(s, "03.  시뮬레이션 조건 — 플랫폼 하나, 표적 둘", "대공 표적은 정면으로 접근하고, 대함 표적은 정북 방향을 가로지른다")
    pic(s, "fig16_scenario_map.png", CX0, 1.50, w=5.6)
    x, w = 6.50, 6.22
    rows = [["항목", "플랫폼", "대함 표적 (1)", "대공 표적 (2)"],
            ["위도 · 경도", "32.0°N  126.0°E", "32.125°N  126.03°E", "32.12°N  126.0°E"],
            ["고도", "0 m", "0 m", "300 m"],
            ["속력 (동체 앞)", "0 m/s", "30 m/s", "200 m/s"],
            ["자세각 (r · y · p)", "—", ("0° · 270° · 0° (서쪽)", {"size": 10}), ("0° · 180° · 0° (남쪽)", {"size": 10})],
            ["시간", "60 s,  간격 0.1 s  (600 스텝)", "", ""]]
    gf = table(s, x, 1.50, w, rows, [1.56, 1.34, 1.66, 1.66], row_h=0.44, size=10.5, head_size=10.5, first_col_bold=True, align_center_cols=(1, 2, 3))
    gf.table.cell(5, 1).merge(gf.table.cell(5, 3))
    card(s, x, 4.36, w, 1.44, "플랫폼 기준으로 바꿔 보면  (참고 코드로 계산)",
         ["대공 :  정북 13.3 km,  고각 1.2°  →  정확히 우리 쪽으로 온다 (정면 접근)",
          "대함 :  북 13.9 km · 동 2.8 km,  거리 14.1 km  →  서쪽으로 가며 정북을 가로지른다",
          "두 표적의 출발점은 550 m · 2.8 km 차이로 아주 가깝다"],
         tone=BLUE, body_size=10.5, line_spacing=14)
    callout(s, CX0, 5.98, CW, 0.76, "같은 자리에서 출발해 하나는 남쪽으로 빠르게, 하나는 서쪽으로 천천히  —  그림에서 대비가 분명하도록 짜인 조건", fill=CARD, tone=TEXT, size=12)
    return s


def s22(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    content_chrome(s, "03.  예상 결과 — 정면 접근과 가로지르기", "코드를 돌리기 전에 결과가 어떻게 나와야 하는지 먼저 그려 둔다")
    pic(s, "fig17_expected.png", CX0, 1.50, w=5.6)
    x, w, y = 6.50, 6.22, 1.50
    card(s, x, y, w, 1.28, "대공 — 정면 접근", ["13.3 km → 1.3 km,  고각 1.2° → 13.4°.   60초 뒤엔 거의 머리 위.", "고도는 300 m 그대로여야 한다 (11장의 곡률 처리가 맞았다면)."], tone=ACCENT, fill=PEACH, body_size=10.5, line_spacing=14)
    card(s, x, y + 1.40, w, 1.06, "대함 — 가로지르기", ["서쪽으로 1.8 km.  대공 표적 출발점 바로 북쪽을 지나 동쪽 1.0 km 지점에서 끝."], tone=GREEN, fill=MINT, body_size=10.5, line_spacing=14)
    card(s, x, y + 2.58, w, 1.72, "기동 예시를 넣으면", ["대공 표적에 20 ~ 30 s 동안 3 g 우선회 :", "10초에 84° 돌아 서쪽을 보고, 그 뒤 직진.  반경 1.36 km 원호가 그려진다.", "조건대로 돌린 것과 기동을 넣은 것을 나란히 보여 준다."], tone=BLUE, body_size=10.5, line_spacing=14)
    callout(s, CX0, 5.98, CW, 0.76, "결과가 어때야 하는지 미리 알아야, 나온 결과가 맞는지 판단할 수 있다  →  다음 장의 검증 계획", size=12)
    return s


def s23(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    content_chrome(s, "03.  검증 계획 — 무엇을 확인할 것인가", "좌표변환은 틀려도 그럴듯한 값이 나온다 (과제 2의 교훈).  점이 600개면 눈으로는 더 못 잡는다")
    rows = [["항목", "어떻게 확인하나", "기대값", "참고 코드로 미리 확인한 값"],
            ["① 직진 거리", "속력 × 시간과 비교", "200 × 60 = 12,000 m", ("13,307 m → 1,307 m  (12,000 m)", {"color": GREEN, "bold": True})],
            ["② 속력 보존", "매 스텝 |v_ecef| 와 입력 속력 비교", "200.000 m/s", ("200.000 m/s", {"color": GREEN, "bold": True})],
            ["③ 고도 유지", "수평 비행 60초 뒤 고도", "300 m ± 0.1", ("300.02 m", {"color": GREEN, "bold": True})],
            ["④ 선회 반경", "궤적에서 반경을 재서 V² / (n·g) 와 비교", "3 g : 1,360 m", "코드 완성 후"],
            ["⑤ 위경도 왕복", "LLA → ECEF → LLA 오차", "1 mm 미만", ("0.03 mm", {"color": GREEN, "bold": True})]]
    table(s, CX0, 1.50, CW, rows, [1.70, 3.60, 2.60, 4.20], row_h=0.50, size=11, head_size=11, first_col_bold=True, align_center_cols=(2, 3))
    y = 4.72
    card(s, CX0, y, 5.96, 1.10, "왜 미리 정하나", ["코드가 내놓을 값을 먼저 알고 있어야 코드가 맞는지 알 수 있다.", "①②③⑤ 는 참고 코드만으로 이미 확인했고, ④ 는 기동 코드가 붙으면 잰다."], tone=BLUE, body_size=10.5, line_spacing=14)
    card(s, CX0 + 6.14, y, 5.96, 1.10, "한 번 틀리면 바로 드러나는 실수들", ["경도·위도 인자 순서 바꿈 → 북위 54° 동경 212°.   속도에 플랫폼 위치 더함 → 6,372 km/s.", "수평면을 한 번만 잡음 → 고도 311 m."], tone=ACCENT, fill=PEACH, body_size=10.5, line_spacing=14)
    callout(s, CX0, 5.98, CW, 0.76, "이 다섯 칸이 채워지면 궤적이 \"그림\" 이 아니라 \"결과\" 가 된다", size=12.5)
    return s


def s24(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    content_chrome(s, "03.  결과 출력과 그림 — 위경도로 적고 세 장으로 그린다", "직진만 가정하고 미리 그려 본 모양.  실제 시뮬레이션 결과로 갈아 끼울 자리")
    pic(s, "fig18_charts_mock.png", CX0 + 0.30, 1.50, w=11.5)
    y, h = 4.66, 1.24
    card(s, CX0, y, 3.90, h, "① 평면 궤적", ["경도 가로, 위도 세로.  축 비율을 1 / cos 32° 로", "맞춰야 선회 원이 타원으로 찌그러지지 않는다."], tone=BLUE, body_size=10.5, line_spacing=14)
    card(s, CX0 + 4.10, y, 3.90, h, "② 고도 – 시간", ["대공 300 m, 대함 0 m 에서 평평해야 한다.", "곡률 처리(11장) 의 검증 그림이기도 하다."], tone=GREEN, fill=MINT, body_size=10.5, line_spacing=14)
    card(s, CX0 + 8.20, y, 3.90, h, "③ 거리 · 고각 – 시간", ["플랫폼 기준.  레이다가 실제로 보게 될 값.", "대공 13.3 → 1.3 km,  고각 1.2° → 13.4°."], tone=ACCENT, fill=PEACH, body_size=10.5, line_spacing=14)
    callout(s, CX0, 6.04, CW, 0.70, "CSV (t, id, lat, lon, alt, …)  →  Python matplotlib  →  그림 세 장.   과제 요구 (시뮬레이션 결과를 LLA 로 출력 후 그림으로 표현) 에 대응", fill=CARD, tone=TEXT, size=11.5)
    return s


def s25(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    content_chrome(s, "03.  다음 단계 — 구현 순서", "오늘 본 순서 그대로 코드로 옮긴다.  손계산 검산값을 먼저 만들어 둔다")
    items = ["구조체 셋 정의 :  ST_Maneuver · ST_Target · ST_Scenario  (코딩 규칙대로 ST_ 접두사, typedef)",
             "초기화 :  위경도 → ECEF  (참고 코드 f_Trans_Lla_To_Ecef, 경도가 첫째 인자)",
             "속도 변환 함수 :  동체 → NED → ECEF, 회전 두 번, 플랫폼 위치 인자는 0",
             "기동 판단과 자세 갱신 :  구간 검사, ω = n·g / V, 자세각 += ω × Δt",
             "루프와 CSV :  60 s / 0.1 s, 표적마다 5단계, 한 스텝에 한 줄",
             "그림 세 장 :  Python matplotlib, 축 비율 1 / cos 32°",
             "검증 표 채우기 :  거리 · 속력 · 고도 · 반경 · 왕복"]
    numbered_rows(s, items, CX0 + 0.10, 1.62, 8.30, dy=0.60, size=12.5, dark=False, circle=0.36)
    x, w = 9.28, 3.44
    card(s, x, 1.62, w, 2.32, "일정 (나흘)",
         [("1일차   수식 손계산 · 검산값", {"size": 11, "color": TEXT, "spacing": 16}),
          ("2일차   구조체 · 루프 · 속도 변환", {"size": 11, "color": TEXT, "spacing": 16}),
          ("3일차   그림 · 검증 표", {"size": 11, "color": TEXT, "spacing": 16}),
          ("4일차   발표 자료 (결과편)", {"size": 11, "color": TEXT, "spacing": 16})],
         tone=BLUE)
    card(s, x, 4.06, w, 1.66, "손계산을 먼저 하는 이유", ["코드가 내놓을 값을 미리 알고 있어야", "코드가 맞는지 알 수 있다 (23장)."], tone=ACCENT, fill=PEACH, body_size=10.5, line_spacing=14)
    callout(s, CX0, 5.98, CW, 0.76, "다음 발표 :  코드와 실행 결과, 검증 표, 기동 시나리오 비교", fill=CARD, tone=TEXT, size=12.5)
    return s


def s27(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    content_chrome(s, "부록 A.  참고 코드의 함수와 함정", "제공된 CoordinateTransform.c 에서 이번에 쓰는 것과, 컴파일해서 확인한 주의점")
    rows = [["함수", "이번 과제의 용도", "주의"],
            ["f_Trans_Lla_To_Ecef (Lon, Lat, Alt)", "초기 위치 입력", ("경도가 첫째 인자.  바꾸면 북위 54° 동경 212°", {"color": ACCENT, "bold": True})],
            ["f_Trans_Ecef_To_Lla (X, Y, Z)", "매 스텝 결과 출력 + 다음 수평면 기준", "5회 고정 반복.  우리 위치에서 오차 0.03 mm"],
            ["f_Trans_Body_To_Ned (…, Roll, Yaw, Pitch)", "동체 속력 → NED 속도", ("인자 순서가 Roll, Yaw, Pitch", {"color": ACCENT, "bold": True})],
            ["f_Trans_Ned_To_Ecef (…, Pf_X, Pf_Y, Pf_Z, Lat, Lon)", "NED 속도 → ECEF 속도", ("속도에는 Pf_X·Y·Z = 0.  아니면 6,372 km/s", {"color": ACCENT, "bold": True})],
            ["f_Trans_Ecef_To_Ned (…)", "플랫폼 기준 거리 · 고각 (결과 그림용)", "각도는 전부 라디안.  DEG2RAD 매크로 직접 정의"],
            ["Ant 계열 4개 · Enu 계열 2개 · PfCompensation", "사용 안 함", "표적을 만드는 쪽이라 안테나 좌표계가 없다"]]
    table(s, CX0, 1.50, CW, rows, [4.30, 3.20, 4.60], row_h=0.50, size=10.5, head_size=11)
    callout(s, CX0, 5.20, CW, 1.54,
            "Body → NED 행렬은 Rz(yaw) · Ry(pitch) · Rx(roll) 과 소수점 6자리까지 일치  —  과제 2 의 coord_frames.c 와 같은 규약",
            "그 밖에 :  한글 주석이 UTF-8 이라 MSVC 에서 /utf-8 옵션 필요.   matrixCalcLib 은 9×9 고정 행렬 (과제 2 와 같음).   상수 G_FORCE = 9.80665 가 Define.h 에 있다.",
            fill=CARD, tone=TEXT, size=12)
    return s


def s28(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    content_chrome(s, "부록 B.  용어 정리", "발표에 나온 말들.  질문이 들어오면 이 표를 보면서 답한다")
    rows = [["용어", "뜻", "이번 과제에서"],
            ["하중배수 n (g)", "안쪽으로 당기는 가속도가 중력가속도의 몇 배인가", "기동의 Gravity Value.  a = n·g"],
            ["각속도 ω", "1초에 몇 도(라디안) 도는가", "ω = n·g / V.  매 스텝 자세각 += ω × 0.1 s"],
            ["오일러 적분", "새 값 = 지금 값 + 변화율 × 시간 간격", "p ← p + v × Δt.  스텝 안 속도 일정이라 정확"],
            ["자세각 roll · pitch · yaw", "동체 축이 수평면(NED) 에 대해 돌아간 세 각", "표적의 진행 방향.  yaw 180° = 남쪽"],
            ["ECEF", "지구 중심 원점, 지구와 함께 도는 직교좌표 [m]", "위치·속도를 갱신하는 자리"],
            ["NED", "한 지점의 북 · 동 · 아래 수평 좌표", "표적 자신의 자리 기준.  매 스텝 다시 잡는다"],
            ["타원체고", "WGS-84 타원체 면에서 수직으로 잰 높이", "과제의 \"고도\".  해발고도와 약 20 m 차이"],
            ["회전행렬 R", "한 좌표계의 축을 다른 좌표계 축으로 돌려 읽는 3×3", "위치엔 R·p + t,  속도엔 R·v"]]
    table(s, CX0, 1.50, CW, rows, [2.30, 5.00, 4.80], row_h=0.50, size=10.5, head_size=11, first_col_bold=True)
    callout(s, CX0, 6.10, CW, 0.64, "이상입니다.  질문 환영합니다.       코드 · 문서 :  radar-sw-study / Part3", fill=CARD, tone=TEXT, size=12)
    return s


# ───────────────────────────────────────────── 조립
def main():
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)

    build_cover(prs)                                                                   # 1
    build_index(prs, [("00", "들어가며", "과제 3이 묻는 것\n과제 2와 무엇이 다른가"),
                      ("01", "표적을 움직이는 원리", "좌표계 넷 · 위치와 속도\n시간 적분 · 지구 곡률"),
                      ("02", "기동 모델", "하중배수 g · 원운동\n세 축 · 기동 시간표"),
                      ("03", "구조체와 시뮬레이션 설계", "표적 카드 · 전체 흐름 · 조건\n예상 결과 · 검증 · 다음 단계")],
                "오늘은 개념편입니다.  수학은 고등학교 물리의 등속 원운동 정도면 충분하고, 나올 때마다 다시 설명합니다.")   # 2
    s03(prs)                                                                           # 3
    s04(prs)                                                                           # 4
    build_divider(prs, "01", "표적을\n움직이는 원리",
                  ["이번에 쓰는 좌표계 넷 — LLA · ECEF · NED · 동체",
                   "위치와 속도는 다르게 변환한다",
                   "동체 속도를 ECEF 속도로 — 회전 두 번",
                   "시간 적분 — 새 위치 = 지금 위치 + 속도 × 0.1 초",
                   "지구는 둥글다 — 60초 직진하면 11 m 떠오른다",
                   "한 스텝의 순서 — 다섯 단계"],
                  "이 발표의 약속 :  각도는 도(°) 로 말하고 프로그램 안에서는 라디안.  거리는 m, 시간은 s.")   # 5
    s06(prs)                                                                           # 6
    s07(prs)                                                                           # 7
    s08(prs)                                                                           # 8
    s09(prs)                                                                           # 9
    s10(prs)                                                                           # 10
    s11(prs)                                                                           # 11
    s12(prs)                                                                           # 12
    build_divider(prs, "02", "기동 모델",
                  ["기동이란 — 하중배수 g",
                   "원운동 공식 — 각속도와 선회 반경",
                   "세 축 — Roll · Yaw · Pitch",
                   "기동 시간표 — StartTime · EndTime"],
                  "약속 :  기동 중에도 속력은 바뀌지 않는다.  방향만 바뀐다.")                              # 13
    s14(prs)                                                                           # 14
    s15(prs)                                                                           # 15
    s16(prs)                                                                           # 16
    s17(prs)                                                                           # 17
    build_divider(prs, "03", "구조체와\n시뮬레이션 설계",
                  ["표적 구조체 — 표적 카드 한 장",
                   "시나리오와 전체 흐름",
                   "시뮬레이션 조건 — 플랫폼 하나, 표적 둘",
                   "예상 결과 — 정면 접근과 가로지르기",
                   "검증 계획 — 무엇을 확인할 것인가",
                   "결과 출력과 그림",
                   "다음 단계 — 구현 순서"],
                  "과제 기타 사항 :  구조체 정의 설명 · 궤적 모의 방법 설명 · 결과 설명 · 제공 코드 이용")     # 18
    s19(prs)                                                                           # 19
    s20(prs)                                                                           # 20
    s21(prs)                                                                           # 21
    s22(prs)                                                                           # 22
    s23(prs)                                                                           # 23
    s24(prs)                                                                           # 24
    s25(prs)                                                                           # 25
    build_summary(prs, ["과제 3 = 과제 2의 변환 체인 + 시간 루프 + 기동 모델.  좌표변환 함수 넷은 그대로 쓴다",
                        "위치는 회전 + 평행이동, 속도는 회전만.  속도에 위치를 더하면 안 된다",
                        "새 위치 = 지금 위치 + 속도 × 0.1 초.  ECEF 에서 더한다",
                        "지구는 둥글다.  매 스텝 표적 자리의 수평면을 다시 잡아야 고도가 유지된다",
                        "기동은 원운동.  ω = n·g / V,  R = V² / (n·g).  속력은 그대로",
                        "표적 카드 = 입력 칸 + 상태 칸.  플랫폼도 카드 한 장",
                        "검증 없는 궤적은 그림일 뿐.  거리 · 속력 · 고도 · 반경으로 확인한다"],
                  "질문 환영합니다       다음 발표 :  코드와 실행 결과       코드 · 문서 :  radar-sw-study / Part3")   # 26
    s27(prs)                                                                           # 27
    s28(prs)                                                                           # 28

    # 노트 : 대본 md 에서
    secs = parse_script(io.open(SCRIPT_MD, encoding="utf-8").read())
    slides = list(prs.slides)
    if len(secs) != len(slides):
        print("경고: 슬라이드 %d 장, 대본 %d 장 — 노트를 채우지 않음" % (len(slides), len(secs)))
    else:
        for i, sl in enumerate(slides, 1):
            notes(sl, secs[i][1])

    prs.save(DECK)
    print("written:", DECK, "(%d slides)" % len(slides))


if __name__ == "__main__":
    main()
