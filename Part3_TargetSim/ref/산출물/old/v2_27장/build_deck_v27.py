# -*- coding: utf-8 -*-
"""Chapter 3 발표자료(pptx) 를 처음부터 만든다. 개념 + C 구현 + 결과, 27장.

    python build_deck.py            # ../Chapter3_target_simulation.pptx 를 새로 쓴다

틀은 Chapter 2 발표자료(네이비판 LIG 틀) 와 같다 : 표지 · 목차 · 간지 · 본문(머리띠 + 흰 패널 + 꼬리말) · 정리.
그림은 ../figures/ (make_figures.py) 와 ../results/ (실행 결과), 노트는 ../Chapter3_발표대본.md 에서 채운다.
"""
import io
import os
import sys

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from sync_notes import parse as parse_script          # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.normpath(os.path.join(HERE, ".."))
FIG_DIR = os.path.join(OUT_DIR, "figures")
RES_DIR = os.path.join(OUT_DIR, "results")
DECK = os.path.join(OUT_DIR, "Chapter3_target_simulation.pptx")
SCRIPT_MD = os.path.join(OUT_DIR, "Chapter3_발표대본.md")

# 디자인 토큰 (Chapter 2 네이비판과 같음)
DARK, BG, PANEL, TEXT, MUTED, LINE, WHITE = "13233F", "E8EDF5", "FFFFFF", "1A2233", "5A6478", "B9C0CC", "FFFFFF"
ACCENT, BLUE, GREEN = "C2521C", "2B57A6", "2F855A"
ON_DARK, ON_DARK_DIM, RULE_SOFT = "C9D6EC", "8FA6CC", "D0D5DE"
CARD, PEACH, SOFT, MINT = "EEF1F7", "FBEADF", "F4F6FA", "E6F3EA"
CIRCLE_DK, CIRCLE_LN, CIRCLE_TX, SUM_TX = "1E3157", "34507F", "9DB3D6", "D6DFF0"
KO, EN, MONO = "맑은 고딕", "Arial", "Consolas"

SW, SH = 13.333, 7.5
M = 0.5
INNER_W = SW - 2 * M
BAR_Y, BAR_H = 0.45, 0.85
PANEL_Y, PANEL_H = 1.36, 5.52
FOOT_LINE_Y, FOOT_TEXT_Y = 6.95, 7.00
CX0, CX1 = 0.62, 12.72
CW = CX1 - CX0
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"


# ───────────────────────────────────────────── 도우미
def rgb(s):
    return RGBColor.from_string(s)


def set_background(slide, color):
    csld = slide._element.find(qn("p:cSld"))
    old = csld.find(qn("p:bg"))
    if old is not None:
        csld.remove(old)
    csld.insert(0, parse_xml('<p:bg xmlns:p="%s" xmlns:a="%s"><p:bgPr><a:solidFill><a:srgbClr val="%s"/></a:solidFill>'
                             '<a:effectLst/></p:bgPr></p:bg>' % (P_NS, A_NS, color)))


def txt(slide, text, x, y, w, h, size=12, color=TEXT, bold=False, font=KO, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, wrap=True, spacing=None, space_after=0):
    """text 는 문자열 또는 (문자열, 옵션dict) 목록. 옵션 'runs' 로 한 문단에 여러 런."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    items = text if isinstance(text, list) else [(text, {})]
    for i, item in enumerate(items):
        s, opt = item if isinstance(item, tuple) else (item, {})
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = opt.get("align", align)
        ls = opt.get("spacing", spacing)
        if ls:
            p.line_spacing = Pt(ls)
        sa = opt.get("space_after", space_after)
        if sa:
            p.space_after = Pt(sa)
        sb = opt.get("space_before", 0)
        if sb:
            p.space_before = Pt(sb)
        for rs, ro in (opt.get("runs") or [(s, {})]):
            r = p.add_run()
            r.text = rs
            f = r.font
            f.name = ro.get("font", opt.get("font", font))
            f.size = Pt(ro.get("size", opt.get("size", size)))
            f.bold = ro.get("bold", opt.get("bold", bold))
            f.color.rgb = rgb(ro.get("color", opt.get("color", color)))
    return box


def rect(slide, x, y, w, h, fill, line=None, lw=0.75, rounded=False, radius=0.06):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        sp.adjustments[0] = radius
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = rgb(fill)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = rgb(line)
        sp.line.width = Pt(lw)
    sp.shadow.inherit = False
    sp.text_frame.word_wrap = False
    return sp


def ellipse(slide, x, y, d, fill, line=None, lw=1.0):
    sp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    sp.fill.solid()
    sp.fill.fore_color.rgb = rgb(fill)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = rgb(line)
        sp.line.width = Pt(lw)
    sp.shadow.inherit = False
    return sp


def hline(slide, x1, x2, y, color=LINE, lw=0.75):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y), Inches(x2), Inches(y))
    ln.line.color.rgb = rgb(color)
    ln.line.width = Pt(lw)
    return ln


SLIDENUM_XML = ('<a:fld xmlns:a="%s" id="{1D0B0B0B-0000-4000-8000-00000000C0DE}" type="slidenum">'
                '<a:rPr lang="ko-KR" altLang="en-US" sz="%d" b="0" dirty="0"><a:solidFill><a:srgbClr val="%s"/></a:solidFill>'
                '<a:latin typeface="%s"/><a:ea typeface="%s"/></a:rPr><a:t>1</a:t></a:fld>')


def slide_number(slide, color=TEXT, size=10):
    box = slide.shapes.add_textbox(Inches(6.2), Inches(FOOT_TEXT_Y), Inches(1.0), Inches(0.3))
    tf = box.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p._p.append(parse_xml(SLIDENUM_XML % (A_NS, int(size * 100), color, EN, EN)))


def footer(slide, color=TEXT):
    hline(slide, M, SW - M, FOOT_LINE_Y, LINE, 0.75)
    slide_number(slide, color)


def pic(slide, name, x, y, w=None, h=None):
    path = os.path.join(RES_DIR if name.startswith("results/") else FIG_DIR, name.split("/")[-1])
    iw, ih = Image.open(path).size
    if w is None:
        w = h * iw / float(ih)
    if h is None:
        h = w * ih / float(iw)
    slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    return w, h


def content_chrome(slide, title, subtitle=None):
    set_background(slide, BG)
    rect(slide, M, PANEL_Y, INNER_W, PANEL_H, PANEL)
    rect(slide, M, BAR_Y, INNER_W, BAR_H, DARK)
    paras = [(title, {"size": 19 if subtitle else 20, "bold": True, "color": WHITE, "space_after": 3})]
    if subtitle:
        paras.append((subtitle, {"size": 11, "color": ON_DARK, "spacing": 14}))
    txt(slide, paras, M + 0.20, BAR_Y, 11.6, BAR_H, anchor=MSO_ANCHOR.MIDDLE)
    footer(slide, TEXT)


def card(slide, x, y, w, h, title, lines, tone=BLUE, fill=CARD, title_size=13, body_size=11, pad=0.22,
         title_after=4, line_spacing=15):
    rect(slide, x, y, w, h, fill, WHITE, 1.0, rounded=True, radius=0.06)
    paras = [(title, {"size": title_size, "bold": True, "color": tone, "space_after": title_after})]
    for s in lines:
        paras.append(s if isinstance(s, tuple) else (s, {"size": body_size, "color": TEXT, "spacing": line_spacing}))
    txt(slide, paras, x + pad, y + pad * 0.7, w - 2 * pad, h - pad * 1.2)


def callout(slide, x, y, w, h, text, sub=None, size=13, tone=ACCENT, fill=PEACH, align=PP_ALIGN.CENTER):
    rect(slide, x, y, w, h, fill, WHITE, 1.0, rounded=True, radius=0.08)
    paras = [(text, {"size": size, "bold": True, "color": tone, "align": align, "space_after": 3})]
    if sub:
        paras.append((sub, {"size": 10.5, "color": MUTED, "align": align, "spacing": 14}))
    txt(slide, paras, x + 0.24, y, w - 0.48, h, anchor=MSO_ANCHOR.MIDDLE, align=align)


def code_box(slide, x, y, w, h, lines, size=9.5, spacing=12.5, title=None):
    rect(slide, x, y, w, h, CARD, None, rounded=True, radius=0.04)
    top = y + 0.12
    if title:
        txt(slide, title, x + 0.18, top, w - 0.36, 0.26, size=10.5, bold=True, color=TEXT)
        top += 0.30
    paras = []
    for s in lines:
        col = MUTED if s.lstrip().startswith("//") or s.lstrip().startswith("/*") else TEXT
        paras.append((s, {"size": size, "font": MONO, "color": col, "spacing": spacing}))
    txt(slide, paras, x + 0.18, top, w - 0.36, h - (top - y) - 0.1, font=MONO)


def table(slide, x, y, w, rows, col_w, row_h=0.42, size=11, head_size=11.5, first_col_bold=False,
          align_center_cols=(), hdr_fill=BLUE):
    n, m = len(rows), len(rows[0])
    gf = slide.shapes.add_table(n, m, Inches(x), Inches(y), Inches(w), Inches(row_h * n))
    tb = gf.table
    tb.first_row = True
    tb.horz_banding = False
    tb.vert_banding = False
    for j, cw_ in enumerate(col_w):
        tb.columns[j].width = Inches(cw_)
    for i in range(n):
        tb.rows[i].height = Inches(row_h)
        for j in range(m):
            cell = tb.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.10)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            val = rows[i][j]
            opt = {}
            if isinstance(val, tuple):
                val, opt = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(hdr_fill if i == 0 else (WHITE if i % 2 else SOFT))
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER if (i == 0 or j in align_center_cols) else PP_ALIGN.LEFT
            r = p.add_run()
            r.text = val
            f = r.font
            f.name = opt.get("font", KO)
            f.size = Pt(head_size if i == 0 else opt.get("size", size))
            f.bold = True if i == 0 else opt.get("bold", first_col_bold and j == 0)
            f.color.rgb = rgb(WHITE if i == 0 else opt.get("color", TEXT))
    return gf


def numbered_rows(slide, items, x, y, w, dy=0.64, size=16, dark=True, circle=0.38, accent_from=None):
    for i, s in enumerate(items):
        yy = y + i * dy
        acc = accent_from is not None and i + 1 >= accent_from
        if dark:
            fc, lc, tc, sc = (ACCENT, ACCENT, WHITE, WHITE) if acc else (CIRCLE_DK, CIRCLE_LN, CIRCLE_TX, SUM_TX)
        else:
            fc, lc, tc, sc = (ACCENT, ACCENT, WHITE, TEXT) if acc else (BLUE, BLUE, WHITE, TEXT)
        ellipse(slide, x, yy + 0.04, circle, fc, lc, 1.0)
        txt(slide, str(i + 1), x, yy + 0.04, circle, circle, size=13, bold=True, color=tc, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        txt(slide, s, x + circle + 0.20, yy, w - circle - 0.20, 0.46, size=size, color=sc, anchor=MSO_ANCHOR.MIDDLE)


# ───────────────────────────────────────────── 특수 페이지
def build_cover(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(s, DARK)
    txt(s, "Chapter 3", 10.30, 0.55, 2.5, 0.35, size=11, bold=True, color=WHITE, font=EN, align=PP_ALIGN.RIGHT, wrap=False)
    txt(s, "표적 궤적 모의", 4.30, 2.70, 8.50, 1.00, size=40, bold=True, color=WHITE, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, "표적을 정의하고, 0.1초마다 움직이고, 위경도로 그리기까지", 4.30, 3.72, 8.50, 0.42, size=16, color=ON_DARK, align=PP_ALIGN.RIGHT)
    txt(s, "원리  ·  기동 모델  ·  C 구현  ·  결과", 4.30, 4.22, 8.50, 0.36, size=12, color=ON_DARK_DIM, align=PP_ALIGN.RIGHT)
    txt(s, "레이다 시스템 소프트웨어 스터디  ·  Part 3", 6.80, 4.80, 6.00, 0.35, size=12, color=ON_DARK, align=PP_ALIGN.RIGHT)


def build_index(prs, rows, tail):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(s, BG)
    rect(s, 0, 0, 4.90, SH, DARK)
    txt(s, "INDEX", 1.25, 0.55, 2.2, 0.60, size=30, bold=True, color=WHITE, font=EN, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
    hline(s, 2.65, 12.00, 0.86, LINE, 0.75)
    txt(s, "목차", 0.45, 3.05, 4.00, 1.40, size=30, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    slide_number(s, TEXT)
    y0, dy = 1.15, 1.02
    for i, (num, name, desc) in enumerate(rows):
        y = y0 + i * dy
        txt(s, num, 5.55, y, 1.10, 0.62, size=34, bold=True, color=DARK, font=EN, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
        txt(s, name, 6.80, y, 3.55, 0.62, size=19, color=TEXT, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
        txt(s, desc, 10.45, y, 2.40, 0.62, size=11, color=MUTED, anchor=MSO_ANCHOR.MIDDLE, spacing=14)
        if i < len(rows) - 1:
            hline(s, 5.55, 12.80, y + 0.82, RULE_SOFT, 0.5)
    txt(s, tail, 5.55, 6.35, 7.30, 0.40, size=11.5, color=MUTED)


def build_divider(prs, num, title, items, foot):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(s, BG)
    rect(s, 0, 0, 4.90, SH, DARK)
    txt(s, num, 1.25, 0.55, 2.2, 0.60, size=30, bold=True, color=WHITE, font=EN, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
    hline(s, 2.65, 12.00, 0.86, LINE, 0.75)
    txt(s, title, 0.45, 3.05, 4.00, 1.40, size=30, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    slide_number(s, TEXT)
    dy, y = 0.70, 1.62
    for i, it in enumerate(items):
        txt(s, str(i + 1), 5.55, y, 0.60, 0.50, size=16, bold=True, color=BLUE, font=EN, anchor=MSO_ANCHOR.MIDDLE)
        txt(s, it, 6.30, y, 6.50, 0.50, size=16, bold=True, color=TEXT, anchor=MSO_ANCHOR.MIDDLE)
        hline(s, 5.55, 12.80, y + dy - 0.12, RULE_SOFT, 0.5)
        y += dy
    txt(s, foot, 5.55, 6.42, 7.25, 0.40, size=11, color=MUTED)


def build_summary(prs, items, tail, accent_from=6):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(s, DARK)
    hline(s, 2.00, 12.83, 0.86, LINE, 0.75)
    txt(s, "정리", 0.95, 1.22, 11.5, 0.58, size=30, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    numbered_rows(s, items, 0.90, 2.00, 11.6, dy=0.62, size=16, dark=True, accent_from=accent_from)
    txt(s, tail, 0.95, 6.46, 11.5, 0.42, size=13, bold=True, color=WHITE)
    hline(s, M, SW - M, FOOT_LINE_Y, LINE, 0.75)
    slide_number(s, WHITE)


def new(prs, title, subtitle=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    content_chrome(s, title, subtitle)
    return s


# ───────────────────────────────────────────── 코드 인용 (src 와 같은 내용)
CODE_STRUCT = [
    "typedef struct",
    "{",
    "    DOUBLE64    dGravity;       // 하중배수 [g]. 부호가 방향",
    "    INT32       nTurnType;      // 0 없음 1 Roll 2 Yaw 3 Pitch",
    "    DOUBLE64    dStartTime;     // [s]",
    "    DOUBLE64    dEndTime;       // [s]",
    "} ST_Maneuver;",
    "",
    "typedef struct",
    "{                                                   /* 입력 */",
    "    STRUCT_Coord_Lla        stInitLla;              // 초기 위치 [rad, rad, m]",
    "    DOUBLE64                dVheading;              // 동체 속력 [m/s]",
    "    STRUCT_Coord_Attitude   stInitAtt;              // 초기 자세 [rad]",
    "    INT32                   nManeuverCnt;           // 기동 수",
    "    ST_Maneuver             astManeuver[MAX_MANEUVER];   // 30",
    "                                                    /* 상태 (ECEF 에서 갱신) */",
    "    STRUCT_Coord_Rect       stPos;                  // ECEF 위치 [m]",
    "    STRUCT_Coord_Rect       stVel;                  // ECEF 속도 [m/s]",
    "    STRUCT_Coord_Attitude   stAtt;                  // 현재 자세 [rad]",
    "    STRUCT_Coord_Lla        stLla;                  // 현재 위경도 (출력용)",
    "} ST_Target;",
    "",
    "typedef struct",
    "{",
    "    ST_Target   stPlatform;                         // 속력 0 인 표적 카드",
    "    INT32       nTargetCnt;",
    "    ST_Target   astTarget[MAX_TARGET];              // 10",
    "    DOUBLE64    dSimTime;                           // 60 s",
    "    DOUBLE64    dDt;                                // 0.1 s",
    "} ST_Scenario;",
]

CODE_ASSIGN = [
    "VOID f_SetAssignment(ST_Scenario *pstScn, INT32 bWithManeuver)",
    "{",
    "    memset(pstScn, 0, sizeof(*pstScn));",
    "    pstScn->dSimTime   = 60.0;",
    "    pstScn->dDt        = 0.1;",
    "    pstScn->nTargetCnt = 2;",
    "                             /*     위도    경도    고도   속력  roll   yaw  pitch */",
    "    f_SetTarget(&pstScn->stPlatform,   32.0,   126.0,    0.0,   0.0, 0.0,   0.0, 0.0);",
    "    f_SetTarget(&pstScn->astTarget[0], 32.125, 126.03,   0.0,  30.0, 0.0, 270.0, 0.0);  // 대함",
    "    f_SetTarget(&pstScn->astTarget[1], 32.12,  126.0,  300.0, 200.0, 0.0, 180.0, 0.0);  // 대공",
    "",
    "    if (bWithManeuver)                                 // 발표용 기동 예시",
    "    {",
    "        f_AddManeuver(&pstScn->astTarget[0], -0.1, TURN_YAW,   20.0, 50.0);  // 대함 좌선회",
    "        f_AddManeuver(&pstScn->astTarget[1],  3.0, TURN_YAW,   10.0, 20.0);  // 대공 3 g 우선회",
    "        f_AddManeuver(&pstScn->astTarget[1],  2.0, TURN_PITCH, 35.0, 40.0);  // 대공 기수 들기",
    "    }",
    "}",
    "",
    "// f_SetTarget : 각도는 도로 받아 라디안으로 저장한다 (프로그램 안은 라디안)",
    "pstTgt->stInitLla.Lat  = DEG2RAD(dLatDeg);  ...  pstTgt->stInitAtt.Yaw = DEG2RAD(dYawDeg);",
]

CODE_VEL = [
    "// 동체 속력 [V, 0, 0] 을 ECEF 속도로. 회전 두 번, 평행이동 없음",
    "static STRUCT_Coord_Rect f_BodyVelToEcef(const ST_Target *pstTgt)",
    "{",
    "    STRUCT_Coord_Rect stNed;",
    "",
    "    // 1) 동체 → NED : 자세각으로 회전 (인자 순서 Roll, Yaw, Pitch)",
    "    stNed = f_Trans_Body_To_Ned(pstTgt->dVheading, 0.0, 0.0,",
    "                                pstTgt->stAtt.Roll, pstTgt->stAtt.Yaw, pstTgt->stAtt.Pitch);",
    "",
    "    // 2) NED → ECEF : 표적 자신의 위경도로 회전.",
    "    //    속도는 벡터라 플랫폼 위치 인자(가운데 셋) 에 0",
    "    return f_Trans_Ned_To_Ecef(stNed.x, stNed.y, stNed.z,  0.0, 0.0, 0.0,",
    "                               pstTgt->stLla.Lat, pstTgt->stLla.Lon);",
    "}",
    "",
    "// 초기화 : 입력 칸 → 상태 칸. f_Trans_Lla_To_Ecef 는 인자 순서가 (경도, 위도, 고도)",
    "static VOID f_InitTarget(ST_Target *pstTgt)",
    "{",
    "    pstTgt->stPos = f_Trans_Lla_To_Ecef(pstTgt->stInitLla.Lon,",
    "                                        pstTgt->stInitLla.Lat, pstTgt->stInitLla.Alt);",
    "    pstTgt->stAtt = pstTgt->stInitAtt;",
    "    pstTgt->stLla = pstTgt->stInitLla;",
    "    pstTgt->stVel = f_BodyVelToEcef(pstTgt);",
    "}",
]

CODE_STEP = [
    "static VOID f_StepTarget(ST_Target *pstTgt, DOUBLE64 dTime, DOUBLE64 dDt)",
    "{",
    "    INT32 i;",
    "",
    "    // 1) 기동 판단, 2) 자세 갱신 : 각속도 ω = n g / V",
    "    for (i = 0; i < pstTgt->nManeuverCnt; i++)",
    "    {",
    "        const ST_Maneuver *pstMan = &pstTgt->astManeuver[i];",
    "",
    "        if ((dTime + 1e-6 >= pstMan->dStartTime) && (dTime + 1e-6 < pstMan->dEndTime)",
    "            && (pstTgt->dVheading > 0.0))",
    "        {",
    "            DOUBLE64 dAng = pstMan->dGravity * G_FORCE / pstTgt->dVheading * dDt;",
    "",
    "            if (pstMan->nTurnType == TURN_ROLL)       pstTgt->stAtt.Roll  += dAng;",
    "            else if (pstMan->nTurnType == TURN_YAW)   pstTgt->stAtt.Yaw   += dAng;",
    "            else if (pstMan->nTurnType == TURN_PITCH) pstTgt->stAtt.Pitch += dAng;",
    "            break;                                     // 먼저 적힌 기동 하나만",
    "        }",
    "    }",
    "",
    "    // 3) 속도 변환 (표적 자신의 현재 위경도 기준 NED)",
    "    pstTgt->stVel = f_BodyVelToEcef(pstTgt);",
    "",
    "    // 4) 위치 적분 (ECEF)",
    "    pstTgt->stPos.x += pstTgt->stVel.x * dDt;",
    "    pstTgt->stPos.y += pstTgt->stVel.y * dDt;",
    "    pstTgt->stPos.z += pstTgt->stVel.z * dDt;",
    "",
    "    // 5) LLA 변환. 출력용이자 다음 스텝의 NED 기준",
    "    pstTgt->stLla = f_Trans_Ecef_To_Lla(pstTgt->stPos.x, pstTgt->stPos.y, pstTgt->stPos.z);",
    "}",
]

CODE_LOOP = [
    "// TargetSimUIDlg.cpp  [실행] : 60 초를 0.1 초 간격으로 돌리며 매 스텝 위경도를 모은다",
    "const int nSteps = (int)floor(m_stScn.dSimTime / m_stScn.dDt + 0.5);     // 600",
    "",
    "f_StartScenario(&m_stScn);                     // 위경도 → ECEF, 속도 초기화",
    "record();                                      // t = 0",
    "for (int k = 0; k < nSteps; k++)",
    "{",
    "    f_StepScenario(&m_stScn, k * m_stScn.dDt); // 플랫폼 + 표적 전부 한 스텝",
    "    record();                                  // 위도 · 경도 · 고도 · yaw 를 배열에",
    "}",
    "m_map.SetPaths(&m_paths);                      // 지도에 궤적, 슬라이더 0 ~ 600",
]

CODE_MAP = [
    "// MapView.cpp  위도 1 도가 경도 1 도보다 1/cos(위도) 배 길다 → 축 비율",
    "const double dAsp = 1.0 / cos(dLatC * PI / 180.0);",
    "if (dLatSpan * dAsp * dW / dH > dLonSpan) dLonSpan = dLatSpan * dAsp * dW / dH;",
    "else                                      dLatSpan = dLonSpan / dAsp * dH / dW;",
    "",
    "// 재생 : 100 ms 타이머 한 틱에 배속(1·2·5·10) 만큼 스텝 → 1 배속 = 실시간",
    "SetTimer(TIMER_PLAY, 100, nullptr);",
]


# ───────────────────────────────────────────── 본문
def s03(prs):
    s = new(prs, "00.  과제 3이 묻는 것", "표적을 정의하고, 0.1초마다 움직이고, 위경도로 그려라")
    pic(s, "fig01_overview.png", CX0, 1.50, w=7.4)
    x, w, y = 8.20, 4.52, 1.50
    specs = [("① 표적 구조체를 정의한다", ["표적 최대 10개 · 표적마다 기동 최대 30개", "초기 위치는 위경도·고도, 속도는 동체 속력 하나로"], BLUE, 1.02),
             ("② 위치·속도는 ECEF 에서 갱신한다", ["동체 속력을 ECEF 성분으로 바꾸는 방법도 고민 (추가 항목)"], BLUE, 0.86),
             ("③ 60초 동안 0.1초 간격으로", ["플랫폼 하나 (정지) · 대함 표적 · 대공 표적"], ACCENT, 0.86),
             ("④ 위경도로 출력하고 그림으로", ["플랫폼과 표적의 궤적을 지도 위에"], ACCENT, 0.86)]
    for t, ls, tone, h in specs:
        card(s, x, y, w, h, t, ls, tone=tone, title_size=12.5, body_size=10.5, pad=0.20, line_spacing=14)
        y += h + 0.10
    callout(s, CX0, 5.42, CW, 1.32,
            "과제 2 와 다른 점 :  변환 함수 넷은 제공 코드 그대로,  새로 들어온 것은  속도 · 시간 · 기동",
            "지난번엔 측정값 하나가 체인을 한 번 지나 점 하나.  이번엔 같은 계산을 600번 반복해 선이 된다.")


def s05(prs):
    s = new(prs, "01.  좌표계 넷, 그리고 왜 ECEF 인가", "입력·출력은 위경도,  계산은 ECEF,  수평면은 NED,  속력은 동체")
    pic(s, "fig03_four_frames.png", CX0, 1.50, w=8.4)
    x, w = 9.28, 3.44
    card(s, x, 1.50, w, 3.95, "이번 과제에서의 역할",
         [("LLA", {"size": 11.5, "bold": True, "color": BLUE, "space_before": 2}),
          ("초기값을 받고, 매 스텝 결과를 적는 형식", {"size": 10.5, "color": TEXT, "space_after": 6}),
          ("ECEF", {"size": 11.5, "bold": True, "color": BLUE}),
          ("위치와 속도를 600번 갱신하는 자리", {"size": 10.5, "color": TEXT, "space_after": 6}),
          ("NED", {"size": 11.5, "bold": True, "color": GREEN}),
          ("표적 자리의 수평면. 자세각의 기준이자 속도 방향을 만드는 자리", {"size": 10.5, "color": TEXT, "space_after": 6}),
          ("동체", {"size": 11.5, "bold": True, "color": ACCENT}),
          ("속력 V 가 적힌 자리. 앞으로 V, 옆·아래는 0", {"size": 10.5, "color": TEXT})],
         tone=TEXT)
    callout(s, CX0, 5.72, CW, 1.02,
            "왜 ECEF 인가 :  위도 1° = 111 km, 경도 1° = 94 km (32°N) — 각도에 미터를 더할 수 없다.  ECEF 는 세 축이 미터라 p + v·Δt 한 줄",
            "플랫폼과 무관한 절대 좌표라, 나중에 배가 움직여도 표적 궤적이 흔들리지 않는다.", size=12)


def s06(prs):
    s = new(prs, "01.  위치와 속도는 다르게 변환한다 — 회전 두 번", "위치는 원점이 있어야 말이 되고,  속도는 화살표 하나면 된다.  그래서 속도는 회전만")
    pic(s, "fig06_vel_chain.png", CX0 + 0.30, 1.50, w=11.5)
    y, h = 3.86, 1.42
    card(s, CX0, y, 3.90, h, "위치 :  p_new = R · p + t", ["원점이 다르면 숫자가 다르다.", "회전하고 원점 차이만큼 옮긴다 (과제 2)."], tone=BLUE, body_size=10.5, line_spacing=14)
    card(s, CX0 + 4.10, y, 3.90, h, "속도 :  v_new = R · v", ["어디에 그려도 같은 화살표.", "축이 돌아간 만큼만 돌린다. t 를 더하면 속도가 아니게 된다."], tone=ACCENT, fill=PEACH, body_size=10.5, line_spacing=14)
    card(s, CX0 + 8.20, y, 3.90, h, "검산", ["결과 벡터의 크기는 항상 V (200.000).", "플랫폼 위치를 잘못 더하면 6,372 km/s."], tone=GREEN, fill=MINT, body_size=10.5, line_spacing=14)
    callout(s, CX0, 5.44, CW, 1.30,
            "과제의 \"추가 고민\" (동체 속도 V_heading → ECEF) 의 답 :  자세각으로 한 번, 표적 자신의 위경도로 한 번, 평행이동은 없다",
            "참고 함수 f_Trans_Ned_To_Ecef 는 플랫폼 위치를 더하므로, 속도를 넣을 때는 그 인자에 0 을 준다.", size=12)


def s07(prs):
    s = new(prs, "01.  시간 적분과 지구 곡률", "새 위치 = 지금 위치 + 속도 × 0.1 초.  단, 매 스텝 수평면을 다시 잡지 않으면 60초에 11 m 떠오른다")
    pic(s, "fig08_curvature.png", CX0, 1.50, w=7.4)
    x, w, y = 8.20, 4.52, 1.50
    card(s, x, y, w, 1.06, "시간 적분 :  p(t+Δt) = p(t) + v(t)·Δt", ["ECEF x, y, z 에 각각.  Δt = 0.1 s, 600 스텝.", "한 스텝 안에서 속도가 일정하니 이 한 줄이면 정확."], tone=BLUE, body_size=10.5, line_spacing=14)
    rows = [["실험 (대공 표적 60초 직진)", "60초 뒤 고도"],
            ["A  처음 ECEF 속도를 60초 내내 그대로", ("311.3 m", {"bold": True, "color": ACCENT})],
            ["B  매 스텝 현재 위경도로 NED 를 다시 잡음", ("300.0 m", {"bold": True, "color": GREEN})]]
    table(s, x, y + 1.18, w, rows, [3.12, 1.40], row_h=0.44, size=10.5, head_size=10.5, align_center_cols=(1,))
    card(s, x, y + 2.62, w, 1.14, "왜 11 m 인가", ["접선을 따라 d 만큼 가면 땅은 d² / 2R 만큼 내려간다.", "12,000² / (2 × 6,371,000) = 11.3 m.  실험값과 일치."], tone=ACCENT, fill=PEACH, body_size=10.5, line_spacing=14)
    callout(s, CX0, 5.42, CW, 1.32,
            "그래서 5단계 LLA 변환은 출력용만이 아니라 다음 스텝의 수평면(NED 축) 을 만드는 기준이다",
            "위경도를 구해야 NED 축이 나오고, NED 축이 있어야 속도 방향이 나온다.  이 순서를 지키면 수평 비행이 수평으로 유지된다.")


def s08(prs):
    s = new(prs, "01.  한 스텝의 순서 — 다섯 단계", "표적 하나가 0.1초 뒤로 가는 일.  이 순서를 600번, 표적 수만큼 반복한다")
    pic(s, "fig09_step_loop.png", CX0 + 0.30, 1.50, w=11.5)
    y, h = 4.62, 1.10
    w5 = (CW - 4 * 0.14) / 5
    specs = [("① 에 필요한 것", "기동 목록,  지금 시각 t", ACCENT, PEACH), ("② 에 필요한 것", "각속도 ω = n·g / V,  Δt", ACCENT, PEACH),
             ("③ 에 필요한 것", "자세각 3개,  표적 위경도", GREEN, MINT), ("④ 에 필요한 것", "ECEF 위치·속도,  Δt = 0.1 s", BLUE, CARD),
             ("⑤ 에 필요한 것", "ECEF → LLA 변환 함수", BLUE, CARD)]
    for i, (t, l, tone, fill) in enumerate(specs):
        card(s, CX0 + i * (w5 + 0.14), y, w5, h, t, [l], tone=tone, fill=fill, title_size=11.5, body_size=10.5, pad=0.18)
    callout(s, CX0, 5.90, CW, 0.84, "표적 10개 × 600 스텝 = 6,000번.   순서가 중요한 건 ③ ④ ⑤ — 자세가 정해져야 속도, 속도가 있어야 위치, 위치가 있어야 새 수평면", fill=CARD, tone=TEXT, size=12)


def s10(prs):
    s = new(prs, "02.  하중배수와 원운동 공식", "과제의 Gravity Value 는 \"몇 g 로 도는가\".  속력은 그대로, 방향만 바뀐다")
    pic(s, "fig11_circle.png", CX0, 1.50, w=6.0)
    x, w, y = 6.90, 5.82, 1.50
    rows = [["표적", "하중배수", "각속도", "선회 반경", "비고"],
            ["대공 200 m/s", "1 g", "2.81 °/s", "4,079 m", "180° 에 64 s"],
            ["대공 200 m/s", ("3 g", {"bold": True, "color": ACCENT}), ("8.43 °/s", {"bold": True, "color": ACCENT}), ("1,360 m", {"bold": True, "color": ACCENT}), "180° 에 21 s"],
            ["대공 200 m/s", "5 g", "14.05 °/s", "816 m", "180° 에 13 s"],
            ["대함 30 m/s", "0.1 g", "1.87 °/s", "918 m", "90° 에 48 s"]]
    table(s, x, y, w, rows, [1.50, 0.90, 1.05, 1.05, 1.32], row_h=0.44, size=11, head_size=11, align_center_cols=(1, 2, 3, 4))
    card(s, x, y + 2.42, w, 1.20, "감 잡기 :  1 g 서 있을 때,  0.3 g 승용차 커브,  3 g 롤러코스터 바닥·전투기 급선회,  9 g 전투기 한계",
         ["가속도가 속도에 수직이라 속력은 안 변하고 방향만 변한다.", "각속도가 속력에 반비례 — 배 기동에는 0.1 g 같은 작은 값을 넣어야 그럴듯하다."], tone=BLUE, title_size=11, body_size=10.5, line_spacing=14)
    callout(s, CX0, 5.62, CW, 1.12, "프로그램에서는 한 줄 :  ω = n·g / V  를 구해 놓고,  매 스텝 자세각에  ω × 0.1 s  를 더한다   (3 g 면 한 스텝에 0.84°)",
            "a = V²/R = n·g  →  R = V²/(n·g),  ω = V/R = n·g/V", size=12.5)


def s11(prs):
    s = new(prs, "02.  세 축과 기동 시간표", "TurnType 은 동체의 어느 축을 돌리는가.  기동 하나 = { g, TurnType, StartTime, EndTime }, 표적마다 최대 30개")
    pic(s, "fig12_axes.png", CX0, 1.50, w=8.4)
    x, w = 9.28, 3.44
    rows = [["TurnType", "축", "효과"], ["0", "—", "직진"], ["1  Roll", "x (앞뒤)", "자세만 기움"], ["2  Yaw", "z (위아래)", "좌우 선회"], ["3  Pitch", "y (날개)", "상승 · 강하"]]
    table(s, x, 1.50, w, rows, [1.04, 1.06, 1.34], row_h=0.40, size=10.5, head_size=10.5, align_center_cols=(0, 1, 2))
    card(s, x, 3.64, w, 1.26, "시간표 규칙 셋", ["StartTime ≤ t < EndTime 이면 적용, 밖이면 직진.", "부호 : n < 0 이면 반대 방향.  겹침 : 먼저 적힌 것.  끝난 뒤 : 자세 유지."], tone=ACCENT, fill=PEACH, title_size=12, body_size=10, line_spacing=13)
    callout(s, CX0, 5.14, CW, 0.74, "축을 돌린 뒤 속도가 어느 쪽인지는 따로 계산하지 않는다 — 자세각만 바꿔 두면 6장의 회전 두 번이 새 속도 방향을 만든다", fill=CARD, tone=TEXT, size=12)
    callout(s, CX0, 6.02, CW, 0.72, "과제 조건에는 기동 목록이 없다 → 발표용 예시 기동을 정해 넣었다 (대공 3 g 우선회 10~20 s, 2 g 기수 들기 35~40 s, 대함 −0.1 g 좌선회 20~50 s)", size=11.5)


def s13(prs):
    s = new(prs, "03.  프로그램 구성", "제공 코드는 무수정으로 링크,  C 코어가 순서대로 부르고,  MFC 는 돌려서 보여 주기만 한다")
    pic(s, "fig19_program.png", CX0 + 0.30, 1.50, w=11.5)
    y, h = 4.36, 1.30
    card(s, CX0, y, 3.90, h, "제공 코드 5 파일 = 원본 그대로", ["복사 후 바이트 단위 비교로 확인.", "좌표변환은 이 코드만 부른다."], tone=BLUE, body_size=10.5, line_spacing=14)
    card(s, CX0 + 4.10, y, 3.90, h, "C 코어 250줄", ["구조체 셋 + 함수 다섯.", "MSVC x64 · C17 · 경고 4단계 0건."], tone=ACCENT, fill=PEACH, body_size=10.5, line_spacing=14)
    card(s, CX0 + 8.20, y, 3.90, h, "MFC 화면 (Part 2 CoordUI 와 같은 뼈대)", ["계산 없음.  코어를 600번 부르고 배열에 모아 그린다.", "지도 · 슬라이더 · 재생 · CSV 저장."], tone=GREEN, fill=MINT, body_size=10.5, line_spacing=14)
    callout(s, CX0, 5.84, CW, 0.90, "걸린 것 하나 :  제공 Define.h 의  #define BOOL bool  이 MFC 의 BOOL(int) 과 충돌 → 제공 코드는 두고, MFC 쪽에서 헤더를 읽은 직후 #undef BOOL", fill=CARD, tone=TEXT, size=11.5)


def s14(prs):
    s = new(prs, "03.  표적 구조체 — target_sim.h", "위 다섯 칸이 입력(사람이 적는 값), 아래 네 칸이 상태(프로그램이 갱신).  좌표 타입은 제공 코드의 구조체 그대로")
    pic(s, "fig14_target_card.png", CX0, 1.50, w=5.4)
    code_box(s, 6.24, 1.50, 6.48, 5.25, CODE_STRUCT, size=8.6, spacing=10.9)
    callout(s, CX0, 5.30, 5.4, 1.45, "플랫폼도 속력 0 인 표적 카드",
            "갱신 함수가 하나로 통일되고, 배가 움직이는 시나리오가 와도 카드 값만 바꾸면 된다.  MAX_TARGET 10, MAX_MANEUVER 30.", size=12)


def s15(prs):
    s = new(prs, "03.  과제 조건 넣기 — f_SetAssignment", "각도는 도로 받아 함수 안에서 라디안으로.  체크박스를 켜면 발표용 기동 예시가 붙는다")
    code_box(s, CX0, 1.50, 7.70, 4.55, CODE_ASSIGN, size=9.2, spacing=12.2)
    x, w = 8.52, 4.20
    rows = [["", "플랫폼", "표적 1 대함", "표적 2 대공"],
            ["위도 · 경도", "32.0 · 126.0", "32.125 · 126.03", "32.12 · 126.0"],
            ["고도 m", "0", "0", "300"],
            ["속력 m/s", "0", "30", "200"],
            ["yaw", "0", "270 (서)", "180 (남)"]]
    table(s, x, 1.50, w, rows, [1.00, 0.90, 1.15, 1.15], row_h=0.40, size=10, head_size=10, first_col_bold=True, align_center_cols=(1, 2, 3))
    card(s, x, 3.70, w, 2.35, "기동 예시 (화면의 체크박스)",
         ["대함 : −0.1 g yaw 20~50 s  (음수 = 좌선회)", "대공 :  +3 g yaw 10~20 s,  +2 g pitch 35~40 s", "표적을 더 넣으려면 f_SetTarget 줄을 더 쓴다 (최대 10)."], tone=ACCENT, fill=PEACH, title_size=12, body_size=10.5, line_spacing=15)
    callout(s, CX0, 6.18, CW, 0.56, "시간 60 s · 간격 0.1 s · 표적 2 개  —  과제 시뮬레이션 조건 그대로", fill=CARD, tone=TEXT, size=11.5)


def s16(prs):
    s = new(prs, "03.  속도 변환 — f_BodyVelToEcef", "6장의 회전 두 번이 열 줄.  참고 코드에서 조심할 점 두 가지가 이 화면에 다 있다")
    x, w = CX0, 4.40
    card(s, x, 1.50, w, 1.30, "① 동체 → NED", ["f_Trans_Body_To_Ned 에 [V, 0, 0] 과 자세각.", "인자 순서는 Roll, Yaw, Pitch (헤더에 적힌 대로)."], tone=BLUE, body_size=10.5, line_spacing=14)
    card(s, x, 2.92, w, 1.50, "② NED → ECEF", ["f_Trans_Ned_To_Ecef 의 가운데 셋이 플랫폼 위치.", "속도는 벡터라 여기에 0.  위경도는 표적 자신의 것."], tone=ACCENT, fill=PEACH, body_size=10.5, line_spacing=14)
    card(s, x, 4.54, w, 1.30, "초기화의 함정", ["f_Trans_Lla_To_Ecef 는 (경도, 위도, 고도) 순서.", "바꿔 넣으면 북위 54° 동경 212° 가 나온다."], tone=GREEN, fill=MINT, body_size=10.5, line_spacing=14)
    code_box(s, 5.22, 1.50, 7.50, 5.25, CODE_VEL, size=9.2, spacing=12.6)
    callout(s, x, 5.96, w, 0.78, "검산 :  |v_ecef| = 200.000000 m/s", "실수로 플랫폼 위치를 더하면 6,372 km/s", size=11.5)


def s17(prs):
    s = new(prs, "03.  한 스텝 — f_StepTarget", "8장의 다섯 단계가 함수 하나.  주석 번호가 8장 번호와 같다")
    x, w = CX0, 3.70
    specs = [("① ② 기동 판단 · 자세 갱신", "for 문. 구간 안이면 ω·dt 만큼 그 축을 돌리고 break (먼저 적힌 기동 우선). +1e-6 은 0.1 을 600번 더한 부동소수점 찌꺼기 여유.", ACCENT, PEACH),
             ("③ 속도 변환", "매 스텝 다시 부른다. 자세가 안 바뀌어도 표적의 위경도가 바뀌니까.", GREEN, MINT),
             ("④ 위치 적분", "ECEF 세 성분에 v × dt.  이 세 줄이 시간 적분.", BLUE, CARD),
             ("⑤ LLA 변환", "기록되고, 다음 스텝 ③ 의 수평면 기준이 된다 (지구 곡률 처리).", BLUE, CARD)]
    y = 1.50
    for t, l, tone, fill in specs:
        h = 1.42 if t.startswith("① ②") else 1.06
        card(s, x, y, w, h, t, [l], tone=tone, fill=fill, title_size=11.5, body_size=10, pad=0.18, line_spacing=13)
        y += h + 0.10
    code_box(s, 4.52, 1.50, 8.20, 5.25, CODE_STEP, size=8.4, spacing=10.6)
    txt(s, "속력 0 인 플랫폼도 같은 함수를 지난다 — 기동 판단에 속력 0 이면 건너뛰는 조건이 있고, 속도가 0 이라 자리만 지킨다.", x, 6.36, w, 0.40, size=9.5, color=MUTED)


def s18(prs):
    s = new(prs, "03.  루프와 화면 — TargetSimUI (MFC)", "MFC 는 계산이 없다.  코어를 600번 부르고, 배열에 모으고, 지도에 그리고, 시각을 움직인다")
    code_box(s, CX0, 1.50, 6.90, 2.55, CODE_LOOP, size=8.6, spacing=11.2)
    code_box(s, CX0, 4.18, 6.90, 1.62, CODE_MAP, size=8.6, spacing=11.2)
    pic(s, "results/ui_playing.png", 7.72, 1.50, w=5.0)
    txt(s, "2 배속 재생 중 (t = 6.2 s).  점의 짧은 선이 진행 방향.", 7.72, 4.42, 5.0, 0.30, size=10, color=MUTED, align=PP_ALIGN.CENTER)
    card(s, 7.72, 4.80, 5.0, 1.00, "CSV 저장", ["t, id, 위도, 경도, 고도, yaw 를 한 줄씩 — 과제가 요구한 위경도 출력.", "뒤의 파이썬 그림은 이 파일로 그렸다."], tone=GREEN, fill=MINT, title_size=12, body_size=10, line_spacing=13)
    callout(s, CX0, 5.96, CW, 0.78, "지도의 축 비율 1/cos(위도) 를 안 맞추면 선회 원이 타원으로 보인다.   재생은 1 배속이 실시간, 10 배속까지.", fill=CARD, tone=TEXT, size=12)


def s20(prs):
    s = new(prs, "04.  시뮬레이션 조건 — 플랫폼 하나, 표적 둘", "대공 표적은 정면으로 접근하고, 대함 표적은 정북 방향을 가로지른다")
    pic(s, "fig16_scenario_map.png", CX0, 1.50, w=5.6)
    x, w = 6.50, 6.22
    rows = [["항목", "플랫폼", "대함 표적 (1)", "대공 표적 (2)"],
            ["위도 · 경도", "32.0°N  126.0°E", "32.125°N  126.03°E", "32.12°N  126.0°E"],
            ["고도", "0 m", "0 m", "300 m"],
            ["속력 (동체 앞)", "0 m/s", "30 m/s", "200 m/s"],
            ["자세각 (r · y · p)", "—", ("0° · 270° · 0° (서쪽)", {"size": 10}), ("0° · 180° · 0° (남쪽)", {"size": 10})],
            ["시간", "60 s,  간격 0.1 s  (600 스텝)", "", ""]]
    gf = table(s, x, 1.50, w, rows, [1.56, 1.34, 1.66, 1.66], row_h=0.44, size=10.5, head_size=10.5, first_col_bold=True, align_center_cols=(1, 2, 3))
    gf.table.cell(5, 1).merge(gf.table.cell(5, 3))
    card(s, x, 4.36, w, 1.44, "플랫폼 기준으로 바꿔 보면",
         ["대공 :  정북 13.3 km,  고각 1.2°  →  정확히 우리 쪽으로 온다 (정면 접근)",
          "대함 :  북 13.9 km · 동 2.8 km,  거리 14.1 km  →  서쪽으로 가며 정북을 가로지른다",
          "두 표적의 출발점은 550 m · 2.8 km 차이로 아주 가깝다"], tone=BLUE, body_size=10.5, line_spacing=14)
    callout(s, CX0, 5.98, CW, 0.76, "같은 자리에서 출발해 하나는 남쪽으로 빠르게, 하나는 서쪽으로 천천히  —  그림에서 대비가 분명하도록 짜인 조건", fill=CARD, tone=TEXT, size=12)


def s21(prs):
    s = new(prs, "04.  실행 화면 — 과제 조건, t = 60 s", "왼쪽 위 표적 카드의 입력 칸,  오른쪽 궤적 지도,  왼쪽 아래 슬라이더 시각의 상태 칸")
    w, h = pic(s, "results/ui_assignment.png", CX0, 1.50, w=7.6)
    x, cw = 8.44, 4.28
    card(s, x, 1.50, cw, 1.66, "대공 표적 (주황)", ["60 s 뒤  32.011788°N, 126.0°E,  고도 300.02 m", "플랫폼에서 1,341 m 앞,  고각 1.2° → 12.9°", "정남으로 내려와 삼각형 바로 위에서 끝난다"], tone=ACCENT, fill=PEACH, title_size=12, body_size=10.5, line_spacing=14)
    card(s, x, 3.28, cw, 1.30, "대함 표적 (초록)", ["60 s 뒤  32.125°N, 126.010925°E", "서쪽으로 1.8 km.  선 길이 차이 = 속력 차이"], tone=GREEN, fill=MINT, title_size=12, body_size=10.5, line_spacing=14)
    card(s, x, 4.70, cw, 1.20, "플랫폼 (파랑 삼각형)", ["32.0°N 126.0°E 에서 자리를 지킨다.", "같은 스텝 함수를 지나지만 속력이 0."], tone=BLUE, title_size=12, body_size=10.5, line_spacing=14)
    callout(s, CX0, 6.06, CW, 0.68, "상태 칸의 값이 그대로 [CSV 저장] 으로 나간다  —  과제 요구 \"시뮬레이션 결과를 LLA 로 출력\"", fill=CARD, tone=TEXT, size=12)


def s22(prs):
    s = new(prs, "04.  궤적과 고도 — CSV 를 파이썬으로", "평면 궤적은 축 비율을 맞춰서,  고도는 평평해야 한다 (지구 곡률 처리의 증거)")
    pic(s, "results/assignment_map.png", CX0, 1.50, h=4.60)
    pic(s, "results/assignment_alt.png", 5.80, 1.50, w=6.92)
    card(s, 5.80, 5.52, 6.92, 0.66, "60 s 뒤 고도 300.019 m — 2 cm 의 정체", ["한 스텝에 20 m 접선으로 가서 20² / 2R 씩 뜨고, 600번이면 1.9 cm.  이론값과 같다."], tone=BLUE, title_size=11.5, body_size=10, pad=0.18, title_after=2, line_spacing=13)
    callout(s, CX0, 6.28, CW, 0.48, "매 스텝 수평면을 다시 잡지 않았다면 이 선이 300 → 311 m 로 기울었을 것", size=11.5)


def s23(prs):
    s = new(prs, "04.  기동 시나리오 — 3 g 우선회, 기수 들기, 좌선회", "체크박스를 켜고 실행.  왼쪽은 t = 30 s 화면,  오른쪽은 CSV 로 그린 궤적")
    pic(s, "results/ui_maneuver_t30.png", CX0, 1.50, w=6.6)
    pic(s, "results/maneuver_map.png", 7.70, 1.50, h=3.80)
    y, h = 5.44, 1.30
    card(s, CX0, y, 3.90, h, "대공  3 g 우선회 (10~20 s)", ["10초에 84.3° 돌아 서쪽으로.", "궤적 세 점 외접원 1,359.5 m  vs  V²/(n·g) 1,359.6 m"], tone=ACCENT, fill=PEACH, title_size=12, body_size=10.5, line_spacing=14)
    card(s, CX0 + 4.10, y, 3.90, h, "대공  2 g 기수 들기 (35~40 s)", ["5초에 28.1° 올라간 뒤 그 각으로 20초.", "60 s 에 고도 2,428.7 m"], tone=BLUE, title_size=12, body_size=10.5, line_spacing=14)
    card(s, CX0 + 8.20, y, 3.90, h, "대함  −0.1 g 좌선회 (20~50 s)", ["30초에 −56.2°, 남서쪽으로.", "반경 917.8 m  vs  이론 917.7 m"], tone=GREEN, fill=MINT, title_size=12, body_size=10.5, line_spacing=14)


def s24(prs):
    s = new(prs, "04.  검증 — 다섯 가지로 확인", "좌표변환은 틀려도 그럴듯한 값이 나온다 (과제 2의 교훈).  저장한 CSV 로 계산했다")
    rows = [["항목", "어떻게 확인하나", "기대값", "결과"],
            ["① 이동 거리", "ECEF 위치 차이의 합  vs  속력 × 시간", "200 × 60 = 12,000 m", ("12,000.000 m  (차 1e-9 m)", {"color": GREEN, "bold": True})],
            ["② 속력 보존", "매 스텝 |v_ecef|  vs  입력 속력", "200 m/s", ("최소·최대 200.000000  (편차 1e-14)", {"color": GREEN, "bold": True})],
            ["③ 고도 유지", "수평 비행 60 s 뒤 고도", "300 m", ("300.019 m  (이론 +1.9 cm)", {"color": GREEN, "bold": True})],
            ["④ 선회 반경", "기동 시작·중간·끝 세 점의 외접원  vs  V²/(n·g)", "3 g : 1,359.6 m", ("1,359.5 m  (−0.01 %)", {"color": GREEN, "bold": True})],
            ["⑤ 위경도 왕복", "LLA → ECEF → LLA → ECEF 거리", "1 mm 미만", ("3e-5 m", {"color": GREEN, "bold": True})]]
    table(s, CX0, 1.50, CW, rows, [1.70, 4.20, 2.30, 3.90], row_h=0.52, size=11, head_size=11, first_col_bold=True, align_center_cols=(2, 3))
    y = 4.82
    card(s, CX0, y, 5.96, 1.06, "왜 미리 정하나", ["코드가 내놓을 값을 먼저 알고 있어야 코드가 맞는지 알 수 있다.", "①②③⑤ 는 개념편에서 손계산한 값과 같고, ④ 는 기동 코드가 붙은 뒤 쟀다."], tone=BLUE, body_size=10.5, line_spacing=14)
    card(s, CX0 + 6.14, y, 5.96, 1.06, "한 번 틀리면 바로 드러나는 실수들", ["경도·위도 인자 순서 바꿈 → 북위 54° 동경 212°.   속도에 플랫폼 위치 더함 → 6,372 km/s.", "수평면을 한 번만 잡음 → 고도 311 m."], tone=ACCENT, fill=PEACH, body_size=10.5, line_spacing=14)
    callout(s, CX0, 6.02, CW, 0.72, "다섯 칸이 채워졌으니 궤적이 \"그림\" 이 아니라 \"결과\" 다", size=12.5)


def s26(prs):
    s = new(prs, "부록 A.  참고 코드의 함수와 함정", "제공된 CoordinateTransform.c 에서 이번에 쓴 것과, 컴파일해서 확인한 주의점")
    rows = [["함수", "이번 과제의 용도", "주의"],
            ["f_Trans_Lla_To_Ecef (Lon, Lat, Alt)", "초기 위치 입력", ("경도가 첫째 인자.  바꾸면 북위 54° 동경 212°", {"color": ACCENT, "bold": True})],
            ["f_Trans_Ecef_To_Lla (X, Y, Z)", "매 스텝 출력 + 다음 수평면 기준", "5회 고정 반복.  우리 위치에서 왕복 오차 3e-5 m"],
            ["f_Trans_Body_To_Ned (…, Roll, Yaw, Pitch)", "동체 속력 → NED 속도", ("인자 순서가 Roll, Yaw, Pitch", {"color": ACCENT, "bold": True})],
            ["f_Trans_Ned_To_Ecef (…, Pf_X, Pf_Y, Pf_Z, Lat, Lon)", "NED 속도 → ECEF 속도", ("속도에는 Pf_X·Y·Z = 0.  아니면 6,372 km/s", {"color": ACCENT, "bold": True})],
            ["Ant 계열 4개 · Enu 계열 2개 · PfCompensation · Ecef_To_Ned", "사용 안 함", "표적을 만드는 쪽이라 안테나 좌표계가 없다"]]
    table(s, CX0, 1.50, CW, rows, [4.30, 3.00, 4.80], row_h=0.52, size=10.5, head_size=11)
    callout(s, CX0, 4.90, CW, 1.84,
            "각도는 전부 라디안 (DEG2RAD 매크로는 target_sim.h 에 정의)  ·  Body → NED 행렬은 Rz(yaw)·Ry(pitch)·Rx(roll) 과 소수점 6자리까지 일치 (과제 2 와 같은 규약)",
            "그 밖에 :  한글 주석이 UTF-8 이라 MSVC /utf-8 옵션 (vcxproj 에 포함).   Define.h 의 #define BOOL bool 은 MFC 와 충돌 → MFC 쪽에서 #undef.   G_FORCE = 9.80665 는 Define.h 에 있다.",
            fill=CARD, tone=TEXT, size=11.5)


def s27(prs):
    s = new(prs, "부록 B.  용어 정리", "발표에 나온 말들.  질문이 들어오면 이 표를 보면서 답한다")
    rows = [["용어", "뜻", "이번 과제에서"],
            ["하중배수 n (g)", "안쪽으로 당기는 가속도가 중력가속도의 몇 배인가", "기동의 Gravity Value.  a = n·g"],
            ["각속도 ω", "1초에 몇 도(라디안) 도는가", "ω = n·g / V.  매 스텝 자세각 += ω × 0.1 s"],
            ["오일러 적분", "새 값 = 지금 값 + 변화율 × 시간 간격", "p ← p + v × Δt.  스텝 안 속도 일정이라 정확"],
            ["자세각 roll · pitch · yaw", "동체 축이 수평면(NED) 에 대해 돌아간 세 각", "표적의 진행 방향.  yaw 180° = 남쪽"],
            ["ECEF", "지구 중심 원점, 지구와 함께 도는 직교좌표 [m]", "위치·속도를 갱신하는 자리"],
            ["NED", "한 지점의 북 · 동 · 아래 수평 좌표", "표적 자신의 자리 기준.  매 스텝 다시 잡는다"],
            ["타원체고", "WGS-84 타원체 면에서 수직으로 잰 높이", "과제의 \"고도\".  해발고도와 약 20 m 차이"],
            ["회전행렬 R", "한 좌표계의 축을 다른 좌표계 축으로 돌려 읽는 3×3", "위치엔 R·p + t,  속도엔 R·v"]]
    table(s, CX0, 1.50, CW, rows, [2.30, 5.00, 4.80], row_h=0.50, size=10.5, head_size=11, first_col_bold=True)
    callout(s, CX0, 6.10, CW, 0.64, "이상입니다.  질문 환영합니다.       코드 · 문서 :  radar-sw-study / Part3_TargetSim", fill=CARD, tone=TEXT, size=12)


# ───────────────────────────────────────────── 조립
def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)

    build_cover(prs)                                                                                       # 1
    build_index(prs, [("00", "들어가며", "과제 3이 묻는 것"),
                      ("01", "표적을 움직이는 원리", "좌표계 넷 · 위치와 속도\n시간 적분 · 지구 곡률 · 한 스텝"),
                      ("02", "기동 모델", "하중배수 · 원운동\n세 축 · 기동 시간표"),
                      ("03", "C 로 구현", "구성 · 구조체 · 과제 조건\n속도 변환 · 한 스텝 · 화면"),
                      ("04", "결과", "조건 · 실행 화면 · 궤적\n기동 시나리오 · 검증")],
                "수학은 고등학교 물리의 등속 원운동 정도면 충분하고, 나올 때마다 다시 설명합니다.")                        # 2
    s03(prs)                                                                                               # 3
    build_divider(prs, "01", "표적을\n움직이는 원리",
                  ["좌표계 넷, 그리고 왜 ECEF 인가", "위치와 속도는 다르게 변환한다 — 회전 두 번",
                   "시간 적분과 지구 곡률 — 60초 직진하면 11 m 떠오른다", "한 스텝의 순서 — 다섯 단계"],
                  "이 발표의 약속 :  각도는 도(°) 로 말하고 프로그램 안에서는 라디안.  거리는 m, 시간은 s.")           # 4
    s05(prs); s06(prs); s07(prs); s08(prs)                                                                 # 5-8
    build_divider(prs, "02", "기동 모델", ["하중배수와 원운동 공식", "세 축과 기동 시간표"],
                  "약속 :  기동 중에도 속력은 바뀌지 않는다.  방향만 바뀐다.")                                     # 9
    s10(prs); s11(prs)                                                                                     # 10-11
    build_divider(prs, "03", "C 로 구현",
                  ["프로그램 구성 — 제공 코드 · C 코어 · MFC 화면", "표적 구조체 — target_sim.h", "과제 조건 넣기 — f_SetAssignment",
                   "속도 변환 — f_BodyVelToEcef", "한 스텝 — f_StepTarget", "루프와 화면 — TargetSimUI"],
                  "제공받은 좌표변환 코드는 한 글자도 고치지 않았다.  MFC 는 계산을 하지 않는다.")                    # 12
    s13(prs); s14(prs); s15(prs); s16(prs); s17(prs); s18(prs)                                             # 13-18
    build_divider(prs, "04", "결과", ["시뮬레이션 조건", "실행 화면", "궤적과 고도", "기동 시나리오", "검증"],
                  "화면의 값 = CSV 의 값 = 파이썬 그림의 값.  같은 코드가 만든 같은 숫자다.")                      # 19
    s20(prs); s21(prs); s22(prs); s23(prs); s24(prs)                                                       # 20-24
    build_summary(prs, ["과제 3 = 과제 2의 변환 함수 + 시간 루프 + 기동 모델.  제공 코드는 한 줄도 안 고쳤다",
                        "위치는 회전 + 평행이동, 속도는 회전만.  속도에 위치를 더하면 안 된다",
                        "새 위치 = 지금 위치 + 속도 × 0.1 초.  ECEF 에서 더한다",
                        "지구는 둥글다.  매 스텝 표적 자리의 수평면을 다시 잡아야 고도가 유지된다",
                        "기동은 원운동.  ω = n·g / V,  R = V² / (n·g).  코드가 0.01 % 로 재현했다",
                        "표적 카드 = 입력 칸 + 상태 칸.  플랫폼도 카드 한 장",
                        "검증 없는 궤적은 그림일 뿐.  거리 · 속력 · 고도 · 반경 · 왕복 다섯 가지로 확인했다"],
                  "질문 환영합니다       코드 · 문서 :  radar-sw-study / Part3_TargetSim")                        # 25
    s26(prs); s27(prs)                                                                                     # 26-27

    secs = parse_script(io.open(SCRIPT_MD, encoding="utf-8").read())
    slides = list(prs.slides)
    if len(secs) != len(slides):
        print("경고: 슬라이드 %d 장, 대본 %d 장 — 노트를 채우지 않음" % (len(slides), len(secs)))
    else:
        for i, sl in enumerate(slides, 1):
            sl.notes_slide.notes_text_frame.text = "\n".join(secs[i][1])
    prs.save(DECK)
    print("written:", DECK, "(%d slides)" % len(slides))


if __name__ == "__main__":
    main()
