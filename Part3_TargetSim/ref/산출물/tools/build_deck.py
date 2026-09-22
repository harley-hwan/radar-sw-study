# -*- coding: utf-8 -*-
"""Chapter 3 발표자료(.pptx)를 만든다.

    python build_deck.py              # ../Chapter3_TargetSim.pptx 생성
    python build_deck.py <출력.pptx>

회사 템플릿(ref/템플릿/Presentation_Template_Blank.pptx)의 색 · 글꼴 · 배치를 그대로 쓴다.
표지와 간지 배경은 템플릿에서 꺼낸 그림(figures/bg_cover.png)이다.
슬라이드 노트는 여기서 넣지 않는다 — 대본 md 를 원본으로 두고 sync_notes.py 가 밀어 넣는다.
"""
import os
import sys

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.normpath(os.path.join(HERE, ".."))
FIG = os.path.join(OUT_DIR, "figures")
TEMPLATE = os.path.normpath(os.path.join(HERE, "..", "..", "..", "..", "ref", "템플릿",
                                         "Presentation_Template_Blank.pptx"))
DECK = os.path.join(OUT_DIR, "Chapter3_TargetSim.pptx")

# ── 템플릿에서 뽑은 색 ──────────────────────────────────────────────────
NAVY = RGBColor(0x1B, 0x27, 0x40)
NAVY2 = RGBColor(0x1E, 0x2C, 0x4F)
SLATE = RGBColor(0x4A, 0x5A, 0x7A)
MUTED = RGBColor(0x87, 0x94, 0xAC)
LIGHT = RGBColor(0xAE, 0xB9, 0xCC)
PALE = RGBColor(0xD2, 0xD9, 0xE5)
BG = RGBColor(0xED, 0xF0, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLUE = RGBColor(0x25, 0x63, 0xEB)
RED = RGBColor(0xDC, 0x26, 0x26)
GREEN = RGBColor(0x05, 0x96, 0x05 + 0x64)
OK = RGBColor(0x05, 0x96, 0x69)
AMBER = RGBColor(0xD9, 0x77, 0x06)
VIOLET = RGBColor(0x7C, 0x3A, 0xED)
CODE_BG = RGBColor(0xF7, 0xF9, 0xFC)

KO = "맑은 고딕"
MONO = "Consolas"

SW, SH = 13.333, 7.5
MX, MW = 0.70, 11.93            # 본문 좌우 여백과 폭 (템플릿 머리말 띠와 같은 자리)
HEAD_Y, HEAD_H = 0.62, 0.62
FOOT_Y = 6.92

TITLE = "Part 3  표적 궤적 모의"
AUTHOR = "Janghwan Kim (Harley)"

prs = Presentation(TEMPLATE)
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)

# 템플릿에 들어 있던 예시 17장을 지운다. 레이아웃(DEFAULT)과 테마는 그대로 남는다.
_ids = prs.slides._sldIdLst
for _sld in list(_ids):
    prs.part.drop_rel(_sld.rId)
    _ids.remove(_sld)
BLANK = prs.slide_layouts[0]

_page = 0
_titles = []


# ── 낱개 도구 ───────────────────────────────────────────────────────────
def txt(slide, x, y, w, h, text, size=12, color=NAVY, bold=False, font=KO,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=1.0, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = space
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = font
        r.font.color.rgb = color
    return tb


def rect(slide, x, y, w, h, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    sh.text_frame.word_wrap = True
    return sh


def hline(slide, x, y, w, color=PALE, lw=0.75):
    ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Emu(0))
    ln.shadow.inherit = False
    ln.fill.background()
    ln.line.color.rgb = color
    ln.line.width = Pt(lw)
    return ln


def picture(slide, name, box, align="center"):
    """box = (x, y, w, h). 그림 비율을 지키며 상자 안에 넣는다."""
    path = os.path.join(FIG, name)
    iw, ih = Image.open(path).size
    bx, by, bw, bh = box
    s = min(bw / iw, bh / ih)
    w, h = iw * s, ih * s
    x = bx + (bw - w) / 2 if align == "center" else bx
    y = by + (bh - h) / 2
    return slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))


# ── 쪽 틀 ───────────────────────────────────────────────────────────────
def new_slide():
    return prs.slides.add_slide(BLANK)


def content(chapter, heading, lead=None):
    """본문 쪽 — 머리말 띠 + 제목 + (한 줄 설명) + 꼬리말. 본문 상자의 y 를 돌려준다."""
    global _page
    _page += 1
    slide = new_slide()
    _titles.append(heading)

    rect(slide, MX, HEAD_Y, MW, HEAD_H, fill=NAVY)
    txt(slide, MX + 0.28, HEAD_Y, MW - 0.56, HEAD_H, chapter, size=13.5, color=WHITE,
        bold=True, anchor=MSO_ANCHOR.MIDDLE)

    txt(slide, MX + 0.25, 1.36, MW - 0.5, 0.44, heading, size=20, color=NAVY, bold=True)
    y = 1.94
    if lead:
        txt(slide, MX + 0.25, 1.90, MW - 0.5, 0.30, lead, size=11.5, color=SLATE)
        y = 2.34

    hline(slide, MX, FOOT_Y, MW, PALE)
    txt(slide, 6.07, 6.98, 1.20, 0.30, str(_page), size=9, color=MUTED,
        align=PP_ALIGN.CENTER, font="Arial")
    return slide, y


def cover():
    global _page
    _page += 1
    slide = new_slide()
    _titles.append("표지")
    slide.shapes.add_picture(os.path.join(FIG, "bg_cover.png"), 0, 0,
                             Inches(SW), Inches(SH))
    txt(slide, 9.00, 0.72, 3.30, 0.30, "Confidential", size=10, color=WHITE, bold=True,
        font="Arial", align=PP_ALIGN.RIGHT)
    txt(slide, 4.60, 2.74, 7.40, 0.50, "레이다 SW 학습 과제 3", size=14, color=LIGHT,
        align=PP_ALIGN.RIGHT)
    txt(slide, 4.60, 3.22, 7.40, 0.75, "표적 궤적 모의", size=38, color=WHITE, bold=True,
        align=PP_ALIGN.RIGHT)
    hline(slide, 8.60, 4.28, 3.40, LIGHT)
    txt(slide, 4.60, 4.44, 7.40, 0.32, "Target Trajectory Simulation", size=13,
        color=LIGHT, font="Arial", align=PP_ALIGN.RIGHT)
    txt(slide, 4.60, 5.60, 7.40, 0.32, AUTHOR, size=12, color=WHITE,
        align=PP_ALIGN.RIGHT)
    return slide


def section(no, ko, en, points):
    global _page
    _page += 1
    slide = new_slide()
    _titles.append("%s. %s" % (no, ko))
    slide.shapes.add_picture(os.path.join(FIG, "bg_cover.png"), 0, 0,
                             Inches(SW), Inches(SH))
    hline(slide, 0.82, 0.78, 11.70, SLATE)
    hline(slide, 0.82, 6.93, 11.70, SLATE)
    txt(slide, 0.94, 2.55, 3.00, 1.10, no, size=64, color=SLATE, bold=True, font="Arial")
    txt(slide, 2.60, 2.72, 8.00, 0.66, ko, size=32, color=WHITE, bold=True)
    txt(slide, 2.66, 3.50, 8.00, 0.34, en, size=13, color=LIGHT, font="Arial")
    for i, p in enumerate(points):
        txt(slide, 2.66, 4.30 + i * 0.42, 8.60, 0.34, "·   " + p, size=12.5, color=LIGHT)
    txt(slide, 6.07, 6.98, 1.20, 0.30, str(_page), size=9, color=MUTED,
        align=PP_ALIGN.CENTER, font="Arial")
    return slide


# ── 본문에서 자주 쓰는 덩어리 ───────────────────────────────────────────
def card(slide, x, y, w, h, title, body, accent=BLUE, tsize=12.5, bsize=10.5):
    rect(slide, x, y, w, h, fill=WHITE, line=PALE, lw=1.0)
    rect(slide, x, y, 0.055, h, fill=accent)
    txt(slide, x + 0.30, y + 0.20, w - 0.55, 0.30, title, size=tsize, color=accent, bold=True)
    txt(slide, x + 0.30, y + 0.58, w - 0.55, h - 0.72, body, size=bsize, color=SLATE,
        space=1.28)


def code(slide, x, y, w, h, lines, size=10.0, lead=1.30, title=None):
    """소스 코드 상자. lines 는 (글자, 종류) 목록 — 종류: '', 'c'(주석), 'k'(강조)."""
    rect(slide, x, y, w, h, fill=CODE_BG, line=PALE, lw=1.0)
    if title:
        txt(slide, x + 0.22, y + 0.14, w - 0.44, 0.26, title, size=9.5, color=MUTED,
            font="Arial")
        y0 = y + 0.50
    else:
        y0 = y + 0.20
    tb = slide.shapes.add_textbox(Inches(x + 0.22), Inches(y0), Inches(w - 0.44),
                                  Inches(h - (y0 - y) - 0.16))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(lines):
        text, kind = item if isinstance(item, tuple) else (item, "")
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = lead
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = text if text else " "
        r.font.size = Pt(size)
        r.font.name = MONO
        if kind == "c":
            r.font.color.rgb = MUTED
        elif kind == "k":
            r.font.color.rgb = BLUE
            r.font.bold = True
        elif kind == "r":
            r.font.color.rgb = RED
            r.font.bold = True
        else:
            r.font.color.rgb = NAVY
    return tb


def banner(slide, x, y, w, h, text, fill=NAVY, color=WHITE, size=12):
    rect(slide, x, y, w, h, fill=fill)
    txt(slide, x + 0.30, y, w - 0.60, h, text, size=size, color=color, bold=True,
        anchor=MSO_ANCHOR.MIDDLE, space=1.3)


def table(slide, x, y, w, cols, rows, widths, hsize=10.5, bsize=10.5, rh=0.42):
    """가벼운 표 — 머리줄은 남색 띠, 본문은 줄무늬."""
    rect(slide, x, y, w, rh, fill=NAVY)
    cx = x
    for c, cw in zip(cols, widths):
        txt(slide, cx + 0.16, y, cw - 0.2, rh, c, size=hsize, color=WHITE, bold=True,
            anchor=MSO_ANCHOR.MIDDLE)
        cx += cw
    for i, row in enumerate(rows):
        ry = y + rh + i * rh
        if i % 2 == 0:
            rect(slide, x, ry, w, rh, fill=BG)
        cx = x
        for j, (v, cw) in enumerate(zip(row, widths)):
            bold = j == 0
            col = NAVY if j == 0 else SLATE
            fnt = KO
            if v.startswith("~"):
                v, col, fnt, bold = v[1:], OK, MONO, True
            elif v.startswith("`"):
                v, fnt = v[1:], MONO
            txt(slide, cx + 0.16, ry, cw - 0.2, rh, v, size=bsize, color=col, bold=bold,
                font=fnt, anchor=MSO_ANCHOR.MIDDLE)
            cx += cw
    return y + rh + len(rows) * rh


CH1 = "01. 표적 구조체 설계"
CH2 = "02. 표적 궤적 모의 구현"
CH3 = "03. 시뮬레이션 결과"
CH0 = "들어가며"


def body_box(y):
    return (MX, y, MW, 6.80 - y)


# ═══════════════════════════════════════════════════════════════════════
# 1. 표지
# ═══════════════════════════════════════════════════════════════════════
cover()

# ═══════════════════════════════════════════════════════════════════════
# 2. 목차
# ═══════════════════════════════════════════════════════════════════════
_page += 1
s = new_slide()
_titles.append("목차")
rect(s, 0, 0, 4.60, SH, fill=NAVY)
txt(s, 0.85, 0.55, 2.00, 0.45, "INDEX", size=26, color=WHITE, bold=True, font="Arial")
hline(s, 2.30, 0.79, 2.30, SLATE)
hline(s, 4.60, 0.79, 7.95, PALE)
txt(s, 0.50, 3.05, 3.60, 1.20,
    "표적 궤적 모의\nTargetSim", size=18, color=WHITE, bold=True,
    align=PP_ALIGN.CENTER, space=1.45)
txt(s, 0.50, 6.60, 3.60, 0.30, AUTHOR, size=10, color=LIGHT, align=PP_ALIGN.CENTER)

_idx = [("01", "표적 구조체 설계",
         "표적 하나를 어떤 숫자로 적을 것인가\n기동을 어떻게 표현할 것인가"),
        ("02", "표적 궤적 모의 구현",
         "동체 속도를 ECEF 속도로 바꾸는 방법\n한 스텝을 전진시키는 방법"),
        ("03", "시뮬레이션 결과",
         "명세 조건 60초 결과와 검산\n정확도 확인과 기동 시연")]
for i, (n, t, d) in enumerate(_idx):
    y = 1.95 + i * 1.60
    txt(s, 5.25, y, 1.00, 0.55, n, size=30, color=NAVY, bold=True, font="Arial")
    txt(s, 6.40, y + 0.08, 2.90, 0.40, t, size=15, color=NAVY, bold=True)
    txt(s, 9.45, y - 0.02, 3.20, 0.95, d, size=10.5, color=SLATE, space=1.45)
    if i < 2:
        hline(s, 5.25, y + 1.22, 7.35, PALE)
txt(s, 8.37, 6.98, 1.20, 0.30, str(_page), size=9, color=MUTED, align=PP_ALIGN.CENTER,
    font="Arial")

# ═══════════════════════════════════════════════════════════════════════
# 3. 과제가 요구한 것
# ═══════════════════════════════════════════════════════════════════════
s, y = content(CH0, "과제가 요구한 것", "「과제 3 내용」 문서의 요구사항을 네 덩어리로 묶고, 발표도 그 순서로 갑니다.")
_req = [("표적 구조체 정의",
         "표적 최대 10개, 표적마다 기동 최대 30개.\n위치는 위도 · 경도 · 고도, 속도는 동체\n속도(V_heading) 한 개로 입력받을 것.", BLUE),
        ("표적 상태 갱신",
         "위치와 속도는 ECEF 좌표계에서 갱신할 것.\n추가 고민 — V_heading 을 ECEF 속도\nVx, Vy, Vz 로 바꾸는 방법.", AMBER),
        ("시뮬레이션 조건",
         "60초 동안 0.1초 간격.\n플랫폼 1 + 대함 표적 1 + 대공 표적 1.\n제공받은 좌표변환 코드를 그대로 사용.", VIOLET),
        ("결과 제시",
         "궤적을 LLA 로 출력하고 그림으로 표현.\n구조체 정의 · 구현 방법 · 결과를\n코드와 함께 설명.", OK)]
for i, (t, d, c) in enumerate(_req):
    card(s, MX + i * 3.03, y + 0.25, 2.80, 2.55, t, d, accent=c)
banner(s, MX, y + 3.20, MW, 0.95,
       "코드는 Part 2 에서 쓰던 좌표변환 라이브러리를 한 줄도 고치지 않고 그대로 가져다 썼습니다.\n"
       "새로 만든 것은 TargetSim.c 하나 — 표적을 시간에 따라 움직이는 부분입니다.")

# ═══════════════════════════════════════════════════════════════════════
# 4. 왜 표적을 모의하는가
# ═══════════════════════════════════════════════════════════════════════
s, y = content(CH0, "표적을 왜 가짜로 만드나요?",
               "레이다 소프트웨어를 만들려면 먼저 \"표적이 이렇게 움직였다\" 는 데이터가 있어야 합니다.")
picture(s, "fig01_why.png", body_box(y))

# ═══════════════════════════════════════════════════════════════════════
# 5. 무엇을 만들었나
# ═══════════════════════════════════════════════════════════════════════
s, y = content(CH0, "만든 것 — 계산은 C, 화면은 MFC",
               "계산 엔진과 화면을 나눠 두면 화면이 바뀌어도 계산은 건드릴 일이 없습니다.")
picture(s, "fig02_system.png", body_box(y))

# ═══════════════════════════════════════════════════════════════════════
# 6. 섹션 01
# ═══════════════════════════════════════════════════════════════════════
section("01", "표적 구조체 설계", "Target Structure Design",
        ["표적 하나를 숫자 몇 개로 적어야 충분한가",
         "\"기동\" 이라는 말을 코드가 알아들을 수 있는 값으로 바꾸기",
         "설정(입력)과 상태(출력)를 나눈 이유"])

# ═══════════════════════════════════════════════════════════════════════
# 7. 표적 하나를 숫자로
# ═══════════════════════════════════════════════════════════════════════
s, y = content(CH1, "표적 하나를 숫자로 적으면",
               "지금 어디에 있고, 어디를 보고 있고, 얼마나 빠른가 — 이 세 가지면 다음 순간을 계산할 수 있습니다.")
picture(s, "fig03_state7.png", body_box(y))

# ═══════════════════════════════════════════════════════════════════════
# 8. 위치 — LLA
# ═══════════════════════════════════════════════════════════════════════
s, y = content(CH1, "위치는 위도 · 경도 · 고도로 받는다",
               "명세가 정한 입력 형식이기도 하고, 사람이 지도에서 바로 찍어 볼 수 있는 값이기도 합니다.")
picture(s, "fig04_lla.png", body_box(y))

# ═══════════════════════════════════════════════════════════════════════
# 9. 속도 — V_heading
# ═══════════════════════════════════════════════════════════════════════
s, y = content(CH1, "속도는 크기 하나만 받는다",
               "방향은 자세각이 이미 들고 있으므로, 속도까지 방향을 갖게 하면 두 개가 서로 어긋납니다.")
picture(s, "fig05_speed.png", body_box(y))

# ═══════════════════════════════════════════════════════════════════════
# 10. 기동
# ═══════════════════════════════════════════════════════════════════════
s, y = content(CH1, "기동을 어떻게 적을 것인가",
               "\"12초부터 16초까지 오른쪽으로 8 G 로 꺾어라\" — 이 한 문장을 네 개의 값으로 쪼갭니다.")
picture(s, "fig06_maneuver.png", body_box(y))

# ═══════════════════════════════════════════════════════════════════════
# 11. G → 각속도
# ═══════════════════════════════════════════════════════════════════════
s, y = content(CH1, "G 값 하나가 회전 속도를 정한다",
               "왜 각도(°/s) 대신 G 로 받을까요? 조종사와 기체가 실제로 견디는 한계가 G 이기 때문입니다.")
picture(s, "fig07_gforce.png", body_box(y))

# ═══════════════════════════════════════════════════════════════════════
# 12. 회전축 세 가지
# ═══════════════════════════════════════════════════════════════════════
s, y = content(CH1, "회전축 세 가지 — 실제로 돌려 보면",
               "같은 표적에 같은 2 G 를 10초씩, 축만 바꿔 걸어 본 결과입니다.")
picture(s, "fig08_turntype.png", body_box(y))

# ═══════════════════════════════════════════════════════════════════════
# 13. 구조체 전체
# ═══════════════════════════════════════════════════════════════════════
s, y = content(CH1, "구조체 전체 그림",
               "입력 쪽과 출력 쪽을 완전히 갈라 두었습니다.")
picture(s, "fig09_struct.png", body_box(y))

# ═══════════════════════════════════════════════════════════════════════
# 14. 코드 ① 기동 · 초기 설정
# ═══════════════════════════════════════════════════════════════════════
s, y = content(CH1, "코드로 보는 구조체 ① — 입력",
               "TargetSimCore/include/TargetSim.h")
code(s, MX, y + 0.02, 7.30, 4.42, [
    ("// 기동 회전축. 값은 과제 명세를 그대로 따른다.", "c"),
    ("typedef enum", "k"),
    "{",
    "    TGT_TURN_NONE = 0,   TGT_TURN_ROLL  = 1,",
    "    TGT_TURN_YAW  = 2,   TGT_TURN_PITCH = 3",
    "} EN_TurnType;",
    "",
    ("// 기동 1건. gravityValue 의 부호가 회전 방향이다.", "c"),
    ("typedef struct", "k"),
    "{",
    "    EN_TurnType   enTurnType;",
    "    FLOAT64       gravityValue;   // [G]",
    "    FLOAT64       startTime;      // [s]",
    "    FLOAT64       endTime;        // [s]",
    "} ST_TargetManeuver;",
    "",
    ("typedef struct", "k"),
    "{",
    "    STRUCT_Coord_Lla       st_InitLla;",
    "    STRUCT_Coord_Attitude  st_InitAtt;",
    "    FLOAT64                headingSpeed;   // V_heading",
    "    INT32                  nManeuverNum;",
    "    ST_TargetManeuver      st_Maneuver[TGT_MAX_MANEUVER_NUM];",
    "} ST_TargetInit;",
], size=9.0, lead=1.12)
card(s, 8.35, y + 0.02, 4.28, 1.38, "구간은 [시작, 종료)",
     "한 기동의 끝과 다음 기동의 시작이 같아도\n겹침이 아니라 이어 붙는 것으로 봅니다.", accent=BLUE,
     bsize=10)
card(s, 8.35, y + 1.54, 4.28, 1.38, "겹치면 아예 거부",
     "같은 표적에 같은 시각을 덮는 기동이 둘이면\nTGT_ERR_MANEUVER_OVERLAP 을 돌려줍니다.", accent=AMBER,
     bsize=10)
card(s, 8.35, y + 3.06, 4.28, 1.38, "10 × 30 은 배열로 고정",
     "동적 할당을 쓰지 않아 실행 중 메모리 실패가\n생길 자리가 없습니다.", accent=OK, bsize=10)

# ═══════════════════════════════════════════════════════════════════════
# 15. 코드 ② 상태 · 표본
# ═══════════════════════════════════════════════════════════════════════
s, y = content(CH1, "코드로 보는 구조체 ② — 출력",
               "설정은 처음 한 번만 읽고, 스텝마다 바뀌는 것은 상태뿐입니다.")
code(s, MX, y + 0.05, 7.30, 4.25, [
    ("// 한 시각의 상태. 위치와 속도는 ECEF 에서 갱신하고,", "c"),
    ("// LLA 는 출력과 다음 스텝의 NED 기준으로 함께 보관한다.", "c"),
    ("typedef struct", "k"),
    "{",
    "    FLOAT64                simTime;      // [s]",
    ("    STRUCT_Coord_Rect      st_PosEcef;   // [m]", "k"),
    ("    STRUCT_Coord_Rect      st_VelEcef;   // [m/s]", "k"),
    "    STRUCT_Coord_Lla       st_Lla;",
    "    STRUCT_Coord_Attitude  st_Att;",
    "} ST_TargetState;",
    "",
    ("// 한 시각의 플랫폼 · 표적 상태.", "c"),
    ("typedef struct", "k"),
    "{",
    "    FLOAT64         simTime;",
    "    INT32           nStepIndex;",
    "    INT32           nTargetNum;",
    "    ST_TargetState  st_Platform;",
    "    ST_TargetState  st_Target[TGT_MAX_TARGET_NUM];",
    "} ST_SimSample;",
], size=9.4, lead=1.16)
card(s, 8.35, y + 0.05, 4.28, 2.05, "왜 ECEF 와 LLA 를 둘 다 들고 있나",
     "계산은 ECEF 에서 하는 쪽이 정확하고,\n출력과 화면은 LLA 라야 읽힙니다.\n"
     "스텝마다 한 번만 바꿔 두면 쓸 때마다\n다시 변환하지 않아도 됩니다.", accent=BLUE, bsize=10)
card(s, 8.35, y + 2.30, 4.28, 2.00, "플랫폼도 같은 구조체",
     "플랫폼은 기동이 없는 표적일 뿐입니다.\n같은 ST_TargetState 를 쓰고 각속도만\n"
     "0 으로 넣어 같은 경로로 전진시킵니다.", accent=OK, bsize=10)

# ═══════════════════════════════════════════════════════════════════════
# 16. 섹션 02
# ═══════════════════════════════════════════════════════════════════════
section("02", "표적 궤적 모의 구현", "Trajectory Propagation",
        ["기수 방향 속력 하나를 ECEF 속도 세 성분으로 바꾸기",
         "0.1초를 한 번 전진시키는 방법 — 중점법",
         "숫자가 깨졌을 때 상태를 지키는 방법"])

# ═══════════════════════════════════════════════════════════════════════
# 17. 좌표계 셋
# ═══════════════════════════════════════════════════════════════════════
s, y = content(CH2, "쓰는 좌표계는 셋뿐입니다",
               "Part 2 에서 여섯 개를 봤지만, 표적을 움직이는 데 실제로 필요한 것은 이 셋입니다.")
picture(s, "fig10_frames.png", body_box(y))

# ═══════════════════════════════════════════════════════════════════════
# 18. 왜 ECEF
# ═══════════════════════════════════════════════════════════════════════
s, y = content(CH2, "왜 ECEF 에서 갱신하는가",
               "명세가 정한 조건이기도 하지만, 실제로 계산해 보면 이유가 분명합니다.")
picture(s, "fig11_whyecef.png", body_box(y))

# ═══════════════════════════════════════════════════════════════════════
# 19. V_heading → ECEF
# ═══════════════════════════════════════════════════════════════════════
s, y = content(CH2, "동체 속도를 ECEF 속도로 — 과제의 추가 고민",
               "\"표적 동체 속도 값(V_heading)을 ECEF 좌표계 Vx, Vy, Vz 성분으로 변환하는 방법\" 에 대한 답입니다.")
picture(s, "fig12_velocity.png", body_box(y))

# ═══════════════════════════════════════════════════════════════════════
# 20. 코드 ③ GetVelEcef
# ═══════════════════════════════════════════════════════════════════════
s, y = content(CH2, "코드 ③ — 두 줄이면 끝난다")
code(s, MX, y + 0.05, 12.63 - MX, 3.22, [
    ("// 동체 속도 (V, 0, 0) 을 자세로 한 번, 위경도로 한 번 돌려 ECEF 속도로 만든다.", "c"),
    ("// 회전만 거치므로 속력은 그대로 보존된다.", "c"),
    ("static STRUCT_Coord_Rect f_Tgt_GetVelEcef(const STRUCT_Coord_Attitude *st_Att,", "k"),
    "        const STRUCT_Coord_Lla *st_Lla, FLOAT64 headingSpeed)",
    "{",
    "    STRUCT_Coord_Rect  st_VelNed;",
    "",
    ("    // 참고 소스는 인자를 Roll, Yaw, Pitch 순서로 받는다.", "c"),
    ("    st_VelNed = f_Trans_Body_To_Ned(headingSpeed, 0.0, 0.0,", "k"),
    "            st_Att->Roll, st_Att->Yaw, st_Att->Pitch);",
    "",
    ("    // 방향만 옮길 때는 기준점을 원점으로 준다.", "c"),
    ("    return f_Trans_Ned_To_Ecef(st_VelNed.x, st_VelNed.y, st_VelNed.z,", "k"),
    "            0.0, 0.0, 0.0, st_Lla->Lat, st_Lla->Lon);",
    "}",
], size=9.4, lead=1.16, title="TargetSimCore/src/TargetSim.c")
card(s, MX, y + 3.40, 3.85, 1.42, "왜 기준점이 0 인가",
     "f_Trans_Ned_To_Ecef 는 회전에 더해\n기준점만큼 평행이동까지 합니다. 속도는\n\"위치\"가 아니라 \"방향\"이라 빼야 합니다.", accent=RED, bsize=9.5)
card(s, MX + 4.05, y + 3.40, 3.85, 1.42, "인자 순서를 왜 적어 두었나",
     "제공받은 함수가 Roll, Yaw, Pitch 순서로\n받습니다. 자세각 구조체의 필드 순서와\n달라, 헷갈리는 자리라 주석을 남겼습니다.", accent=AMBER, bsize=9.5)
card(s, MX + 8.10, y + 3.40, 3.83, 1.42, "속력 보존 확인",
     "회전행렬은 길이를 바꾸지 않습니다.\n60초 601개 표본에서 속력을 재보면\n200.000000000 m/s 가 601번 나옵니다.", accent=OK, bsize=9.5)

# ═══════════════════════════════════════════════════════════════════════
# 21. 중점법
# ═══════════════════════════════════════════════════════════════════════
s, y = content(CH2, "0.1초를 어떻게 전진시킬 것인가",
               "속도가 계속 바뀌는데 한 방향으로만 밀면 안쪽 곡선을 못 따라갑니다.")
picture(s, "fig13_midpoint.png", body_box(y))

# ═══════════════════════════════════════════════════════════════════════
# 22. 코드 ④ Propagate
# ═══════════════════════════════════════════════════════════════════════
s, y = content(CH2, "코드 ④ — 한 스텝 전진",
               "TargetSimCore/src/TargetSim.c   f_Tgt_Propagate (발췌)")
code(s, MX, y + 0.05, 7.60, 4.25, [
    ("    halfStep = 0.5 * stepTime;", ""),
    "",
    ("    // ① 반 스텝 뒤의 자세", "c"),
    "    st_AttMid.Roll  = st_State->st_Att.Roll  + (st_Rate->Roll  * halfStep);",
    "    st_AttMid.Pitch = st_State->st_Att.Pitch + (st_Rate->Pitch * halfStep);",
    "    st_AttMid.Yaw   = st_State->st_Att.Yaw   + (st_Rate->Yaw   * halfStep);",
    "",
    ("    // ② 반 스텝 뒤의 위치 — 스텝 시작 속도로 예측", "c"),
    "    st_VelStart = f_Tgt_GetVelEcef(&st_State->st_Att, &st_State->st_Lla, headingSpeed);",
    "    st_PosMid.x = st_State->st_PosEcef.x + (st_VelStart.x * halfStep);",
    ("    ...", "c"),
    "",
    ("    // ③ 그 자리 · 그 자세에서 구한 속도가 이 스텝의 대표 속도", "c"),
    ("    st_LlaMid = f_Trans_Ecef_To_Lla(st_PosMid.x, st_PosMid.y, st_PosMid.z);", "k"),
    ("    st_VelMid = f_Tgt_GetVelEcef(&st_AttMid, &st_LlaMid, headingSpeed);", "k"),
    "",
    ("    // ④ 그 속도로 한 스텝을 옮긴다", "c"),
    ("    st_Next.st_PosEcef.x = st_State->st_PosEcef.x + (st_VelMid.x * stepTime);", "k"),
    ("    ...", "c"),
], size=9.0, lead=1.14)
card(s, 8.50, y + 0.05, 4.13, 2.05, "왜 반 스텝 위치까지 다시 구하나",
     "자세만 반 스텝 돌리고 위치는 그대로 두면,\n지구가 둥근 만큼 방향이 뒤처집니다.\n"
     "그러면 수평으로 날아도 고도가 매 스텝\nd²/2R 씩 올라갑니다.", accent=BLUE, bsize=9.8)
card(s, 8.50, y + 2.30, 4.13, 2.00, "짐벌락은 알고 남겨 두었다",
     "오일러각을 직접 적분하므로 Pitch 가\n±90° 근처면 Yaw 축이 속도 방향과 겹쳐\n"
     "Yaw 기동이 궤적을 바꾸지 못합니다.\n코드에 주석으로 명시해 두었습니다.", accent=AMBER,
     bsize=9.8)

# ═══════════════════════════════════════════════════════════════════════
# 23. 한 스텝 전체 + 안전장치
# ═══════════════════════════════════════════════════════════════════════
s, y = content(CH2, "한 스텝에 벌어지는 일 전부",
               "그리고 계산이 깨졌을 때 상태를 지키는 방법.")
picture(s, "fig14_stepflow.png", body_box(y))

# ═══════════════════════════════════════════════════════════════════════
# 24. 섹션 03
# ═══════════════════════════════════════════════════════════════════════
section("03", "시뮬레이션 결과", "Simulation Results",
        ["명세 조건 그대로 60초를 돌린 결과",
         "손으로 계산한 값과 맞는지 숫자로 검산",
         "갱신 간격을 바꿔 가며 확인한 정확도"])

# ═══════════════════════════════════════════════════════════════════════
# 25. 시뮬레이션 조건
# ═══════════════════════════════════════════════════════════════════════
s, y = content(CH3, "시뮬레이션 조건", "「과제 3 내용」 2항의 값을 그대로 넣었습니다.")
end = table(s, MX, y + 0.10, MW,
            ["", "위도 [°]", "경도 [°]", "고도 [m]", "속력 [m/s]", "Roll [°]", "Pitch [°]", "Yaw [°]"],
            [["플랫폼", "32.0", "126.0", "0", "0", "0", "0", "0"],
             ["대함 표적", "32.125", "126.03", "0", "30", "0", "0", "270"],
             ["대공 표적", "32.12", "126.0", "300", "200", "0", "0", "180"]],
            [2.13, 1.40, 1.40, 1.40, 1.40, 1.40, 1.40, 1.40], rh=0.46)
card(s, MX, end + 0.35, 3.85, 1.35, "시뮬레이션 시간 60 초",
     "갱신 간격 0.1 초 → 601개 표본\n(t = 0 과 t = 60 을 모두 포함)", accent=BLUE, bsize=10.5)
card(s, MX + 4.05, end + 0.35, 3.85, 1.35, "Yaw 270° 는 서쪽",
     "Yaw 는 북쪽이 0, 시계 방향이 +.\n대함은 서쪽, 대공은 남쪽(180°)으로 갑니다.", accent=AMBER,
     bsize=10.5)
card(s, MX + 8.10, end + 0.35, 3.83, 1.35, "기동 없는 등속 직진",
     "먼저 기동 없이 돌려 손계산과 맞춰 보고,\n그 다음 기동을 얹었습니다.", accent=OK, bsize=10.5)
banner(s, MX, end + 1.82, MW, 0.62,
       "플랫폼 속력이 0 이라 f_Tgt_Propagate 의 나눗셈 분모가 되는 자리가 없는지 따로 확인했습니다 — 각속도 계산은 기동이 있을 때만 들어갑니다.",
       fill=BG, color=SLATE, size=11)

# ═══════════════════════════════════════════════════════════════════════
# 26. 결과 ① 궤적
# ═══════════════════════════════════════════════════════════════════════
s, y = content(CH3, "결과 ① — 60초 궤적",
               "플랫폼을 원점으로 잡은 동-북 평면입니다. 점 하나가 10초입니다.")
picture(s, "fig15_result_map.png", body_box(y))

# ═══════════════════════════════════════════════════════════════════════
# 27. 결과 ② 검산
# ═══════════════════════════════════════════════════════════════════════
s, y = content(CH3, "결과 ② — 손으로 검산해 보면",
               "속력 × 시간, 고도 유지, 위도 유지 — 답을 미리 아는 항목으로 맞춰 봤습니다.")
picture(s, "fig16_verify.png", body_box(y))

# ═══════════════════════════════════════════════════════════════════════
# 28. 결과 ③ 정확도
# ═══════════════════════════════════════════════════════════════════════
s, y = content(CH3, "결과 ③ — 0.1초는 충분한 간격인가",
               "직진은 틀릴 수가 없으니, 가장 많이 휘는 선회 구간으로 확인했습니다.")
picture(s, "fig17_convergence.png", body_box(y))

# ═══════════════════════════════════════════════════════════════════════
# 29. 결과 ④ 기동
# ═══════════════════════════════════════════════════════════════════════
s, y = content(CH3, "결과 ④ — 기동을 얹으면",
               "같은 초기값에 기동 10줄만 추가했습니다. 초기값 · 시간 · 간격은 명세 그대로입니다.")
picture(s, "fig18_zigzag.png", body_box(y))

# ═══════════════════════════════════════════════════════════════════════
# 30. UI
# ═══════════════════════════════════════════════════════════════════════
s, y = content(CH3, "TargetSim — 값을 바꿔 가며 볼 수 있게",
               "표에서 숫자를 고치면 바로 다시 돌고, 그림에서 시작점과 기수를 끌 수도 있습니다.")
picture(s, "fig19_ui.png", (MX, y - 0.05, 7.55, 4.50))
card(s, 8.55, y + 0.00, 4.08, 1.40, "고치면 곧바로 다시 돈다",
     "표의 칸을 고치면 0.25초 뒤 자동으로\n다시 계산합니다. 601 표본 × 3 객체가\n5 ms 안에 끝나 기다릴 일이 없습니다.", accent=BLUE, bsize=9.8)
card(s, 8.55, y + 1.53, 4.08, 1.40, "그림에서 끌어서 배치",
     "시작점과 기수 손잡이를 마우스로 끌면\n위도 · 경도 · Yaw 가 표에 그대로\n반영됩니다.", accent=AMBER, bsize=9.8)
card(s, 8.55, y + 3.06, 4.08, 1.44, "결과는 표 · CSV · 그림",
     "시각마다의 LLA 를 표로 보고, 슬라이더로\n되감아 보고, CSV 로 내보냅니다.\n시나리오는 .tsim 파일로 저장합니다.", accent=OK, bsize=9.8)

# ═══════════════════════════════════════════════════════════════════════
# 31. 정리
# ═══════════════════════════════════════════════════════════════════════
s, y = content(CH3, "정리", "요구사항과 결과를 한 장에 놓고 보겠습니다.")
end = table(s, MX, y + 0.05, MW,
            ["과제 요구사항", "이렇게 했습니다", "확인"],
            [["표적 10개 · 기동 30개", "고정 배열로 상한을 두고 범위 밖은 오류 코드로 거부", "~완료"],
             ["위치는 LLA, 속도는 V_heading", "ST_TargetInit 에 그대로. 방향은 자세각이 담당", "~완료"],
             ["기동 = G · 축 · 시작 · 종료", "ST_TargetManeuver 네 필드. 구간 겹침도 검사", "~완료"],
             ["ECEF 에서 상태 갱신", "위치 · 속도 모두 ECEF, 출력용 LLA 를 함께 보관", "~완료"],
             ["V_heading → Vx, Vy, Vz  (추가 고민)", "동체 → NED → ECEF 두 번 회전. 속력 보존 확인", "~완료"],
             ["제공 좌표변환 코드 사용", "CoordinateTransform.c / matrixCalcLib.c 원본 그대로", "~완료"],
             ["결과를 LLA 로 출력하고 그림으로", "표 · CSV · 궤적 그림 · 고도 그림", "~완료"]],
            [5.30, 5.60, 1.03], rh=0.365, bsize=9.8, hsize=10)
card(s, MX, end + 0.26, 5.85, 1.22, "남겨 둔 것",
     "· Pitch 가 ±90° 근처면 짐벌락 — 오일러각 대신 사원수를 쓰면 풀립니다\n"
     "· 중력 · 항력 없이 기수 방향으로만 나아가는 운동학 모델입니다", accent=AMBER, bsize=10)
card(s, MX + 6.08, end + 0.26, 5.85, 1.22, "다음에 이어 붙일 것",
     "· 이 궤적에 잡음을 얹으면 곧바로 추적 필터 입력이 됩니다\n"
     "· Part 2 의 좌표변환을 붙이면 안테나 기준 거리 · 방위 · 고각이 나옵니다", accent=BLUE,
     bsize=10)

# ═══════════════════════════════════════════════════════════════════════
# 마무리
# ═══════════════════════════════════════════════════════════════════════
_page += 1
s = new_slide()
_titles.append("감사합니다")
s.shapes.add_picture(os.path.join(FIG, "bg_cover.png"), 0, 0, Inches(SW), Inches(SH))
hline(s, 0.82, 0.78, 11.70, SLATE)
hline(s, 0.82, 6.93, 11.70, SLATE)
txt(s, 0, 3.05, SW, 0.80, "감사합니다", size=34, color=WHITE, bold=True,
    align=PP_ALIGN.CENTER)
txt(s, 0, 4.05, SW, 0.34, "질문 주시면 답변드리겠습니다", size=13, color=LIGHT,
    align=PP_ALIGN.CENTER)
txt(s, 6.07, 6.98, 1.20, 0.30, str(_page), size=9, color=MUTED, align=PP_ALIGN.CENTER,
    font="Arial")


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else DECK
    prs.core_properties.title = TITLE
    prs.core_properties.author = AUTHOR
    prs.save(out)
    print("생성: %s (%d 장)" % (out, len(prs.slides._sldIdLst)))
    for i, t in enumerate(_titles, 1):
        print("  %2d. %s" % (i, t))
