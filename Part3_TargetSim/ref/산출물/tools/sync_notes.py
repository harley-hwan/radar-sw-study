# -*- coding: utf-8 -*-
"""발표대본 md 를 슬라이드 노트로 밀어 넣는다 (대본이 원본, 노트가 사본).

    python sync_notes.py            # 산출물 폴더의 pptx / md 를 그대로 쓴다
    python sync_notes.py <발표자료.pptx> <대본.md>
    python sync_notes.py --check    # 쓰지 않고 어긋난 장만 알려 준다

대본의 "## N. 제목" 이 N 번째 슬라이드에 그대로 대응한다. 장 수와 슬라이드 수가 다르면
그대로 멈춘다. build_deck.py 도 같은 parse() 를 써서 처음 만들 때 노트를 채운다.
"""
import io
import os
import re
import sys

from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.normpath(os.path.join(HERE, ".."))
DECK = os.path.join(OUT_DIR, "Chapter3_target_simulation.pptx")
SCRIPT_MD = os.path.join(OUT_DIR, "Chapter3_발표대본.md")


def parse(md):
    """## N. 제목  ->  {번호: (제목, [문단, ...])}"""
    out = {}
    for blk in re.split(r"(?m)^## ", md)[1:]:
        m = re.match(r"(\d+)\. (.+?)\n\n\[.+?\]\n(.*?)(?:\n---\n|\Z)", blk, re.S)
        if not m:
            continue
        paras = [p.strip() for p in m.group(3).strip().split("\n\n") if p.strip()]
        out[int(m.group(1))] = (m.group(2).strip(), paras)
    return out


def run(deck_path, md_path, check_only=False):
    prs = Presentation(deck_path)
    secs = parse(io.open(md_path, encoding="utf-8").read())
    slides = list(prs.slides)
    if len(slides) != len(secs):
        sys.exit("슬라이드 %d 장, 대본 %d 장 — 수가 달라서 멈춤" % (len(slides), len(secs)))

    changed = []
    for num, slide in enumerate(slides, 1):
        want = "\n".join(secs[num][1])
        have = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        if have == want:
            continue
        changed.append((num, secs[num][0]))
        if not check_only:
            slide.notes_slide.notes_text_frame.text = want

    for num, title in changed:
        print("  %2d. %s" % (num, title))
    if check_only:
        print("어긋난 장 %d 개" % len(changed))
        return 1 if changed else 0
    if changed:
        prs.save(deck_path)
        print("노트 %d 장 갱신: %s" % (len(changed), deck_path))
    else:
        print("이미 같음")
    return 0


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if a != "--check"]
    check = "--check" in sys.argv[1:]
    if not args:
        args = [DECK, SCRIPT_MD]
    if len(args) != 2:
        sys.exit("사용법: python sync_notes.py [발표자료.pptx 대본.md] [--check]")
    sys.exit(run(args[0], args[1], check))
