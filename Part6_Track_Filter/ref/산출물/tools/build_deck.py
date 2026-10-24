# 과제 6 발표자료 생성 (틀: 과제 4 발표자료의 레이아웃)
# python build_deck.py <과제4 pptx> <그림 폴더> <소스 폴더 src/TargetSim> <출력 pptx> <대본 md> [결과 데이터 폴더, 기본: 그림 폴더 옆 data]
# 프로그램 화면은 그림 폴더의 Windows 캡처 (ui_<이름>.png + .json) 가 있으면 그것을, 없으면 리눅스 미리보기 (pv_<이름>.png + .json) 를 씀.
# .json 은 컨트롤 위치 (그림 부분을 자를 때 씀). 출처는 그림 아래 표시 (src_tag) 로 구분
import sys, os, re, json, copy, tempfile, unicodedata
import pandas as pd
from PIL import Image
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

from notes import NOTES

TEMPLATE, FIG, SRC, OUT, SCRIPT_MD = sys.argv[1:6]
DATA = sys.argv[6] if len(sys.argv) > 6 else os.path.join(os.path.dirname(os.path.abspath(FIG)), 'data')
CROP = tempfile.mkdtemp(prefix='chapter6_crop_')

NAVY = '1B2740'; SUB = '4A5A7A'; MUTED = '8794AC'; LIGHT = 'AEB9CC'; BORDER = 'D2D9E5'
BLUE = '2563EB'; RED = 'DC2626'; ORANGE = 'D97706'; GREEN = '059669'; PURPLE = '7C3AED'; NUM = 'B45309'; WHITE = 'FFFFFF'
F_GRAY = 'F7F9FC'; F_BLUE = 'EBF2FE'; F_RED = 'FDEEEE'; F_ORANGE = 'FEF5E7'; F_GREEN = 'E8F6F0'; F_PURPLE = 'F3E8FF'
KO = '맑은 고딕'; EN = 'Arial'; MONO = 'Consolas'; MATH = 'Cambria Math'
MARK = {1: RED, 2: ORANGE, 3: GREEN, 4: BLUE, 5: PURPLE}
MARK_F = {1: F_RED, 2: F_ORANGE, 3: F_GREEN, 4: F_BLUE, 5: F_PURPLE}
SEC = {1: '01. 추적 필터', 2: '02. 구현', 3: '03. 결과: 과제 값', 4: '04. 파라미터 조정'}

# ------------------------------------------------------------------ 결과 데이터 (슬라이드 숫자와 대본 숫자를 같은 곳에서)
RUN = pd.read_csv(os.path.join(DATA, 'runs.csv')).set_index('run')
SWEEP = pd.read_csv(os.path.join(DATA, 'sweep.csv'))


def rv(run, col):
    return float(RUN.loc[run, col])


def sweep_row(sc, sig, q, tau=10.0, kind='q_sigma'):
    s = SWEEP[(SWEEP.scenario == sc) & (SWEEP.sweep == kind) & ((SWEEP.sigma_az_deg - sig).abs() < 1e-9) & ((SWEEP.q - q).abs() < 1e-9) & ((SWEEP.tau - tau).abs() < 1e-9)]
    assert len(s) == 1, (sc, sig, q, tau, kind)
    return s.iloc[0]


def setting(run):
    return rv(run, 'sigma_az_deg'), rv(run, 'q'), rv(run, 'tau')


SEEDS = {run: int(sweep_row(run.split('_')[0], *setting(run)).diverged_seeds) for run in RUN.index}
SWEEP_RMS = {run: float(sweep_row(run.split('_')[0], *setting(run)).rms) for run in RUN.index}

prs = Presentation(TEMPLATE)
L = {l.name: l for l in prs.slide_layouts}

# ------------------------------------------------------------------ template preparation
# 과제 4 슬라이드 제거 (레이아웃, 마스터만 남김)
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.get(qn('r:id')))
    sldIdLst.remove(sldId)

# 표지, 끝 배경 그림을 과제 6 그림으로
for lname, img in (('COVER', 'cover_bg.png'), ('END', 'end_bg.png')):
    lay = L[lname]
    blip = lay._element.find('.//' + qn('a:blip'))
    part = lay.part.related_part(blip.get(qn('r:embed')))
    part._blob = open(os.path.join(FIG, img), 'rb').read()

SLIDE_NUM_SP = {}
for lname in ('INDEX', 'CONTENT', 'DIVIDER'):
    for sh in L[lname].placeholders:
        if sh.placeholder_format.type == 13:
            SLIDE_NUM_SP[lname] = sh._element


def add_slide(lname):
    s = prs.slides.add_slide(L[lname])
    if lname in SLIDE_NUM_SP:
        s.shapes._spTree.append(copy.deepcopy(SLIDE_NUM_SP[lname]))
    return s


# ------------------------------------------------------------------ text helpers
def rgb(h):
    return RGBColor.from_string(h)


def _style_run(r, size, color, bold=False, font=KO, italic=False, sub=False, sup=False):
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = rgb(color)
    r.font.name = font
    rPr = r._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = etree.SubElement(rPr, qn('a:ea'))
    ea.set('typeface', font if font not in (EN, MATH) else KO)
    if sub:
        rPr.set('baseline', '-25000')
    if sup:
        rPr.set('baseline', '30000')


TOKEN = re.compile(r'(_\{[^}]*\}|\^\{[^}]*\}|\*\*[^*]+\*\*)')


def add_runs(p, text, size, color, bold=False, font=KO):
    # 간단한 표기: _{아래첨자}, ^{위첨자}, **굵게**
    for part in TOKEN.split(text):
        if not part:
            continue
        sub = part.startswith('_{'); sup = part.startswith('^{'); bb = part.startswith('**')
        t = part[2:-1] if (sub or sup) else (part[2:-2] if bb else part)
        r = p.add_run()
        r.text = t
        _style_run(r, size, color, bold or bb, font, sub=sub, sup=sup)


def tb(slide, x, y, w, h, paras, size=12, color=NAVY, bold=False, font=KO, align='l', anchor='t', spacing=None, after=0, wrap=True, margin=0.0):
    # paras: str 또는 [str | (str, {size,color,bold,font,align,after})]
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    m = Inches(margin)
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = m
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, str):
        paras = [paras]
    for i, item in enumerate(paras):
        opt = {}
        if isinstance(item, tuple):
            item, opt = item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[opt.get('align', align)]
        if spacing:
            p.line_spacing = spacing
        p.space_after = Pt(opt.get('after', after))
        add_runs(p, item, opt.get('size', size), opt.get('color', color), opt.get('bold', bold), opt.get('font', font))
    return box


def rect(slide, x, y, w, h, fill=WHITE, line=None, radius=None, lw=0.75, dash=False, shape=None):
    st = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    sp = slide.shapes.add_shape(st, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius and st == MSO_SHAPE.ROUNDED_RECTANGLE:
        sp.adjustments[0] = radius
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = rgb(fill)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = rgb(line); sp.line.width = Pt(lw)
        if dash:
            ln = sp.line._get_or_add_ln()
            pd_ = etree.SubElement(ln, qn('a:prstDash')); pd_.set('val', 'dash')
    sp.shadow.inherit = False
    sp.text_frame.text = ''
    return sp


def label_box(slide, x, y, w, h, text, size=11, color=NAVY, fill=WHITE, line=None, bold=True, radius=0.25, align='c'):
    sp = rect(slide, x, y, w, h, fill, line, radius=radius)
    tf = sp.text_frame
    tf.margin_left = tf.margin_right = Inches(0.04); tf.margin_top = tf.margin_bottom = Inches(0.0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = True
    lines = text.split('\n')
    for i, t in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER}[align]
        add_runs(p, t, size if i == 0 else size - 1.5, color if i == 0 else SUB, bold if i == 0 else False)
    return sp


def line(slide, x1, y1, x2, y2, color=LIGHT, width=1.0, dash=False, arrow=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color); c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        pd_ = etree.SubElement(ln, qn('a:prstDash')); pd_.set('val', 'dash')
    if arrow:
        te = etree.SubElement(ln, qn('a:tailEnd')); te.set('type', 'triangle'); te.set('w', 'med'); te.set('len', 'med')
    return c


def marker(slide, x, y, n, d=0.26, size=9.5, color=None, text=None):
    sp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    sp.fill.solid(); sp.fill.fore_color.rgb = rgb(color or MARK[n]); sp.line.fill.background(); sp.shadow.inherit = False
    tf = sp.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text or str(n); _style_run(r, size, WHITE, True)
    return sp


def fig_path(name):
    return name if os.path.isabs(name) else os.path.join(FIG, name)


def img_ratio(name):
    with Image.open(fig_path(name)) as im:
        return im.size[1] / im.size[0]


def picture(slide, name, x, y, w=None, h=None, border=False):
    pic = slide.shapes.add_picture(fig_path(name), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)
    if border:
        pic.line.color.rgb = rgb(BORDER); pic.line.width = Pt(0.75)
    return pic


def callout(slide, x, y, w, n, head, body, hsize=11.5, bsize=10.5, color=None):
    marker(slide, x, y + 0.02, n, d=0.28, color=color)
    paras = [(head, dict(size=hsize, bold=True, color=NAVY, after=2))]
    for b in body:
        paras.append((b, dict(size=bsize, color=SUB, after=1)))
    return tb(slide, x + 0.42, y, w - 0.42, 0.3 + 0.22 * len(body), paras)


# 결과 그림의 출처 표시: 파랑 = TargetSimUI 실행 화면 캡처 (Windows), 보라 = TargetSimUI 의 그리기 코드를 리눅스에서 실행한 미리보기,
# 초록 = 결과 CSV (data 폴더) 로 그린 그림
SRC_KIND = {'ui': ('프로그램 화면', BLUE), 'pv': ('화면 미리보기', PURPLE), 'data': ('결과 데이터', GREEN)}


def src_tag(slide, x, y, kind, text, w, size=9.5):
    name, col = SRC_KIND[kind]
    label_box(slide, x, y, 0.92, 0.22, name, size=8.5, color=WHITE, fill=col, radius=0.5)
    return tb(slide, x + 1.0, y - 0.01, w - 1.0, 0.24, text, size=size, color=SUB, anchor='m')


def content(sec, title, subtitle):
    s = add_slide('CONTENT')
    ph = {p.placeholder_format.idx: p for p in s.placeholders}
    for idx, txt, size, color, bold in ((102, SEC[sec], 13.5, WHITE, True), (103, title, 21, NAVY, True), (104, subtitle, 12.5, SUB, False)):
        tf = ph[idx].text_frame
        tf.text = ''
        p = tf.paragraphs[0]
        add_runs(p, txt, size, color, bold)
    return s


def divider(title, desc, items):
    s = add_slide('DIVIDER')
    tb(s, 0.94, 1.0, 10.0, 0.52, title, size=24, bold=True)
    tb(s, 0.94, 1.58, 11.58, 0.33, desc, size=13, color=NAVY)
    y = 2.25
    for num, name, page in items:
        tb(s, 0.94, y, 0.56, 0.44, num, size=14, bold=True)
        tb(s, 1.55, y, 6.6, 0.44, name, size=14)
        est = 1.55 + 0.02 + sum(0.19 if unicodedata.east_asian_width(ch) in 'WF' else 0.105 for ch in name)
        line(s, est + 0.12, y + 0.28, 11.6, y + 0.28, color=BORDER, width=0.75, dash=True)
        tb(s, 11.74, y, 0.78, 0.44, str(page), size=13, color=SUB, align='r')
        y += 0.445
    DIVIDER_ITEMS.extend((num, page) for num, _, page in items)
    return s


def chips(slide, x, y, w, items, h=0.36, gap=0.12):
    # 수치 한 줄: [(이름, 값)] -> 같은 폭의 칸
    cw = (w - gap * (len(items) - 1)) / len(items)
    for i, (name, value) in enumerate(items):
        xx = x + i * (cw + gap)
        rect(slide, xx, y, cw, h, F_GRAY, BORDER, radius=0.25)
        box = tb(slide, xx, y, cw, h, '', anchor='m', align='c')
        p = box.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        add_runs(p, name + '   ', 9.5, SUB)
        add_runs(p, value, 11, NAVY, bold=True)


def table(slide, x, y, widths, rows, row_h=0.3, size=10, fills=None, colors=None, bold_cols=(), aligns=None, fonts=None, header=True):
    t = slide.shapes.add_table(len(rows), len(widths), Inches(x), Inches(y), Inches(sum(widths)), Inches(row_h * len(rows))).table
    for j, w in enumerate(widths):
        t.columns[j].width = Inches(w)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(row_h)
        for j, txt in enumerate(row):
            cell = t.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.06); cell.margin_top = cell.margin_bottom = Inches(0.0)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            head = header and i == 0
            fill = NAVY if head else ((fills[i] if fills else None) or WHITE)
            cell.fill.solid(); cell.fill.fore_color.rgb = rgb(fill)
            tf = cell.text_frame; tf.text = ''; tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[(aligns[j] if aligns else 'c')]
            col = WHITE if head else ((colors[i] if colors else None) or NAVY)
            font = (fonts[j] if fonts else None) or KO
            add_runs(p, txt, size, col, bold=head or (j in bold_cols), font=KO if head else font)
    tblpr = slide.shapes[-1]._element.graphic.graphicData.tbl.tblPr
    tblpr.set('firstRow', '0'); tblpr.set('bandRow', '0')
    return t


# ------------------------------------------------------------------ program screens (Windows capture or Linux preview)
VIEWS_ALL = ('IDC_VIEW3D', 'IDC_PROFILE_LAT', 'IDC_PROFILE_LON')


def ui_shot(name):
    for kind, prefix in (('ui', 'ui_'), ('pv', 'pv_')):
        png = os.path.join(FIG, prefix + name + '.png')
        if os.path.exists(png):
            meta = os.path.join(FIG, prefix + name + '.json')
            assert os.path.exists(meta), f'{meta}: 컨트롤 위치 파일이 없음'
            with open(meta, encoding='utf-8') as f:
                return kind, png, json.load(f)['rects']
    raise FileNotFoundError(f'{name}: figures 에 ui_{name}.png (Windows 캡처) 도 pv_{name}.png (미리보기) 도 없음')


def ui_crop(name, keys, pad=0):
    # 화면 그림에서 keys 컨트롤을 모두 덮는 사각형만 잘라 임시 PNG 로
    kind, png, rects = ui_shot(name)
    boxes = [rects[k] for k in keys]
    x0 = min(b[0] for b in boxes) - pad; y0 = min(b[1] for b in boxes) - pad
    x1 = max(b[0] + b[2] for b in boxes) + pad; y1 = max(b[1] + b[3] for b in boxes) + pad
    out = os.path.join(CROP, f'{kind}_{name}_{len(keys)}.png')
    with Image.open(png) as im:
        im.crop((x0, y0, x1, y1)).save(out)
    return kind, out


def ui_source_note():
    kinds = {ui_shot(n)[0] for n in UI_SHOTS}
    if kinds == {'ui'}:
        return 'Visual Studio 2022 Release x64 로 빌드해 실행한 화면입니다.'
    return ('이 발표의 프로그램 화면은 TargetSimUI 의 실제 그리기 코드를 리눅스에서 그대로 실행해 그린 미리보기입니다. '
            '그래서 글꼴 모양은 Windows 화면과 조금 다를 수 있습니다.')


UI_SHOTS = ('air_tuned_all', 'ship_given_all', 'air_given_all', 'air_given_perf', 'ship_tuned_all')
UI_TAG = {'ui': 'TargetSimUI 실행 화면', 'pv': 'TargetSimUI 그리기 코드를 리눅스에서 실행한 미리보기'}

# ------------------------------------------------------------------ code panels
TYPES = {'UINT32', 'INT32', 'FLOAT32', 'FLOAT64', 'DOUBLE64', 'VOID', 'matrix', 'const', 'static', 'UCHAR8', 'INT64', 'CHAR', 'typedef', 'struct'}
KEYWORDS = {'if', 'else', 'switch', 'case', 'break', 'for', 'while', 'return', 'default'}
CTOKEN = re.compile(r'(//.*$|\b[A-Za-z_][A-Za-z0-9_]*\b|\b\d+(?:\.\d+)?(?:e\d+)?U?\b|\s+|.)')
DECL = re.compile(r'^\s*(?:FLOAT64|INT32|UINT32|matrix)\s+[\w\[\]]+;\s*(?://.*)?$')        # 선언만 있는 줄
NODECL = lambda t: not DECL.match(t)
SEP = None                                                                              # 구간 사이 줄 (⋮)


def src_lines(rel):
    with open(os.path.join(SRC, rel), encoding='utf-8') as f:
        return f.read().split('\n')


def find_line(lines, text, start=0):
    for i in range(start, len(lines)):
        if text in lines[i]:
            return i
    raise ValueError(text)


def dwidth(s):
    return sum(2 if unicodedata.east_asian_width(c) in 'WF' else 1 for c in s)


def tidy(lines, maxw=78):
    # 탭 정렬을 패널 안에서 다시 맞춤: 앞 들여쓰기는 탭 하나 = 2칸, 줄 안의 탭 묶음은 열로 정렬
    rows = []
    for raw in lines:
        lead = len(raw) - len(raw.lstrip('\t'))
        body = raw.lstrip('\t')
        cmt = ''
        if '//' in body and not body.startswith('//'):
            k = body.index('//'); cmt = body[k:]; body = body[:k].rstrip('\t ')
        rows.append((lead, re.split(r'\t+', body) if body else [''], cmt))
    width = {}
    for lead, f, c in rows:
        for i, x in enumerate(f[:-1]):
            width[(lead, i)] = max(width.get((lead, i), 0), dwidth(x))
    out = []
    for lead, f, c in rows:
        s = '  ' * lead
        for i, x in enumerate(f):
            s += x + (' ' * (width[(lead, i)] - dwidth(x) + 2) if i < len(f) - 1 else '')
        if c:
            s += '  ' + c
        while dwidth(s) > maxw:
            s = s[:-2]
            if not s.endswith('…'):
                s = s[:-1] + '…'
        out.append(s)
    return out


def code_runs(p, num, text, size):
    if num is None:
        r = p.add_run(); r.text = '     ...'; _style_run(r, size, MUTED, font=MONO)
        return
    r = p.add_run(); r.text = f'{num:>3}  '; _style_run(r, size, LIGHT, font=MONO)
    for m in CTOKEN.finditer(text):
        t = m.group(0)
        if t.startswith('//'):
            c = GREEN
        elif t in TYPES or t.startswith('STRUCT_') or t.startswith('ST_'):
            c = BLUE
        elif t in KEYWORDS:
            c = BLUE
        elif re.fullmatch(r'[A-Z][A-Z0-9_]{2,}', t):
            c = PURPLE
        elif re.fullmatch(r'\d+(?:\.\d+)?(?:e\d+)?U?', t):
            c = NUM
        else:
            c = NAVY
        r = p.add_run(); r.text = t; _style_run(r, size, c, font=MONO)


def code_panel(slide, x, y, w, rel, segments, title=None, highlights=None, size=8.0, lh=0.15, maxw=None, skip=frozenset({''}), keep=None):
    # segments: [(첫 줄 글자, 줄 수)] -> 구간 사이에 ⋮. highlights: {표시 번호: (첫 줄 글자, 줄 수)}
    lines = src_lines(rel)
    idx = []
    i0 = 0
    for k, (start_text, count) in enumerate(segments):
        i0 = find_line(lines, start_text, i0)                                          # 앞 구간 다음부터 찾음
        part = [i for i in range(i0, i0 + count) if not (skip and lines[i].strip() in skip) and (keep is None or keep(lines[i]))]
        if k > 0:
            idx.append(SEP)
        idx += part
    shown = tidy([lines[i] if i is not SEP else '' for i in idx], maxw=maxw or int((w - 0.5) * 72 / (size * 0.6)) - 5)
    if title:
        box = tb(slide, x, y - 0.29, w, 0.26, title[0] + '    ', size=11, bold=True).text_frame.paragraphs[0]
        r = box.add_run(); r.text = title[1]; _style_run(r, 9.5, MUTED, font=MONO)
    h = 0.12 + lh * len(idx) + 0.06
    rect(slide, x, y, w, h, WHITE, BORDER)
    marks = []
    if highlights:
        for n, (txt, nl) in highlights.items():
            parts = txt if isinstance(txt, tuple) else (txt,)
            k = next(j for j, i in enumerate(idx) if i is not SEP and all(t in lines[i] for t in parts))
            yy = y + 0.08 + lh * k
            rect(slide, x + 0.03, yy, w - 0.06, lh * nl, MARK_F[n])
            rect(slide, x + 0.03, yy, 0.04, lh * nl, MARK[n])
            marks.append((n, yy, nl))
    box = slide.shapes.add_textbox(Inches(x + 0.1), Inches(y + 0.08), Inches(w - 0.2), Inches(lh * len(idx)))
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for j, (i, s) in enumerate(zip(idx, shown)):
        p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        p.line_spacing = Pt(lh * 72)
        code_runs(p, None if i is SEP else i + 1, s, size)
    placed = []
    for n, yy, nl in sorted(marks, key=lambda m: m[1]):
        my = yy + (lh * nl - 0.22) / 2
        mx = x + w - 0.16
        while any(abs(my - py) < 0.24 and abs(mx - px) < 0.24 for px, py in placed):
            mx -= 0.27
        placed.append((mx, my))
        marker(slide, mx, my, n, d=0.22, size=8.5)
    return h


def legend_row(slide, x, y, items, size=10):
    for n, t in items:
        marker(slide, x, y + 0.03, n, d=0.22, size=8.5)
        w = 0.12 + sum(0.145 if unicodedata.east_asian_width(c) in 'WF' else 0.072 for c in t)
        tb(slide, x + 0.28, y, w, 0.28, t, size=size, color=NAVY)
        x += w + 0.42


# ------------------------------------------------------------------ slides
slides_meta = []      # (노트 키, 슬라이드) 대본 문서용
PAGE = {}             # 노트 키 -> 쪽 번호
DIVIDER_ITEMS = []    # (번호, 적은 쪽)


def note(slide, key):
    PAGE[key] = len(prs.slides)
    slide.notes_slide.notes_text_frame.text = NOTES[key].replace('{ui_source}', ui_source_note())
    slides_meta.append((key, slide))


# 1 표지
s = add_slide('COVER')
tb(s, 9.0, 0.72, 3.3, 0.3, 'Confidential', size=10, bold=True, color=WHITE, font=EN, align='r')
tb(s, 4.6, 2.74, 7.4, 0.5, '레이다 SW 학습 과제 6', size=14, color=LIGHT, align='r')
tb(s, 4.6, 3.22, 7.4, 0.75, '추적 필터', size=38, bold=True, color=WHITE, align='r')
line(s, 8.6, 4.28, 12.0, 4.28, color=MUTED, width=0.75)
tb(s, 4.6, 4.44, 7.4, 0.32, 'Tracking Filter: CMKF with Singer Model', size=13, color=LIGHT, font=EN, align='r')
tb(s, 4.6, 5.6, 7.4, 0.32, 'Janghwan Kim (Harley)', size=12, color=WHITE, font=EN, align='r')
note(s, 'cover')

# 2 목차
s = add_slide('INDEX')
tb(s, 0.5, 3.0, 3.6, 1.3, [('추적 필터', dict(after=4)), 'Singer CMKF'], size=18, bold=True, color=WHITE, align='c')
tb(s, 0.5, 6.6, 3.6, 0.3, 'Janghwan Kim (Harley)', size=10, color=LIGHT, font=EN, align='c')
INDEX = [('01', '추적 필터', 3), ('02', '구현', 10), ('03', '결과: 과제 값', 16), ('04', '파라미터 조정', 22)]
y = 1.97
for i, (n, name, page) in enumerate(INDEX):
    tb(s, 5.25, y, 1.0, 0.55, n, size=30, bold=True, font=EN)
    tb(s, 6.4, y, 2.6, 0.55, name, size=15, bold=True, anchor='m')
    est = sum(0.24 if unicodedata.east_asian_width(c) in 'WF' else 0.12 for c in name)
    line(s, 6.45 + est + 0.2, y + 0.33, 11.95, y + 0.33, color=BORDER, width=0.75, dash=True)
    tb(s, 12.04, y, 0.56, 0.55, str(page), size=14, font=EN, anchor='m', align='r')
    if i < len(INDEX) - 1:
        line(s, 5.25, y + 0.72, 12.6, y + 0.72, color=BORDER, width=0.75)
    y += 0.84
note(s, 'index')

# ================================================================== 01 추적 필터
s = divider('01. 추적 필터', '트랙의 위치 · 속도 · 가속도를 스캔마다 예측하고 플롯으로 고치는 칼만 필터 (CMKF) 와, 표적 기동을 나타내는 Singer 모델을 봅니다.',
            [('1.1', '추적 필터가 하는 일', 4), ('1.2', '칼만 필터 한 주기: 예측과 갱신', 5), ('1.3', 'CMKF: 측정치 변환과 측정잡음 공분산 R', 6),
             ('1.4', 'Singer 모델과 상태천이 행렬 F', 7), ('1.5', '공정잡음 공분산 Q', 8), ('1.6', '발산 판단: NIS 와 게이트', 9)])
note(s, 'div1')

# 1.1 추적 필터가 하는 일
s = content(1, '1.1 추적 필터가 하는 일', '과제 5 는 트랙을 만들고 상태만 관리했습니다. 과제 6 은 그 트랙의 위치 · 속도 · 가속도를 스캔마다 예측하고 플롯으로 고치는 추적 필터를 넣습니다.')
flow = [('측정 모의', '과제 4 · 스캔마다 플롯', F_GRAY, NAVY), ('트랙 개시 · 관리', '과제 5 · 2-Point, 스코어', F_GRAY, NAVY),
        ('추적 필터', '과제 6 · 이번 과제', F_BLUE, NAVY), ('트랙 연관', '이번에는 무조건 연관 (가정)', F_GRAY, '7A879F')]
for i, (a, b, f, c) in enumerate(flow):
    x = 0.7 + i * 3.08
    rect(s, x, 2.48, 2.69, 0.64, f, BLUE if i == 2 else 'E3E8F0', lw=1.2 if i == 2 else 0.75)
    tb(s, x, 2.53, 2.69, 0.3, a, size=12.5, bold=True, color=c, align='c')
    tb(s, x, 2.83, 2.69, 0.24, b, size=10, color=SUB if i == 2 else MUTED, align='c')
    if i < 3:
        line(s, x + 2.74, 2.8, x + 3.03, 2.8, color=LIGHT, width=1.0, arrow=True)
line(s, 0.7, 3.3, 12.63, 3.3, color='E3E8F0', width=0.75)
W = 5.45
picture(s, 'fig_role.png', 0.7, 3.42, w=W)
src_tag(s, 0.7, 3.42 + W * img_ratio('fig_role.png') + 0.04, 'data', '대공 표적 조정 값 실행 22 ~ 34 s (plots · tracks.csv)', W)
tb(s, 6.5, 3.45, 6.1, 0.3, '스캔마다 하는 일', size=12, bold=True)
ROLE = [('① 예측', BLUE, '지난 추정으로 이번 스캔의 위치를 미리 계산', '표적 운동 모델: Singer 모델 (1.4)'),
        ('② 갱신', BLUE, '이번 플롯과 예측의 차이 (잔차) 를 칼만 게인 K 만큼 반영', 'K 는 예측 오차 P 와 플롯 오차 R 로 정함 (1.2)'),
        ('결과', NAVY, '트랙 = 상태 9 개 + 오차 공분산 P (9 × 9)', '위치 · 속도 · 가속도 × 북 · 동 · 아래 (VC-NED)')]
y = 3.85
for tag, col, head, body in ROLE:
    label_box(s, 6.5, y, 0.95, 0.3, tag, size=10.5, color=col, line=col)
    tb(s, 7.6, y - 0.03, 5.0, 0.62, [(head, dict(size=11, bold=True, after=2)), (body, dict(size=10, color=SUB))])
    y += 0.72
rect(s, 6.5, 6.02, 6.13, 0.82, F_ORANGE, None, radius=0.1)
tb(s, 6.68, 6.07, 5.85, 0.74, [('과제 5 와 달라진 점: 탐지 플롯을 트랙에 무조건 연관 (가정)', dict(size=10.5, bold=True, after=2)),
                               (f'Assigned_Flag 1 → {rv("air_tuned", "init_time"):.1f} s 개시, {rv("air_tuned", "confirm_time"):.1f} s 확정 뒤 끝까지 유지 '
                                '(과제 5 는 늘 0 → 두 스캔 뒤 삭제)', dict(size=9.6, color=NAVY))])
note(s, 's1_1')

# 1.2 칼만 필터 한 주기
s = content(1, '1.2 칼만 필터 한 주기: 예측과 갱신', '스캔마다 예측과 갱신을 한 번씩 합니다. 과제 자료의 다섯 식이 이 두 단계이고, 오차 공분산 (타원) 이 예측에서 커졌다가 갱신에서 줄어듭니다.')
picture(s, 'fig_kf_cycle.png', 0.6, 2.42, w=5.9)
EQ = [('예측', '추정치 예측', 'X̂_{k|k−1} = F_{k} X̂_{k−1|k−1}', '지난 추정을 T (0.1 s) 만큼 앞으로 (F: Singer 모델)'),
      ('예측', '오차 공분산 예측', 'P_{k|k−1} = F_{k} P_{k−1|k−1} F_{k}^{T} + Q_{k}', '모르는 기동 Q 만큼 커짐'),
      ('갱신', '칼만 게인', 'K_{k} = P_{k|k−1} H^{T} ( H P_{k|k−1} H^{T} + R_{k} )^{−1}', '예측 오차 P 와 플롯 오차 R 의 비율'),
      ('갱신', '추정치', 'X̂_{k|k} = X̂_{k|k−1} + K_{k} ( z_{k} − H X̂_{k|k−1} )', '잔차 (플롯 − 예측) 의 K 만큼 고침'),
      ('갱신', '추정 오차 공분산', 'P_{k|k} = ( I − K_{k} H ) P_{k|k−1}', '플롯을 반영해 줄어듦')]
y = 2.5
for grp, name, eq, desc in EQ:
    col = BLUE if grp == '예측' else GREEN
    label_box(s, 6.75, y + 0.02, 0.62, 0.26, grp, size=9.5, color=col, line=col)
    box = tb(s, 7.5, y - 0.02, 5.13, 0.3, '')
    p = box.text_frame.paragraphs[0]
    add_runs(p, name + '   ', 10.5, NAVY, bold=True)
    add_runs(p, desc, 9.5, SUB)
    tb(s, 7.5, y + 0.27, 5.13, 0.34, eq, size=12.5, font=MATH)
    y += 0.7
rect(s, 0.7, 6.22, 11.93, 0.62, F_BLUE, None, radius=0.2)
tb(s, 0.9, 6.22, 11.6, 0.62, [('K 가 크면 플롯을, 작으면 예측 (운동 모델) 을 따름.  R 이 작거나 (플롯을 믿음) Q 가 클수록 (모델을 덜 믿음) K 가 커짐  →  4 장에서 조정할 손잡이',
                               dict(size=11.5))], anchor='m')
note(s, 's1_2')

# 1.3 CMKF
s = content(1, '1.3 CMKF: 측정치 변환과 측정잡음 공분산 R', '레이다는 거리 · 방위각 · 고각을 잽니다. CMKF 는 이 값을 VC-NED 직교 좌표로 바꾼 위치를 측정치로 써서 측정 식이 선형 (H) 이 되게 합니다.')
tb(s, 0.7, 2.47, 5.8, 0.3, '측정치 변환 (CM: Converted Measurement)', size=12, bold=True)
CHIP = [('r · θ · φ', '구좌표 (플롯)'), ('안테나', '직교 좌표'), ('동체', '좌표'), ('z = [x y z]^{T}', 'VC-NED')]
for i, (a, b) in enumerate(CHIP):
    x = 0.7 + i * 1.5
    rect(s, x, 2.85, 1.24, 0.6, F_BLUE if i in (0, 3) else F_GRAY, BLUE if i in (0, 3) else 'E3E8F0', radius=0.12)
    tb(s, x, 2.89, 1.24, 0.28, a, size=11, bold=True, align='c', font=MATH if i in (0, 3) else KO)
    tb(s, x, 3.16, 1.24, 0.24, b, size=9, color=SUB, align='c')
    if i < 3:
        line(s, x + 1.27, 3.15, x + 1.47, 3.15, color=SUB, width=1.1, arrow=True)
tb(s, 0.7, 3.5, 5.8, 0.25, '과제 3 좌표 변환 함수 그대로 (수신 쓰레드 f_Lnk_FillPlotCoord)', size=9.5, color=SUB)
tb(s, 0.7, 3.9, 5.8, 0.3, '측정 행렬 H: 상태 9 개 가운데 위치 3 개만 꺼냄', size=12, bold=True)
picture(s, 'fig_hmatrix.png', 0.7, 4.2, w=5.6)
rect(s, 0.7, 6.1, 5.75, 0.74, F_BLUE, None, radius=0.12)
tb(s, 0.88, 6.12, 5.5, 0.7, [('잔차 ν = z_{k} − H X̂_{k|k−1} = 플롯 위치 − 예측 위치', dict(size=10.5, bold=True, after=2)),
                             ('측정 식이 선형이라 칼만 필터 식을 그대로 씀. 대신 R 을 직교 좌표로', dict(size=10, color=SUB))], anchor='m')
tb(s, 6.85, 2.47, 5.78, 0.3, '측정잡음 공분산 R  (과제 5 의 1.4 와 같은 식)', size=12, bold=True)
rect(s, 6.85, 2.85, 5.78, 0.92, F_GRAY, 'E3E8F0', radius=0.08)
tb(s, 6.85, 2.9, 5.78, 0.42, 'R_{k}^{NED} = M J R_{k}^{rθφ} J^{T} M^{T}', size=15, bold=True, align='c', font=MATH)
tb(s, 6.95, 3.36, 5.6, 0.36, 'M = C_{Body}^{NED} C_{Ant}^{Body} (안테나 → 동체 → NED 회전),   R_{k}^{rθφ} = diag( σ_{rng}^{2}, σ_{azi}^{2}, σ_{ele}^{2} )',
   size=9.5, color=SUB, align='c')
for i, (a, b) in enumerate((('σ_{rng}', '10 m'), ('σ_{azi}', '0.3°'), ('σ_{ele}', '0.3°'))):
    xx = 6.85 + i * 1.97
    rect(s, xx, 3.9, 1.84, 0.66, F_BLUE, None, radius=0.12)
    tb(s, xx, 3.92, 1.84, 0.24, a + '  (과제 값)', size=9.5, color=SUB, align='c')
    tb(s, xx, 4.13, 1.84, 0.4, b, size=18, bold=True, color=BLUE, align='c', font=EN)
tb(s, 6.85, 4.68, 5.78, 0.28, '자코비안 J = ∂(x, y, z) / ∂(r, θ, φ)    θ 방위각, φ 고각', size=10.5, bold=True)
table(s, 6.85, 4.98, [1.75, 2.0, 2.03], [['sinθ cosφ', 'r cosθ cosφ', '−r sinθ sinφ'], ['sinφ', '0', 'r cosφ'], ['cosθ cosφ', '−r sinθ cosφ', '−r cosθ sinφ']],
      row_h=0.27, size=10.5, fills=[F_GRAY, WHITE, F_GRAY], fonts=[MATH] * 3, header=False)
tb(s, 6.85, 5.9, 5.78, 0.94, [('0.3° 는 13.5 ~ 14 km (두 표적의 시작 거리) 에서 약 70 m', dict(size=10.5, bold=True, after=2)),
                              ('→ R 은 거리 방향 10 m, 옆 · 위 방향 70 m 인 납작한 타원', dict(size=10, color=SUB, after=1)),
                              ('거리 · 각도가 플롯마다 달라 R 도 플롯마다 계산 (f_Trk_MeasCov)', dict(size=10, color=SUB))])
note(s, 's1_3')

# 1.4 Singer 모델과 F
s = content(1, '1.4 Singer 모델과 상태천이 행렬 F', '표적 가속도를 시간이 지나면 점점 잊히는 랜덤 과정 (1 차 마르코프) 으로 봅니다. τ 는 기동이 이어지는 시간, q 는 기동 가속도의 크기입니다.')
picture(s, 'fig_singer.png', 0.7, 2.45, w=5.7)
PARAM = [('τ  기동 시상수 [s]', '과제 값 10 s', ['가속도가 이어지는 시간', '작으면 등속 (CV), 크면 등가속 (CA) 에 가까움']),
         ('q  기동 가속도 표준편차 [m/s²]', '과제 값 0.5', ['얼마나 세게 기동하나', '공정잡음 Q 의 크기를 정함 (1.5)'])]
for i, (h, v, body) in enumerate(PARAM):
    xx = 0.7 + i * 2.93
    rect(s, xx, 5.45, 2.8, 1.38, F_GRAY, None, radius=0.08)
    tb(s, xx + 0.15, 5.52, 2.55, 0.28, h, size=11, bold=True)
    tb(s, xx + 0.15, 5.81, 2.55, 0.28, v, size=11, bold=True, color=BLUE)
    tb(s, xx + 0.15, 6.12, 2.55, 0.68, body, size=9.8, color=SUB, after=1)
tb(s, 6.75, 2.47, 5.88, 0.3, 'F = blockdiag( f, f, f )    x · y · z 축마다 같은 f', size=12, bold=True)
F_SYM = [['1', 'T', 'τ^{2} ( −1 + T/τ + e^{−T/τ} )'], ['0', '1', 'τ ( 1 − e^{−T/τ} )'], ['0', '0', 'e^{−T/τ}']]
table(s, 6.75, 2.85, [0.75, 0.75, 4.38], F_SYM, row_h=0.42, size=12.5, fills=[F_GRAY, WHITE, F_GRAY], fonts=[MATH] * 3, header=False)
tb(s, 6.75, 4.2, 5.88, 0.3, 'T 0.1 s, τ 10 s 이면', size=11.5, bold=True)
F_NUM = [['1', '0.1', '0.00498   (≈ T^{2}/2)'], ['0', '1', '0.0995   (≈ T)'], ['0', '0', '0.990']]
table(s, 6.75, 4.52, [0.75, 0.75, 4.38], F_NUM, row_h=0.36, size=12, fills=[F_GRAY, WHITE, F_GRAY], fonts=[EN] * 3, header=False)
tb(s, 6.75, 5.7, 5.88, 1.1, [('위치 += 속도 × T + 가속도 × 약 T^{2}/2,   속도 += 가속도 × 약 T', dict(size=10.5, after=3)),
                             ('가속도 × 0.990: 스캔마다 1 % 씩 줄어 τ (10 s) 뒤 37 %', dict(size=10.5, after=3)),
                             ('구현: e^{−T/τ} − 1 은 expm1 로 계산 (τ 가 커도 자리 손실 없음, 2.3)', dict(size=9.8, color=SUB))])
note(s, 's1_4')

# 1.5 Q
s = content(1, '1.5 공정잡음 공분산 Q', 'Q 는 한 스캔 사이 모르는 기동이 만드는 오차입니다. 과제 자료의 q_{k} 식을 쓰되, q12 의 오타와 작은 T/τ 에서의 계산 오차를 바로잡았습니다.')
rect(s, 0.7, 2.47, 5.75, 0.55, F_GRAY, 'E3E8F0', radius=0.12)
tb(s, 0.7, 2.47, 5.75, 0.55, 'Q_{k} = (2q^{2}/τ) · blockdiag( q_{k}, q_{k}, q_{k} )', size=14, bold=True, align='c', anchor='m', font=MATH)
tb(s, 0.7, 3.1, 5.75, 0.28, 'q_{k} 의 원소   (x = T/τ 로 줄여 씀)', size=11, bold=True)
QK = [('q_{11} = (τ^{5}/2)( 1 − e^{−2x} + 2x + 2x^{3}/3 − 2x^{2} − 4x e^{−x} )', NAVY),
      ('q_{12} = (τ^{4}/2)( 1 + e^{−2x} − 2e^{−x} + 2x e^{−x} − 2x + x^{2} )', BLUE),
      ('q_{13} = (τ^{3}/2)( 1 − e^{−2x} − 2x e^{−x} )', NAVY),
      ('q_{22} = (τ^{3}/2)( 4e^{−x} − 3 − e^{−2x} + 2x )', NAVY),
      ('q_{23} = (τ^{2}/2)( e^{−2x} + 1 − 2e^{−x} )', NAVY),
      ('q_{33} = (τ/2)( 1 − e^{−2x} )', NAVY)]
tb(s, 0.7, 3.4, 5.75, 1.75, [(t, dict(color=c, bold=(c == BLUE), after=2)) for t, c in QK], size=10.5, font=MATH)
rect(s, 0.7, 5.2, 5.75, 1.64, F_RED, None, radius=0.06)
tb(s, 0.85, 5.25, 5.5, 0.5, [('① q12 오타: 자료의 2(T/τ)·**T**·e^{−T/τ} 는 이 항만 [s] 단위가 남음', dict(size=10.5, bold=True, after=1)),
                             ('→ 표준식 2(T/τ)e^{−T/τ} 로 구현 (위 파란 줄).  T 0.1 s, τ 10 s 일 때:', dict(size=9.8, color=SUB))])
table(s, 0.85, 5.82, [1.45, 0.9, 0.9, 0.9, 1.3], [['', 'q11', 'q22', 'q12', 'q12² ≤ q11·q22'],
                                                 ['표준식 (구현)', '4.97e−7', '3.31e−4', '1.24e−5', '만족'],
                                                 ['자료 식 그대로', '4.97e−7', '3.31e−4', '−89.1', '깨짐: 공분산 아님']],
      row_h=0.32, size=9.5, fills=[None, WHITE, WHITE], colors=[None, NAVY, RED], fonts=[KO, EN, EN, EN, KO])
tb(s, 6.75, 2.47, 5.88, 0.3, '② 작은 T/τ 에서 자리 손실 → 급수', size=12, bold=True)
picture(s, 'fig_qprec.png', 6.75, 2.8, w=5.88)
tb(s, 6.75, 2.8 + 5.88 * img_ratio('fig_qprec.png') + 0.08, 5.88, 0.7,
   [('T/τ < 0.05 이면 같은 식을 테일러 급수 5 항으로 계산 (TRK_SERIES_X)', dict(size=10.5, bold=True, after=2)),
    ('검증: 60 자리 기준값 · 수치 적분과 상대 오차 1e−9 미만, 모든 T/τ 에서 Q 가 양의 정부호', dict(size=9.8, color=SUB))])
note(s, 's1_5')

# 1.6 발산 판단
s = content(1, '1.6 발산 판단: NIS 와 게이트', '과제는 트랙이 발산하거나 유지되지 않으면 파라미터를 조정하라고 합니다. 발산을 잔차의 통계적 크기 NIS 로 판단하는 규칙을 정했습니다.')
rect(s, 0.7, 2.47, 5.75, 0.82, F_GRAY, 'E3E8F0', radius=0.1)
tb(s, 0.7, 2.5, 5.75, 0.4, 'NIS = ν^{T} S^{−1} ν', size=16, bold=True, align='c', font=MATH)
tb(s, 0.7, 2.9, 5.75, 0.34, 'ν = z_{k} − H X̂_{k|k−1},    S = H P_{k|k−1} H^{T} + R_{k}', size=11, align='c', font=MATH, color=SUB)
picture(s, 'fig_chi2.png', 0.7, 3.38, w=5.75)
rect(s, 0.7, 6.2, 5.75, 0.64, F_BLUE, None, radius=0.15)
tb(s, 0.85, 6.2, 5.5, 0.64, [('R · Q 가 맞으면 NIS 는 자유도 3 카이제곱 (평균 3), 99 % 값 11.34 를 게이트로', dict(size=10, bold=True, after=1)),
                             ('정상이면 게이트 밖 1 % → 5 번 연속은 100 억 번에 1 번 꼴 (0.01^{5})', dict(size=9.8, color=SUB))], anchor='m')
tb(s, 6.75, 2.47, 5.88, 0.3, '대공 표적 · 과제 값: 첫 선회 중 NIS', size=12, bold=True)
picture(s, 'fig_runs.png', 6.75, 2.8, w=5.88)
yy = 2.8 + 5.88 * img_ratio('fig_runs.png') + 0.04
src_tag(s, 6.75, yy, 'data', 'air_given/tracks.csv 의 nis · gate_run', 5.88)
RULE = [('게이트 밖', 'NIS > 11.34  (χ²(3) 99 %, TRK_GATE_CHI2)'),
        ('발산 판정', '게이트 밖이 5 번 이어질 때마다 1 회 (TRK_DIVERGE_NUM)'),
        ('처리', '갱신은 가정대로 계속, 트랙은 지우지 않고 횟수만 셈 (화면 빨간 띠)')]
yy += 0.36
for i, (a, b) in enumerate(RULE):
    label_box(s, 6.75, yy + 0.03, 0.95, 0.26, a, size=9.5, color=RED if i < 2 else NAVY, line=RED if i < 2 else NAVY)
    tb(s, 7.82, yy, 4.8, 0.32, b, size=10, anchor='m')
    yy += 0.34
note(s, 's1_6')

# ================================================================== 02 구현
s = divider('02. 구현', '과제 5 프로그램 (TargetSim) 의 트랙 처리에 Singer CMKF 를 넣고, 시나리오 · 파라미터 입력과 LLA 결과 화면을 만들었습니다.',
            [('2.1', '전체 구조와 바뀐 곳', 11), ('2.2', '스캔 처리 순서와 연관 가정 (f_Trk_ProcessScan)', 12), ('2.3', '예측: Singer F · Q 와 f_Trk_Predict', 13),
             ('2.4', '갱신: CMKF 와 발산 판정 (f_Trk_Update)', 14), ('2.5', '결과 화면 (TargetSimUI)', 15)])
note(s, 'div2')

# 2.1 전체 구조
s = content(2, '2.1 전체 구조와 바뀐 곳', '과제 5 의 구조 (송신 · 수신 쓰레드, 스캔마다 UDP 메시지) 는 그대로 두고, 트랙 처리에 추적 필터를 넣고 시나리오와 결과 화면을 바꿨습니다.')
rect(s, 0.7, 2.5, 11.93, 3.38, F_GRAY, 'E3E8F0', radius=0.03)
tb(s, 0.9, 2.56, 5, 0.28, 'TargetSimCore.dll  (C)', size=11, bold=True, color=SUB, font=EN)
label_box(s, 10.55, 2.58, 1.9, 0.26, '주황 = 과제 6 에서 바뀜', size=9.5, color=ORANGE, line=ORANGE)
B = [(1.0, 2.92, '궤적 계산  (과제 3)', '과제 3 그대로 (시나리오는 화면에서)', 'TargetSim.c', NAVY, WHITE),
     (4.75, 2.92, '측정 모의  (과제 4)', '표적 플롯만 (클러터 0), 잡음 seed 지정', 'PlotSim.c', ORANGE, F_ORANGE),
     (8.5, 2.92, '송신 쓰레드  (신호처리 역할)', '스캔의 플롯 + R σ (파라미터) 를 메시지로', 'PlotLink.c', BLUE, F_BLUE),
     (8.5, 4.68, '수신 쓰레드  (데이터 처리 역할)', '좌표 변환, 플롯마다 R 계산', 'PlotLink.c', BLUE, F_BLUE),
     (4.75, 4.68, '추적 필터 + 트랙 관리', '예측 · 연관 (가정) · CMKF 갱신 · 발산 판정', 'TrackMgr.c', ORANGE, F_ORANGE),
     (1.0, 4.68, '결과 버퍼', '+ 잔차 · NIS · P · 발산 판정 횟수', 'PlotLink.c', ORANGE, F_ORANGE)]
for x, y, a, b, f_, c, f in B:
    rect(s, x, y, 3.3, 0.98, f, c, radius=0.1, lw=1.1)
    tb(s, x, y + 0.08, 3.3, 0.3, a, size=12, bold=True, color=c, align='c')
    tb(s, x, y + 0.41, 3.3, 0.28, b, size=10, color=NAVY, align='c')
    tb(s, x, y + 0.67, 3.3, 0.24, f_, size=8.5, color=MUTED, align='c', font=MONO)
for x1, y1, x2, y2 in ((4.32, 3.41, 4.73, 3.41), (8.07, 3.41, 8.48, 3.41), (8.48, 5.17, 8.07, 5.17), (4.73, 5.17, 4.32, 5.17)):
    line(s, x1, y1, x2, y2, color=SUB, width=1.25, arrow=True)
line(s, 10.15, 3.92, 10.15, 4.66, color=BLUE, width=2.0, arrow=True)
tb(s, 10.27, 4.06, 2.3, 0.5, ['UDP', '스캔마다 메시지 1 개'], size=10, color=BLUE, bold=True)
line(s, 2.65, 5.68, 2.65, 6.03, color=SUB, width=1.25, arrow=True)
rect(s, 0.7, 6.05, 11.93, 0.78, F_ORANGE, ORANGE, radius=0.1, lw=1.1)
tb(s, 0.95, 6.09, 11.5, 0.72, [('TargetSimUI.exe  (MFC 화면)', dict(size=12, bold=True, after=2)),
                              ('시나리오 (대함 · 대공) · 필터 파라미터 입력,  3D (LLA) · 위도-고도 · 경도-고도 · 평면 · 추적 성능 보기,  참값 대비 오차,  트랙 표 (가속도 · NIS · 발산 판정)',
                               dict(size=10.5, color=SUB))])
note(s, 's2_1')

# 2.2 스캔 처리 순서
s = content(2, '2.2 스캔 처리 순서와 연관 가정 (f_Trk_ProcessScan)', '수신 쓰레드가 플롯 메시지를 받을 때마다 아래 여섯 단계를 차례로 합니다. 과제 5 의 개시 · 관리 사이에 예측과 연관 (가정) · 갱신이 들어갔습니다.')
STEP = [('삭제 트랙 빼기', '앞 스캔에서 삭제 (4) 가 된 트랙을 표에서 지움', False),
        ('예측', '모든 트랙을 이번 스캔 시각까지 (F, Q)\n트랙-플롯 통계적 거리 d^{2} = ν^{T}S^{−1}ν 계산', True),
        ('연관 (가정) · 갱신', 'd^{2} 가 가장 작은 짝부터, 문턱 없이 연관\nCMKF 갱신, NIS 로 게이트 밖 · 발산 판정 셈', True),
        ('트랙 관리', '스코어 로직 (과제 5). 연관되면 Assigned_Flag 1', False),
        ('2-Point 개시', '트랙이 가져가지 않은 플롯끼리 (과제 5)', False),
        ('개시 후보 보관', '이번 플롯을 다음 스캔 개시 후보로 (과제 5)', False)]
y = 2.5
for i, (h, b, new) in enumerate(STEP):
    hh = 0.8 if new else 0.6
    rect(s, 0.7, y, 5.3, hh - 0.08, F_ORANGE if new else F_GRAY, ORANGE if new else None, radius=0.1)
    marker(s, 0.82, y + 0.1, i + 1, d=0.3, size=10, color=ORANGE if new else MUTED)
    tb(s, 1.25, y + 0.06, 4.7, 0.28, h + ('   과제 6' if new else ''), size=11.5, bold=True)
    tb(s, 1.25, y + 0.33, 4.7, hh - 0.4, b.split('\n'), size=9.8, color=SUB)
    y += hh
code_panel(s, 6.25, 2.78, 6.38, 'TargetSimCore/src/TrackMgr.c', [('// 3. 연관 (가정) 과 갱신', 35)], title=('③ 연관 (가정) · 갱신', 'TrackMgr.c'),
           highlights={1: ('(st_Table->st_Track[nTrack].Assigned_Flag == 0) && (st_Cur[nCur].isUsed == 0)', 2),
                       2: ('f_Trk_Update(st_Track, &st_Cur[bestCur]);', 1), 3: (('Assigned_Flag', '= 1;'), 1)},
           size=7.2, lh=0.12)
note(s, 's2_2')

# 2.3 예측 코드
s = content(2, '2.3 예측: Singer F · Q 와 f_Trk_Predict', 'F 와 Q 는 과제 자료의 식 그대로 계산합니다. 다만 숫자 자리 손실을 막으려고 F 는 expm1, Q 는 T/τ 가 작을 때 급수를 씁니다.')
h1 = code_panel(s, 0.7, 2.78, 5.6, 'TargetSimCore/src/TrackMgr.c', [('VOID f_Trk_SingerF(', 12)], title=('Singer F 의 블록 f', 'TrackMgr.c'),
                highlights={1: ('expm1(-x)', 1)}, size=7.6, lh=0.13, keep=NODECL)
code_panel(s, 0.7, 2.78 + h1 + 0.45, 5.6, 'TargetSimCore/src/TrackMgr.c', [('VOID f_Trk_Predict(', 3), ('if (dt > 0.0)', 12)], title=('예측 f_Trk_Predict', 'TrackMgr.c'),
           highlights={2: ('Matrix_Product2(&mat_F, &st_Track->mat_EstX)', 1), 3: ('Matrix_Product3(&mat_F, &st_Track->mat_EstP', 2)}, size=7.6, lh=0.13, keep=NODECL)
code_panel(s, 6.55, 2.78, 6.08, 'TargetSimCore/src/TrackMgr.c', [('VOID f_Trk_SingerQ(', 1), ('x = dt / tau;', 4), ('if (x < TRK_SERIES_X)', 23)],
           title=('Singer Q: 급수 / 닫힌 식', 'TrackMgr.c'), highlights={4: ('if (x < TRK_SERIES_X)', 1), 5: ('bracket[1] = 1.0 + e2', 1)},
           size=7.0, lh=0.12, keep=NODECL)
legend_row(s, 0.7, 6.58, [(1, 'e^(−T/τ) − 1 을 expm1 로'), (2, 'X = F X'), (3, 'P = F P Fᵀ + Q'), (4, 'T/τ < 0.05 → 급수'), (5, 'q12 표준식 (오타 고침)')], size=9.5)
note(s, 's2_3')

# 2.4 갱신 코드
s = content(2, '2.4 갱신: CMKF 와 발산 판정 (f_Trk_Update)', '연관된 플롯으로 칼만 게인 · 추정치 · 추정 오차 공분산을 계산하고, 그때의 잔차로 NIS 를 구해 게이트 밖 횟수와 발산 판정을 셉니다.')
code_panel(s, 0.7, 2.78, 5.95, 'TargetSimCore/src/TrackMgr.c', [('// CMKF 갱신: K = P H^T', 4), ('f_Trk_Innovation(st_Track, st_Plot, &mat_V, &mat_S);', 19)],
           title=('CMKF 갱신', 'TrackMgr.c'),
           highlights={1: ('Matrix_Product2(&st_Track->mat_EstP, &mat_Ht)', 2), 2: ('Matrix_Product2(&mat_K, &mat_V)', 2), 3: ('Matrix_Identity(&mat_I);', 5)},
           size=7.8, lh=0.135, keep=NODECL)
code_panel(s, 6.85, 2.78, 5.78, 'TargetSimCore/src/TrackMgr.c', [('// 잔차 기록, 게이트 밖 횟수', 30)], title=('NIS · 게이트 밖 · 발산 판정', 'TrackMgr.c'),
           highlights={4: ('f_Trk_Mahalanobis(&mat_V, &mat_Si)', 1), 5: ('if (st_Track->nis > TRK_GATE_CHI2)', 13)}, size=7.8, lh=0.135)
legend_row(s, 0.7, 6.58, [(1, 'K = P Hᵀ S⁻¹'), (2, 'X = X + K ν'), (3, 'P = (I − K H) P, 대칭 맞춤'), (4, 'NIS = νᵀ S⁻¹ ν'), (5, '> 11.34 게이트 밖, 5 번 이어지면 발산 판정')], size=9.5)
note(s, 's2_4')

# 2.5 결과 화면
s = content(2, '2.5 결과 화면 (TargetSimUI)', '과제가 요구한 3D (LLA) 와 위도-고도 · 경도-고도 2D 궤적에 참값 · 플롯 · 트랙을 함께 그리고, 같은 ID 의 트랙은 선으로 잇습니다. 시나리오와 필터 파라미터는 화면에서 바꿉니다.')
kind, png, rects = ui_shot('air_tuned_all')
with Image.open(png) as im:
    iw, ih = im.size
W = 7.0
picture(s, png, 0.7, 2.45, w=W, border=True)
sc = W / iw
# 표시 위치 (컨트롤 기준 픽셀): 콤보는 펼침 단추 위, 그림은 안쪽 오른쪽 위, 파라미터는 고각 칸 오른쪽 빈 곳, 트랙 표는 빈 줄
def combo_anchor(x, y, w, h):
    # 펼침 단추 위. 미리보기 기록의 콤보 높이는 펼친 목록까지라 닫힌 칸 (24 px) 으로, Windows 캡처는 닫힌 칸 높이 그대로
    hc = 24 if h > 40 else h
    return x + w - 0.6 * hc, y + 0.5 * hc


UIMARK = [(1, 'IDC_SCN_COMBO', combo_anchor), (2, 'IDC_VIEW_MODE', combo_anchor),
          (3, 'IDC_VIEW3D', lambda x, y, w, h: (x + w - 24, y + 18)), (4, 'IDC_PROFILE_LAT', lambda x, y, w, h: (x + w - 24, y + 48)),
          (5, 'IDC_PRM_SIGEL', lambda x, y, w, h: (x + w + 30, y + h / 2)), (6, 'IDC_TRK_LIST', lambda x, y, w, h: (x + w - 40, y + h - 40))]
for n, key, anchor in UIMARK:
    px, py = anchor(*rects[key])
    marker(s, 0.7 + px * sc - 0.13, 2.45 + py * sc - 0.13, n, d=0.26, size=9.5, color=NAVY)
UITEXT = [('시나리오', ['대함 · 대공 (바꾸면 궤적 · 플롯부터 다시)']),
          ('보기', ['3D + 2D (기본) · 3D · 2D · 평면 · 추적 성능 (오차 · NIS)']),
          ('3D 궤적 (LLA)', ['경도 · 위도 · 고도 눈금, 끌어서 회전 · 휠로 확대']),
          ('2D 궤적', ['위도-고도 (위), 경도-고도 (아래)']),
          ('추적 필터 파라미터', ['R σ · q · τ 입력, 과제 값 · 조정 값 단추, 적용 · 다시 실행']),
          ('트랙 표', ['지금 스캔의 트랙: 상태 · 위치 · 가속도 · NIS · 발산 판정 · 오차'])]
y = 2.45
for i, (h, b) in enumerate(UITEXT):
    callout(s, 8.3, y, 4.33, i + 1, h, b, hsize=11, bsize=9.8, color=NAVY)
    y += 0.6
rect(s, 8.3, 6.12, 4.33, 0.72, F_GRAY, None, radius=0.1)
box = tb(s, 8.42, 6.15, 4.15, 0.66, '', anchor='m')
p = box.text_frame.paragraphs[0]; p.space_after = Pt(1)
add_runs(p, '그림: ', 9.8, NAVY, bold=True)
for t, c in (('━ 참값  ', BLUE), ('● 플롯  ', NAVY), ('━■ 트랙', GREEN)):
    add_runs(p, t, 9.8, c, bold=True)
add_runs(box.text_frame.add_paragraph(), '트랙 기호: □ 임시 ■ 확정 ⬚ 메모리 ⊠ 삭제, 같은 ID 끼리 선으로', 9.3, SUB)
src_tag(s, 0.7, 2.45 + ih * sc + 0.04, kind, UI_TAG[kind] + ' (대공 · 조정 값 · 스캔 1159)', W)
note(s, 's2_5')

# ================================================================== 03 결과: 과제 값
s = divider('03. 결과: 과제 값', '과제가 정한 값 (R σ 10 m · 0.3° · 0.3°, q 0.5, τ 10) 으로 대함 · 대공 시나리오를 하나씩 실행한 결과입니다.',
            [('3.1', '시나리오와 시뮬레이션 조건', 17), ('3.2', '대함 표적: 트랙이 끝까지 유지됨', 18), ('3.3', '대공 표적: 트랙은 이어지지만 오차가 큼', 19),
             ('3.4', f'대공 표적 추적 성능: 발산 판정 {int(rv("air_given", "divergence"))} 번', 20), ('3.5', '발산 원인: 필터가 가정한 기동 · 잡음이 실제보다 작음', 21)])
note(s, 'div3')

# 3.1 시나리오
s = content(3, '3.1 시나리오와 시뮬레이션 조건', '대함 · 대공 표적을 하나씩 따로 실행합니다. 클러터 없이 표적 플롯만 만들고, 탐지 플롯은 모두 트랙에 연관된다고 가정합니다.')
H = 4.05
picture(s, 'fig_scenarios.png', 0.7, 2.45, h=H)
src_tag(s, 0.7, 2.45 + H + 0.04, 'data', 'truth.csv 의 참값 궤적과 turn 열 (VC-NED, km)', H / img_ratio('fig_scenarios.png'))
COND = [('대함 표적', '32.125°N 126.03°E, 고도 0 m · 서쪽 30 m/s\n±1 G 선회 4 번 (2.4 s 동안 45°)'),
        ('대공 표적', '32.12°N 126.0°E, 고도 2,000 m · 남쪽 200 m/s\n±4 G 선회 6 번 (8 s 동안 90°, 마지막은 150°)'),
        ('레이다', '32.0°N 126.0°E 에 정지, 북쪽을 봄 (VC-NED 원점)'),
        ('시간', '116 s, 스캔 간격 0.1 s → 1,160 스캔'),
        ('플롯', 'SNR 20 표적 플롯만 (클러터 0), 실제 σ 4.74 m · 0.48° · 0.49°'),
        ('연관', '탐지 플롯을 트랙에 무조건 연관 (가정)'),
        ('필터', 'CMKF + Singer (VC-NED), R σ 10 m · 0.3° · 0.3°, q 0.5, τ 10'),
        ('개시 · 관리', '2-Point (V_{max} 800 m/s), 스코어 5 / 0 / 5 (과제 5)'),
        ('통계', '갱신 20 번 (2 s) 뒤부터 참값 대비 위치 오차')]
y = 2.47
for a, b in COND:
    hh = 0.56 if '\n' in b else 0.4
    rect(s, 6.1, y, 1.15, hh - 0.06, F_BLUE, None, radius=0.15)
    tb(s, 6.1, y, 1.15, hh - 0.06, a, size=10.5, bold=True, color=BLUE, align='c', anchor='m')
    tb(s, 7.38, y, 5.25, hh - 0.06, b.split('\n'), size=10, color=NAVY, anchor='m')
    y += hh
note(s, 's3_1')


def result_screen(slide, shot, keys, tag):
    kind, path = ui_crop(shot, keys)
    w = 11.93
    h = w * img_ratio(path)
    if h > 3.75:
        h = 3.75; w = h / img_ratio(path)
    x = 0.7 + (11.93 - w) / 2
    picture(slide, path, x, 2.45, w=w, border=True)
    src_tag(slide, x, 2.45 + h + 0.04, kind, f'{UI_TAG[kind]} · {tag}', w)
    return 2.45 + h + 0.04


def run_chips(slide, run, y, extra=None):
    items = [('위치 오차 RMS', f'{rv(run, "rms"):.1f} m'), ('최대', f'{rv(run, "max_err"):.1f} m'), ('NIS 평균', f'{rv(run, "mean_nis"):.2f}'),
             ('게이트 밖', f'{rv(run, "gate_out_pct"):.1f} %'), ('발산 판정', f'{int(rv(run, "divergence"))} 번')]
    chips(slide, 0.7, y, 11.93, items + (extra or []))


# 3.2 대함 과제 값
s = content(3, '3.2 대함 표적: 트랙이 끝까지 유지됨', f'과제 값으로 대함 표적을 끝까지 (115.9 s) 돌린 화면입니다. 트랙 하나 (ID 100) 가 {rv("ship_given", "init_time"):.1f} s 에 개시, '
            f'{rv("ship_given", "confirm_time"):.1f} s 에 확정된 뒤 끝까지 유지되고 발산 판정은 없습니다.')
y = result_screen(s, 'ship_given_all', VIEWS_ALL, '대함 · 과제 값 · 스캔 1159: 3D (LLA) + 위도-고도 + 경도-고도')
run_chips(s, 'ship_given', 6.48)
note(s, 's3_2')

# 3.3 대공 과제 값
s = content(3, '3.3 대공 표적: 트랙은 이어지지만 오차가 큼', '과제 값으로 대공 표적을 돌린 화면입니다. 연관 가정 덕분에 트랙 하나가 끝까지 이어지지만, 4 G 선회마다 참값과 수평으로 140 ~ 310 m 벌어집니다.')
y = result_screen(s, 'air_given_all', VIEWS_ALL, '대공 · 과제 값 · 스캔 1159: 3D (LLA) + 위도-고도 + 경도-고도')
run_chips(s, 'air_given', 6.48)
note(s, 's3_3')

# 3.4 대공 추적 성능
s = content(3, f'3.4 대공 표적 추적 성능: 발산 판정 {int(rv("air_given", "divergence"))} 번', '추적 성능 보기 (위: 참값 대비 위치 오차, 아래: NIS 로그 눈금). 선회 (회색 띠) 마다 오차와 NIS 가 커지고, '
            '게이트 밖이 5 번 이상 이어진 구간 (빨간 띠) 이 생깁니다.')
y = result_screen(s, 'air_given_perf', ('IDC_FILTER_PLOT',), '대공 · 과제 값 · 추적 성능 보기')
NOTE34 = [('선회마다 오차 140 ~ 310 m', '최대 314 m (세 번째 선회)'),
          ('NIS 수십 ~ 170', f'게이트 밖 {rv("air_given", "gate_out_pct"):.1f} %'),
          ('선회 뒤에도 몇 초 이어짐', '가속도 추정이 늦게 빠짐')]
for i, (a, b) in enumerate(NOTE34):
    xx = 0.7 + i * 4.02
    rect(s, xx, 6.46, 3.89, 0.4, F_GRAY, BORDER, radius=0.2)
    box = tb(s, xx + 0.12, 6.46, 3.7, 0.4, '', anchor='m')
    pp = box.text_frame.paragraphs[0]
    add_runs(pp, a + '   ', 10.5, NAVY, bold=True)
    add_runs(pp, b, 9, SUB)
note(s, 's3_4')

# 3.5 발산 원인
s = content(3, '3.5 발산 원인: 필터가 가정한 기동 · 잡음이 실제보다 작음', '과제 값의 q 0.5 m/s² 는 대공 표적의 4 G (39 m/s²) 선회를 거의 없는 일로 봅니다. R 의 0.3° 도 실제 플롯 잡음 (0.48°) 보다 작습니다.')
H = 3.55
picture(s, 'fig_lag_given.png', 0.7, 2.45, h=H)
WL = H / img_ratio('fig_lag_given.png')
src_tag(s, 0.7, 2.45 + H + 0.04, 'data', 'air_given 의 truth · plots · tracks.csv (3 ~ 17.5 s)', WL)
rect(s, 0.7, 6.36, WL, 0.48, F_BLUE, None, radius=0.15)
tb(s, 0.82, 6.36, WL - 0.2, 0.48, 'q 가 작으면 P · K 가 작아 직진하는 예측을 믿음 → 선회를 늦게 따라감', size=10, bold=True, anchor='m')
XR = 0.7 + WL + 0.25
WR = 12.63 - XR
picture(s, 'fig_accel_given.png', XR, 2.45, w=WR)
yy = 2.45 + WR * img_ratio('fig_accel_given.png') + 0.04
src_tag(s, XR, yy, 'data', '참값: truth.csv 위치의 2 차 차분, 트랙: tracks.csv 의 an · ae · ad', WR)
table(s, XR, yy + 0.36, [1.45, 2.2, WR - 3.65], [['', '과제 값 (필터가 가정)', '실제'],
                                              ['기동 가속도', 'q 0.5 m/s²', '대공 4 G = 39 m/s², 대함 1 G = 9.8 m/s²'],
                                              ['플롯 각도 잡음', 'R σ 0.3°', 'SNR 20: 0.48° (방위), 0.49° (고각)']],
      row_h=0.32, size=10, fills=[None, F_ORANGE, F_ORANGE], bold_cols=(0,))
note(s, 's3_5')

# ================================================================== 04 파라미터 조정
s = divider('04. 파라미터 조정', '트랙이 끊기지 않고 발산하지 않도록 측정잡음 (R σ) 과 공정잡음 (q, τ) 파라미터를 30 회 몬테카를로 실험으로 골랐습니다.',
            [('4.1', '조정 방법: 30 회 몬테카를로', 23), ('4.2', '측정잡음 R: σ 0.3° → 0.5°', 24), ('4.3', '공정잡음 q · τ', 25),
             ('4.4', '대함 조정 결과', 26), ('4.5', '대공 조정 결과', 27), ('4.6', '과제 값과 조정 값 비교 (대공)', 28), ('4.7', '정리', 29)])
note(s, 'div4')

# 4.1 조정 방법
s = content(4, '4.1 조정 방법: 30 회 몬테카를로', '한 번의 실행은 잡음 하나에 대한 결과라 우연이 섞입니다. 플롯 잡음 seed 를 30 개로 바꿔 돌린 평균과 발산이 난 회차 수로 값을 골랐습니다.')
n_set = len(SWEEP)
n_run = n_set * int(SWEEP.seeds.iloc[0])
PROC = [('① 설정', 'q 0.1 ~ 80 (10 단계)\n× R σ 0.3° · 0.5°\n+ τ 1 ~ 60 s (6 단계)'),
        ('② 30 회 반복', f'플롯 잡음 seed 30 개\n(seed 0 = 프로그램 기본)\n{n_set // 2} 설정 × 2 시나리오 → {n_run:,} 번'),
        ('③ 지표', '위치 오차 RMS (갱신 20 번 뒤)\nNIS 평균, 게이트 밖 %\n발산 판정 횟수'),
        ('④ 고르는 규칙', '30 회 모두 발산 판정 0\n그 가운데 RMS 가 가장 작은 값')]
for i, (h, b) in enumerate(PROC):
    xx = 0.7 + i * 3.04
    last = i == 3
    rect(s, xx, 2.5, 2.8, 1.62, F_BLUE if last else F_GRAY, BLUE if last else 'E3E8F0', radius=0.08, lw=1.2 if last else 0.75)
    tb(s, xx + 0.15, 2.58, 2.5, 0.3, h, size=12, bold=True, color=BLUE if last else NAVY)
    tb(s, xx + 0.15, 2.93, 2.55, 1.15, b.split('\n'), size=10, color=NAVY, after=2)
    if i < 3:
        line(s, xx + 2.83, 3.31, xx + 3.01, 3.31, color=SUB, width=1.1, arrow=True)
tb(s, 0.7, 4.22, 11.93, 0.28, '실행은 실제 코드 그대로: PlotLink.c UDP 송수신 쓰레드 + TrackMgr.c (검증 프로그램 tools/verify/harness_main.c, 결과 data/sweep.csv)',
   size=9.8, color=SUB)
tb(s, 0.7, 4.62, 6, 0.3, '고른 값', size=12, bold=True)
TUNE = [['', 'R σ (거리 · 방위 · 고각)', 'q [m/s²]', 'τ [s]', '30 회 중 발산 판정이 난 회차', '위치 오차 RMS (30 회 평균)'],
        ['과제 값', '10 m · 0.3° · 0.3°', '0.5', '10', f'대함 {SEEDS["ship_given"]} · 대공 {SEEDS["air_given"]}',
         f'대함 {SWEEP_RMS["ship_given"]:.1f} m · 대공 {SWEEP_RMS["air_given"]:.1f} m'],
        ['조정 값 대함', '10 m · 0.5° · 0.5°', f'{rv("ship_tuned", "q"):g}', f'{rv("ship_tuned", "tau"):g}', f'{SEEDS["ship_tuned"]}', f'{SWEEP_RMS["ship_tuned"]:.1f} m'],
        ['조정 값 대공', '10 m · 0.5° · 0.5°', f'{rv("air_tuned", "q"):g}', f'{rv("air_tuned", "tau"):g}', f'{SEEDS["air_tuned"]}', f'{SWEEP_RMS["air_tuned"]:.1f} m']]
table(s, 0.7, 4.95, [1.6, 2.4, 1.15, 0.95, 2.6, 3.23], TUNE, row_h=0.36, size=10.5, fills=[None, F_ORANGE, F_BLUE, F_BLUE], bold_cols=(0,))
tb(s, 0.7, 6.48, 11.93, 0.3, '화면의 [조정 값] 단추가 시나리오에 맞는 값을 넣음 (Scenario.cpp s_Tuned)', size=10, color=SUB)
note(s, 's4_1')

# 4.2 R 조정
s = content(4, '4.2 측정잡음 R: σ 0.3° → 0.5°', '실제 플롯의 각도 잡음은 0.48° 인데 R 은 0.3° 로 가정해 플롯을 너무 믿습니다. 그래서 q 를 아무리 바꿔도 NIS 가 3 으로 내려오지 않습니다.')
W = 7.25
picture(s, 'fig_plotnoise.png', 0.7, 2.45, w=W)
yy = 2.45 + W * img_ratio('fig_plotnoise.png') + 0.04
src_tag(s, 0.7, yy, 'data', 'plots.csv − truth.csv (두 시나리오 플롯 2,320 개)', W)
rect(s, 0.7, yy + 0.42, W, 6.84 - yy - 0.42, F_GRAY, None, radius=0.06)
tb(s, 0.9, yy + 0.5, W - 0.3, 6.84 - yy - 0.55,
   [('왜 0.5° 인가', dict(size=11.5, bold=True, after=3)),
    ('· 실제 각도 잡음 (방위 0.47°, 고각 0.49°) 보다 조금 크게, 보수적으로 잡은 값', dict(size=10.3, after=2)),
    ('· 거리는 실제 4.7 m 보다 과제 값 10 m 가 넓어 그대로 둠', dict(size=10.3, after=2)),
    ('· 과제 5 와 같은 이유: 받는 쪽은 실제 잡음을 모름. 결과 (NIS) 를 보고 R 을 맞춤', dict(size=10.3))])
XR = 8.2; WR = 12.63 - XR
picture(s, 'fig_sweep_nis.png', XR, 2.45, w=WR)
yy = 2.45 + WR * img_ratio('fig_sweep_nis.png') + 0.04
src_tag(s, XR, yy, 'data', 'sweep.csv (설정마다 30 회 평균)', WR)
ship_hi = sweep_row('ship', 0.3, 80.0); air_hi = sweep_row('air', 0.3, 80.0)
tb(s, XR, yy + 0.4, WR, 1.9, [(f'σ 0.3°: q 를 80 까지 키워도 NIS {min(ship_hi.mean_nis, air_hi.mean_nis):.1f} 이상', dict(size=10.5, bold=True, after=2)),
                              (f'게이트 밖 {min(ship_hi.gate_out_pct, air_hi.gate_out_pct):.0f} % 남음 → q 로는 안 풀림', dict(size=10, color=SUB, after=6)),
                              ('σ 0.5°: NIS 2.0 ~ 2.5 (q 가 충분할 때)', dict(size=10.5, bold=True, after=2)),
                              ('게이트 밖 0.2 ~ 0.4 %, 3 보다 조금 작음 = 보수적', dict(size=10, color=SUB))])
note(s, 's4_2')

# 4.3 q · τ
s = content(4, '4.3 공정잡음 q · τ', 'q 가 작으면 선회를 못 따라가 발산하고, 크면 플롯 잡음을 따라가 오차가 커집니다. 대함은 q 0.5, 대공은 q 20 에서 발산 없이 오차가 가장 작습니다.')
W = 7.3
picture(s, 'fig_sweep_q.png', 0.7, 2.42, w=W)
src_tag(s, 0.7, 2.42 + W * img_ratio('fig_sweep_q.png') + 0.03, 'data', 'sweep.csv (설정마다 seed 30 개)', W)
XR = 8.3; WR = 12.63 - XR
tb(s, XR, 2.45, WR, 0.28, 'τ 의 영향 (조정한 R σ, q 에서)', size=11.5, bold=True)
picture(s, 'fig_sweep_tau.png', XR, 2.78, w=WR)
yy = 2.78 + WR * img_ratio('fig_sweep_tau.png') + 0.03
src_tag(s, XR, yy, 'data', 'sweep.csv τ 실험', WR)
tau_rng = {sc: SWEEP[(SWEEP.scenario == sc) & (SWEEP.sweep == 'tau')].rms for sc in ('ship', 'air')}
tb(s, XR, yy + 0.38, WR, 2.2, [(f'τ 1 ~ 60 s 에서 RMS 차이 대함 {tau_rng["ship"].max() - tau_rng["ship"].min():.1f} m, 대공 {tau_rng["air"].max() - tau_rng["air"].min():.1f} m',
                                 dict(size=10.3, bold=True, after=2)),
                                ('→ τ 는 과제 값 10 s 그대로', dict(size=10, color=SUB, after=8)),
                                ('q 를 고른 기준', dict(size=11.5, bold=True, after=3)),
                                ('· 대함: 선회가 짧아 (2.4 s) 작은 q 가 유리 → 0.5', dict(size=10, after=2)),
                                ('· 대공: 4 G (39 m/s²) 가 약 2σ 안에 들도록 → 20', dict(size=10, after=2)),
                                ('· 과제 값 q 0.5 는 대공에서 30 회 모두 발산', dict(size=10))])
note(s, 's4_3')

# 4.4 대함 조정 결과
s = content(4, '4.4 대함 조정 결과: R σ 0.5°', f'R 의 각도 σ 만 0.5° 로 바꿔 NIS 가 실제 잡음과 맞아졌습니다 (평균 {rv("ship_given", "mean_nis"):.2f} → {rv("ship_tuned", "mean_nis"):.2f}). '
            '궤적과 위치 오차는 과제 값과 거의 같습니다.')
y = result_screen(s, 'ship_tuned_all', VIEWS_ALL, '대함 · 조정 값 (R σ 10 m · 0.5° · 0.5°, q 0.5, τ 10) · 스캔 1159')


def delta_chips(slide, sc, y):
    g, t = sc + '_given', sc + '_tuned'
    items = [('RMS', f'{rv(g, "rms"):.1f} → {rv(t, "rms"):.1f} m'), ('최대', f'{rv(g, "max_err"):.0f} → {rv(t, "max_err"):.0f} m'),
             ('NIS 평균', f'{rv(g, "mean_nis"):.2f} → {rv(t, "mean_nis"):.2f}'), ('게이트 밖', f'{rv(g, "gate_out_pct"):.1f} → {rv(t, "gate_out_pct"):.1f} %'),
             ('발산 판정', f'{int(rv(g, "divergence"))} → {int(rv(t, "divergence"))} 번'), ('30 회 중 발산', f'{SEEDS[g]} → {SEEDS[t]}')]
    chips(slide, 0.7, y, 11.93, items, gap=0.1)


delta_chips(s, 'ship', 6.48)
note(s, 's4_4')

# 4.5 대공 조정 결과
s = content(4, '4.5 대공 조정 결과: R σ 0.5°, q 20', f'선회에서도 트랙이 참값을 따라가며 발산 판정이 사라졌습니다 ({int(rv("air_given", "divergence"))} 번 → {int(rv("air_tuned", "divergence"))} 번). '
            f'위치 오차 RMS 는 {rv("air_given", "rms"):.0f} m → {rv("air_tuned", "rms"):.0f} m 로 줄었습니다.')
y = result_screen(s, 'air_tuned_all', VIEWS_ALL, '대공 · 조정 값 (R σ 10 m · 0.5° · 0.5°, q 20, τ 10) · 스캔 1159')
delta_chips(s, 'air', 6.48)
note(s, 's4_5')

# 4.6 비교
s = content(4, '4.6 과제 값과 조정 값 비교 (대공)', '과제 값 (주황) 은 선회마다 오차가 커지고 NIS 가 게이트를 넘지만, 조정 값 (파랑) 은 선회 중에도 참값을 따라가고 NIS 가 게이트 아래에 머뭅니다.')
W = 7.3
picture(s, 'fig_compare_air.png', 0.7, 2.45, w=W)
yy = 2.45 + W * img_ratio('fig_compare_air.png') + 0.04
src_tag(s, 0.7, yy, 'data', 'air_given · air_tuned 의 tracks.csv (err, nis)', W)
H = yy - 2.45 - 0.04
WL = H / img_ratio('fig_lag_both.png')
picture(s, 'fig_lag_both.png', 12.63 - WL, 2.45, w=WL)
src_tag(s, 12.63 - WL, yy, 'data', '첫 선회 평면 (truth · plots · tracks.csv)', WL)
rect(s, 0.7, yy + 0.38, 11.93, 6.84 - yy - 0.38, F_BLUE, None, radius=0.12)
tb(s, 0.9, yy + 0.38, 11.6, 6.84 - yy - 0.38, [('q ↑ → 예측 P ↑ → 칼만 게인 K ↑ : 선회하는 플롯을 바로 따라감.  대신 직진에서는 플롯 잡음을 조금 더 따라가므로 RMS 가 가장 작은 q 20 을 고름',
                                                dict(size=10.8, bold=True))], anchor='m')
note(s, 's4_6')

# 4.7 정리
s = content(4, '4.7 정리', '과제 값으로는 대공 표적이 선회마다 발산 판정을 받았고, R σ 와 q 를 조정해 두 시나리오 모두 발산 판정 없이 트랙을 유지했습니다.')
SUMROW = [['실행', 'R σ (방위 · 고각)', 'q', '위치 오차 RMS', '최대 오차', 'NIS 평균', '게이트 밖', '발산 판정', '30 회 중 발산']]
for run, name in (('ship_given', '대함 · 과제 값'), ('ship_tuned', '대함 · 조정 값'), ('air_given', '대공 · 과제 값'), ('air_tuned', '대공 · 조정 값')):
    SUMROW.append([name, f'{rv(run, "sigma_az_deg"):g}° · {rv(run, "sigma_el_deg"):g}°', f'{rv(run, "q"):g}', f'{rv(run, "rms"):.1f} m', f'{rv(run, "max_err"):.1f} m',
                   f'{rv(run, "mean_nis"):.2f}', f'{rv(run, "gate_out_pct"):.1f} %', f'{int(rv(run, "divergence"))} 번', f'{SEEDS[run]} 회'])
table(s, 0.7, 2.5, [1.85, 1.6, 0.7, 1.45, 1.3, 1.1, 1.15, 1.15, 1.63], SUMROW, row_h=0.36, size=10.5,
      fills=[None, F_ORANGE, F_BLUE, F_ORANGE, F_BLUE], bold_cols=(0,))
src_tag(s, 0.7, 4.36, 'data', 'runs.csv (프로그램 기본 잡음 한 번), sweep.csv (30 회)', 11.93)
rect(s, 0.7, 4.75, 5.85, 2.09, F_GRAY, None, radius=0.06)
tb(s, 0.9, 4.82, 5.5, 2.0, [('결론', dict(size=12, bold=True, after=4)),
                            ('· 과제 값 q 0.5 는 4 G 선회를 못 따라감 → 대공 발산 판정 11 번', dict(size=11, after=5)),
                            ('· R σ 를 실제 플롯 잡음 수준 (0.5°) 으로 → NIS 가 정상 범위', dict(size=11, after=5)),
                            ('· q 를 표적 기동에 맞게 (대공 20) → 발산 0, RMS 137 → 37 m', dict(size=11, after=5)),
                            ('· 30 회 반복에서도 두 시나리오 모두 발산 판정 0', dict(size=11))])
rect(s, 6.78, 4.75, 5.85, 2.09, WHITE, BORDER, radius=0.06)
tb(s, 6.98, 4.82, 5.5, 2.0, [('한계와 다음 단계', dict(size=12, bold=True, after=4)),
                             ('· 가장 좋은 q 가 대함 0.5 · 대공 20 으로 다름 → 표적을 모르면 IMM 같은 다중 모델', dict(size=11, after=5)),
                             ('· 무조건 연관 가정 → 클러터가 있으면 게이트 · 연관 기법 필요 (NIS 가 그 기준)', dict(size=11, after=5)),
                             ('· 대함 조정 값: 직진 NEES 3.5 (맞으면 3), 선회 중 · 직후 10 ~ 19 → P 가 실제 오차보다 작은 구간', dict(size=11))])
note(s, 's4_7')

# 끝
s = add_slide('END')
tb(s, 4.6, 3.05, 7.4, 0.8, '감사합니다', size=36, bold=True, color=WHITE, align='r')
line(s, 8.6, 4.02, 12.0, 4.02, color=MUTED, width=0.75)
tb(s, 4.6, 4.16, 7.4, 0.34, '질문 주시면 편하게 답변드리겠습니다', size=13, color=LIGHT, align='r')
note(s, 'end')

# ------------------------------------------------------------------ checks: 목차 · 간지 쪽 번호, 대본 숫자 = 결과 데이터
for num, name, page in INDEX:
    key = 'div' + str(int(num))
    assert PAGE[key] == page, f'목차 {num} {name}: 적은 쪽 {page}, 실제 {PAGE[key]}'
SECTION_KEYS = {k: v for k, v in PAGE.items() if re.fullmatch(r's\d_\d', k)}
for num, page in DIVIDER_ITEMS:
    key = 's' + num.replace('.', '_')
    assert SECTION_KEYS.get(key) == page, f'간지 {num}: 적은 쪽 {page}, 실제 {SECTION_KEYS.get(key)}'


def need(key, *texts):
    for t in texts:
        assert t in NOTES[key], f'대본 {key} 에 결과 데이터 값 "{t}" 가 없음 (notes.py 를 데이터에 맞게 고칠 것)'


need('s1_1', f'{rv("air_tuned", "init_time"):.1f} 초에 개시', f'{rv("air_tuned", "confirm_time"):.1f} 초에 확정')
need('s3_2', f'RMS {rv("ship_given", "rms"):.1f} 미터', f'NIS 평균 {rv("ship_given", "mean_nis"):.2f}', f'게이트 밖 {rv("ship_given", "gate_out_pct"):.1f} 퍼센트')
need('s3_3', f'RMS 가 {rv("air_given", "rms"):.1f} 미터', f'최대 {rv("air_given", "max_err"):.1f} 미터', f'NIS 평균은 {rv("air_given", "mean_nis"):.2f}',
     f'{rv("air_given", "gate_out_pct"):.1f} 퍼센트', f'발산 판정이 {int(rv("air_given", "divergence"))} 번')
need('s3_5', f'NIS 평균이 3 이 아니라 {rv("ship_given", "mean_nis"):.2f}')
need('s4_3', f'대함에서 30 번 중 {SEEDS["ship_given"]} 번', f'q {rv("ship_tuned", "q"):g} 에서 {SWEEP_RMS["ship_tuned"]:.0f} 미터', f'q {rv("air_tuned", "q"):g} 에서 {SWEEP_RMS["air_tuned"]:.0f} 미터')
need('s4_4', f'{rv("ship_given", "mean_nis"):.2f} 에서 {rv("ship_tuned", "mean_nis"):.2f}', f'{rv("ship_given", "gate_out_pct"):.1f} 퍼센트에서 {rv("ship_tuned", "gate_out_pct"):.1f} 퍼센트',
     f'{rv("ship_given", "rms"):.1f} 미터에서 {rv("ship_tuned", "rms"):.1f} 미터', f'{SEEDS["ship_given"]} 번에서 {SEEDS["ship_tuned"]} 번')
need('s4_5', f'{int(rv("air_given", "divergence"))} 번에서 {int(rv("air_tuned", "divergence"))} 번', f'{rv("air_given", "rms"):.1f} 미터에서 {rv("air_tuned", "rms"):.1f} 미터',
     f'{rv("air_given", "max_err"):.1f} 미터에서 {rv("air_tuned", "max_err"):.1f} 미터', f'NIS 평균은 {rv("air_tuned", "mean_nis"):.2f}')
need('s4_1', f'{n_run:,} 번')
need('s4_7', f'모두 {int(rv("air_given", "divergence"))} 번')

prs.save(OUT)
print('saved', OUT, len(prs.slides), 'slides')

# 대본 문서
TITLES = []
for key, sl in slides_meta:
    t = ''
    for sh in sl.placeholders:
        if sh.placeholder_format.idx == 103:
            t = sh.text_frame.text
    if not t:
        texts = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and not sh.is_placeholder and sh.text_frame.text.strip()]
        t = {'cover': '표지', 'index': '목차', 'end': '마무리'}.get(key, texts[0] if texts else key)
    TITLES.append(t)
with open(SCRIPT_MD, 'w', encoding='utf-8') as f:
    f.write('# 과제 6 추적 필터 (Singer CMKF) — 발표 대본\n\n')
    f.write('슬라이드 노트와 같은 내용입니다. 발표 시간은 25 ~ 30 분을 기준으로 썼습니다.\n\n')
    for i, ((key, sl), t) in enumerate(zip(slides_meta, TITLES), 1):
        f.write(f'## {i}. {t}\n\n{NOTES[key].replace("{ui_source}", ui_source_note())}\n\n')
print('saved', SCRIPT_MD)
