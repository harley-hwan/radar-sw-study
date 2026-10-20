# -*- coding: utf-8 -*-
"""과제 5 발표자료 만들기.

    python build_deck.py <틀.pptx> <그림 폴더> <화면 캡처 폴더> <출력.pptx>
    (예: 틀 = ../../../../Part4_Plot_Clutter_Sim/ref/산출물/Chapter4_PlotClutterSim.pptx, 그림 · 캡처 = ../figures)

틀은 과제 4 발표자료 (Chapter4_PlotClutterSim.pptx). 장은 모두 지우고 틀의 레이아웃
(COVER, INDEX, DIVIDER, CONTENT, END) 과 장 모양 (표지 · 목차 · 간지 글자, 머리띠, 쪽 번호) 만 가져온다.
표지와 끝 장 바탕 그림은 그림 폴더의 cover_bg.png 로 바꾼다.
코드 발췌는 src 의 실제 파일에서 줄을 찾아 읽음. 그래서 줄 번호가 늘 파일과 같다.
노트는 같은 폴더의 Chapter5_발표대본.md 에서 읽는다.
"""
import copy
import io
import os
import re
import sys

from lxml import etree
from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.normpath(os.path.join(HERE, '..'))
SRC = os.path.normpath(os.path.join(HERE, '..', '..', '..', 'src', 'TargetSim'))
SCRIPT_MD = os.path.join(OUT_DIR, 'Chapter5_발표대본.md')

KO = '맑은 고딕'
EN = 'Arial'
CODE = 'Consolas'
NAVY = '1B2740'
SUB = '4A5A7A'
MUTE = '8794AC'
LNUM = 'AEB9CC'
LINE = 'D2D9E5'
BOXF = 'F7F9FC'
WHITE = 'FFFFFF'
STRONG = ['DC2626', 'D97706', '059669', '2563EB', '7C3AED']
LIGHT = ['FDEEEE', 'FEF3E2', 'E6F5EE', 'EAF1FD', 'F2ECFD']
TENT = 'EA580C'

FONT_REG = ImageFont.truetype(r'C:\Windows\Fonts\malgun.ttf', 100)
FONT_BOLD = ImageFont.truetype(r'C:\Windows\Fonts\malgunbd.ttf', 100)


def text_width(text, size, bold=False):
    """글자 폭 [in]"""
    f = FONT_BOLD if bold else FONT_REG
    return f.getlength(text) / 100.0 * size / 72.0


# 노트 (발표 대본 md)

def read_notes(path):
    md = io.open(path, encoding='utf-8').read()
    out = {}
    for blk in re.split(r'(?m)^## ', md)[1:]:
        head = re.match(r'(\d+)\.\s+(.+)', blk)
        if not head:
            continue
        body = []
        for raw in blk.split('\n')[1:]:
            line = raw.strip()
            if not line or line.startswith(('---', '#', '|', '>')) or re.match(r'\[(.+?)\]$', line):
                continue
            body.append(line)
        out[int(head.group(1))] = '\n'.join(body)
    return out


# XML 도우미

def new_id(slide):
    ids = [int(e.get('id')) for e in slide.shapes._spTree.iter(qn('p:cNvPr'))]
    return max(ids + [1]) + 1


def clone(slide, proto):
    el = copy.deepcopy(proto)
    el.find('.//' + qn('p:cNvPr')).set('id', str(new_id(slide)))
    slide.shapes._spTree.append(el)
    return el


def el_set_text(el, text):
    """첫 단락 첫 run 의 글자만 바꾸고 나머지 run, 단락은 지움"""
    body = el.find(qn('p:txBody'))
    paras = body.findall(qn('a:p'))
    for p in paras[1:]:
        body.remove(p)
    runs = paras[0].findall(qn('a:r'))
    for r in runs[1:]:
        paras[0].remove(r)
    runs[0].find(qn('a:t')).text = text


def el_set_pos(el, x=None, y=None, w=None, h=None):
    off = el.find('.//' + qn('a:off'))
    ext = el.find('.//' + qn('a:ext'))
    if x is not None:
        off.set('x', str(int(Inches(x))))
    if y is not None:
        off.set('y', str(int(Inches(y))))
    if w is not None:
        ext.set('cx', str(int(Inches(w))))
    if h is not None:
        ext.set('cy', str(int(Inches(h))))


def strip_style(shape):
    st = shape._element.find(qn('p:style'))
    if st is not None:
        shape._element.remove(st)


def set_run(run, text, size, color=NAVY, bold=False, font=KO):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._r.get_or_add_rPr()
    for tag in ('a:latin', 'a:ea', 'a:cs'):
        e = rpr.find(qn(tag))
        if e is None:
            e = etree.SubElement(rpr, qn(tag))
        e.set('typeface', font)


def body_zero(tf, anchor='t', wrap=True):
    bp = tf._txBody.find(qn('a:bodyPr'))
    for k in ('lIns', 'tIns', 'rIns', 'bIns'):
        bp.set(k, '0')
    bp.set('wrap', 'square' if wrap else 'none')
    bp.set('anchor', {'t': 't', 'm': 'ctr', 'b': 'b'}[anchor])


def para_spacing(p, lnspc_pct=None, lnspc_pt=None, after_pt=None):
    ppr = p._p.get_or_add_pPr()
    if lnspc_pct is not None or lnspc_pt is not None:
        ln = etree.SubElement(ppr, qn('a:lnSpc'))
        if lnspc_pt is not None:
            etree.SubElement(ln, qn('a:spcPts')).set('val', str(int(lnspc_pt * 100)))
        else:
            etree.SubElement(ln, qn('a:spcPct')).set('val', str(int(lnspc_pct * 1000)))
    if after_pt is not None:
        sa = etree.SubElement(ppr, qn('a:spcAft'))
        etree.SubElement(sa, qn('a:spcPts')).set('val', str(int(after_pt * 100)))


def text(slide, x, y, w, h, paras, anchor='t', align='l', wrap=True, lnspc=None, after=None):
    """paras: [[(글자, 크기, 색, 굵게, 글꼴), ...], ...] 또는 단락마다 str"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    body_zero(tf, anchor, wrap)
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[align]
        if lnspc is not None or after is not None:
            para_spacing(p, lnspc_pct=lnspc, after_pt=after)
        if isinstance(para, str):
            para = [(para, 11, NAVY, False, KO)]
        for spec in para:
            t, size, color, bold, font = (list(spec) + [11, NAVY, False, KO][len(spec) - 1:])[:5]
            set_run(p.add_run(), t, size, color, bold, font)
    return tb


def rect(slide, x, y, w, h, fill=None, line=None, lw=0.75, shape=MSO_SHAPE.RECTANGLE, dash=None):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    strip_style(sp)
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = RGBColor.from_string(line)
        sp.line.width = Pt(lw)
        if dash:
            ln = sp._element.spPr.find(qn('a:ln'))
            etree.SubElement(ln, qn('a:prstDash')).set('val', dash)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sp.adjustments[0] = min(0.5, 0.08 / max(min(w, h), 0.01))
    return sp


def hline(slide, x0, y, x1, color=LINE, lw=0.75, dash=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x0), Inches(y), Inches(x1), Inches(y))
    strip_style(c)
    c.line.color.rgb = RGBColor.from_string(color)
    c.line.width = Pt(lw)
    if dash:
        ln = c._element.spPr.find(qn('a:ln'))
        etree.SubElement(ln, qn('a:prstDash')).set('val', dash)
    return c


def arrow_line(slide, x0, y0, x1, y1, color=MUTE, lw=1.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x0), Inches(y0), Inches(x1), Inches(y1))
    strip_style(c)
    c.line.color.rgb = RGBColor.from_string(color)
    c.line.width = Pt(lw)
    ln = c._element.spPr.find(qn('a:ln'))
    t = etree.SubElement(ln, qn('a:tailEnd'))
    t.set('type', 'triangle')
    t.set('w', 'med')
    t.set('len', 'med')
    return c


def badge(slide, n, x, y, d=0.28, color_idx=None):
    col = STRONG[(n - 1) % len(STRONG)] if color_idx is None else STRONG[color_idx]
    sp = rect(slide, x, y, d, d, fill=col, shape=MSO_SHAPE.OVAL)
    tf = sp.text_frame
    body_zero(tf, 'm')
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    set_run(p.add_run(), str(n), 9.5, WHITE, True)
    return sp


def callout(slide, n, x, y, w, title, lines, size=11.0, h=None):
    badge(slide, n, x, y)
    paras = [[(title, size, NAVY, True)]] + [[(l, size, SUB, False)] for l in lines]
    hh = h if h is not None else 0.24 * len(paras)
    return text(slide, x + 0.42, y - 0.02, w - 0.42, hh, paras, lnspc=105)


def picture(slide, path, x, y, w=None, h=None, border=False):
    pic = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)
    if border:
        pic.line.color.rgb = RGBColor.from_string(LINE)
        pic.line.width = Pt(0.75)
    return pic


def label(slide, x, y, w, title, sub=None):
    paras = [[(title, 11.5, NAVY, True)] + ([('    ' + sub, 9.5, MUTE, False, CODE)] if sub else [])]
    return text(slide, x, y, w, 0.24, paras, anchor='m')


# 코드 발췌

TYPES = {'VOID', 'INT32', 'UINT32', 'INT64', 'UINT16', 'FLOAT64', 'FLOAT32', 'DOUBLE64', 'UCHAR8', 'CHAR', 'matrix', 'const', 'static',
         'return', 'if', 'else', 'for', 'while', 'switch', 'case', 'break', 'default', 'typedef', 'struct', 'sizeof', 'SOCKET', 'HANDLE',
         'DWORD', 'ULONGLONG', 'TSCORE_API', '__stdcall', 'size_t'}
TOKEN = re.compile(r'(//.*$)|("[^"]*")|(\b\d+(?:\.\d+)?(?:[eE][-+]?\d+)?[UuLlFf]*\b)|([A-Za-z_][A-Za-z0-9_]*)|(\s+)|(.)')


def colorize(code):
    """C 한 줄 -> [(글자, 색, 글꼴)]"""
    out = []
    for m in TOKEN.finditer(code):
        com, s, num, ident, ws, other = m.groups()
        if com:
            out.append((com, '059669', KO))
        elif s:
            out.append((s, 'B45309', CODE))
        elif num:
            out.append((num, 'B45309', CODE))
        elif ident:
            if ident in TYPES or re.match(r'^(ST|STRUCT)_[A-Za-z]', ident) and not ident.isupper():
                col = '2563EB'
            elif re.match(r'^[A-Z][A-Z0-9_]{2,}$', ident) and ident not in TYPES:
                col = '7C3AED'
            else:
                col = NAVY
            out.append((ident, col, CODE))
        else:
            out.append((ws or other, NAVY, CODE))
    # 이어진 같은 색 · 글꼴 합치기
    merged = []
    for t, c, f in out:
        if merged and merged[-1][1] == c and merged[-1][2] == f:
            merged[-1] = (merged[-1][0] + t, c, f)
        else:
            merged.append((t, c, f))
    return merged


def detab(line):
    lead = len(line) - len(line.lstrip('\t'))
    rest = line[lead:]
    rest = re.sub(r'\t+', '  ', rest)
    return '  ' * lead + rest.rstrip()


def read_lines(rel):
    return io.open(os.path.join(SRC, rel), encoding='utf-8').read().split('\n')


def find(lines, needle, start=0):
    for i in range(start, len(lines)):
        if needle in lines[i]:
            return i
    raise KeyError(needle)


def excerpt(rel, parts, dedent=0):
    """parts: [(시작 글자, 끝 글자 또는 줄 수)] -> [(줄 번호, 글자)]. 조각 사이에 '...' 줄."""
    lines = read_lines(rel)
    out = []
    pos = 0
    for k, (a, b) in enumerate(parts):
        i0 = find(lines, a, pos)
        i1 = i0 + b - 1 if isinstance(b, int) else find(lines, b, i0)
        if k > 0 and out and out[-1][0] is not None and i0 > out[-1][0]:
            out.append((None, '...'))
        for i in range(i0, i1 + 1):
            t = detab(lines[i])
            if dedent and t.startswith(' ' * dedent):
                t = t[dedent:]
            out.append((i + 1, t))
        pos = i1 + 1
    return out


def annotate(rows, notes):
    """발표용 줄 주석. notes: [(찾을 글자, 주석)] 또는 [(찾을 글자, 주석, 'above')], 코드 순서대로.
    줄 끝에 '// 주석' 을 붙이고, 'above' 면 그 줄 위에 주석만 있는 줄 (번호 없음) 을 넣음.
    찾을 글자는 탭을 공백 두 칸으로 바꾼 뒤의 글자로 찾음."""
    out = list(rows)
    pos = 0
    for note in notes:
        needle, com = note[0], note[1]
        i = row_of(out, needle, pos)
        num, code = out[i]
        if len(note) > 2 and note[2] == 'above':
            ind = len(code) - len(code.lstrip(' '))
            out.insert(i, (None, ' ' * ind + '// ' + com))
            i += 1
        else:
            out[i] = (num, code + '  // ' + com)
        pos = i + 1
    return out


def row_width(lab, code, size):
    """코드 한 줄의 폭 [in]. Consolas 글자 폭 0.5498 em, 주석 (맑은 고딕) 은 글꼴로 잼"""
    wd = len(lab) * 0.5498 * size / 72.0
    for t, c, f in colorize(code):
        wd += len(t) * 0.5498 * size / 72.0 if f == CODE else FONT_REG.getlength(t) / 100.0 * size / 72.0
    return wd


def code_box(slide, x, y, w, rows, bands=(), size=9.0, step=11.5, markers=True, pad=0.09):
    """rows: [(줄 번호, 글자)], bands: [(첫 행, 행 수, 색 번호, 표시 번호)]
    PowerPoint 가 고정 줄 간격을 1 pt 단위로 맞춤. 그래서 step 도 정수 pt 로 맞춰야 띠가 줄과 어긋나지 않음"""
    step = float(round(step))
    n = len(rows)
    h = pad * 2 + n * step / 72.0
    rect(slide, x, y, w, h, fill=WHITE, line=LINE)
    for (r0, cnt, ci, mark) in bands:
        by = y + pad + r0 * step / 72.0 - 0.012
        bh = cnt * step / 72.0
        rect(slide, x + 0.03, by, w - 0.06, bh, fill=LIGHT[ci])
        rect(slide, x + 0.03, by, 0.04, bh, fill=STRONG[ci])
        if markers and mark:
            badge(slide, mark, x + w - 0.11, by + min(bh, step / 72.0) / 2 - 0.11, d=0.22, color_idx=ci)
    tb = slide.shapes.add_textbox(Inches(x + 0.1), Inches(y + pad), Inches(w - 0.25), Inches(h - 2 * pad))
    tf = tb.text_frame
    body_zero(tf, 't', wrap=False)
    width = max(len(str(r[0])) for r in rows if r[0] is not None)
    badge_rows = {b[0] for b in bands if markers and b[3]}
    for i, (num, code) in enumerate(rows):
        room = w - 0.25 - (0.26 if i in badge_rows else 0.0)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para_spacing(p, lnspc_pt=step)
        lab = (str(num).rjust(width) if num is not None else ' ' * width) + '  '
        set_run(p.add_run(), lab, size, LNUM, False, CODE)
        if num is None and code.strip() == '...':
            set_run(p.add_run(), code, size, MUTE, False, CODE)
        else:
            for t, c, f in colorize(code):
                set_run(p.add_run(), t, size, c, False, f)
        wd = row_width(lab, code, size)
        if wd > room:
            print('  [폭 넘침] %.2f > %.2f in : %s' % (wd, room, code.strip()[:70]))
    return h


def row_of(rows, needle, start=0):
    for i in range(start, len(rows)):
        if rows[i][0] is not None and needle in rows[i][1]:
            return i
    raise KeyError(needle)


# 덱

class Deck:
    def __init__(self, template, cover_png):
        self.prs = Presentation(template)
        self.lay = {l.name: l for l in self.prs.slide_layouts}
        sl = list(self.prs.slides)
        g = lambda s, name: [sh._element for sh in sl[s - 1].shapes if sh.name == name][0]
        self.proto = {
            'cover': {n: g(1, n) for n in ('Text 0', 'Text 1', 'Text 2', 'Shape 3', 'Text 4', 'Text 5')},
            'index': {n: g(2, n) for n in ('Text 0', 'Text 1', 'Text 2', 'Text 3', 'Shape 4', 'Text 5', 'Shape 6')},
            'index_num': g(2, 'Slide Number Placeholder 0'),
            'div': {n: g(3, n) for n in ('Text 0', 'Text 1', 'Text 2', 'Text 3', 'Shape 4', 'Text 5')},
            'div_num': g(3, 'Slide Number Placeholder 0'),
            'content_ph': [sh._element for sh in sl[14].placeholders],
            'content_num': g(15, 'Slide Number Placeholder 0'),
            'end': {n: g(24, n) for n in ('Text 0', 'Shape 1', 'Text 2')},
        }
        self.proto = copy.deepcopy(self.proto)
        # 장 모두 지우기
        lst = self.prs.slides._sldIdLst
        for sid in list(lst):
            self.prs.part.drop_rel(sid.rId)
            lst.remove(sid)
        # 표지 · 끝 바탕 그림 바꾸기
        blob = open(cover_png, 'rb').read()
        for name in ('COVER', 'END'):
            lay = self.lay[name]
            for rel in lay.part.rels.values():
                if rel.reltype.endswith('/image'):
                    rel.target_part._blob = blob
        self.notes = read_notes(SCRIPT_MD)

    def _new(self, layout):
        s = self.prs.slides.add_slide(self.lay[layout])
        for ph in list(s.placeholders):
            ph._element.getparent().remove(ph._element)
        return s

    def _notes(self, s):
        n = len(self.prs.slides)
        if n in self.notes:
            s.notes_slide.notes_text_frame.text = self.notes[n]

    def cover(self, kicker, title, en, author):
        s = self._new('COVER')
        P = self.proto['cover']
        clone(s, P['Text 0'])
        el_set_text(clone(s, P['Text 1']), kicker)
        el_set_text(clone(s, P['Text 2']), title)
        clone(s, P['Shape 3'])
        el_set_text(clone(s, P['Text 4']), en)
        el_set_text(clone(s, P['Text 5']), author)
        self._notes(s)
        return s

    def index(self, left_title, left_sub, author, items):
        s = self._new('INDEX')
        P = self.proto['index']
        e = clone(s, P['Text 0'])
        el_set_text(e, left_title)
        body = e.find(qn('p:txBody'))
        p2 = copy.deepcopy(body.findall(qn('a:p'))[0])
        p2.findall(qn('a:r'))[0].find(qn('a:t')).text = left_sub
        body.append(p2)
        el_set_text(clone(s, P['Text 1']), author)
        y0, dy = 1.97, 0.84
        for i, (num, name, page) in enumerate(items):
            y = y0 + i * dy
            el_set_pos(el_set_text_ret(clone(s, P['Text 2']), num), y=y)
            nw = text_width(name, 15, True) + 0.1
            el_set_pos(el_set_text_ret(clone(s, P['Text 3']), name), y=y, w=nw)
            el_set_pos(clone(s, P['Shape 4']), x=6.40 + nw + 0.12, y=y + 0.33, w=12.0 - (6.40 + nw + 0.12) - 0.05)
            el_set_pos(el_set_text_ret(clone(s, P['Text 5']), str(page)), y=y)
            if i < len(items) - 1:
                el_set_pos(clone(s, P['Shape 6']), y=y + 0.72)
        clone(s, self.proto['index_num'])
        self._notes(s)
        return s

    def divider(self, title, desc, items):
        s = self._new('DIVIDER')
        P = self.proto['div']
        el_set_text(clone(s, P['Text 0']), title)
        el_set_text(clone(s, P['Text 1']), desc)
        for i, (num, name, page) in enumerate(items):
            y = 2.25 + i * 0.445
            el_set_pos(el_set_text_ret(clone(s, P['Text 2']), num), y=y)
            nw = text_width(name, 14) + 0.08
            el_set_pos(el_set_text_ret(clone(s, P['Text 3']), name), y=y, w=nw)
            el_set_pos(clone(s, P['Shape 4']), x=1.55 + nw + 0.12, y=y + 0.28, w=11.60 - (1.55 + nw + 0.12))
            el_set_pos(el_set_text_ret(clone(s, P['Text 5']), str(page)), y=y)
        clone(s, self.proto['div_num'])
        self._notes(s)
        return s

    def content(self, section, title, subtitle, path=None, tag=None):
        """tag: 이 장이 답하는 과제 항목. 머리띠 오른쪽 끝에 씀"""
        s = self._new('CONTENT')
        if text_width(subtitle, 12.5) > 11.43:
            print('  [부제 넘침] %.2f in : %s' % (text_width(subtitle, 12.5), title))
        for el, t in zip(self.proto['content_ph'], (section, title, subtitle)):
            el_set_text(clone(s, el), t)
        if tag:
            text(s, 8.45, 0.62, 4.0, 0.62, [[('과제  ', 10, '8FA3C4', False), (tag, 10.5, 'D5DEEC', True)]], anchor='m', align='r')
        if path:
            text(s, 7.30, 1.33, 5.33, 0.48, [[(path, 10, MUTE, False, CODE)]], anchor='m', align='r')
        clone(s, self.proto['content_num'])
        self._notes(s)
        return s

    def end(self, title, sub):
        s = self._new('END')
        P = self.proto['end']
        el_set_text(clone(s, P['Text 0']), title)
        clone(s, P['Shape 1'])
        el_set_text(clone(s, P['Text 2']), sub)
        self._notes(s)
        return s

    def save(self, path):
        self.prs.save(path)


def el_set_text_ret(el, t):
    el_set_text(el, t)
    return el


def table(slide, x, y, w, col_w, rows, head_size=9.5, size=9.0, row_h=0.27):
    nr, nc = len(rows), len(rows[0])
    gt = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
    tbl = gt.table
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i in range(nr):
        tbl.rows[i].height = Inches(row_h)
        for j in range(nc):
            c = tbl.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.06)
            c.margin_top = c.margin_bottom = Inches(0.02)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            c.fill.fore_color.rgb = RGBColor.from_string(NAVY if i == 0 else (WHITE if i % 2 else BOXF))
            tf = c.text_frame
            val = rows[i][j]
            spec = val if isinstance(val, tuple) else (val,)
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            set_run(p.add_run(), spec[0], head_size if i == 0 else size, WHITE if i == 0 else (spec[1] if len(spec) > 1 else NAVY),
                    i == 0 or (len(spec) > 2 and spec[2]), spec[3] if len(spec) > 3 else KO)
    return gt


def grid_table(slide, x, y, col_w, row_h, cells):
    """칸마다 바탕색과 여러 줄 글자를 따로 주는 표. cells[i][j] = (바탕색, [(글자, 크기, 색, 굵게, 글꼴), ...]), row_h 는 행마다"""
    nr, nc = len(cells), len(cells[0])
    gt = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(sum(col_w)), Inches(sum(row_h)))
    tbl = gt.table
    tbl.first_row = False
    tbl.horz_banding = False
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i in range(nr):
        tbl.rows[i].height = Inches(row_h[i])
        for j in range(nc):
            fill, paras = cells[i][j]
            c = tbl.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.05)
            c.margin_top = c.margin_bottom = Inches(0.02)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            c.fill.fore_color.rgb = RGBColor.from_string(fill)
            tf = c.text_frame
            for k, spec in enumerate(paras):
                p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                p.alignment = PP_ALIGN.CENTER
                t, size, color, bold, font = (list(spec) + [10, NAVY, False, KO][len(spec) - 1:])[:5]
                set_run(p.add_run(), t, size, color, bold, font)
    return gt


def check_fit(paras, w):
    """한 줄 글자 폭이 상자 폭을 넘는지 확인 (넘으면 알림만)"""
    for para in paras:
        wd = sum(text_width(sp[0], sp[1], len(sp) > 3 and sp[3]) for sp in para)
        if wd > w:
            print('  [글자 넘침] %.2f > %.2f in : %s' % (wd, w, ''.join(sp[0] for sp in para)[:60]))


def build(template, fig, shots, out):
    F = lambda n: os.path.join(fig, n)
    S = lambda n: os.path.join(shots, n)
    d = Deck(template, F('cover_bg.png'))
    sec1, sec2, sec3, sec4 = '01. 트랙 개시', '02. 트랙 관리', '03. 구현', '04. 결과'

    # 1 표지, 2 목차
    d.cover('레이다 SW 학습 과제 5', '트랙 개시 · 관리', 'Track Initiation and Management', 'Janghwan Kim (Harley)')
    d.index('트랙 개시 · 관리', 'TrackMgr', 'Janghwan Kim (Harley)', [('01', '트랙 개시', 3), ('02', '트랙 관리', 9), ('03', '구현', 14), ('04', '결과', 22)])

    # 3 간지 01
    d.divider('01. 트랙 개시', '트랙이 무엇인지 보고, 플롯 두 개로 트랙을 시작하는 2 Point 개시와 속도 게이트를 설명합니다.',
              [('1.1', '트랙과 트랙 개시', 4), ('1.2', '1-Point 와 2-Point 개시', 5), ('1.3', '2-Point 초기 상태변수와 오차 공분산', 6),
               ('1.4', '속도 게이팅', 7), ('1.5', '측정 공분산을 VC-NED 로 옮기기', 8)])

    # 4  1.1
    s = d.content(sec1, '1.1 트랙과 트랙 개시', '트랙은 같은 표적의 플롯을 시간 순서로 이어 위치와 속도를 추정한 것입니다. 이번 과제는 트랙의 개시와 관리입니다.',
                  tag='1 · 2 · 3 개요')
    text(s, 0.70, 2.46, 5.0, 0.30, [[('레이다 처리 흐름에서의 위치', 12, NAVY, True)]], anchor='m')
    chain = [('표적 궤적 모의', '과제 3 · 시각별 참값 위치', 0), ('측정 모의', '과제 4 · 플롯 · 클러터 생성', 0), ('트랙 개시 · 관리', '과제 5 · 이번 과제', 1),
             ('연관 · 추적 필터', '다음 과제 · 플롯을 트랙에 붙여 추정', 2)]
    for i, (t, sub, kind) in enumerate(chain):
        x = 0.70 + i * 3.08
        rect(s, x, 2.82, 2.69, 0.66, fill='EEF4FF' if kind == 1 else BOXF, line='2563EB' if kind == 1 else LINE, lw=1.25 if kind == 1 else 0.75)
        text(s, x, 2.89, 2.69, 0.30, [[(t, 12.5, '7A879F' if kind == 2 else NAVY, True)]], align='c', anchor='m')
        text(s, x, 3.18, 2.69, 0.24, [[(sub, 10.5, '8794AC' if kind != 1 else SUB, False)]], align='c', anchor='m')
        if i < 3:
            arrow_line(s, x + 2.72, 3.15, x + 3.05, 3.15)
    hline(s, 0.70, 3.70, 12.63)
    text(s, 0.70, 3.80, 6.9, 0.28, [[('스캔마다 들어오는 플롯과 트랙', 11.5, NAVY, True)]], anchor='m')
    picture(s, F('fig_track_concept.png'), 0.70, 4.12, w=6.9)
    tags = [('플롯', '레이다가 표적을 한 번 잰 결과. 거리 · 방위각 · 고각', '1F2937'), ('클러터', '표적이 아닌 것에서 나온 플롯. 같이 섞여 들어온다', '6B7280'),
            ('트랙 개시', '같은 표적으로 보이는 플롯으로 트랙을 처음 만듦', '2563EB'), ('연관', '이번 스캔 플롯을 기존 트랙에 이어 붙이는 것 (2.1)', '7C3AED'),
            ('트랙 관리', '연관 결과로 확정 · 메모리 · 삭제를 정함', '16A34A')]
    for i, (t, sub, col) in enumerate(tags):
        y = 4.12 + i * 0.52
        sp = rect(s, 7.95, y, 0.92, 0.30, fill=WHITE, line=col, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        body_zero(sp.text_frame, 'm')
        sp.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        set_run(sp.text_frame.paragraphs[0].add_run(), t, 10.5, col, True)
        text(s, 9.02, y + 0.04, 3.6, 0.3, [[(sub, 11, NAVY)]], anchor='t')

    # 5  1.2
    s = d.content(sec1, '1.2 1-Point 와 2-Point 개시', '플롯 하나로 바로 시작하면 속도를 모르고, 연속 두 스캔 (또는 드웰) 의 플롯으로 시작하면 위치 변화로 속도를 구할 수 있습니다.',
                  tag='1.1 · 1.2 · 3-1')
    picture(s, F('fig_1p2p.png'), 0.70, 2.55, w=7.3)
    callout(s, 1, 8.30, 2.60, 4.33, '1-Point 개시', ['장점: 플롯 하나로 바로 개시 (빠름)', '단점: 속도를 몰라 초기 속도 오차가 크다'])
    callout(s, 2, 8.30, 3.70, 4.33, '2-Point 개시  (이번 과제)', ['장점: 두 플롯의 위치 변화로 속도를 구해', '1-Point 보다 속도 오차가 작다', '단점: 연속 두 스캔 (또는 드웰) 이 필요해 늦다'])
    callout(s, 3, 8.30, 5.05, 4.33, '속도 게이트', ['두 플롯이 같은 표적인지 먼저 거른다 (1.4)'])

    # 6  1.3
    s = d.content(sec1, '1.3 2-Point 초기 상태변수와 오차 공분산', '상태변수는 두 플롯으로 바로 계산하고, 오차 공분산 P 에는 그 값이 얼마나 틀릴 수 있는지를 적습니다.',
                  tag='1.2 · 3-1')
    PB, PG, PO, PZ = 'DCE7FD', 'DDF3EA', 'FDECD7', 'F3F4F7'
    # ① 상태변수
    badge(s, 1, 0.70, 2.47, d=0.26)
    text(s, 1.04, 2.44, 5.8, 0.32, [[('상태변수 X', 11.5, NAVY, True), ('   두 플롯으로 바로 계산 (세 축 x 위치 · 속도 · 가속도 = 9 개)', 9.5, SUB)]], anchor='m')
    text(s, 0.70, 2.80, 6.15, 0.24, [[('예 (x 축)   앞 스캔 플롯 5,000 m,  이번 스캔 플롯 5,030 m,  스캔 간격 T = 0.1 s', 9.5, SUB)]], anchor='m')
    C = lambda t, col=NAVY, b=False: (t, col, b, EN)
    st = [('성분', '식', '예 (x 축)'),
          ('위치  x', C('z_k'), C('5,030 m')),
          ('속도  x\'', C('( z_k − z_k-1 ) / T'), C('( 5,030 − 5,000 ) / 0.1 = 300 m/s')),
          ('가속도  x\'\'', ('0  (과제 조건)', NAVY), C('0'))]
    table(s, 0.70, 3.08, 6.15, [1.25, 2.05, 2.85], st, head_size=9.5, size=10, row_h=0.29)
    # ② 한 축의 P (3 x 3)
    badge(s, 2, 0.70, 4.47, d=0.26)
    text(s, 1.04, 4.44, 5.8, 0.32, [[('오차 공분산 P', 11.5, NAVY, True), ('   한 축의 3 x 3  (예: 플롯 오차 ±10 m, T = 0.1 s)', 9.5, SUB)]], anchor='m')
    H = lambda t: ('2B3A55', [(t, 9.5, WHITE, True)])
    H2 = lambda t, sub: ('2B3A55', [(t, 10, WHITE, True), (sub, 8.5, 'C9D3E6', False)])
    F2 = lambda fill, f, v: (fill, [(f, 10.5, NAVY, True, EN), (v, 9.5, SUB, False, EN)])
    Z = (PZ, [('0', 10, MUTE, False, EN)])
    pc = [[H(''), H('위치'), H('속도'), H('가속도')],
          [H2('위치', '= 이번 플롯'), F2(PB, 'R_k', '100  (±10 m)'), F2(PG, 'R_k / T', '1,000'), Z],
          [H2('속도', '= 두 플롯 차 / T'), F2(PG, 'R_k / T', '1,000'), F2(PO, '( R_k + R_k-1 ) / T²', '20,000  (±141 m/s)'), Z],
          [H2('가속도', '= 0'), Z, Z, Z]]
    grid_table(s, 0.70, 4.82, [1.40, 1.45, 2.20, 1.10], [0.27, 0.56, 0.56, 0.44], pc)
    # ③ 9 x 9 전체
    badge(s, 3, 7.15, 2.47, d=0.26)
    text(s, 7.49, 2.44, 5.1, 0.32, [[('9 x 9 전체 P', 11.5, NAVY, True), ('   ② 의 블록이 축마다 대각선에 하나씩', 9.5, SUB)]], anchor='m')
    picture(s, F('fig_pmatrix.png'), 7.15, 2.80, w=5.45)

    # 7  1.4  (처음 판 구성으로 되돌림. 아래 설명 상자 대신 그림을 키움)
    s = d.content(sec1, '1.4 속도 게이팅', '두 플롯이 같은 표적이면 한 스캔 동안 움직인 거리가 V_max · T 보다 작아야 합니다. 이 조건을 만족하는 (게이트를 통과한) 쌍으로만 개시합니다.',
                  tag='속도 게이팅 · 3-1')
    rect(s, 0.70, 2.45, 11.93, 0.62, fill=BOXF, line=LINE)
    text(s, 0.95, 2.45, 5.6, 0.62, [[('| z_(k,i)  −  z_(k-1,j) |  /  ΔT   <   V_max', 15, NAVY, True, EN)]], anchor='m')
    text(s, 6.40, 2.45, 6.1, 0.62, [[('V_max = 800 m/s,  ΔT = 0.1 s  ->  앞 플롯에서 반경 80 m 안', 11.5, '2563EB', True)],
                                    [('i: 이번 스캔 플롯,  j: 앞 스캔 플롯.  거리는 VC-NED 직교좌표로 잰다', 10, SUB)]], anchor='m')
    picture(s, F('fig_gate.png'), 0.70, 3.25, w=11.93)

    # 8  1.5
    s = d.content(sec1, '1.5 측정 공분산을 VC-NED 로 옮기기', '트랙은 VC-NED (플랫폼 중심, 북 · 동 · 아래로 고정된 축) 에서 만듭니다. 거리 · 방위각 · 고각의 오차를 VC-NED 위치 오차 R 로 바꿉니다.',
                  tag='1.2 · 3-1')
    steps = [('측정 표준편차 (과제 4, SNR 20)', 'σ_r = 4.74 m,   σ_az = 0.478°,   σ_el = 0.494°', 1),
             ('안테나 직교좌표로', 'R_ant = J · diag( σ_r², σ_az², σ_el² ) · Jᵀ', 2),
             ('VC-NED 로', 'R = M · R_ant · Mᵀ', 3)]
    for i, (t, f, n) in enumerate(steps):
        y = 2.50 + i * 0.95
        badge(s, n, 0.74, y + 0.05)
        text(s, 1.16, y, 6.2, 0.30, [[(t, 11.5, NAVY, True)]], anchor='m')
        rect(s, 1.16, y + 0.33, 6.2, 0.44, fill=BOXF, line=LINE)
        text(s, 1.30, y + 0.33, 6.0, 0.44, [[(f, 12.5, NAVY, False, EN)]], anchor='m')
    text(s, 1.16, 5.35, 6.2, 0.75, [[('J', 10.5, NAVY, True), (': 구좌표 (r, az, el) -> 안테나 직교 (x, y, z) 변환의 미분', 10.5, SUB)],
                                    [('M', 10.5, NAVY, True), (': 안테나 -> 동체 -> VC-NED 회전 (플랫폼 Roll · Pitch · Yaw)', 10.5, SUB)]], lnspc=115)
    picture(s, F('fig_rcov.png'), 7.55, 2.42, w=5.08)

    # 9 간지 02
    d.divider('02. 트랙 관리', '트랙의 상태와 연관을 보고, 만든 트랙을 확정할지 지울지 정하는 M/N 방식과 트랙 스코어 방식을 설명합니다.',
              [('2.1', '트랙 상태와 연관', 10), ('2.2', 'M/N 방식', 11), ('2.3', '트랙 스코어 방식', 12), ('2.4', 'Assigned_Flag 에 따른 스코어 변화', 13)])

    # 10  2.1
    s = d.content(sec2, '2.1 트랙 상태와 연관', '트랙은 임시로 시작해 확정, 메모리, 삭제 상태를 오갑니다. 상태를 바꾸는 근거는 스캔마다의 연관 결과 (Assigned_Flag) 입니다.',
                  tag='2 · 3-2 · 3-3')
    label(s, 0.70, 2.44, 6.2, '연관이란')
    picture(s, F('fig_assoc.png'), 0.70, 2.74, w=6.2)
    rect(s, 0.70, 5.45, 6.2, 1.08, fill=BOXF, line=LINE)
    text(s, 0.88, 5.45, 5.9, 1.08, anchor='m', paras=[[('연관  ', 10.5, NAVY, True), ('이번 스캔 플롯 가운데 하나를 기존 트랙에 이어 붙이는 것', 10.5, NAVY)],
                                    [('Assigned_Flag  ', 10.5, NAVY, True), ('트랙마다 스캔마다 남기는 연관 결과. 1 성공, 0 실패', 10.5, NAVY)],
                                    [('이번 과제  ', 10.5, '2563EB', True), ('연관은 다음 과제라 매 스캔 0 으로 둔다 (과제 조건)', 10.5, NAVY)]], lnspc=135)
    label(s, 7.25, 2.44, 5.38, '트랙 상태 네 가지', 'TrackState')
    states = [('1  임시', TENT, 'FFF1E6', '막 만든 트랙. 표적인지 아직 확인하는 중', '클러터 두 개가 우연히 이어진 트랙일 수도 있다'),
              ('2  확정', '16A34A', 'E8F7EE', '연관이 이어져 표적으로 인정한 트랙', '보통 운용자에게 알리는 것은 확정 트랙'),
              ('3  메모리', '9333EA', 'F5EBFF', '확정 트랙이 이번 스캔에 플롯을 놓친 상태', '바로 지우지 않고 다시 잡히기를 기다린다'),
              ('4  삭제', '6B7280', 'F1F2F4', '더 쫓지 않기로 한 트랙', '다음 스캔을 시작할 때 트랙 표에서 지운다')]
    for i, (name, col, fc, l1, l2) in enumerate(states):
        y = 2.78 + i * 1.0
        rect(s, 7.25, y, 5.38, 0.86, fill=fc, line=col, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        text(s, 7.42, y, 1.15, 0.86, [[(name, 12, col, True)]], anchor='m')
        text(s, 8.55, y, 4.0, 0.86, [[(l1, 10.5, NAVY, True)], [(l2, 10, SUB)]], anchor='m', lnspc=115)

    # 11  2.2
    s = d.content(sec2, '2.2 M/N 방식', '최근 N 번의 연관 결과 가운데 성공 횟수 M 으로 확정을, 연속 실패 횟수로 삭제를 정합니다.', tag='2.1')
    picture(s, F('fig_mn.png'), 0.70, 2.45, w=6.85)
    callout(s, 1, 7.75, 2.60, 4.88, '임시 트랙', ['N 번 중 M 번 트랙 연관에 성공하면 확정', 'M 번 연속 결합에 실패하면 삭제'])
    callout(s, 2, 7.75, 3.65, 4.88, '확정 트랙', ['트랙 연관에 실패하면 메모리 트랙으로 바꿈', '정한 시간 이상 메모리 상태면 삭제'])
    callout(s, 3, 7.75, 4.70, 4.88, '특징', ['성공 · 실패 횟수만 세면 되어 단순하다', '최근 N 번의 성공 · 실패 기록을 들고 있어야 한다'])

    # 12  2.3
    s = d.content(sec2, '2.3 트랙 스코어 방식', '트랙마다 TrackScore 와 MemoryScore 를 두고, 점수가 문턱을 넘으면 상태 (TrackState) 를 바꿉니다.', tag='2.2 · 3-2')
    picture(s, F('fig_state.png'), 0.70, 2.45, w=7.0)
    callout(s, 1, 7.95, 2.55, 4.68, 'TrackState (과제 값)', ['1 임시,  2 확정,  3 메모리,  4 삭제'])
    callout(s, 2, 7.95, 3.30, 4.68, '문턱 (과제 값)', ['임시 트랙 확정 5,  임시 트랙 삭제 0', '메모리 트랙 삭제 5'])
    callout(s, 3, 7.95, 4.30, 4.68, '메모리 트랙이 다시 연관되면', ['바로 확정 트랙으로, MemoryScore = 0'])
    callout(s, 4, 7.95, 5.05, 4.68, '과제에 없어 정한 값', ['개시 직후 TrackScore 2 (두 플롯으로 만든 트랙)', '한 번에 올리고 내리는 양 1'])

    # 13  2.4
    s = d.content(sec2, '2.4 Assigned_Flag 에 따른 스코어 변화', '왼쪽은 연관 결과를 넣은 예, 오른쪽은 과제 조건 (매 스캔 0) 입니다. 맨 아래 계산 줄이 스캔마다 점수가 바뀐 과정입니다.',
                  tag='3-2 · 3-3')
    picture(s, F('fig_score.png'), 0.70, 2.42, w=11.93)

    # 14 간지 03
    d.divider('03. 구현', '쓰레드 · UDP 구조와 플롯 메시지를 보고, 송신 · 수신과 트랙 개시 · 관리 코드를 차례로 봅니다.',
              [('3.1', '쓰레드와 UDP 구조', 15), ('3.2', '플롯 메시지 STRUCT_PlotMsg', 16), ('3.3', '송신: 플롯 메시지 만들기 (f_Lnk_BuildMsg)', 17),
               ('3.4', '수신: 좌표 변환과 측정 공분산 (f_Lnk_ProcessMsg)', 18), ('3.5', '2 Point 개시: 쌍 고르기 (f_Trk_ProcessScan)', 19),
               ('3.6', '2 Point 초기화 (f_Trk_TwoPointInit)', 20), ('3.7', '트랙 스코어 로직 (f_Trk_UpdateState)', 21)])

    # 15  3.1
    s = d.content(sec3, '3.1 쓰레드와 UDP 구조', '과제 4 의 측정 모의 결과를 신호처리가 보낸 플롯으로 보고, 송신 쓰레드가 스캔마다 UDP 로 보내면 수신 쓰레드가 받아 처리합니다.',
                  tag='3-3')
    picture(s, F('fig_arch.png'), 0.70, 2.42, w=11.93)
    text(s, 0.70, 5.27, 11.93, 0.30, [[('왜 TCP 가 아니라 UDP 인가', 11.5, NAVY, True), ('   과제는 UDP 또는 TCP/IP 중 하나를 고르게 함', 9.5, SUB)]], anchor='m')
    U = lambda t: (t, '1D4ED8', False)
    ut = [('', 'UDP  (이번 과제)', 'TCP'),
          ('메시지 단위', U('보낸 메시지 1 개 = 받는 메시지 1 개 (스캔 1 개)'), '바이트가 이어서 와서, 메시지를 직접 잘라야 함'),
          ('지연', U('연결 · 재전송 없이 바로 전달'), '하나가 늦으면 뒤 스캔도 함께 기다림'),
          ('손실', U('생길 수 있음 -> 스캔 번호로 확인 (1,160 / 1,160, 빠짐 0)'), '없음 (잃으면 다시 보냄)')]
    table(s, 0.70, 5.60, 11.93, [1.55, 5.40, 4.98], ut, head_size=10, size=10.5, row_h=0.29)

    # 16  3.2
    s = d.content(sec3, '3.2 플롯 메시지 STRUCT_PlotMsg', '참고 소스 코드의 Idd.h 는 STRUCT_PlotInfo 의 Y_Body 에서 끊겨 있어 나머지와 메시지 구조체를 이어 썼습니다.',
                  'TargetSimCore/include/Idd.h', tag='3-3')
    rows = excerpt('TargetSimCore/include/Idd.h', [('typedef struct STRUCT_PlotInfo', '}STRUCT_PlotInfo;')])
    rows = annotate(rows, [('Id;', '플롯 ID'), ('EmittedDwellNum;', '드웰 번호'), ('EmittedScanNum;', '스캔 번호'), ('EmittedBeamNum;', '빔 번호'),
                           ('EmittedTimeStamp_s;', '탐지 시각 (초)'), ('EmittedTimeStamp_ns;', '탐지 시각 (나노초)'), ('BeamType;', '빔 종류'),
                           ('Range_Ant;', '안테나 기준 거리 [m]'), ('Azimuth_Ant;', '안테나 기준 방위각 [rad]'), ('Elevation_Ant;', '안테나 기준 고각 [rad]'),
                           ('Range_VcNed;', 'VC-NED 기준 거리'), ('Azimuth_VcNed;', 'VC-NED 기준 방위각'), ('Elevation_VcNed;', 'VC-NED 기준 고각'),
                           ('X_Ant;', '안테나 직교 x [m]'), ('Y_Ant;', '안테나 직교 y'), ('Z_Ant;', '안테나 직교 z'),
                           ('X_Body;', '동체 직교 x [m]'), ('Y_Body;', '동체 직교 y'), ('Z_Body;', '동체 직교 z'),
                           ('X_VcNed;', 'VC-NED 북 [m]'), ('Y_VcNed;', 'VC-NED 동'), ('Z_VcNed;', 'VC-NED 아래')])
    i_r = row_of(rows, 'Range_Ant')
    i_z = row_of(rows, 'Z_Body')
    label(s, 0.70, 2.44, 5.7, '플롯 하나', 'STRUCT_PlotInfo')
    code_box(s, 0.70, 2.72, 5.75, rows, bands=[(i_r, 3, 0, 1), (i_z, 4, 1, 2)], size=8.3, step=11)
    rows2 = excerpt('TargetSimCore/include/Idd.h', [('typedef struct STRUCT_PlotMsg', '}STRUCT_PlotMsg;')])
    rows2 = annotate(rows2, [('MsgId;', '메시지 종류 (플롯 / 끝)'), ('ScanNum;', '스캔 번호'), ('TimeStamp_ns;', '스캔 시각 (나노초)'),
                             ('st_Platform;', '플랫폼 위치 · 자세'), ('st_Plot[IDD_MAX_PLOT];', '플롯 배열 (최대 64 개)')])
    i_h = row_of(rows2, 'MsgId')
    i_s = row_of(rows2, 'Sigma_Range')
    i_p = row_of(rows2, 'st_Platform;')
    label(s, 6.75, 2.44, 5.9, '스캔 하나의 메시지', 'STRUCT_PlotMsg')
    h2 = code_box(s, 6.75, 2.72, 5.88, rows2, bands=[(i_h, 6, 3, 3), (i_s, 3, 2, 4), (i_p, 2, 4, 5)], size=8.3, step=11)
    y = 2.72 + h2 + 0.14
    notes = [(1, '송신은 안테나 구좌표 세 칸만 채운다'), (2, 'Y_Body 다음부터 이어 쓴 칸 (받는 쪽이 계산해 채움)'), (3, '머리: 종류 · 크기 · 스캔 번호 · 시각 · 플롯 수'),
             (4, '측정 표준편차: 받는 쪽 오차 공분산 계산에 씀'), (5, '플랫폼 위치 · 자세, 플롯 배열 (최대 64 개)')]
    for k, (n, t) in enumerate(notes):
        badge(s, n, 6.80, y + k * 0.29, d=0.22)
        text(s, 7.12, y + k * 0.29 - 0.01, 5.5, 0.25, [[(t, 10, NAVY)]], anchor='m')

    # 17  3.3
    s = d.content(sec3, '3.3 송신: 플롯 메시지 만들기 (f_Lnk_BuildMsg)', '표적 플롯과 클러터를 섞어 방위각 순서로 담고, 플롯마다 ID 를 매겨 안테나 구좌표만 채웁니다.',
                  'TargetSimCore/src/PlotLink.c', tag='3-3')
    rows = excerpt('TargetSimCore/src/PlotLink.c', [
        ('static UINT32 f_Lnk_BuildMsg', 1),
        ('for (nIndex = 0; (nIndex < st_Scan->nTargetNum)', 4),
        ('for (nIndex = 0; (nIndex < st_Scan->nClutterNum)', 4),
        ('// 방위각 오름차순', 1),
        ('st_Msg->MsgId', 'st_Msg->PlotNum'),
        ('for (nIndex = 0; nIndex < nPlotNum; nIndex++)', 4),
        ('st_Info->Range_Ant', 'st_Info->Elevation_Ant'),
        ('// 플롯 ID 는 10000', '*pt_PlotId = ')])
    rows = annotate(rows, [('f_Lnk_BuildMsg(', '스캔 하나의 메시지 만들기'), ('nTargetNum)', '표적 플롯부터'),
                           ('st_Scan->st_Plot[nIndex];', '거리 · 방위각 · 고각'), ('= nIndex + 1;', '출처: 표적 번호 (송신 쪽 기록용)'),
                           ('nClutterNum)', '이어서 클러터를 같은 배열에'), ('st_Scan->st_Clutter[nIndex];', '클러터 플롯'), ('= 0;', '출처: 0 은 클러터'),
                           ('st_Msg->MsgId', '메시지 종류: 플롯'), ('st_Msg->MsgSize', '보낼 바이트 수'), ('st_Msg->ScanNum', '스캔 번호'),
                           ('st_Msg->TimeStamp_s', '스캔 시각 (초)'), ('st_Msg->TimeStamp_ns', '스캔 시각 (나노초)'), ('st_Msg->PlotNum', '표적 + 클러터 수'),
                           ('nIndex < nPlotNum;', '정렬된 순서대로 플롯마다'), ('&st_Msg->st_Plot[nIndex];', '메시지의 플롯 칸'), ('st_Info->Id', '플롯 ID'),
                           ('st_Info->Range_Ant', '안테나 기준 거리'), ('st_Info->Azimuth_Ant', '방위각'), ('st_Info->Elevation_Ant', '고각'),
                           ('*pt_PlotId = ', '다음 플롯 ID')])
    b1 = row_of(rows, 'nTargetNum)')
    b2 = row_of(rows, '방위각 오름차순')
    b3 = row_of(rows, 'st_Msg->MsgId')
    b4 = row_of(rows, 'st_Info->Id')
    b5 = row_of(rows, 'Range_Ant')
    b6 = row_of(rows, '// 플롯 ID 는')
    b7 = row_of(rows, 'nClutterNum)')
    code_box(s, 0.70, 2.45, 7.85, rows, bands=[(b1, 4, 0, 1), (b7, 4, 0, None), (b2, 1, 1, 2), (b3, 6, 2, 3), (b4, 1, 3, 4), (b5, 3, 3, None), (b6, 2, 4, 5)],
             size=7.6, step=9)
    callout(s, 1, 8.85, 2.50, 3.78, '표적 플롯 + 클러터를 한 배열에', ['표적 플롯 뒤에 클러터를 이어 붙임', '출처 (표적 k / 클러터) 는 송신 쪽에만 기록'], size=10.5)
    callout(s, 2, 8.85, 3.45, 3.78, '방위각 순서로 정렬', ['빔이 지나가는 순서처럼', '배열 위치로 표적을 알 수 없게'], size=10.5)
    callout(s, 3, 8.85, 4.40, 3.78, '메시지 머리', ['스캔 번호, 시각, 플롯 수'], size=10.5)
    callout(s, 4, 8.85, 5.05, 3.78, '플롯 ID 와 안테나 구좌표', ['거리 · 방위각 · 고각만 채움'], size=10.5)
    callout(s, 5, 8.85, 5.70, 3.78, '플롯 ID 10000 ~ 1000000', ['1000000 다음은 10000 으로'], size=10.5)

    # 18  3.4
    s = d.content(sec3, '3.4 수신: 좌표 변환과 측정 공분산 (f_Lnk_ProcessMsg)', '받은 플롯을 VC-NED 로 바꾸고 측정 공분산을 구한 뒤, 트랙 관리와 개시를 한 번에 돌립니다.',
                  'TargetSimCore/src/PlotLink.c', tag='3-3')
    rows = excerpt('TargetSimCore/src/PlotLink.c', [
        ('static VOID f_Lnk_ProcessMsg', 1),
        ('const FLOAT64\t\t\ttime = ', 1),
        ('st_Sigma.r\t= st_Msg->Sigma_Range;', 'st_Att.Yaw'),
        ('for (nPlot = 0; nPlot < nPlotNum; nPlot++)', 4),
        ('st_Plot[nPlot].pos[0]', 'st_Plot[nPlot].pos[2]'),
        ('f_Trk_MeasCov(st_Plot[nPlot].cov', 1),
        ('f_Trk_MoveOrigin(&s_Link.st_Table', 1),
        ('f_Trk_ProcessScan(&s_Link.st_Table', 1)])
    rows = annotate(rows, [('f_Lnk_ProcessMsg(', '받은 메시지 하나 처리'), ('time = ', '스캔 시각 [s]'), ('st_Sigma.r', '측정 표준편차: 거리'),
                           ('st_Sigma.az', '방위각'), ('st_Sigma.el', '고각'), ('st_Att.Roll', '플랫폼 자세'),
                           ('nPlot < nPlotNum;', '플롯마다'), ('&st_Msg->st_Plot[nPlot];', '받은 플롯'), ('f_Lnk_FillPlotCoord(', '좌표 변환 (아래 세 줄)'),
                           ('pos[0]', '북'), ('pos[1]', '동'), ('pos[2]', '아래'), ('f_Trk_MeasCov(', '측정 공분산 R (1.5)'),
                           ('f_Trk_MoveOrigin(', '앞 스캔 값을 이번 원점 기준으로'), ('f_Trk_ProcessScan(', '트랙 관리 + 개시')])
    i1 = row_of(rows, 'st_Sigma.r')
    i2 = row_of(rows, 'f_Lnk_FillPlotCoord(')
    i3 = row_of(rows, 'pos[0]')
    i4 = row_of(rows, 'f_Trk_MeasCov(')
    i5 = row_of(rows, 'f_Trk_MoveOrigin(')
    i6 = row_of(rows, 'f_Trk_ProcessScan(')
    hm = code_box(s, 0.70, 2.45, 7.85, rows, bands=[(i1, 6, 0, 1), (i2, 1, 1, 2), (i3, 3, 2, 3), (i4, 1, 2, None), (i5, 1, 4, 4), (i6, 1, 3, 5)],
                  size=7.9, step=9)
    rowsc = excerpt('TargetSimCore/src/PlotLink.c', [('st_Ant\t= f_Trans_Ant_Sph_To_XYZ(st_Info->Range_Ant', 'st_Ned\t= f_Trans_Body_To_Ned(st_Body.x')], dedent=2)
    rowsc = annotate(rowsc, [('st_Ant', '구좌표 -> 안테나 직교'), ('st_Body', '안테나 -> 동체'), ('st_Ned', '동체 -> VC-NED (플랫폼 자세)')])
    label(s, 0.70, 2.45 + hm + 0.08, 7.85, '2  f_Lnk_FillPlotCoord 의 변환 세 줄', 'PlotLink.c')
    code_box(s, 0.70, 2.45 + hm + 0.36, 11.93, rowsc, size=7.4, step=9)
    callout(s, 1, 8.85, 2.50, 3.78, '머리에서 꺼냄', ['시각, 측정 표준편차, 플랫폼 자세'], size=10.5)
    callout(s, 2, 8.85, 3.15, 3.78, '안테나 -> 동체 -> VC-NED', ['아래 세 줄. 비어 있던 좌표 칸도 채움'], size=10.5)
    callout(s, 3, 8.85, 3.80, 3.78, '트랙 개시용 플롯', ['위치 VC-NED, 공분산 f_Trk_MeasCov (1.5)'], size=10.5)
    callout(s, 4, 8.85, 4.45, 3.78, '원점 옮기기', ['플랫폼이 움직이면 앞 스캔 플롯 · 트랙을', '이번 스캔 VC-NED 로 (이번 시나리오는 0)'], size=10.5)
    callout(s, 5, 8.85, 5.35, 3.78, '트랙 관리 + 2 Point 개시', ['결과는 스캔 단위로 쌓아 둠'], size=10.5)

    # 19  3.5
    s = d.content(sec3, '3.5 2 Point 개시: 쌍 고르기 (f_Trk_ProcessScan)', '모든 트랙을 Assigned_Flag 0 으로 판정한 뒤, 게이트를 통과한 쌍 가운데 가장 느린 쌍부터 하나씩 골라 트랙을 만듭니다.',
                  'TargetSimCore/src/TrackMgr.c', tag='3-1 · 3-3')
    rows = excerpt('TargetSimCore/src/TrackMgr.c', [
        ('st_Track->Assigned_Flag\t= 0;', 2),
        ('while (isSearching != 0)', 3),
        ('for (nPrev = 0; nPrev < st_Table->nPrevPlotNum; nPrev++)', 'bestCur\t\t= nCur;'),
        ('st_Table->st_PrevPlot[bestPrev].isUsed\t= 1;', 2),
        ('newId = (st_Table->nTrackNum < TRK_MAX_TRACK)', 1),
        ('f_Trk_TwoPointInit(st_Track, &st_Table->st_PrevPlot[bestPrev]', 1),
        ('st_Track->TrackId\t\t= newId;', 'st_Track->TrackScore\t= TRK_SCORE_INIT;'),
        ('st_Track->PlotId[0]', 'st_Track->PlotId[1]')], dedent=2)
    rows = annotate(rows, [('Assigned_Flag', '연관 기능이 없어 늘 실패 (과제 조건)'), ('f_Trk_UpdateState(', '스코어로 상태 판정 (3.7)'),
                           ('while (isSearching', '고를 쌍이 없을 때까지 반복'), ('TRK_VMAX;', 'V_max 에서 시작 = 속도 게이트'),
                           ('for (nPrev', '앞 스캔 플롯'), ('for (nCur', '이번 스캔 플롯 (모든 쌍)'), ('isUsed == 0) && (', '아직 안 쓴 플롯끼리만'),
                           ('f_Trk_PairSpeed(', '|z_k - z_k-1| / T'), ('if (speed < bestSpeed)', '지금까지 가장 느린 쌍이면'),
                           ('bestSpeed  = speed;', '그 쌍을 기억'), ('[bestPrev].isUsed', '고른 두 플롯은 다시 쓰지 않음'),
                           ('f_Trk_NewId(', '트랙 ID (100 ~ 10000)'), ('f_Trk_TwoPointInit(', '초기 상태변수 · 공분산 (3.6)'),
                           ('TrackState', '임시 트랙 (1)'), ('TrackScore', '개시 스코어 2'), ('PlotId[0]', '개시에 쓴 플롯 ID')])
    a1 = row_of(rows, 'Assigned_Flag')
    a2 = row_of(rows, 'TRK_VMAX;')
    a3 = row_of(rows, 'for (nPrev')
    a4 = row_of(rows, '[bestPrev].isUsed')
    a5 = row_of(rows, 'f_Trk_NewId')
    a6 = row_of(rows, 'f_Trk_TwoPointInit')
    a7 = row_of(rows, 'TrackId')
    code_box(s, 0.70, 2.45, 7.85, rows, bands=[(a1, 2, 0, 1), (a2, 1, 1, 2), (a3, 13, 1, None), (a4, 2, 2, 3), (a5, 1, 3, 4), (a6, 1, 4, 5),
                                               (a7, 3, 4, None)], size=7.6, step=9)
    callout(s, 1, 8.85, 2.50, 3.78, '트랙 관리 먼저', ['연관이 없어 모든 트랙 Assigned_Flag = 0', 'f_Trk_UpdateState 로 판정 (3.7)'], size=10.5)
    callout(s, 2, 8.85, 3.45, 3.78, '속도 게이트', ['bestSpeed 를 V_max 로 시작', '-> 800 m/s 이상인 쌍은 고르지 않음'], size=10.5)
    callout(s, 3, 8.85, 4.40, 3.78, '쓴 플롯 표시', ['한 플롯은 한 쌍에만'], size=10.5)
    callout(s, 4, 8.85, 5.05, 3.78, '트랙 ID', ['100 부터 차례로, 10000 다음은 100', '쓰는 중인 ID 는 건너뜀'], size=10.5)
    callout(s, 5, 8.85, 5.95, 3.78, '초기화 (3.6)', ['임시 (1), TrackScore 2, 개시 플롯 ID 2 개'], size=10.5)

    # 20  3.6
    s = d.content(sec3, '3.6 2 Point 초기화 (f_Trk_TwoPointInit)', '9 x 1 상태변수와 9 x 9 오차 공분산을 0 으로 만든 뒤, 축마다 위치 · 속도 칸만 채웁니다. 가속도 칸은 0 으로 남습니다.',
                  'TargetSimCore/src/TrackMgr.c', tag='3-1')
    rows = excerpt('TargetSimCore/src/TrackMgr.c', [('VOID f_Trk_TwoPointInit(', 'st_P->e[(3 * nAxis) + 1][(3 * nAxisB) + 1]')])
    rows.append((rows[-1][0] + 1, '    }'))
    rows.append((rows[-1][0] + 1, '  }'))
    rows.append((rows[-1][0] + 1, '}'))
    rows = annotate(rows, [('f_Trk_TwoPointInit(', '앞 플롯, 이번 플롯'), ('dt = ', '스캔 간격 T [s]'), ('*st_X = ', '상태변수 X (9 x 1)'),
                           ('*st_P = ', '오차 공분산 P (9 x 9)'), ('nAxis;', '축 a'), ('nAxisB;', '축 b'),
                           ('Matrix_initialize(st_X', '9 x 1, 모두 0'), ('Matrix_initialize(st_P', '9 x 9, 모두 0'),
                           ('for (nAxis = 0', '축 a: 0 북, 1 동, 2 아래'), ('st_X->e[(3 * nAxis)][0]', '위치 = 이번 플롯'),
                           ('st_X->e[(3 * nAxis) + 1][0]', '속도 = 위치 차 / T'), ('st_X->e[(3 * nAxis) + 2][0]', '가속도 = 0'),
                           ('for (nAxisB = 0', '축 b'), ('st_P->e[(3 * nAxis)][(3 * nAxisB)]', '위치 · 위치 = R_k'),
                           ('st_P->e[(3 * nAxis)][(3 * nAxisB) + 1]', '위치 · 속도 = R_k / T'), ('st_P->e[(3 * nAxis) + 1][(3 * nAxisB)]', '속도 · 위치 = R_k / T'),
                           ('st_P->e[(3 * nAxis) + 1][(3 * nAxisB) + 1]', '속도 · 속도 = (R_k + R_k-1) / T²', 'above')])
    c1 = row_of(rows, 'Matrix_initialize(st_X')
    c2 = row_of(rows, 'st_X->e[(3 * nAxis)][0]')
    c3 = row_of(rows, 'st_P->e[(3 * nAxis)][(3 * nAxisB)]')
    code_box(s, 0.70, 2.45, 7.85, rows, bands=[(c1, 2, 0, 1), (c2, 3, 1, 2), (c3, 5, 2, 3)], size=7.6, step=10)
    callout(s, 1, 8.85, 2.50, 3.78, '0 으로 채운 행렬', ['참고 소스 코드의 Matrix_initialize', '가속도 칸은 손대지 않아 0'], size=10.5)
    callout(s, 2, 8.85, 3.45, 3.78, '상태변수 (축 a 마다)', ['위치 = 이번 플롯  (3a 번째 칸)', '속도 = 위치 차 / 시간 차  (3a+1)', '가속도 = 0  (3a+2)'], size=10.5)
    callout(s, 3, 8.85, 4.60, 3.78, '오차 공분산 (1.3 의 3 x 3 표)', ['위치 · 위치 = R_k', '위치 · 속도 = R_k / T', '속도 · 속도 = ( R_k + R_k-1 ) / T²'],
            size=10.5)

    # 21  3.7
    s = d.content(sec3, '3.7 트랙 스코어 로직 (f_Trk_UpdateState)', '상태마다 case 하나씩입니다. 임시는 TrackScore 로, 확정 · 메모리는 MemoryScore 로 다음 상태를 정합니다.',
                  'TargetSimCore/src/TrackMgr.c', tag='3-2')
    rows = excerpt('TargetSimCore/src/TrackMgr.c', [('VOID f_Trk_UpdateState(', 'case TRK_STATE_CONFIRMED:')])
    rows = rows[:-1]
    rows = annotate(rows, [('f_Trk_UpdateState(', '스캔마다 트랙 하나씩'), ('isAssigned = ', '연관 결과 1 / 0'),
                           ('switch (', '지금 상태에 따라'), ('case TRK_STATE_TENTATIVE:', '임시 트랙'),
                           ('TrackScore + ', '성공 +1, 실패 -1', 'above'), ('>= TRK_SCORE_CONFIRM', '5 이상이면'), ('= TRK_STATE_CONFIRMED;', '확정'),
                           ('<= TRK_SCORE_DELETE', '0 이하면'), ('= TRK_STATE_DELETED;', '삭제')])
    d1 = row_of(rows, 'TrackScore + ')
    d2 = row_of(rows, '>= TRK_SCORE_CONFIRM')
    d3 = row_of(rows, '<= TRK_SCORE_DELETE')
    label(s, 0.70, 2.44, 6.0, '임시 트랙')
    code_box(s, 0.70, 2.72, 6.05, rows, bands=[(d1 - 1, 2, 0, 1), (d2, 4, 2, 2), (d3, 4, 4, None)], size=7.6, step=9)
    rows2 = excerpt('TargetSimCore/src/TrackMgr.c', [('case TRK_STATE_CONFIRMED:', 'default:')])
    rows2 = rows2[:-1]
    rows2 = annotate(rows2, [('case TRK_STATE_CONFIRMED:', '확정 트랙'), ('isAssigned == 0', '연관 실패면'), ('= TRK_STATE_MEMORY;', '메모리로'),
                             ('MemoryScore + 1;', 'MemoryScore +1'), ('case TRK_STATE_MEMORY:', '메모리 트랙'), ('isAssigned != 0', '연관 성공이면'),
                             ('= TRK_STATE_CONFIRMED;', '바로 확정'), ('MemoryScore  = 0;', 'MemoryScore 0 으로'), ('else', '연관 실패면'),
                             ('MemoryScore + 1;', 'MemoryScore +1'), ('>= TRK_MEMORY_DELETE', '5 이상이면'), ('= TRK_STATE_DELETED;', '삭제')])
    e1 = row_of(rows2, 'TRK_STATE_MEMORY;')
    e3 = row_of(rows2, '= 0;')
    e4 = row_of(rows2, '>= TRK_MEMORY_DELETE')
    label(s, 6.95, 2.44, 5.7, '확정 트랙, 메모리 트랙')
    code_box(s, 6.95, 2.72, 5.68, rows2, bands=[(e1, 2, 4, 3), (e3 - 1, 2, 2, 4), (e4, 4, 3, 5)], size=7.6, step=9)
    y = 6.42
    marks = [(1, '성공 +1, 실패 -1'), (2, '5 이상 확정, 0 이하 삭제'), (3, '확정 -> 메모리, MemoryScore +1'), (4, '메모리 -> 바로 확정, MemoryScore 0'), (5, 'MemoryScore 5 이상 삭제')]
    for k, (n, t) in enumerate(marks):
        x = 0.70 + k * 2.40
        badge(s, n, x, y, d=0.22)
        text(s, x + 0.28, y - 0.03, 2.1, 0.5, [[(t, 9.5, NAVY)]], anchor='t')

    # 22 간지 04
    d.divider('04. 결과', '2D · 3D 전시, 트랙 수명과 스코어 확인, 거리에 따른 개시 결과를 봅니다.',
              [('4.1', 'True · Plot · Track 2D · 3D 전시', 23), ('4.2', '트랙 수명과 스코어 확인', 24), ('4.3', '거리에 따른 개시 결과', 25)])

    # 23  4.1
    s = d.content(sec4, '4.1 True · Plot · Track 2D · 3D 전시', '1,160 스캔에서 트랙 355 개를 개시했고 (클러터끼리 만든 트랙 3), 아래는 스캔 546 (54.6 s) 의 화면입니다.',
                  tag='3-3')
    # 그림 카드 한 장 (스캔 슬라이더 줄 + 2D + 3D). 2D 는 휠로 1.3 배 확대
    picture(s, S('ui_view.png'), 1.565, 2.42, w=10.2, border=True)
    leg = [('참값', '선: 궤적,  원: 출발점,  마름모: 지금 위치 (표적 색)'), ('플롯', '옅은 점: 지난 스캔,  흰 테 큰 점: 지금 스캔,  회색: 클러터'),
           ('트랙', '빈 주황 사각형: 임시,  회색 X: 삭제,  숫자: ID'), ('속도 선', '트랙 속도로 1 초 동안 가는 거리'),
           ('3D', '고도 3 배,  옅은 선: 궤적의 바닥 그림자 (사이가 고도)')]
    for k, (a, b) in enumerate(leg):
        x = 0.70 + (k % 2) * 6.05
        y = 6.10 + (k // 2) * 0.25
        text(s, x, y, 1.0, 0.25, [[(a, 10.5, NAVY, True)]], anchor='m')
        text(s, x + 0.85, y, 5.1, 0.25, [[(b, 10.5, NAVY)]], anchor='m')

    # 24  4.2
    s = d.content(sec4, '4.2 트랙 수명과 스코어 확인', '과제 조건 (Assigned_Flag 늘 0) 에서 트랙은 세 스캔 동안만 보입니다. 확정 · 메모리 상태는 스코어 확인 창으로 확인했습니다.',
                  tag='3-2 · 3-3')
    label(s, 0.70, 2.44, 6.9, '스캔 545 의 트랙 표')
    picture(s, S('table_s545.png'), 0.70, 2.74, w=6.9, border=True)
    callout(s, 1, 0.75, 4.30, 6.85, '나이가 다른 트랙 세 개', ['245: 이번 스캔에 개시, TrackScore 2', '244: 한 스캔 전 개시, 1  /  243: 두 스캔 전 개시, 0 -> 삭제 (4)'], size=10.5)
    text(s, 7.95, 6.30, 4.68, 0.30, [[('스코어 확인 창: 1 1 1 0 0 1 1 을 넣은 결과', 10, SUB)]], align='c')
    callout(s, 2, 0.75, 5.25, 6.85, '모든 트랙이 같은 수명', ['TrackScore 2 -> 1 -> 0 (삭제), 다음 스캔에 표에서 빠짐', '355 개 모두 같은 순서 (끝에 만든 2 개는 1 까지)'], size=10.5)
    picture(s, S('score.png'), 7.95, 2.50, h=3.75, border=True)

    # 25  4.3
    s = d.content(sec4, '4.3 거리에 따른 개시 결과', '멀리 있는 표적일수록 같은 표적 쌍도 속도 게이트를 통과하기 어려워, 트랙이 덜 만들어집니다.',
                  tag='1.2 · 속도 게이팅')
    picture(s, F('fig_rate.png'), 0.70, 2.50, w=6.9)
    # 가까운 순. 표적 2 는 13.4 km 에서 4.8 km 까지 다가오므로 거리 구간으로 나눔. 구간마다 머문 스캔 쌍 수가 달라 쌍당 개시 비율로 견줌
    tab = [('표적', '거리 [km]', '스캔 쌍', '게이트 통과', '개시 트랙', '개시 비율'),
           ('2', '4.8 ~ 6.5', '308', '45 %', '97', '31 %'),
           ('2', '6.5 ~ 8.0', '87', '32 %', '21', '24 %'),
           ('2', '8.0 ~ 10.5', '414', '23 %', '71', '17 %'),
           ('2', '10.5 ~ 13.4', '350', '15 %', '46', '13 %'),
           ('1', '13.6 ~ 14.2', '1,159', '12 %', '117', '10 %')]
    table(s, 7.80, 3.05, 4.83, [0.50, 1.00, 0.70, 0.90, 0.83, 0.90], tab, head_size=10, size=10.5, row_h=0.36)

    # 26 끝
    d.end('감사합니다', '질문 주시면 답변드리겠습니다')
    d.save(out)
    print('saved', out, len(d.prs.slides))


if __name__ == '__main__':
    build(*sys.argv[1:5])
