# 과제 5 발표자료 생성 (틀: 과제 4 발표자료의 레이아웃)
# python build_deck.py <과제4 pptx> <그림 폴더> <소스 폴더 src/TargetSim> <출력 pptx> <대본 md>
import sys, os, re, copy, unicodedata
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

from notes import NOTES

TEMPLATE, FIG, SRC, OUT, SCRIPT_MD = sys.argv[1:6]

NAVY = '1B2740'; SUB = '4A5A7A'; MUTED = '8794AC'; LIGHT = 'AEB9CC'; BORDER = 'D2D9E5'
BLUE = '2563EB'; RED = 'DC2626'; ORANGE = 'D97706'; GREEN = '059669'; PURPLE = '7C3AED'; NUM = 'B45309'; WHITE = 'FFFFFF'
F_GRAY = 'F7F9FC'; F_BLUE = 'EBF2FE'; F_RED = 'FDEEEE'; F_ORANGE = 'FEF5E7'; F_GREEN = 'E8F6F0'; F_PURPLE = 'F3E8FF'
KO = '맑은 고딕'; EN = 'Arial'; MONO = 'Consolas'
FULL_T1 = '━ 표적 1 대함 30 m/s'
FULL_T2 = '━ 표적 2 대공 200 m/s'
MARK = {1: RED, 2: ORANGE, 3: GREEN, 4: BLUE, 5: PURPLE}
MARK_F = {1: F_RED, 2: F_ORANGE, 3: F_GREEN, 4: F_BLUE, 5: F_PURPLE}

prs = Presentation(TEMPLATE)
L = {l.name: l for l in prs.slide_layouts}

# ------------------------------------------------------------------ template preparation
# 과제 4 슬라이드 제거 (레이아웃, 마스터만 남김)
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.get(qn('r:id')))
    sldIdLst.remove(sldId)

# 표지, 끝 배경 그림을 과제 5 그림으로
for lname, img in (('COVER', 'cover_bg.png'), ('END', 'end_bg.png')):
    lay = L[lname]
    blip = lay._element.find('.//' + qn('a:blip'))
    part = lay.part.related_part(blip.get(qn('r:embed')))
    part._blob = open(os.path.join(FIG, img), 'rb').read()

SLIDE_NUM_SP = {}
for lname in ('INDEX', 'CONTENT', 'DIVIDER'):
    for sh in L[lname].placeholders:
        if sh.placeholder_format.type == 13:
            SLIDE_NUM_SP[lname] = sh._element


def add_slide(lname):
    s = prs.slides.add_slide(L[lname])
    if lname in SLIDE_NUM_SP:
        s.shapes._spTree.append(copy.deepcopy(SLIDE_NUM_SP[lname]))
    return s


# ------------------------------------------------------------------ text helpers
def rgb(h):
    return RGBColor.from_string(h)


def _style_run(r, size, color, bold=False, font=KO, italic=False, sub=False, sup=False):
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = rgb(color)
    r.font.name = font
    rPr = r._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = etree.SubElement(rPr, qn('a:ea'))
    ea.set('typeface', font if font != EN else KO)
    if sub:
        rPr.set('baseline', '-25000')
    if sup:
        rPr.set('baseline', '30000')


TOKEN = re.compile(r'(_\{[^}]*\}|\^\{[^}]*\}|\*\*[^*]+\*\*)')


def add_runs(p, text, size, color, bold=False, font=KO):
    # 간단한 표기: _{아래첨자}, ^{위첨자}, **굵게**
    for part in TOKEN.split(text):
        if not part:
            continue
        sub = part.startswith('_{'); sup = part.startswith('^{'); bb = part.startswith('**')
        t = part[2:-1] if (sub or sup) else (part[2:-2] if bb else part)
        r = p.add_run()
        r.text = t
        _style_run(r, size, color, bold or bb, font, sub=sub, sup=sup)


def tb(slide, x, y, w, h, paras, size=12, color=NAVY, bold=False, font=KO, align='l', anchor='t', spacing=None, after=0, wrap=True, margin=0.0):
    # paras: str 또는 [str | (str, {size,color,bold,font,align,after})]
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    m = Inches(margin)
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = m
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, str):
        paras = [paras]
    for i, item in enumerate(paras):
        opt = {}
        if isinstance(item, tuple):
            item, opt = item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[opt.get('align', align)]
        if spacing:
            p.line_spacing = spacing
        p.space_after = Pt(opt.get('after', after))
        add_runs(p, item, opt.get('size', size), opt.get('color', color), opt.get('bold', bold), opt.get('font', font))
    return box


def rect(slide, x, y, w, h, fill=WHITE, line=None, radius=None, lw=0.75, dash=False, shape=None):
    st = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    sp = slide.shapes.add_shape(st, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius and st == MSO_SHAPE.ROUNDED_RECTANGLE:
        sp.adjustments[0] = radius
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = rgb(fill)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = rgb(line); sp.line.width = Pt(lw)
        if dash:
            ln = sp.line._get_or_add_ln()
            pd = etree.SubElement(ln, qn('a:prstDash')); pd.set('val', 'dash')
    sp.shadow.inherit = False
    sp.text_frame.text = ''
    return sp


def label_box(slide, x, y, w, h, text, size=11, color=NAVY, fill=WHITE, line=None, bold=True, radius=0.25, align='c'):
    sp = rect(slide, x, y, w, h, fill, line, radius=radius)
    tf = sp.text_frame
    tf.margin_left = tf.margin_right = Inches(0.04); tf.margin_top = tf.margin_bottom = Inches(0.0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = True
    lines = text.split('\n')
    for i, t in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER}[align]
        add_runs(p, t, size if i == 0 else size - 1.5, color if i == 0 else SUB, bold if i == 0 else False)
    return sp


def line(slide, x1, y1, x2, y2, color=LIGHT, width=1.0, dash=False, arrow=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color); c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        pd = etree.SubElement(ln, qn('a:prstDash')); pd.set('val', 'dash')
    if arrow:
        te = etree.SubElement(ln, qn('a:tailEnd')); te.set('type', 'triangle'); te.set('w', 'med'); te.set('len', 'med')
    return c


def marker(slide, x, y, n, d=0.26, size=9.5):
    sp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    sp.fill.solid(); sp.fill.fore_color.rgb = rgb(MARK[n]); sp.line.fill.background(); sp.shadow.inherit = False
    tf = sp.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = str(n); _style_run(r, size, WHITE, True)
    return sp


def picture(slide, name, x, y, w=None, h=None, border=False):
    path = os.path.join(FIG, name)
    pic = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)
    if border:
        pic.line.color.rgb = rgb(BORDER); pic.line.width = Pt(0.75)
    return pic


def callout(slide, x, y, w, n, head, body, hsize=11.5, bsize=10.5):
    marker(slide, x, y + 0.02, n, d=0.28)
    paras = [(head, dict(size=hsize, bold=True, color=NAVY, after=2))]
    for b in body:
        paras.append((b, dict(size=bsize, color=SUB, after=1)))
    return tb(slide, x + 0.42, y, w - 0.42, 0.3 + 0.22 * len(body), paras)


def section_label(text):
    return text


def content(sec, title, subtitle):
    s = add_slide('CONTENT')
    ph = {p.placeholder_format.idx: p for p in s.placeholders}
    for idx, txt, size, color, bold in ((102, sec, 13.5, WHITE, True), (103, title, 21, NAVY, True), (104, subtitle, 12.5, SUB, False)):
        tf = ph[idx].text_frame
        tf.text = ''
        p = tf.paragraphs[0]
        add_runs(p, txt, size, color, bold)
    return s


def divider(title, desc, items):
    s = add_slide('DIVIDER')
    tb(s, 0.94, 1.0, 10.0, 0.52, title, size=24, bold=True)
    tb(s, 0.94, 1.58, 11.58, 0.33, desc, size=13, color=NAVY)
    y = 2.25
    for num, name, page in items:
        tb(s, 0.94, y, 0.56, 0.44, num, size=14, bold=True)
        t = tb(s, 1.55, y, 6.6, 0.44, name, size=14)
        est = 1.55 + 0.02 + sum(0.19 if unicodedata.east_asian_width(ch) in 'WF' else 0.105 for ch in name)
        line(s, est + 0.12, y + 0.28, 11.6, y + 0.28, color=BORDER, width=0.75, dash=True)
        tb(s, 11.74, y, 0.78, 0.44, str(page), size=13, color=SUB, align='r')
        y += 0.445
    return s


# ------------------------------------------------------------------ code panels
TYPES = {'UINT32', 'INT32', 'FLOAT32', 'FLOAT64', 'DOUBLE64', 'VOID', 'matrix', 'const', 'static', 'UCHAR8', 'INT64', 'CHAR', 'typedef', 'struct'}
KEYWORDS = {'if', 'else', 'switch', 'case', 'break', 'for', 'while', 'return', 'default'}
CTOKEN = re.compile(r'(//.*$|\b[A-Za-z_][A-Za-z0-9_]*\b|\b\d+(?:\.\d+)?(?:e\d+)?U?\b|\s+|.)')


def src_lines(rel):
    with open(os.path.join(SRC, rel), encoding='utf-8') as f:
        return f.read().split('\n')


def find_line(lines, text, start=0):
    for i in range(start, len(lines)):
        if text in lines[i]:
            return i
    raise ValueError(text)


def dwidth(s):
    return sum(2 if unicodedata.east_asian_width(c) in 'WF' else 1 for c in s)


def tidy(lines, maxw=78):
    # 탭 정렬을 패널 안에서 다시 맞춤: 앞 들여쓰기는 탭 하나 = 2칸, 줄 안의 탭 묶음은 열로 정렬
    rows = []
    for raw in lines:
        lead = len(raw) - len(raw.lstrip('\t'))
        body = raw.lstrip('\t')
        cmt = ''
        if '//' in body and not body.startswith('//'):
            k = body.index('//'); cmt = body[k:].replace('// @var ', '// '); body = body[:k].rstrip('\t ')
        rows.append((lead, re.split(r'\t+', body) if body else [''], cmt))
    width = {}
    for lead, f, c in rows:
        for i, x in enumerate(f[:-1]):
            width[(lead, i)] = max(width.get((lead, i), 0), dwidth(x))
    out = []
    for lead, f, c in rows:
        s = '  ' * lead
        for i, x in enumerate(f):
            s += x + (' ' * (width[(lead, i)] - dwidth(x) + 2) if i < len(f) - 1 else '')
        if c:
            s += '  ' + c
        while dwidth(s) > maxw:
            s = s[:-2]
            if not s.endswith('…'):
                s = s[:-1] + '…'
        out.append(s)
    return out


def code_runs(p, num, text, size):
    r = p.add_run(); r.text = f'{num:>3}  '; _style_run(r, size, LIGHT, font=MONO)
    for m in CTOKEN.finditer(text):
        t = m.group(0)
        if t.startswith('//'):
            c = GREEN
        elif t in TYPES or t.startswith('STRUCT_') or t.startswith('ST_'):
            c = BLUE
        elif t in KEYWORDS:
            c = BLUE
        elif re.fullmatch(r'[A-Z][A-Z0-9_]{2,}', t):
            c = PURPLE
        elif re.fullmatch(r'\d+(?:\.\d+)?(?:e\d+)?U?', t):
            c = NUM
        else:
            c = NAVY
        r = p.add_run(); r.text = t; _style_run(r, size, c, font=MONO)


def code_panel(slide, x, y, w, rel, start_text, count, title=None, highlights=None, size=9.0, lh=0.165, maxw=None, skip=None, start_after=0, keep=None):
    # highlights: {표시 번호: (첫 줄 글자, 줄 수)}
    lines = src_lines(rel)
    i0 = find_line(lines, start_text, start_after)
    idx = [i for i in range(i0, i0 + count) if not (skip and lines[i].strip() in skip) and (keep is None or keep(lines[i]))]
    shown = tidy([lines[i] for i in idx], maxw=maxw or int((w - 0.5) * 72 / (size * 0.6)) - 5)
    if title:
        tb(slide, x, y - 0.29, w, 0.26, [title[0] + '    ' + title[1]], size=11, bold=True)
        # 파일 이름은 옅게
        box = slide.shapes[-1].text_frame.paragraphs[0]
        box.runs[0].text = title[0] + '    '
        r = box.add_run(); r.text = title[1]; _style_run(r, 9.5, MUTED, font=MONO)
    h = 0.12 + lh * len(idx) + 0.06
    rect(slide, x, y, w, h, WHITE, BORDER)
    marks = []
    if highlights:
        for n, (txt, nl) in highlights.items():
            k = next(j for j, i in enumerate(idx) if txt in lines[i])
            yy = y + 0.08 + lh * k
            rect(slide, x + 0.03, yy, w - 0.06, lh * nl, MARK_F[n])
            rect(slide, x + 0.03, yy, 0.04, lh * nl, MARK[n])
            marks.append((n, yy))
    box = slide.shapes.add_textbox(Inches(x + 0.1), Inches(y + 0.08), Inches(w - 0.2), Inches(lh * len(idx)))
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for j, (i, s) in enumerate(zip(idx, shown)):
        p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        p.line_spacing = Pt(lh * 72)
        code_runs(p, i + 1, s, size)
    for n, yy in marks:
        marker(slide, x + w - 0.16, yy + (lh - 0.22) / 2, n, d=0.22, size=8.5)
    return h


# ------------------------------------------------------------------ slides
slides_meta = []      # (번호, 제목) 대본 문서용


def note(slide, key):
    slide.notes_slide.notes_text_frame.text = NOTES[key]
    slides_meta.append((key, slide))


# 1 표지
s = add_slide('COVER')
tb(s, 9.0, 0.72, 3.3, 0.3, 'Confidential', size=10, bold=True, color=WHITE, font=EN, align='r')
tb(s, 4.6, 2.74, 7.4, 0.5, '레이다 SW 학습 과제 5', size=14, color=LIGHT, align='r')
tb(s, 4.6, 3.22, 7.4, 0.75, '트랙 개시 · 관리', size=38, bold=True, color=WHITE, align='r')
line(s, 8.6, 4.28, 12.0, 4.28, color=MUTED, width=0.75)
tb(s, 4.6, 4.44, 7.4, 0.32, 'Track Initiation and Track Management', size=13, color=LIGHT, font=EN, align='r')
tb(s, 4.6, 5.6, 7.4, 0.32, 'Janghwan Kim (Harley)', size=12, color=WHITE, font=EN, align='r')
note(s, 'cover')

# 2 목차
s = add_slide('INDEX')
tb(s, 0.5, 3.0, 3.6, 1.3, [('트랙 개시 · 관리', dict(after=4)), 'TrackMgr'], size=18, bold=True, color=WHITE, align='c')
tb(s, 0.5, 6.6, 3.6, 0.3, 'Janghwan Kim (Harley)', size=10, color=LIGHT, font=EN, align='c')
INDEX = [('01', '트랙 개시', 3), ('02', '트랙 관리', 9), ('03', '구현', 13), ('04', '결과', 19)]
y = 1.97
for i, (n, name, page) in enumerate(INDEX):
    tb(s, 5.25, y, 1.0, 0.55, n, size=30, bold=True, font=EN)
    tb(s, 6.4, y, 1.6, 0.55, name, size=15, bold=True, anchor='m')
    line(s, 6.45 + 0.24 * len(name) + 0.2, y + 0.33, 11.95, y + 0.33, color=BORDER, width=0.75, dash=True)
    tb(s, 12.04, y, 0.56, 0.55, str(page), size=14, font=EN, anchor='m', align='r')
    if i < len(INDEX) - 1:
        line(s, 5.25, y + 0.72, 12.6, y + 0.72, color=BORDER, width=0.75)
    y += 0.84
note(s, 'index')

# 3 간지 01
s = divider('01. 트랙 개시', '플롯으로 트랙을 처음 만드는 방법과, 그때 필요한 초기 상태 · 오차 공분산 · 속도 게이트를 설명합니다.',
            [('1.1', '플롯에서 트랙으로', 4), ('1.2', '1-Point 와 2-Point 트랙 개시', 5), ('1.3', '2-Point 초기 상태와 오차 공분산', 6),
             ('1.4', '측정잡음 공분산 R (σr 10 m, σaz · σel 0.2°)', 7), ('1.5', '속도 게이팅 (V_max = 800 m/s)', 8)])
note(s, 'div1')

# 4 플롯에서 트랙으로
s = content('01. 트랙 개시', '1.1 플롯에서 트랙으로', '레이다는 스캔마다 점 (플롯) 만 줍니다. 이 점들을 시간 순서로 묶어 "같은 물체" 로 관리하는 단위가 트랙입니다.')
tb(s, 0.7, 2.46, 5.0, 0.3, '레이다 처리 흐름에서의 위치', size=12, bold=True)
flow = [('표적 궤적 모의', '과제 3 · 시각별 참값', F_GRAY, NAVY), ('측정 모의', '과제 4 · 플롯 · 클러터', F_GRAY, NAVY),
        ('트랙 개시 · 관리', '과제 5 · 이번 과제', F_BLUE, NAVY), ('추적 필터 · 트랙 연관', '다음 과제', F_GRAY, '7A879F')]
for i, (a, b, f, c) in enumerate(flow):
    x = 0.7 + i * 3.08
    rect(s, x, 2.82, 2.69, 0.66, f, BLUE if i == 2 else 'E3E8F0', lw=1.2 if i == 2 else 0.75)
    tb(s, x, 2.88, 2.69, 0.3, a, size=12.5, bold=True, color=c, align='c')
    tb(s, x, 3.18, 2.69, 0.24, b, size=10.5, color=SUB if i == 2 else '8794AC', align='c')
    if i < 3:
        line(s, x + 2.74, 3.15, x + 3.03, 3.15, color=LIGHT, width=1.0, arrow=True)
label_box(s, 7.42, 3.56, 0.94, 0.26, '이번 과제', size=10, color=BLUE, line=BLUE)
tb(s, 9.94, 3.56, 2.7, 0.26, '연관 · 필터는 다음 과제에서', size=10.5, color=SUB)
line(s, 0.7, 3.94, 12.63, 3.94, color='E3E8F0', width=0.75)
picture(s, 'fig_plot_track.png', 0.75, 3.99, w=7.85)
cards = [('플롯 (Plot)', '한 스캔에서 탐지된 측정치 하나.\n거리 · 방위각 · 고각. 표적인지 모름'),
         ('트랙 (Track)', '같은 물체라고 묶은 플롯들.\nID · 위치 · 속도 · 상태를 가짐'),
         ('이번 과제 범위', '트랙을 처음 만들기 (개시)\n트랙 상태 올리고 내리기 (관리)')]
for i, (h, b) in enumerate(cards):
    yy = 4.1 + i * 0.93
    label_box(s, 9.15, yy, 1.25, 0.28, h.split(' (')[0], size=10.5, color=[NAVY, ORANGE, BLUE][i], line=[NAVY, ORANGE, BLUE][i])
    tb(s, 9.15, yy + 0.33, 3.5, 0.55, b.split('\n'), size=10.5, color=NAVY)
note(s, 's4')

# 5 1-Point / 2-Point
s = content('01. 트랙 개시', '1.2 1-Point 와 2-Point 트랙 개시', '트랙을 플롯 몇 개로 시작하느냐의 차이입니다. 빨리 시작할지, 속도를 알고 시작할지를 고르는 문제입니다.')
picture(s, 'fig_1p2p.png', 0.8, 2.45, w=11.7)
rect(s, 0.7, 6.3, 11.93, 0.46, F_BLUE, None, radius=0.2)
tb(s, 0.9, 6.3, 11.6, 0.46, ['이번 과제: **2-Point 개시** — 연속 두 스캔의 플롯 쌍으로 초기 상태 변수 (mat_EstX) 와 초기 오차 공분산 (mat_EstP) 을 계산'],
   size=12, color=NAVY, anchor='m')
note(s, 's5')

# 6 초기 상태와 오차 공분산
s = content('01. 트랙 개시', '1.3 2-Point 초기 상태와 오차 공분산', '상태 변수는 축마다 위치 · 속도 · 가속도, 모두 9 개. 두 플롯으로 위치와 속도를 채우고 가속도 항은 과제 조건대로 0 으로 둡니다.')
tb(s, 0.7, 2.46, 6.8, 0.3, '상태 변수 (VC-NED: 북 · 동 · 아래, 레이다 원점)', size=12, bold=True)
rect(s, 0.7, 2.8, 6.8, 0.5, F_GRAY, 'E3E8F0', radius=0.15)
tb(s, 0.7, 2.8, 6.8, 0.5, "X_{k} = [ x  x'  x''   y  y'  y''   z  z'  z'' ]^{T}", size=14, bold=True, align='c', anchor='m', font='Cambria Math')
tb(s, 0.7, 3.48, 6.8, 0.3, '초기 상태 mat_EstX  (축마다 같은 식, T = 두 플롯의 시간 차 0.1 s)', size=12, bold=True)
rows = [('위치', "x̂ = z_{k}", '이번 스캔 플롯 그대로', BLUE), ('속도', "x̂' = ( z_{k} − z_{k−1} ) / T", '위치 변화량 ÷ 시간', GREEN), ('가속도', "x̂'' = 0", '두 점으로는 못 구함 (과제 조건)', MUTED)]
for i, (a, b, c, col) in enumerate(rows):
    yy = 3.82 + i * 0.42
    label_box(s, 0.7, yy, 0.8, 0.3, a, size=10.5, color=col, line=col)
    tb(s, 1.65, yy - 0.02, 3.2, 0.34, b, size=13, font='Cambria Math', anchor='m')
    tb(s, 4.75, yy - 0.02, 2.8, 0.34, c, size=10.5, color=SUB, anchor='m')
tb(s, 0.7, 5.15, 6.8, 0.3, '초기 오차 공분산 mat_EstP  (축 a, b 마다, R = 플롯 위치 측정잡음)', size=12, bold=True)
prow = [('P_{pp} = R_{k}', BLUE), ('P_{pv} = P_{vp} = R_{k} / T', GREEN), ('P_{vv} = ( R_{k} + R_{k−1} ) / T^{2}', RED), ('가속도 행 · 열 = 0', MUTED)]
for i, (t, col) in enumerate(prow):
    xx = 0.7 + (i % 2) * 3.45; yy = 5.5 + (i // 2) * 0.42
    rect(s, xx, yy, 3.3, 0.34, WHITE, col, radius=0.2)
    tb(s, xx, yy, 3.3, 0.34, t, size=12, color=col, bold=True, align='c', anchor='m', font='Cambria Math' if i < 3 else KO)
picture(s, 'fig_pmatrix.png', 7.95, 2.45, h=4.05)
tb(s, 7.95, 6.5, 4.4, 0.3, 'mat_EstP 9 × 9 구조 (같은 색 = 같은 식, 회색 = 가속도 0)', size=10, color=SUB, align='c')
note(s, 's6')

# 7 측정잡음 공분산 R
s = content('01. 트랙 개시', '1.4 측정잡음 공분산 R', 'P 를 만들려면 플롯 하나의 위치 오차 R 이 필요합니다. 과제가 정해 준 표준편차로 R 을 만들어 VC-NED 좌표로 바꿉니다.')
stats = [('σ range', '10', 'm'), ('σ azimuth', '0.2', '°'), ('σ elevation', '0.2', '°')]
for i, (a, b, u) in enumerate(stats):
    xx = 0.7 + i * 2.0
    rect(s, xx, 2.5, 1.85, 0.84, F_BLUE, None, radius=0.12)
    tb(s, xx, 2.54, 1.85, 0.26, a, size=11, color=SUB, align='c', font=EN)
    tb(s, xx, 2.76, 1.85, 0.52, [b + ' ' + u], size=24, bold=True, color=BLUE, align='c', font=EN)
# R 의 표준편차는 과제가 정해 준 값. 과제 4 플롯 모의의 SNR 기반 표준편차 (PlotSim.c f_Pcs_CalcSigma, SNR 20) 와 따로 정해짐
for i, (tag, col, txt) in enumerate((('과제 5 지정값', BLUE, '트랙 초기화 (초기 오차 공분산) 용 R 의 표준편차'),
                                     ('과제 4 SNR', '6B7280', '플롯 모의 잡음은 SNR 20 으로 계산: 4.74 m · 0.48° · 0.49° (R 과 별개)'))):
    yy = 3.41 + i * 0.235
    label_box(s, 0.7, yy, 1.08, 0.21, tag, size=9, color=col, line=col)
    tb(s, 1.88, yy, 4.7, 0.21, txt, size=10, color=NAVY if i == 0 else SUB, anchor='m')
rect(s, 6.85, 2.5, 5.78, 1.32, WHITE, BORDER, radius=0.06)
tb(s, 7.0, 2.55, 5.5, 0.28, 'R 이 뜻하는 것', size=11.5, bold=True)
MEANING = (('R = 플롯이 흩어지는 정도', '같은 표적을 여러 번 잴 때 흩어지는 크기 · 모양'),
           ('거리는 정확, 방향은 멀수록 퍼짐', '거리 10 m 그대로, 옆 · 위아래 = 거리 × 0.2°'),
           ('그래서 옆으로 긴 타원', '거리 · 방향이 플롯마다 달라 R 도 플롯마다 계산'))
for i, (h, b) in enumerate(MEANING):
    yy = 2.88 + i * 0.3
    marker(s, 7.0, yy + 0.03, i + 1, d=0.22, size=8.5)
    p = tb(s, 7.32, yy, 5.25, 0.28, '', size=10.5, anchor='m').text_frame.paragraphs[0]
    add_runs(p, h + '   ', 10.5, NAVY, True)
    add_runs(p, b, 9.5, SUB)
picture(s, 'fig_rcov.png', 0.7, 3.95, w=7.6)
# R 계산 행렬식: 위치 변환 (이전 과제 함수) 과 공분산 변환을 나란히. 회전은 같은 행렬, 구좌표 → 직교 함수 s 만 기울기 G (f_Trk_MeasCov 의 mat_J)
p = tb(s, 8.55, 3.97, 4.1, 0.28, '', size=12, anchor='m').text_frame.paragraphs[0]
add_runs(p, 'R 계산 행렬식   ', 12, NAVY, True)
add_runs(p, '이전 과제 좌표 변환과 같은 행렬', 9.5, MUTED)
picture(s, 'fig_rmatrix.png', 8.48, 4.24, w=4.15)
note(s, 's7')

# 8 속도 게이팅
s = content('01. 트랙 개시', '1.5 속도 게이팅', '앞 스캔 플롯과 이번 스캔 플롯을 아무렇게나 짝지으면 클러터끼리도 트랙이 됩니다. 표적이 낼 수 있는 최대 속도로 짝을 거릅니다.')
rect(s, 0.7, 2.5, 6.2, 1.22, F_BLUE, None, radius=0.1)
tb(s, 0.7, 2.58, 6.2, 0.6, '| z_{k,i} − z_{k−1,j} | / ΔT  <  V_{max}', size=22, bold=True, color=NAVY, align='c', anchor='m', font='Cambria Math')
tb(s, 0.7, 3.2, 6.2, 0.4, 'V_{max} = 800 m/s (과제 조건),  ΔT = 0.1 s  →  앞 플롯 중심 반경 80 m', size=12, color=BLUE, align='c', bold=True)
tb(s, 0.7, 3.92, 6.2, 0.3, '구현에서 정한 짝짓기 규칙', size=12, bold=True)
callout(s, 0.7, 4.3, 6.2, 1, '게이트를 통과한 쌍이 여럿이면 가장 가까운 쌍부터', ['두 플롯의 속력 (위치 변화량 ÷ ΔT) 이 작은 순서로 선택'])
callout(s, 0.7, 5.0, 6.2, 2, '한 플롯은 한 트랙에만 씀', ['개시에 쓴 플롯은 다음 스캔의 개시 후보에서 뺌'])
callout(s, 0.7, 5.7, 6.2, 3, '짝이 없던 이번 스캔 플롯은 다음 스캔 후보로', ['오른쪽 그림의 d 처럼 다음 스캔에서 다시 짝을 찾음'])
picture(s, 'fig_gate.png', 7.3, 2.45, w=5.3)
note(s, 's8')

# 9 간지 02
s = divider('02. 트랙 관리', '개시한 트랙을 계속 믿을지 지울지 정하는 규칙입니다. 트랙 상태 네 가지와 M/N 방식, 트랙 스코어 방식을 봅니다.',
            [('2.1', '트랙 상태 네 가지', 10), ('2.2', 'M/N 방식', 11), ('2.3', '트랙 스코어 방식 (이번 과제)', 12)])
note(s, 'div2')

# 10 트랙 상태
s = content('02. 트랙 관리', '2.1 트랙 상태 네 가지', 'TrackState 1 임시 · 2 확정 · 3 메모리 · 4 삭제. 매 스캔 트랙 연관 결과 (Assigned_Flag: 1 성공, 0 실패) 로 상태가 바뀝니다.')
picture(s, 'fig_states.png', 0.7, 2.6, w=8.0)
rect(s, 8.95, 2.5, 3.68, 3.32, F_GRAY, None, radius=0.06)
tb(s, 9.1, 2.56, 3.4, 0.3, '트랙 연관과 Assigned_Flag', size=12, bold=True)
picture(s, 'fig_assoc.png', 9.0, 2.9, w=3.58)
tb(s, 9.1, 4.55, 3.42, 1.25, ['트랙 연관: 이번 스캔 플롯 가운데 기존 트랙에 이어지는 플롯을 찾아 붙이는 과정',
                              '트랙 위치 · 속도로 이번 위치를 예측하고, 그 주변 연관 게이트 안의 가장 가까운 플롯을 할당',
                              ('이번 과제는 연관 기능이 없어 항상 0', dict(color=ORANGE))], size=9.5, color=NAVY, after=3)
cards = [('1 임시', '막 개시한 트랙. 아직 진짜 표적인지 모름', ORANGE, F_ORANGE), ('2 확정', '연관이 꾸준히 성공해 믿을 만해진 트랙', GREEN, F_GREEN),
         ('3 메모리', '확정이었다가 이번 스캔에 놓친 트랙. 바로 지우지 않고 기다림', PURPLE, F_PURPLE), ('4 삭제', '더 볼 필요 없는 트랙. 한 스캔 보인 뒤 표에서 빠짐', '6B7280', F_GRAY)]
for i, (a, b, c, f) in enumerate(cards):
    xx = 0.7 + i * 3.02
    rect(s, xx, 5.95, 2.86, 0.85, f, None, radius=0.12)
    tb(s, xx + 0.15, 6.0, 2.6, 0.28, a, size=12, bold=True, color=c)
    tb(s, xx + 0.15, 6.28, 2.6, 0.5, b, size=10.5, color=NAVY)
note(s, 's10')

# 11 M/N
s = content('02. 트랙 관리', '2.2 M/N 방식', '최근 N 번 가운데 M 번 연관되면 확정합니다. 연관 성공 · 실패를 횟수로 세는 방식입니다.')
picture(s, 'fig_mn.png', 0.9, 2.45, w=11.2)
for i, (h, items, c, f) in enumerate([('임시 트랙', ['N 번 중 M 번 트랙 연관 성공 → 트랙 확정', 'M 번 연속 결합 실패 → 트랙 삭제'], ORANGE, F_ORANGE),
                                      ('확정 트랙', ['트랙 연관 실패 → 메모리 트랙 전환', '특정 시간 이상 메모리 상태 유지 → 트랙 삭제'], GREEN, F_GREEN)]):
    xx = 0.7 + i * 4.05
    rect(s, xx, 5.3, 3.85, 1.5, f, None, radius=0.08)
    tb(s, xx + 0.2, 5.4, 3.5, 0.3, h, size=12.5, bold=True, color=c)
    tb(s, xx + 0.2, 5.8, 3.5, 0.9, items, size=11, color=NAVY, after=4)
rect(s, 8.8, 5.3, 3.83, 1.5, WHITE, BORDER, radius=0.08)
tb(s, 9.0, 5.4, 3.5, 1.35, [('특징', dict(size=12.5, bold=True, after=4)), ('성공 · 실패 횟수로 판단하는 단순한 규칙', dict(size=11, color=NAVY, after=3)),
                            ('트랙마다 최근 N 번의 연관 이력을 관리', dict(size=11, color=NAVY))])
note(s, 's11')

# 12 트랙 스코어
s = content('02. 트랙 관리', '2.3 트랙 스코어 방식 (이번 과제)', '연관에 성공하면 점수를 올리고 실패하면 내립니다. 점수가 문턱치를 넘는지로 상태를 바꿉니다. 과제는 이 방식을 씁니다.')
picture(s, 'fig_score.png', 0.62, 2.5, w=7.75)
tb(s, 0.7, 5.15, 7.6, 0.3, 'Flag 1 1 1 0 0 1 0 0 0 0 0 을 넣은 예 (UI 스코어 확인 창과 같은 계산)', size=10, color=SUB)
rect(s, 8.55, 2.5, 4.08, 4.3, WHITE, BORDER, radius=0.04)
tb(s, 8.75, 2.58, 3.8, 0.3, '과제 조건과 규칙', size=12.5, bold=True)
rules = [('임시 (1)', ORANGE, ['성공 TrackScore +1, 실패 −1', 'TrackScore ≥ **5** → 확정', 'TrackScore ≤ **0** → 삭제']),
         ('확정 (2)', GREEN, ['실패 → 메모리, MemoryScore +1']),
         ('메모리 (3)', PURPLE, ['성공 → 바로 확정, MemoryScore = 0', '실패 → MemoryScore +1', 'MemoryScore ≥ **5** → 삭제'])]
yy = 2.98
for h, c, items in rules:
    label_box(s, 8.75, yy, 1.05, 0.28, h, size=10.5, color=c, line=c)
    tb(s, 9.95, yy - 0.02, 2.65, 0.25 * len(items) + 0.05, items, size=10.5, color=NAVY, after=2)
    yy += 0.3 + 0.27 * len(items)
line(s, 8.75, yy + 0.02, 12.45, yy + 0.02, color=BORDER, width=0.75)
tb(s, 8.75, yy + 0.08, 3.8, 0.75, ['개시 직후 TrackScore = 2', '(플롯 두 개로 개시한 트랙)'], size=10, color=SUB, after=1)
rect(s, 0.7, 5.55, 7.6, 1.25, F_GRAY, None, radius=0.08)
tb(s, 0.9, 5.62, 7.3, 1.15, [('과제 적용 조건', dict(size=12, bold=True, after=3)),
                             ('· 임시 트랙 확정 문턱치 5, 임시 트랙 삭제 문턱치 0, 메모리 트랙 삭제 문턱치 5', dict(size=10.5, after=2)),
                             ('· 메모리 트랙이 다시 연관되면 바로 확정 트랙으로 전환, MemoryScore = 0', dict(size=10.5, after=2)),
                             ('· TrackState: 1 임시, 2 확정, 3 메모리, 4 삭제', dict(size=10.5))])
note(s, 's12')

# 13 간지 03
s = divider('03. 구현', '과제 4 프로그램 (TargetSim) 에 플롯 메시지 송수신과 트랙 개시 · 관리를 넣었습니다. 코드마다 맡은 역할과 처리 순서, 계산 방법을 봅니다.',
            [('3.1', '전체 구조와 데이터 흐름', 14), ('3.2', '플롯 메시지 송수신 순서', 15), ('3.3', '트랙 자료구조와 스캔 처리 순서 (f_Trk_ProcessScan)', 16),
             ('3.4', '2-Point 개시와 R 계산 (f_Trk_TwoPointInit)', 17), ('3.5', '트랙 스코어 로직 (f_Trk_UpdateState)', 18)])
note(s, 'div3')

# 14 전체 구조
s = content('03. 구현', '3.1 전체 구조와 데이터 흐름', '한 프로그램 안에 신호처리 역할의 송신 쓰레드와 데이터 처리 역할의 수신 쓰레드를 두고, 스캔마다 플롯 메시지를 UDP 로 주고받습니다.')
rect(s, 0.7, 2.5, 11.93, 3.38, F_GRAY, 'E3E8F0', radius=0.03)
tb(s, 0.9, 2.56, 5, 0.28, 'TargetSimCore.dll  (C)', size=11, bold=True, color=SUB, font=EN)
B = [(1.0, 2.92, '궤적 계산  (과제 3)', '0.1 s 마다 플랫폼 · 표적 위치 계산', 'TargetSim.c', NAVY, WHITE),
     (4.75, 2.92, '측정 모의  (과제 4)', '스캔마다 표적 플롯 + 클러터 5 개 생성', 'PlotSim.c', NAVY, WHITE),
     (8.5, 2.92, '송신 쓰레드  (신호처리 역할)', '스캔의 플롯을 메시지로 묶어 보냄', 'PlotLink.c', BLUE, F_BLUE),
     (8.5, 4.68, '수신 쓰레드  (데이터 처리 역할)', '좌표 변환, 플롯마다 R 계산', 'PlotLink.c', BLUE, F_BLUE),
     (4.75, 4.68, '트랙 개시 · 관리', '트랙 표 갱신 (개시 · 상태 관리)', 'TrackMgr.c', ORANGE, F_ORANGE),
     (1.0, 4.68, '결과 버퍼', '스캔마다 플롯 · 트랙 결과 저장', 'PlotLink.c', NAVY, WHITE)]
for x, y, a, b, f_, c, f in B:
    rect(s, x, y, 3.3, 0.98, f, c, radius=0.1, lw=1.1)
    tb(s, x, y + 0.08, 3.3, 0.3, a, size=12, bold=True, color=c, align='c')
    tb(s, x, y + 0.41, 3.3, 0.28, b, size=10, color=NAVY, align='c')
    tb(s, x, y + 0.67, 3.3, 0.24, f_, size=8.5, color=MUTED, align='c', font=MONO)
for x1, y1, x2, y2 in ((4.32, 3.41, 4.73, 3.41), (8.07, 3.41, 8.48, 3.41), (8.48, 5.17, 8.07, 5.17), (4.73, 5.17, 4.32, 5.17)):
    line(s, x1, y1, x2, y2, color=SUB, width=1.25, arrow=True)
line(s, 10.15, 3.92, 10.15, 4.66, color=BLUE, width=2.0, arrow=True)
tb(s, 10.27, 4.06, 2.3, 0.5, ['UDP', '스캔마다 메시지 1 개'], size=10, color=BLUE, bold=True)
line(s, 2.65, 5.68, 2.65, 6.03, color=SUB, width=1.25, arrow=True)
rect(s, 0.7, 6.05, 11.93, 0.78, WHITE, NAVY, radius=0.1, lw=1.1)
tb(s, 0.95, 6.09, 11.5, 0.72, [('TargetSimUI.exe  (MFC 화면)', dict(size=12, bold=True, after=2)),
                              ('결과 버퍼에서 처리된 스캔을 가져와  True · Plot · Clutter · Track 2D / 3D 그림,  지금 스캔의 트랙 표,  송수신 · 트랙 요약으로 보여 줌', dict(size=10.5, color=SUB))])
note(s, 's14')

# 15 플롯 메시지 송수신 순서
s = content('03. 구현', '3.2 플롯 메시지 송수신 순서', '송신 쓰레드는 스캔마다 플롯을 메시지 하나로 묶어 보내고, 수신 쓰레드는 받은 플롯을 트랙 계산에 쓸 수 있는 형태로 바꿔 넘깁니다.')


def flow_lane(slide, y, title, func, steps, col, fill, w_total=8.25, gap=0.3, h=1.12):
    tb(slide, 0.7, y, 6.0, 0.3, [title], size=12, bold=True, color=col)
    t = slide.shapes[-1].text_frame.paragraphs[0]
    r = t.add_run(); r.text = '    ' + func; _style_run(r, 9.5, MUTED, font=MONO)
    n = len(steps)
    w = (w_total - gap * (n - 1)) / n
    yy = y + 0.38
    for i, (num, head, body, fn) in enumerate(steps):
        x = 0.7 + i * (w + gap)
        rect(slide, x, yy, w, h, fill, col, radius=0.08, lw=1.0)
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.1), Inches(yy + 0.1), Inches(0.25), Inches(0.25))
        dot.fill.solid(); dot.fill.fore_color.rgb = rgb(col); dot.line.fill.background(); dot.shadow.inherit = False
        dtf = dot.text_frame; dtf.margin_left = dtf.margin_right = dtf.margin_top = dtf.margin_bottom = 0; dtf.vertical_anchor = MSO_ANCHOR.MIDDLE
        dp = dtf.paragraphs[0]; dp.alignment = PP_ALIGN.CENTER; dr = dp.add_run(); dr.text = num; _style_run(dr, 9.5, WHITE, True)
        tb(slide, x + 0.42, yy + 0.08, w - 0.5, 0.3, [head], size=11.5, bold=True, color=col)
        tb(slide, x + 0.1, yy + 0.4, w - 0.2, 0.45, body.split('\n'), size=9.8, color=NAVY)
        if fn:
            tb(slide, x + 0.1, yy + h - 0.27, w - 0.2, 0.22, fn, size=8, color=MUTED, font=MONO)
        if i < n - 1:
            line(slide, x + w + 0.03, yy + h / 2, x + w + gap - 0.03, yy + h / 2, color=col, width=1.25, arrow=True)
    return yy, w, h


TX = [('1', '플롯 모으기', '표적 플롯 2 + 클러터 5\n방위각 순서로 섞음', 'f_Lnk_BuildMsg'),
      ('2', '메시지 만들기', '플롯마다 ID 부여\n거리 · 방위각 · 고각 + R σ', 'f_Lnk_BuildMsg'),
      ('3', '보내기', '스캔 순서대로\n메시지 1 개씩 UDP 전송', 'f_Lnk_TxThread')]
RX = [('4', '좌표 변환', '구좌표 →\nVC-NED 직교 좌표', 'f_Lnk_FillPlotCoord'),
      ('5', 'R 계산', '플롯마다\n측정잡음 공분산', 'f_Trk_MeasCov'),
      ('6', '트랙 처리', '개시 · 상태 관리\n(다음 장)', 'f_Trk_ProcessScan'),
      ('7', '결과 저장', '플롯 · 트랙 표 저장\n→ 화면이 가져감', 'f_Lnk_ProcessMsg')]
y_tx, w_tx, h_tx = flow_lane(s, 2.5, '송신 쓰레드  (신호처리 역할)', 'PlotLink.c  f_Lnk_TxThread', TX, BLUE, F_BLUE, gap=0.4)
y_rx, w_rx, h_rx = flow_lane(s, 4.62, '수신 쓰레드  (데이터 처리 역할)', 'PlotLink.c  f_Lnk_RxThread', RX, GREEN, F_GREEN, gap=0.3)
# ③ → ④ (UDP)
x3 = 0.7 + 2 * (w_tx + 0.4) + w_tx / 2
ym = y_rx + h_rx / 2
line(s, x3, y_tx + h_tx, x3, 4.44, color=BLUE, width=1.6)
line(s, x3, 4.44, 0.48, 4.44, color=BLUE, width=1.6)
line(s, 0.48, 4.44, 0.48, ym, color=BLUE, width=1.6)
line(s, 0.48, ym, 0.68, ym, color=BLUE, width=1.6, arrow=True)
tb(s, x3 - 2.7, y_tx + h_tx + 0.08, 2.6, 0.25, 'UDP  (스캔마다 메시지 1 개)', size=10, bold=True, color=BLUE, align='r')
# 메시지 내용 카드
cx, cy, cw = 9.35, 2.5, 3.28
rect(s, cx, cy, cw, 3.85, WHITE, BORDER, radius=0.04)
tb(s, cx + 0.2, cy + 0.12, cw - 0.4, 0.3, '메시지 한 개에 담는 내용', size=12, bold=True)
tb(s, cx + 0.2, cy + 0.42, cw - 0.4, 0.25, 'STRUCT_PlotMsg (Idd.h)', size=8.5, color=MUTED, font=MONO)
ITEMS = [('스캔 번호 · 시각', ''), ('플랫폼 위치 · 자세', '좌표 변환에 사용'), ('R 표준편차', 'σr 10 m, σaz 0.2°, σel 0.2°'),
         ('플롯 7 개', '플롯 ID (10000 ~ 1000000)\n거리 · 방위각 · 고각')]
yy = cy + 0.82
for head, body in ITEMS:
    rect(s, cx + 0.2, yy + 0.06, 0.07, 0.2, BLUE)
    paras = [(head, dict(size=10.5, bold=True, after=1))] + [(b, dict(size=9.5, color=SUB)) for b in body.split('\n') if b]
    tb(s, cx + 0.38, yy, cw - 0.55, 0.62, paras)
    yy += 0.33 + 0.2 * len([b for b in body.split('\n') if b]) + 0.12
tb(s, cx + 0.2, cy + 3.42, cw - 0.4, 0.3, '표적 / 클러터 구분 정보는 넣지 않음', size=9.5, color=ORANGE, bold=True)
note(s, 's15')

# 16 자료구조와 처리 순서
s = content('03. 구현', '3.3 트랙 자료구조와 스캔 처리 순서', '수신 쓰레드가 메시지를 받을 때마다 f_Trk_ProcessScan 이 아래 네 단계를 차례로 수행합니다. 트랙 변수 이름은 과제 명세 그대로입니다.')
code_panel(s, 0.7, 2.75, 6.3, 'TargetSimCore/include/TrackMgr.h', 'typedef struct', 13, title=('트랙 하나 ST_Track', 'TrackMgr.h'), start_after=find_line(src_lines('TargetSimCore/include/TrackMgr.h'), '트랙 하나. 변수 이름은'),
           size=8.6, lh=0.15)
CONST = ('TRK_VMAX', 'TRK_SCORE_CONFIRM', 'TRK_SCORE_DELETE', 'TRK_MEMORY_DELETE', 'TRK_SCORE_INIT', 'TRK_R_SIGMA_RANGE', 'TRK_R_SIGMA_AZIMUTH', 'TRK_R_SIGMA_ELEVATION')
code_panel(s, 0.7, 5.25, 6.3, 'TargetSimCore/include/TrackMgr.h', '#define TRK_VMAX', 40, title=('과제 조건 상수', 'TrackMgr.h'), keep=lambda t: t.startswith('#define') and t.split()[1] in CONST, size=8.6, lh=0.15)
steps = [('삭제 트랙 정리', '앞 스캔에서 삭제 (4) 가 된 트랙을\n트랙 표에서 지움'),
         ('트랙 관리  →  3.5', '모든 트랙에 Assigned_Flag = 0 (연관 기능 없음)\n스코어 로직으로 상태 갱신'),
         ('2-Point 개시  →  3.4', '앞 · 이번 스캔 플롯 쌍 중 속도 게이트 통과 쌍으로\n가까운 쌍부터 새 트랙 생성, 트랙 ID 부여'),
         ('개시 후보 보관', '이번 스캔 플롯을 다음 스캔의 개시 후보로 보관\n(트랙을 만드는 데 쓴 플롯은 제외)')]
for i, (h, b) in enumerate(steps):
    yy = 2.5 + i * 1.08
    rect(s, 7.35, yy, 5.28, 0.95, F_BLUE if i in (1, 2) else F_GRAY, None, radius=0.1)
    tb(s, 7.45, yy + 0.08, 0.45, 0.45, ['①②③④'[i]], size=20, bold=True, color=BLUE)
    tb(s, 8.0, yy + 0.08, 4.5, 0.3, h, size=12, bold=True)
    tb(s, 8.0, yy + 0.4, 4.6, 0.55, b.split('\n'), size=10, color=SUB)
note(s, 's16')

# 17 2-Point 개시 코드
s = content('03. 구현', '3.4 2-Point 개시와 R 계산', '두 플롯으로 트랙의 첫 상태 mat_EstX 와 오차 공분산 mat_EstP 를 채웁니다. 시간 차를 구한 뒤 x · y · z 축마다 같은 계산을 반복합니다.')
code_panel(s, 0.7, 2.75, 7.55, 'TargetSimCore/src/TrackMgr.c', 'VOID f_Trk_TwoPointInit(', 26, title=('2-Point 초기화', 'TrackMgr.c'),
           highlights={1: ('st_X->e[(3 * nAxis)][0]', 1), 2: ('st_X->e[(3 * nAxis) + 1][0]', 1), 3: ('st_X->e[(3 * nAxis) + 2][0]', 1), 4: ('st_P->e[(3 * nAxis)][(3 * nAxisB)]', 4)},
           size=7.8, lh=0.148, skip={''})
code_panel(s, 8.5, 2.75, 4.13, 'TargetSimCore/src/TrackMgr.c', 'VOID f_Trk_GetMeasSigma(', 6, title=('R 표준편차 (과제 값)', 'TrackMgr.c'),
           highlights={5: ('st_Sigma->r', 3)}, size=7.4, lh=0.148, maxw=60)
callout(s, 8.5, 4.0, 4.13, 1, '위치 = 이번 플롯 z_{k}', [])
callout(s, 8.5, 4.4, 4.13, 2, '속도 = (z_{k} − z_{k−1}) / T', [])
callout(s, 8.5, 4.8, 4.13, 3, '가속도 = 0 (과제 조건)', [])
callout(s, 8.5, 5.2, 4.13, 4, 'P = R_{k},  R_{k}/T,  (R_{k}+R_{k−1})/T^{2}', ['위치 · 위치-속도 · 속도 칸, 가속도 칸은 0'])
callout(s, 8.5, 5.95, 4.13, 5, 'σ 10 m, 0.2°, 0.2°', ['이 값으로 플롯마다 R 계산 (1.4 의 식)'])
note(s, 's17')

# 18 스코어 로직 코드
s = content('03. 구현', '3.5 트랙 스코어 로직 (f_Trk_UpdateState)', '트랙마다 매 스캔 한 번 불려, 현재 상태와 연관 결과 (Assigned_Flag) 에 따라 점수와 상태를 바꿉니다. 과제 2.2 규칙을 그대로 옮겼습니다.')
code_panel(s, 0.7, 2.75, 5.85, 'TargetSimCore/src/TrackMgr.c', 'VOID f_Trk_UpdateState(', 22, title=('임시 트랙', 'TrackMgr.c'),
           highlights={1: ('st_Track->TrackScore = st_Track->TrackScore +', 1), 2: ('if (st_Track->TrackScore >= TRK_SCORE_CONFIRM)', 3), 3: ('else if (st_Track->TrackScore <= TRK_SCORE_DELETE)', 3)},
           size=7.6, lh=0.158, skip={''}, maxw=82)
code_panel(s, 6.78, 2.75, 5.85, 'TargetSimCore/src/TrackMgr.c', 'case TRK_STATE_CONFIRMED:', 24, title=('확정 · 메모리 트랙', 'TrackMgr.c'),
           highlights={4: ('TRK_STATE_MEMORY;', 2), 5: ('TRK_STATE_CONFIRMED;', 2), 3: ('MemoryScore >= TRK_MEMORY_DELETE', 3)},
           size=7.6, lh=0.158, skip={''}, maxw=82)
legend = [(1, 'TrackScore ±1'), (2, '≥ 5 → 확정'), (3, '삭제 (TrackScore ≤ 0 · MemoryScore ≥ 5)'), (4, '실패 → 메모리, MemoryScore +1'), (5, '성공 → 바로 확정, MemoryScore 0')]
xx = 0.7
for n, t in legend:
    marker(s, xx, 6.55, n, d=0.22, size=8.5)
    w = 0.12 + sum(0.145 if unicodedata.east_asian_width(c) in 'WF' else 0.072 for c in t)
    tb(s, xx + 0.28, 6.52, w, 0.28, t, size=10, color=NAVY)
    xx += w + 0.42
note(s, 's18')

# 19 간지 04
s = divider('04. 결과', '프로그램을 실행한 결과 화면과 수치로 트랙 개시 · 관리 동작을 확인합니다.',
            [('4.1', '시뮬레이션 조건과 결과 화면', 20), ('4.2', 'True · Plot · Track 동시 표시 (스캔 701)', 21), ('4.3', '한 스캔 확대: 트랙이 만들어지는 모습', 22),
             ('4.4', '트랙 수명과 스코어 로직 확인', 23), ('4.5', '전체 궤적 (시뮬레이션 마지막 화면)', 24)])
note(s, 'div4')

# 20 조건과 결과 화면
s = content('04. 결과', '4.1 시뮬레이션 조건과 결과 화면', 'Visual Studio 2022 Release x64 로 빌드해 실행한 화면입니다. 1,160 개 플롯 메시지를 모두 보내고 받아 처리했습니다 (손실 0).')
conds = [('시나리오', '과제 3 지그재그 응용 · 플랫폼 정지\n표적 1 대함 30 m/s, 표적 2 대공 200 m/s'), ('시간', '116 s, 스캔 간격 0.1 s → 1,160 스캔'),
         ('플롯', 'SNR 20 · 표적 2 + 클러터 5 = 스캔당 7 개'), ('개시', '2-Point, V_max 800 m/s'), ('R', 'σr 10 m, σaz 0.2°, σel 0.2°'),
         ('관리', '스코어 5 / 0 / 5, 개시 스코어 2'), ('연관', '없음 → Assigned_Flag 0 고정'), ('전송', 'UDP 루프백 · 보냄 = 받음 = 1,160')]
yy = 2.5
for a, b in conds:
    hh = 0.56 if '\n' in b else 0.4
    rect(s, 0.7, yy, 1.0, hh - 0.06, F_BLUE, None, radius=0.15)
    tb(s, 0.7, yy, 1.0, hh - 0.06, a, size=10.5, bold=True, color=BLUE, align='c', anchor='m')
    tb(s, 1.82, yy, 3.3, hh - 0.06, b.split('\n'), size=10, color=NAVY, anchor='m')
    yy += hh
picture(s, 'ui_s701_both.png', 5.63, 2.5, w=7.0, border=True)
tb(s, 5.35, 7.0 - 0.18, 7.28, 0.25, '', size=9)
note(s, 's20')

# 21 2D 3D (스캔 701)
s = content('04. 결과', '4.2 True · Plot · Track 동시 표시 (스캔 701)', '스캔 701 (70.1 s) 화면. 참값 궤적, 지금까지의 플롯 · 클러터, 지금 스캔의 트랙 심볼을 한 그림에 겹쳐 그립니다.')
picture(s, 'ui_crop_2d.png', 0.7, 2.5, w=5.85, border=True)
picture(s, 'ui_crop_3d.png', 6.78, 2.5, w=5.85, border=True)
tb(s, 0.7, 2.5 + 5.85 * 765 / 1200 + 0.05, 5.85, 0.25, '2D: 위경도 격자 위 동-북 평면', size=10, color=SUB)
tb(s, 6.78, 2.5 + 5.85 * 765 / 1200 + 0.05, 5.85, 0.25, '3D: 동-북-고도 (고도 2 배), 끌어서 회전', size=10, color=SUB)
leg = [('━ 참값 궤적 · ○ 시작점', NAVY), ('● 플롯 (큰 점 = 지금 스캔)', NAVY), ('· 클러터', '6B7280'), ('□ 임시', ORANGE), ('■ 확정', GREEN), ('⬚ 메모리', PURPLE), ('⊠ 삭제', '6B7280')]
xx = 0.7
for t, c in leg:
    w = 0.13 * len(t) + 0.3
    tb(s, xx, 6.56, w, 0.26, t, size=10.5, color=c, bold=True)
    xx += w
note(s, 's21')

# 22 한 스캔 확대
s = content('04. 결과', '4.3 한 스캔 확대: 트랙이 만들어지는 모습 (스캔 1154)', '표적 2 가 레이다 가까이 온 스캔 1154 를 확대했습니다. 두 스캔의 플롯이 게이트 안이면 트랙을 만들고, 연관이 없는 트랙은 두 스캔 뒤 삭제됩니다.')
picture(s, 'fig_zoom.png', 0.62, 2.4, w=6.6)
callout(s, 7.55, 2.5, 5.08, 1, '속도 게이트: 앞 스캔 플롯에서 반경 80 m', ['최대 속도 800 m/s × 스캔 간격 0.1 s'])
callout(s, 7.55, 3.15, 5.08, 2, '이번 스캔 플롯까지 31 m → 게이트 안', ['이 두 플롯으로 트랙 451 개시 (임시, 점수 2)'])
callout(s, 7.55, 3.8, 5.08, 3, '위치 = 이번 플롯,  속도 = 31 m ÷ 0.1 s ≈ 310 m/s', ['플롯의 측정 잡음 때문에 참값 200 m/s 와 차이가 남'])
callout(s, 7.55, 4.45, 5.08, 4, '트랙 450: 두 스캔 전 (스캔 1152) 에 개시', ['연관이 없어 점수 2 → 1 → 0, 이번 스캔에 삭제'])
tb(s, 7.55, 5.32, 5.08, 0.28, '스캔 1154 의 트랙 표', size=11, bold=True)
TROWS = [('트랙 ID', 'TrackState', 'TrackScore', 'Assigned_Flag', '개시 스캔'), ('451', '1 임시', '2', '0', '1154'), ('450', '4 삭제', '0', '0', '1152')]
cw = [0.8, 1.1, 1.0, 1.18, 1.0]
tbl = s.shapes.add_table(len(TROWS), len(cw), Inches(7.55), Inches(5.62), Inches(sum(cw)), Inches(0.3 * len(TROWS))).table
for j, w in enumerate(cw):
    tbl.columns[j].width = Inches(w)
for i, row in enumerate(TROWS):
    tbl.rows[i].height = Inches(0.3)
    for j, t in enumerate(row):
        cell = tbl.cell(i, j)
        cell.margin_left = cell.margin_right = Inches(0.06); cell.margin_top = cell.margin_bottom = Inches(0.01)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid(); cell.fill.fore_color.rgb = rgb(NAVY if i == 0 else (F_ORANGE if i == 1 else F_GRAY))
        tf = cell.text_frame; tf.text = ''
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        add_runs(p, t, 10, WHITE if i == 0 else (ORANGE if i == 1 else '6B7280'), bold=(i == 0 or j <= 1))
tblpr = s.shapes[-1]._element.graphic.graphicData.tbl.tblPr
tblpr.set('firstRow', '0'); tblpr.set('bandRow', '0')
note(s, 's22')

# 23 수명과 스코어
s = content('04. 결과', '4.4 트랙 수명과 스코어 로직 확인', '트랙 355 개 (표적 1 117 · 표적 2 235 · 클러터끼리 3) 가 모두 같은 일생을 거칩니다: 개시 → 두 스캔 뒤 삭제. 확정 · 메모리 전환은 스코어 확인 창에 Flag 를 넣어 확인했습니다.')
# 트랙의 일생: 처음 10 초 동안 생긴 트랙 22 개 (make_figures.py fig_lifeline) 와 첫 줄 트랙 100 의 스캔별 변화
picture(s, 'fig_lifeline.png', 0.66, 2.45, w=7.45)
tb(s, 0.7, 5.34, 7.4, 0.26, '트랙 100 의 경우  (표적 2, 위 그림의 첫 줄)', size=11, bold=True)
LIFE = (('스캔 10', 'TrackScore 2', ['두 플롯으로 개시', '(10067 + 10074)'], 1), ('스캔 11', 'TrackScore 1', 'Assigned_Flag 0 → −1', 1),
        ('스캔 12', 'TrackScore 0', '0 이하라 삭제 상태 (4)', 4), ('스캔 13', '표에 없음', '처리 시작 때 표에서 뺌', 0))
cw = (7.45 - 3 * 0.22) / 4
for i, (scan_txt, score_txt, body, state) in enumerate(LIFE):
    xx, yy = 0.66 + i * (cw + 0.22), 5.66
    col = {1: ORANGE, 4: '6B7280', 0: LIGHT}[state]
    rect(s, xx, yy, cw, 1.1, WHITE, col, radius=0.08, lw=1.4)
    if state:
        rect(s, xx + 0.12, yy + 0.11, 0.2, 0.2, WHITE, col, lw=1.4)
        if state == 4:
            line(s, xx + 0.12, yy + 0.11, xx + 0.32, yy + 0.31, color=col, width=1.0)
            line(s, xx + 0.12, yy + 0.31, xx + 0.32, yy + 0.11, color=col, width=1.0)
    tb(s, xx + 0.42, yy + 0.07, cw - 0.5, 0.26, scan_txt, size=11, bold=True, anchor='m')
    tb(s, xx + 0.12, yy + 0.4, cw - 0.2, 0.26, score_txt, size=10.5, bold=True, color=col if state else SUB)
    tb(s, xx + 0.12, yy + 0.68, cw - 0.2, 0.36, body, size=9, color=SUB)
    if i < len(LIFE) - 1:
        line(s, xx + cw + 0.02, yy + 0.55, xx + cw + 0.2, yy + 0.55, color=LIGHT, width=1.2, arrow=True)
picture(s, 'ui_score_p2.png', 8.3, 2.5, w=4.33, border=True)
tb(s, 8.3, 2.5 + 4.33 * 602 / 751 + 0.06, 4.33, 0.7, ['스코어 확인 창: Flag 1 1 1 0 0 1 1', '임시 → 확정 (5) → 메모리 → 다시 확정 (MemoryScore 0)',
                                                      '트랙 ID 100 ~ 454 · 플롯 ID 10000 ~ 18119 (과제 범위 안)'], size=10, color=SUB)
note(s, 's23')

# 24 전체 궤적
s = content('04. 결과', '4.5 전체 궤적 (시뮬레이션 마지막 화면)', '스캔 1159 (115.9 s), 116 초를 끝까지 돌린 화면입니다. 두 표적의 참값 궤적 전체와 그동안 받은 플롯 · 클러터가 모두 그려져 있습니다.')
picture(s, 'ui_crop_full_2d.png', 0.7, 2.5, w=5.85, border=True)
picture(s, 'ui_crop_full_3d.png', 6.78, 2.5, w=5.85, border=True)
tb(s, 0.7, 2.5 + 5.85 * 765 / 1200 + 0.05, 5.85, 0.25, '2D: 위경도 격자 위 동-북 평면', size=10, color=SUB)
tb(s, 6.78, 2.5 + 5.85 * 765 / 1200 + 0.05, 5.85, 0.25, '3D: 동-북-고도 (고도 2 배)', size=10, color=SUB)
leg = [(FULL_T1, BLUE), (FULL_T2, RED), ('○ 시작점', NAVY), ('● 플롯', NAVY), ('· 클러터', '6B7280'), ('□ 임시 트랙', ORANGE)]
xx = 0.7
for t, c in leg:
    w = sum(0.16 if ord(ch) > 0x1100 else 0.085 for ch in t) + 0.38
    tb(s, xx, 6.56, w, 0.26, t, size=10.5, color=c, bold=True)
    xx += w
note(s, 's21full')

# 25 끝
s = add_slide('END')
tb(s, 4.6, 3.05, 7.4, 0.8, '감사합니다', size=36, bold=True, color=WHITE, align='r')
line(s, 8.6, 4.02, 12.0, 4.02, color=MUTED, width=0.75)
tb(s, 4.6, 4.16, 7.4, 0.34, '질문 주시면 편하게 답변드리겠습니다', size=13, color=LIGHT, align='r')
note(s, 'end')

prs.save(OUT)
print('saved', OUT, len(prs.slides), 'slides')

# 대본 문서
TITLES = []
for key, sl in slides_meta:
    t = ''
    for sh in sl.placeholders:
        if sh.placeholder_format.idx == 103:
            t = sh.text_frame.text
    if not t:
        texts = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and not sh.is_placeholder and sh.text_frame.text.strip()]
        t = {'cover': '표지', 'index': '목차', 'end': '마무리'}.get(key, texts[0] if texts else key)
    TITLES.append(t)
with open(SCRIPT_MD, 'w', encoding='utf-8') as f:
    f.write('# 과제 5 트랙 개시 · 관리 — 발표 대본\n\n')
    f.write('슬라이드 노트와 같은 내용입니다. 발표 시간은 20 ~ 25 분을 기준으로 썼습니다.\n\n')
    for i, ((key, sl), t) in enumerate(zip(slides_meta, TITLES), 1):
        f.write(f'## {i}. {t}\n\n{NOTES[key]}\n\n')
print('saved', SCRIPT_MD)
