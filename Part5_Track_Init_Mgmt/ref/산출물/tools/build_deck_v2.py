# -*- coding: utf-8 -*-
"""과제 5 발표자료 만들기 (2 판: 처음 보는 사람도 따라올 수 있게 그림 · 표 중심으로 다시 짬).

    python build_deck_v2.py <틀.pptx> <그림 폴더> <출력.pptx>
    (예: 틀 = ../../../../Part4_Plot_Clutter_Sim/ref/산출물/Chapter4_PlotClutterSim.pptx, 그림 = ../figures)
    환경변수 SRC 로 소스 폴더 (TargetSim) 를 바꿀 수 있음. 글꼴은 맑은 고딕이 없으면 Noto Sans CJK 로 폭을 잼.

틀은 과제 4 발표자료 (Chapter4_PlotClutterSim.pptx). 장은 모두 지우고 틀의 레이아웃
(COVER, INDEX, DIVIDER, CONTENT, END) 과 장 모양 (표지 · 목차 · 간지 글자, 머리띠, 쪽 번호) 만 가져온다.
표지와 끝 장 바탕 그림은 그림 폴더의 cover_bg.png 로 바꾼다.
코드 발췌는 src 의 실제 파일에서 줄을 찾아 읽음. 그래서 줄 번호가 늘 파일과 같다.
노트는 같은 폴더의 Chapter5_발표대본.md 에서 읽는다.
"""
import copy
import io
import os
import re
import sys

from lxml import etree
from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.normpath(os.path.join(HERE, '..'))
SRC = os.environ.get('SRC') or os.path.normpath(os.path.join(HERE, '..', '..', '..', 'src', 'TargetSim'))
SCRIPT_MD = os.path.join(OUT_DIR, 'Chapter5_발표대본.md')

KO = '맑은 고딕'
EN = 'Arial'
CODE = 'Consolas'
NAVY = '1B2740'
SUB = '4A5A7A'
MUTE = '8794AC'
LNUM = 'AEB9CC'
LINE = 'D2D9E5'
BOXF = 'F7F9FC'
WHITE = 'FFFFFF'
STRONG = ['DC2626', 'D97706', '059669', '2563EB', '7C3AED']
LIGHT = ['FDEEEE', 'FEF3E2', 'E6F5EE', 'EAF1FD', 'F2ECFD']
TENT = 'EA580C'

def _font(cands):
    for c in cands:
        try:
            return ImageFont.truetype(c, 100)
        except (OSError, IOError):
            continue
    raise RuntimeError('font not found: %s' % cands)


FONT_REG = _font([r'C:\Windows\Fonts\malgun.ttf', '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'])
FONT_BOLD = _font([r'C:\Windows\Fonts\malgunbd.ttf', '/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc'])


def text_width(text, size, bold=False):
    """글자 폭 [in]"""
    f = FONT_BOLD if bold else FONT_REG
    return f.getlength(text) / 100.0 * size / 72.0


# 노트 (발표 대본 md)

def read_notes(path):
    md = io.open(path, encoding='utf-8').read()
    out = {}
    for blk in re.split(r'(?m)^## ', md)[1:]:
        head = re.match(r'(\d+)\.\s+(.+)', blk)
        if not head:
            continue
        body = []
        for raw in blk.split('\n')[1:]:
            line = raw.strip()
            if not line or line.startswith(('---', '#', '|', '>')) or re.match(r'\[(.+?)\]$', line):
                continue
            body.append(line)
        out[int(head.group(1))] = '\n'.join(body)
    return out


# XML 도우미

def new_id(slide):
    ids = [int(e.get('id')) for e in slide.shapes._spTree.iter(qn('p:cNvPr'))]
    return max(ids + [1]) + 1


def clone(slide, proto):
    el = copy.deepcopy(proto)
    el.find('.//' + qn('p:cNvPr')).set('id', str(new_id(slide)))
    slide.shapes._spTree.append(el)
    return el


def el_set_text(el, text):
    """첫 단락 첫 run 의 글자만 바꾸고 나머지 run, 단락은 지움"""
    body = el.find(qn('p:txBody'))
    paras = body.findall(qn('a:p'))
    for p in paras[1:]:
        body.remove(p)
    runs = paras[0].findall(qn('a:r'))
    for r in runs[1:]:
        paras[0].remove(r)
    runs[0].find(qn('a:t')).text = text


def el_set_pos(el, x=None, y=None, w=None, h=None):
    off = el.find('.//' + qn('a:off'))
    ext = el.find('.//' + qn('a:ext'))
    if x is not None:
        off.set('x', str(int(Inches(x))))
    if y is not None:
        off.set('y', str(int(Inches(y))))
    if w is not None:
        ext.set('cx', str(int(Inches(w))))
    if h is not None:
        ext.set('cy', str(int(Inches(h))))


def strip_style(shape):
    st = shape._element.find(qn('p:style'))
    if st is not None:
        shape._element.remove(st)


def set_run(run, text, size, color=NAVY, bold=False, font=KO):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._r.get_or_add_rPr()
    for tag in ('a:latin', 'a:ea', 'a:cs'):
        e = rpr.find(qn(tag))
        if e is None:
            e = etree.SubElement(rpr, qn(tag))
        e.set('typeface', font)


def body_zero(tf, anchor='t', wrap=True):
    bp = tf._txBody.find(qn('a:bodyPr'))
    for k in ('lIns', 'tIns', 'rIns', 'bIns'):
        bp.set(k, '0')
    bp.set('wrap', 'square' if wrap else 'none')
    bp.set('anchor', {'t': 't', 'm': 'ctr', 'b': 'b'}[anchor])


def para_spacing(p, lnspc_pct=None, lnspc_pt=None, after_pt=None):
    ppr = p._p.get_or_add_pPr()
    if lnspc_pct is not None or lnspc_pt is not None:
        ln = etree.SubElement(ppr, qn('a:lnSpc'))
        if lnspc_pt is not None:
            etree.SubElement(ln, qn('a:spcPts')).set('val', str(int(lnspc_pt * 100)))
        else:
            etree.SubElement(ln, qn('a:spcPct')).set('val', str(int(lnspc_pct * 1000)))
    if after_pt is not None:
        sa = etree.SubElement(ppr, qn('a:spcAft'))
        etree.SubElement(sa, qn('a:spcPts')).set('val', str(int(after_pt * 100)))


def text(slide, x, y, w, h, paras, anchor='t', align='l', wrap=True, lnspc=None, after=None):
    """paras: [[(글자, 크기, 색, 굵게, 글꼴), ...], ...] 또는 단락마다 str"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    body_zero(tf, anchor, wrap)
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[align]
        if lnspc is not None or after is not None:
            para_spacing(p, lnspc_pct=lnspc, after_pt=after)
        if isinstance(para, str):
            para = [(para, 11, NAVY, False, KO)]
        for spec in para:
            t, size, color, bold, font = (list(spec) + [11, NAVY, False, KO][len(spec) - 1:])[:5]
            set_run(p.add_run(), t, size, color, bold, font)
    return tb


def rect(slide, x, y, w, h, fill=None, line=None, lw=0.75, shape=MSO_SHAPE.RECTANGLE, dash=None):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    strip_style(sp)
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = RGBColor.from_string(line)
        sp.line.width = Pt(lw)
        if dash:
            ln = sp._element.spPr.find(qn('a:ln'))
            etree.SubElement(ln, qn('a:prstDash')).set('val', dash)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sp.adjustments[0] = min(0.5, 0.08 / max(min(w, h), 0.01))
    return sp


def hline(slide, x0, y, x1, color=LINE, lw=0.75, dash=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x0), Inches(y), Inches(x1), Inches(y))
    strip_style(c)
    c.line.color.rgb = RGBColor.from_string(color)
    c.line.width = Pt(lw)
    if dash:
        ln = c._element.spPr.find(qn('a:ln'))
        etree.SubElement(ln, qn('a:prstDash')).set('val', dash)
    return c


def arrow_line(slide, x0, y0, x1, y1, color=MUTE, lw=1.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x0), Inches(y0), Inches(x1), Inches(y1))
    strip_style(c)
    c.line.color.rgb = RGBColor.from_string(color)
    c.line.width = Pt(lw)
    ln = c._element.spPr.find(qn('a:ln'))
    t = etree.SubElement(ln, qn('a:tailEnd'))
    t.set('type', 'triangle')
    t.set('w', 'med')
    t.set('len', 'med')
    return c


def badge(slide, n, x, y, d=0.28, color_idx=None):
    col = STRONG[(n - 1) % len(STRONG)] if color_idx is None else STRONG[color_idx]
    sp = rect(slide, x, y, d, d, fill=col, shape=MSO_SHAPE.OVAL)
    tf = sp.text_frame
    body_zero(tf, 'm')
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    set_run(p.add_run(), str(n), 9.5, WHITE, True)
    return sp


def callout(slide, n, x, y, w, title, lines, size=11.0, h=None):
    badge(slide, n, x, y)
    paras = [[(title, size, NAVY, True)]] + [[(l, size, SUB, False)] for l in lines]
    hh = h if h is not None else 0.24 * len(paras)
    return text(slide, x + 0.42, y - 0.02, w - 0.42, hh, paras, lnspc=105)


def picture(slide, path, x, y, w=None, h=None, border=False):
    pic = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)
    if border:
        pic.line.color.rgb = RGBColor.from_string(LINE)
        pic.line.width = Pt(0.75)
    return pic


def label(slide, x, y, w, title, sub=None):
    paras = [[(title, 11.5, NAVY, True)] + ([('    ' + sub, 9.5, MUTE, False, CODE)] if sub else [])]
    return text(slide, x, y, w, 0.24, paras, anchor='m')


# 코드 발췌

TYPES = {'VOID', 'INT32', 'UINT32', 'INT64', 'UINT16', 'FLOAT64', 'FLOAT32', 'DOUBLE64', 'UCHAR8', 'CHAR', 'matrix', 'const', 'static',
         'return', 'if', 'else', 'for', 'while', 'switch', 'case', 'break', 'default', 'typedef', 'struct', 'sizeof', 'SOCKET', 'HANDLE',
         'DWORD', 'ULONGLONG', 'TSCORE_API', '__stdcall', 'size_t'}
TOKEN = re.compile(r'(//.*$)|("[^"]*")|(\b\d+(?:\.\d+)?(?:[eE][-+]?\d+)?[UuLlFf]*\b)|([A-Za-z_][A-Za-z0-9_]*)|(\s+)|(.)')


def colorize(code):
    """C 한 줄 -> [(글자, 색, 글꼴)]"""
    out = []
    for m in TOKEN.finditer(code):
        com, s, num, ident, ws, other = m.groups()
        if com:
            out.append((com, '059669', KO))
        elif s:
            out.append((s, 'B45309', CODE))
        elif num:
            out.append((num, 'B45309', CODE))
        elif ident:
            if ident in TYPES or re.match(r'^(ST|STRUCT)_[A-Za-z]', ident) and not ident.isupper():
                col = '2563EB'
            elif re.match(r'^[A-Z][A-Z0-9_]{2,}$', ident) and ident not in TYPES:
                col = '7C3AED'
            else:
                col = NAVY
            out.append((ident, col, CODE))
        else:
            out.append((ws or other, NAVY, CODE))
    # 이어진 같은 색 · 글꼴 합치기
    merged = []
    for t, c, f in out:
        if merged and merged[-1][1] == c and merged[-1][2] == f:
            merged[-1] = (merged[-1][0] + t, c, f)
        else:
            merged.append((t, c, f))
    return merged


def detab(line):
    lead = len(line) - len(line.lstrip('\t'))
    rest = line[lead:]
    rest = re.sub(r'\t+', '  ', rest)
    return '  ' * lead + rest.rstrip()


def read_lines(rel):
    return io.open(os.path.join(SRC, rel), encoding='utf-8').read().split('\n')


def find(lines, needle, start=0):
    for i in range(start, len(lines)):
        if needle in lines[i]:
            return i
    raise KeyError(needle)


def excerpt(rel, parts, dedent=0):
    """parts: [(시작 글자, 끝 글자 또는 줄 수)] -> [(줄 번호, 글자)]. 조각 사이에 '...' 줄."""
    lines = read_lines(rel)
    out = []
    pos = 0
    for k, (a, b) in enumerate(parts):
        i0 = find(lines, a, pos)
        i1 = i0 + b - 1 if isinstance(b, int) else find(lines, b, i0)
        if k > 0 and out and out[-1][0] is not None and i0 > out[-1][0]:
            out.append((None, '...'))
        for i in range(i0, i1 + 1):
            t = detab(lines[i])
            if dedent and t.startswith(' ' * dedent):
                t = t[dedent:]
            out.append((i + 1, t))
        pos = i1 + 1
    return out


def annotate(rows, notes):
    """발표용 줄 주석. notes: [(찾을 글자, 주석)] 또는 [(찾을 글자, 주석, 'above')], 코드 순서대로.
    줄 끝에 '// 주석' 을 붙이고, 'above' 면 그 줄 위에 주석만 있는 줄 (번호 없음) 을 넣음.
    찾을 글자는 탭을 공백 두 칸으로 바꾼 뒤의 글자로 찾음."""
    out = list(rows)
    pos = 0
    for note in notes:
        needle, com = note[0], note[1]
        i = row_of(out, needle, pos)
        num, code = out[i]
        if len(note) > 2 and note[2] == 'above':
            ind = len(code) - len(code.lstrip(' '))
            out.insert(i, (None, ' ' * ind + '// ' + com))
            i += 1
        else:
            out[i] = (num, code + '  // ' + com)
        pos = i + 1
    return out


def row_width(lab, code, size):
    """코드 한 줄의 폭 [in]. Consolas 글자 폭 0.5498 em, 주석 (맑은 고딕) 은 글꼴로 잼"""
    wd = len(lab) * 0.5498 * size / 72.0
    for t, c, f in colorize(code):
        wd += len(t) * 0.5498 * size / 72.0 if f == CODE else FONT_REG.getlength(t) / 100.0 * size / 72.0
    return wd


def code_box(slide, x, y, w, rows, bands=(), size=9.0, step=11.5, markers=True, pad=0.09):
    """rows: [(줄 번호, 글자)], bands: [(첫 행, 행 수, 색 번호, 표시 번호)]
    PowerPoint 가 고정 줄 간격을 1 pt 단위로 맞춤. 그래서 step 도 정수 pt 로 맞춰야 띠가 줄과 어긋나지 않음"""
    step = float(round(step))
    n = len(rows)
    h = pad * 2 + n * step / 72.0
    rect(slide, x, y, w, h, fill=WHITE, line=LINE)
    for (r0, cnt, ci, mark) in bands:
        by = y + pad + r0 * step / 72.0 - 0.012
        bh = cnt * step / 72.0
        rect(slide, x + 0.03, by, w - 0.06, bh, fill=LIGHT[ci])
        rect(slide, x + 0.03, by, 0.04, bh, fill=STRONG[ci])
        if markers and mark:
            badge(slide, mark, x + w - 0.11, by + min(bh, step / 72.0) / 2 - 0.11, d=0.22, color_idx=ci)
    tb = slide.shapes.add_textbox(Inches(x + 0.1), Inches(y + pad), Inches(w - 0.25), Inches(h - 2 * pad))
    tf = tb.text_frame
    body_zero(tf, 't', wrap=False)
    width = max(len(str(r[0])) for r in rows if r[0] is not None)
    badge_rows = {b[0] for b in bands if markers and b[3]}
    for i, (num, code) in enumerate(rows):
        room = w - 0.25 - (0.26 if i in badge_rows else 0.0)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para_spacing(p, lnspc_pt=step)
        lab = (str(num).rjust(width) if num is not None else ' ' * width) + '  '
        set_run(p.add_run(), lab, size, LNUM, False, CODE)
        if num is None and code.strip() == '...':
            set_run(p.add_run(), code, size, MUTE, False, CODE)
        else:
            for t, c, f in colorize(code):
                set_run(p.add_run(), t, size, c, False, f)
        wd = row_width(lab, code, size)
        if wd > room:
            print('  [폭 넘침] %.2f > %.2f in : %s' % (wd, room, code.strip()[:70]))
    return h


def row_of(rows, needle, start=0):
    for i in range(start, len(rows)):
        if rows[i][0] is not None and needle in rows[i][1]:
            return i
    raise KeyError(needle)


# 덱

class Deck:
    def __init__(self, template, cover_png):
        self.prs = Presentation(template)
        self.lay = {l.name: l for l in self.prs.slide_layouts}
        sl = list(self.prs.slides)
        g = lambda s, name: [sh._element for sh in sl[s - 1].shapes if sh.name == name][0]
        self.proto = {
            'cover': {n: g(1, n) for n in ('Text 0', 'Text 1', 'Text 2', 'Shape 3', 'Text 4', 'Text 5')},
            'index': {n: g(2, n) for n in ('Text 0', 'Text 1', 'Text 2', 'Text 3', 'Shape 4', 'Text 5', 'Shape 6')},
            'index_num': g(2, 'Slide Number Placeholder 0'),
            'div': {n: g(3, n) for n in ('Text 0', 'Text 1', 'Text 2', 'Text 3', 'Shape 4', 'Text 5')},
            'div_num': g(3, 'Slide Number Placeholder 0'),
            'content_ph': [sh._element for sh in sl[14].placeholders],
            'content_num': g(15, 'Slide Number Placeholder 0'),
            'end': {n: g(24, n) for n in ('Text 0', 'Shape 1', 'Text 2')},
        }
        self.proto = copy.deepcopy(self.proto)
        # 장 모두 지우기
        lst = self.prs.slides._sldIdLst
        for sid in list(lst):
            self.prs.part.drop_rel(sid.rId)
            lst.remove(sid)
        # 표지 · 끝 바탕 그림 바꾸기
        blob = open(cover_png, 'rb').read()
        for name in ('COVER', 'END'):
            lay = self.lay[name]
            for rel in lay.part.rels.values():
                if rel.reltype.endswith('/image'):
                    rel.target_part._blob = blob
        self.notes = NOTES

    def _new(self, layout):
        s = self.prs.slides.add_slide(self.lay[layout])
        for ph in list(s.placeholders):
            ph._element.getparent().remove(ph._element)
        return s

    def _notes(self, s):
        n = len(self.prs.slides)
        if n in self.notes:
            s.notes_slide.notes_text_frame.text = self.notes[n][1].strip()

    def cover(self, kicker, title, en, author):
        s = self._new('COVER')
        P = self.proto['cover']
        clone(s, P['Text 0'])
        el_set_text(clone(s, P['Text 1']), kicker)
        el_set_text(clone(s, P['Text 2']), title)
        clone(s, P['Shape 3'])
        el_set_text(clone(s, P['Text 4']), en)
        el_set_text(clone(s, P['Text 5']), author)
        self._notes(s)
        return s

    def index(self, left_title, left_sub, author, items):
        s = self._new('INDEX')
        P = self.proto['index']
        e = clone(s, P['Text 0'])
        el_set_text(e, left_title)
        body = e.find(qn('p:txBody'))
        p2 = copy.deepcopy(body.findall(qn('a:p'))[0])
        p2.findall(qn('a:r'))[0].find(qn('a:t')).text = left_sub
        body.append(p2)
        el_set_text(clone(s, P['Text 1']), author)
        y0, dy = 1.97, 0.84
        for i, (num, name, page) in enumerate(items):
            y = y0 + i * dy
            el_set_pos(el_set_text_ret(clone(s, P['Text 2']), num), y=y)
            nw = text_width(name, 15, True) + 0.45
            el_set_pos(el_set_text_ret(clone(s, P['Text 3']), name), y=y, w=nw)
            el_set_pos(clone(s, P['Shape 4']), x=6.40 + nw + 0.12, y=y + 0.33, w=12.0 - (6.40 + nw + 0.12) - 0.05)
            el_set_pos(el_set_text_ret(clone(s, P['Text 5']), str(page)), y=y)
            if i < len(items) - 1:
                el_set_pos(clone(s, P['Shape 6']), y=y + 0.72)
        clone(s, self.proto['index_num'])
        self._notes(s)
        return s

    def divider(self, title, desc, items):
        s = self._new('DIVIDER')
        P = self.proto['div']
        el_set_text(clone(s, P['Text 0']), title)
        el_set_text(clone(s, P['Text 1']), desc)
        for i, (num, name, page) in enumerate(items):
            y = 2.25 + i * 0.445
            el_set_pos(el_set_text_ret(clone(s, P['Text 2']), num), y=y)
            nw = text_width(name, 14) + 0.45
            el_set_pos(el_set_text_ret(clone(s, P['Text 3']), name), y=y, w=nw)
            el_set_pos(clone(s, P['Shape 4']), x=1.55 + nw + 0.12, y=y + 0.28, w=11.60 - (1.55 + nw + 0.12))
            el_set_pos(el_set_text_ret(clone(s, P['Text 5']), str(page)), y=y)
        clone(s, self.proto['div_num'])
        self._notes(s)
        return s

    def content(self, section, title, subtitle, path=None, tag=None):
        """tag: 이 장이 답하는 과제 항목. 머리띠 오른쪽 끝에 씀"""
        s = self._new('CONTENT')
        if text_width(subtitle, 12.5) > 11.43:
            print('  [부제 넘침] %.2f in : %s' % (text_width(subtitle, 12.5), title))
        for el, t in zip(self.proto['content_ph'], (section, title, subtitle)):
            el_set_text(clone(s, el), t)
        if tag:
            text(s, 8.45, 0.62, 4.0, 0.62, [[('과제  ', 10, '8FA3C4', False), (tag, 10.5, 'D5DEEC', True)]], anchor='m', align='r')
        if path:
            text(s, 7.30, 1.33, 5.33, 0.48, [[(path, 10, MUTE, False, CODE)]], anchor='m', align='r')
        clone(s, self.proto['content_num'])
        self._notes(s)
        return s

    def end(self, title, sub):
        s = self._new('END')
        P = self.proto['end']
        el_set_text(clone(s, P['Text 0']), title)
        clone(s, P['Shape 1'])
        el_set_text(clone(s, P['Text 2']), sub)
        self._notes(s)
        return s

    def save(self, path):
        self.prs.save(path)


def el_set_text_ret(el, t):
    el_set_text(el, t)
    return el


def table(slide, x, y, w, col_w, rows, head_size=9.5, size=9.0, row_h=0.27):
    nr, nc = len(rows), len(rows[0])
    gt = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
    tbl = gt.table
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i in range(nr):
        tbl.rows[i].height = Inches(row_h)
        for j in range(nc):
            c = tbl.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.06)
            c.margin_top = c.margin_bottom = Inches(0.02)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            c.fill.fore_color.rgb = RGBColor.from_string(NAVY if i == 0 else (WHITE if i % 2 else BOXF))
            tf = c.text_frame
            val = rows[i][j]
            spec = val if isinstance(val, tuple) else (val,)
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            set_run(p.add_run(), spec[0], head_size if i == 0 else size, WHITE if i == 0 else (spec[1] if len(spec) > 1 else NAVY),
                    i == 0 or (len(spec) > 2 and spec[2]), spec[3] if len(spec) > 3 else KO)
    return gt


def grid_table(slide, x, y, col_w, row_h, cells):
    """칸마다 바탕색과 여러 줄 글자를 따로 주는 표. cells[i][j] = (바탕색, [(글자, 크기, 색, 굵게, 글꼴), ...]), row_h 는 행마다"""
    nr, nc = len(cells), len(cells[0])
    gt = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(sum(col_w)), Inches(sum(row_h)))
    tbl = gt.table
    tbl.first_row = False
    tbl.horz_banding = False
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i in range(nr):
        tbl.rows[i].height = Inches(row_h[i])
        for j in range(nc):
            fill, paras = cells[i][j]
            c = tbl.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.05)
            c.margin_top = c.margin_bottom = Inches(0.02)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            c.fill.fore_color.rgb = RGBColor.from_string(fill)
            tf = c.text_frame
            for k, spec in enumerate(paras):
                p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                p.alignment = PP_ALIGN.CENTER
                t, size, color, bold, font = (list(spec) + [10, NAVY, False, KO][len(spec) - 1:])[:5]
                set_run(p.add_run(), t, size, color, bold, font)
    return gt


def check_fit(paras, w):
    """한 줄 글자 폭이 상자 폭을 넘는지 확인 (넘으면 알림만)"""
    for para in paras:
        wd = sum(text_width(sp[0], sp[1], len(sp) > 3 and sp[3]) for sp in para)
        if wd > w:
            print('  [글자 넘침] %.2f > %.2f in : %s' % (wd, w, ''.join(sp[0] for sp in para)[:60]))




# 장 번호. 목차 · 간지의 쪽 번호와 노트가 모두 여기를 봄
PG = dict(cover=1, index=2, glance=3, d1=4, s11=5, s12=6, s13=7, s14=8, s15=9, d2=10, s21=11, s22=12, s23=13, s24=14,
          d3=15, s31=16, s32=17, s33=18, s34=19, s35=20, s36=21, s37=22, d4=23, s41=24, s42=25, s43=26, wrap=27, end=28)


def chip(slide, x, y, w, h, title, col, fill, lines, title_size=12, size=10.5):
    """색 테두리 상자: 굵은 제목 + 설명 줄"""
    rect(slide, x, y, w, h, fill=fill, line=col, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    paras = [[(title, title_size, col, True)]] + [[(l, size, NAVY)] for l in lines]
    text(slide, x + 0.16, y + 0.08, w - 0.32, h - 0.16, paras, anchor='m', lnspc=112)


def num_step(slide, n, x, y, w, title, lines, size=10.5, color_idx=None):
    badge(slide, n, x, y, color_idx=color_idx)
    paras = [[(title, 11.5, NAVY, True)]] + [[(l, size, SUB)] for l in lines]
    return text(slide, x + 0.42, y - 0.02, w - 0.42, 0.26 * len(paras), paras, lnspc=108)


def build(template, fig, out):
    F = lambda n: os.path.join(fig, n)
    d = Deck(template, F('cover_bg.png'))
    sec0, sec1, sec2, sec3, sec4 = '00. 과제 개요', '01. 트랙 개시', '02. 트랙 관리', '03. 구현', '04. 결과'

    # 1 표지, 2 목차
    d.cover('레이다 SW 학습 과제 5', '트랙 개시 · 관리', 'Track Initiation and Management', 'Janghwan Kim (Harley)')
    d.index('트랙 개시 · 관리', 'TrackMgr', 'Janghwan Kim (Harley)',
            [('00', '과제 개요', PG['glance']), ('01', '트랙 개시', PG['d1']), ('02', '트랙 관리', PG['d2']), ('03', '구현', PG['d3']), ('04', '결과', PG['d4'])])

    # 3 과제 한눈에 보기
    s = d.content(sec0, '이번 과제 한눈에 보기', '레이다가 잰 점 (플롯) 을 이어 표적의 움직임 (트랙) 을 처음 만들고, 그 트랙을 점수로 관리하는 기능을 구현했습니다.',
                  tag='1 · 2 · 3')
    text(s, 0.70, 2.42, 4.4, 0.28, [[('이번 과제에서 하는 일 세 가지', 12, NAVY, True)]], anchor='m')
    jobs = [('① 트랙 개시', '2563EB', 'EAF1FD', ['연속 두 스캔의 플롯 두 개로 트랙을 만든다', '2-Point Track Initiation, 속도 게이트 800 m/s']),
            ('② 트랙 관리', '059669', 'E6F5EE', ['트랙마다 점수를 매겨 확정 · 메모리 · 삭제를 정한다', '트랙 스코어 방식, 문턱 5 / 0 / 5']),
            ('③ 시뮬레이션', 'D97706', 'FEF3E2', ['플롯 메시지를 UDP 로 보내고 받아 처리한다', '2D · 3D 화면에 참값 · 플롯 · 트랙을 함께 표시'])]
    for i, (t, col, fc, lines) in enumerate(jobs):
        chip(s, 0.70, 2.78 + i * 1.22, 4.4, 1.08, t, col, fc, lines, title_size=12.5, size=10.5)
    text(s, 5.40, 2.42, 7.2, 0.28, [[('과제 요구사항과 구현한 곳', 12, NAVY, True), ('   (오른쪽 끝은 이 자료의 장)', 9.5, SUB)]], anchor='m')
    B = lambda t: (t, NAVY, True)
    rows = [('과제 항목', '요구 사항', '구현 (소스)', '장'),
            ('트랙 개시 방식', '2 Point Track Init.', 'f_Trk_TwoPointInit (TrackMgr.c)', '1.2 ~ 1.3, 3.6'),
            ('상태변수', '[x x\' x\'\' y y\' y\'\' z z\' z\'\']  9 개, 가속도 0', 'mat_EstX 9 x 1,  mat_EstP 9 x 9', '1.3, 3.6'),
            ('속도 게이팅', 'V_max = 800 m/s', 'f_Trk_ProcessScan', '1.4, 3.5'),
            ('트랙 ID', '100 ~ 10000', 'f_Trk_NewId', '3.5'),
            ('트랙 관리', '트랙 스코어 방식, 문턱 5 / 0 / 5', 'f_Trk_UpdateState', '2.3, 3.7'),
            ('플롯 메시지', 'STRUCT_PlotMsg, 안테나 구좌표', 'Idd.h, 송신 · 수신 쓰레드 (UDP)', '3.1 ~ 3.4'),
            ('플롯 ID', '10000 ~ 1000000, 표적 + 클러터', 'f_Lnk_BuildMsg (PlotLink.c)', '3.3'),
            ('트랙 연관', 'Assigned_Flag 매 스캔 0 (실패)', 'f_Trk_ProcessScan', '2.1, 3.5'),
            ('결과 화면', 'True · Plot · Track 심볼 2D · 3D', 'TargetSimUI (MFC)', '4.1 ~ 4.3')]
    table(s, 5.40, 2.78, 7.23, [1.35, 2.75, 2.13, 1.0], rows, head_size=9.5, size=9.5, row_h=0.345)

    # 4 간지 01
    d.divider('01. 트랙 개시', '트랙이 무엇인지부터 보고, 플롯 두 개로 트랙을 시작하는 2-Point 개시와 속도 게이트를 설명합니다.',
              [('1.1', '트랙이란', PG['s11']), ('1.2', '1-Point 와 2-Point 개시', PG['s12']), ('1.3', '2-Point 개시의 계산', PG['s13']),
               ('1.4', '속도 게이팅', PG['s14']), ('1.5', '측정 오차를 VC-NED 로 옮기기', PG['s15'])])

    # 5  1.1
    s = d.content(sec1, '1.1 트랙이란', '레이다는 스캔마다 점 (플롯) 을 찍어 줄 뿐입니다. 같은 표적에서 나온 점을 시간 순서로 이어 위치와 속도를 붙인 것이 트랙입니다.',
                  tag='1 · 2 개요')
    text(s, 0.70, 2.42, 5.0, 0.28, [[('레이다 데이터 처리 흐름에서 이번 과제의 자리', 12, NAVY, True)]], anchor='m')
    chain = [('표적 궤적 모의', '과제 3 · 시각별 참값 위치', 0), ('측정 모의', '과제 4 · 플롯 · 클러터 생성', 0), ('트랙 개시 · 관리', '과제 5 · 이번 과제', 1),
             ('연관 · 추적 필터', '다음 과제 · 플롯을 트랙에 붙여 추정', 2)]
    for i, (t, sub, kind) in enumerate(chain):
        x = 0.70 + i * 3.08
        rect(s, x, 2.78, 2.69, 0.66, fill='EEF4FF' if kind == 1 else BOXF, line='2563EB' if kind == 1 else LINE, lw=1.25 if kind == 1 else 0.75)
        text(s, x, 2.85, 2.69, 0.30, [[(t, 12.5, '7A879F' if kind == 2 else NAVY, True)]], align='c', anchor='m')
        text(s, x, 3.14, 2.69, 0.24, [[(sub, 10.5, '8794AC' if kind != 1 else SUB, False)]], align='c', anchor='m')
        if i < 3:
            arrow_line(s, x + 2.72, 3.11, x + 3.05, 3.11)
    hline(s, 0.70, 3.66, 12.63)
    text(s, 0.70, 3.76, 6.9, 0.28, [[('스캔마다 들어오는 플롯이 트랙이 되기까지', 11.5, NAVY, True)]], anchor='m')
    picture(s, F('fig_track_concept.png'), 0.70, 4.08, w=6.9)
    tags = [('스캔', '레이다가 전체 영역을 한 번 훑는 것. 이번 과제는 0.1 초마다 한 번', '374151'),
            ('플롯', '레이다가 표적을 한 번 잰 결과. 거리 · 방위각 · 고각', '1F2937'),
            ('클러터', '표적이 아닌 것에서 나온 플롯. 표적 플롯과 섞여 들어온다', '6B7280'),
            ('트랙 개시', '같은 표적으로 보이는 플롯으로 트랙을 처음 만드는 것', '2563EB'),
            ('트랙 관리', '트랙이 진짜 표적인지 판단해 확정하거나 지우는 것', '16A34A')]
    for i, (t, sub, col) in enumerate(tags):
        y = 4.08 + i * 0.52
        sp = rect(s, 7.95, y, 0.92, 0.30, fill=WHITE, line=col, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        body_zero(sp.text_frame, 'm')
        sp.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        set_run(sp.text_frame.paragraphs[0].add_run(), t, 10.5, col, True)
        text(s, 9.02, y + 0.04, 3.6, 0.3, [[(sub, 10.5, NAVY)]], anchor='t')

    # 6  1.2
    s = d.content(sec1, '1.2 1-Point 와 2-Point 개시', '플롯 하나로 바로 시작할 수도 있고, 연속 두 스캔 (또는 드웰) 의 플롯 두 개를 기다려 시작할 수도 있습니다. 차이는 속도를 아느냐입니다.',
                  tag='1.1 · 1.2 · 3-1')
    picture(s, F('fig_1p2p.png'), 0.70, 2.42, w=7.2)
    rows = [('', '1-Point 개시', '2-Point 개시  (이번 과제)'),
            ('쓰는 플롯', '1 개 (이번 스캔)', '2 개 (앞 스캔 + 이번 스캔)'),
            ('초기 속도', '모름 -> 0 으로 두고 오차를 크게', '두 위치의 차 / 스캔 간격'),
            ('장점', '플롯이 오자마자 바로 개시 (빠름)', '속도를 알아 초기 속도 오차가 작음'),
            ('단점', '초기 속도 오차가 큼', '두 스캔을 기다려야 해 개시가 한 스캔 늦음')]
    table(s, 8.15, 2.50, 4.48, [0.85, 1.72, 1.91], rows, head_size=9.5, size=9.5, row_h=0.46)
    rect(s, 8.15, 5.05, 4.48, 1.35, fill='EEF4FF', line='2563EB', lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, 8.32, 5.12, 4.15, 1.22, [[('이번 과제: 2-Point 개시', 11.5, '2563EB', True)],
                                      [('두 플롯이 정말 같은 표적인지는 속도 게이트 (1.4) 로 먼저 거르고,', 10, NAVY)],
                                      [('통과한 쌍으로만 트랙을 만듭니다.', 10, NAVY)],
                                      [('드웰 = 빔 하나로 한 방향을 잠깐 비추는 단위. 여기서는 스캔 = 드웰', 9, SUB)]], anchor='m', lnspc=112)

    # 7  1.3
    s = d.content(sec1, '1.3 2-Point 개시의 계산: 초기 상태변수와 오차 공분산', '트랙은 상태변수 X (위치 · 속도 · 가속도) 와 그 값이 얼마나 틀릴 수 있는지 적은 오차 공분산 P 를 갖습니다. 둘 다 플롯 두 개로 계산합니다.',
                  tag='1.2 · 3-1')
    PB, PG_, PO, PZ = 'DCE7FD', 'DDF3EA', 'FDECD7', 'F3F4F7'
    badge(s, 1, 0.70, 2.47, d=0.26)
    text(s, 1.04, 2.44, 5.8, 0.32, [[('상태변수 X', 11.5, NAVY, True), ('   세 축 (x, y, z) 마다 위치 · 속도 · 가속도 = 9 개', 9.5, SUB)]], anchor='m')
    text(s, 0.70, 2.80, 6.15, 0.24, [[('예 (x 축)   앞 스캔 플롯 5,000 m,  이번 스캔 플롯 5,030 m,  스캔 간격 T = 0.1 s', 9.5, SUB)]], anchor='m')
    C = lambda t, col=NAVY, b=False: (t, col, b, EN)
    st = [('성분', '식', '예 (x 축)'),
          ('위치  x', C('z_k  (이번 플롯 그대로)'), C('5,030 m')),
          ('속도  x\'', C('( z_k − z_k-1 ) / T'), C('( 5,030 − 5,000 ) / 0.1 = 300 m/s')),
          ('가속도  x\'\'', ('0  (과제 조건)', NAVY), C('0'))]
    table(s, 0.70, 3.08, 6.15, [1.25, 2.25, 2.65], st, head_size=9.5, size=10, row_h=0.29)
    badge(s, 2, 0.70, 4.47, d=0.26)
    text(s, 1.04, 4.44, 5.8, 0.32, [[('오차 공분산 P', 11.5, NAVY, True), ('   한 축의 3 x 3  (예: 플롯 오차 ±10 m, T = 0.1 s)', 9.5, SUB)]], anchor='m')
    H = lambda t: ('2B3A55', [(t, 9.5, WHITE, True)])
    H2 = lambda t, sub: ('2B3A55', [(t, 10, WHITE, True), (sub, 8.5, 'C9D3E6', False)])
    F2 = lambda fill, f, v: (fill, [(f, 10.5, NAVY, True, EN), (v, 9.5, SUB, False, EN)])
    Z = (PZ, [('0', 10, MUTE, False, EN)])
    pc = [[H(''), H('위치'), H('속도'), H('가속도')],
          [H2('위치', '= 이번 플롯'), F2(PB, 'R_k', '100  (±10 m)'), F2(PG_, 'R_k / T', '1,000'), Z],
          [H2('속도', '= 두 플롯 차 / T'), F2(PG_, 'R_k / T', '1,000'), F2(PO, '( R_k + R_k-1 ) / T²', '20,000  (±141 m/s)'), Z],
          [H2('가속도', '= 0'), Z, Z, Z]]
    grid_table(s, 0.70, 4.82, [1.40, 1.45, 2.20, 1.10], [0.27, 0.56, 0.56, 0.44], pc)
    badge(s, 3, 7.15, 2.47, d=0.26)
    text(s, 7.49, 2.44, 5.1, 0.32, [[('9 x 9 전체 P', 11.5, NAVY, True), ('   ② 의 3 x 3 블록이 축마다 대각선에 하나씩', 9.5, SUB)]], anchor='m')
    picture(s, F('fig_pmatrix.png'), 7.15, 2.80, w=4.6)
    rect(s, 7.15, 6.18, 5.48, 0.52, fill='FEF3E2', line='D97706', lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, 7.30, 6.18, 5.2, 0.52, [[('눈여겨볼 점  ', 10, 'D97706', True), ('위치 오차 ±10 m 가 0.1 초로 나뉘면서 속도 오차 ±141 m/s 가 됩니다.', 10, NAVY)],
                                    [('R: 플롯 위치의 측정 오차 (분산). 구하는 법은 1.5', 9, SUB)]], anchor='m', lnspc=110)

    # 8  1.4
    s = d.content(sec1, '1.4 속도 게이팅', '표적이 아무리 빨라도 한 스캔 동안 갈 수 있는 거리는 정해져 있습니다. 두 플롯 사이 거리를 시간으로 나눈 속력이 V_max 보다 작을 때만 같은 표적으로 봅니다.',
                  tag='속도 게이팅 · 3-1')
    rect(s, 0.70, 2.45, 11.93, 0.62, fill=BOXF, line=LINE)
    text(s, 0.95, 2.45, 5.6, 0.62, [[('| z_(k,i)  −  z_(k-1,j) |  /  ΔT   <   V_max', 15, NAVY, True, EN)]], anchor='m')
    text(s, 6.40, 2.45, 6.1, 0.62, [[('V_max = 800 m/s,  ΔT = 0.1 s  ->  앞 플롯에서 반경 80 m 안이면 통과', 11.5, '2563EB', True)],
                                    [('i: 이번 스캔 플롯,  j: 앞 스캔 플롯.  거리는 VC-NED 직교좌표 (북 · 동 · 아래) 로 잰다', 10, SUB)]], anchor='m')
    picture(s, F('fig_gate.png'), 0.70, 3.25, w=11.93)

    # 9  1.5
    s = d.content(sec1, '1.5 측정 오차를 VC-NED 로 옮기기', '플롯은 거리 · 방위각 · 고각으로 들어오지만 트랙은 북 · 동 · 아래 (VC-NED) 좌표로 만듭니다. 그래서 측정 오차도 같은 좌표로 바꿔야 합니다.',
                  tag='1.2 · 3-1')
    steps = [('측정 표준편차 (과제 4, SNR 20)', 'σ_r = 4.74 m,   σ_az = 0.478°,   σ_el = 0.494°', 1),
             ('안테나 직교좌표의 오차로', 'R_ant = J · diag( σ_r², σ_az², σ_el² ) · Jᵀ', 2),
             ('VC-NED 의 오차로', 'R = M · R_ant · Mᵀ', 3)]
    for i, (t, f, n) in enumerate(steps):
        y = 2.50 + i * 0.95
        badge(s, n, 0.74, y + 0.05)
        text(s, 1.16, y, 6.2, 0.30, [[(t, 11.5, NAVY, True)]], anchor='m')
        rect(s, 1.16, y + 0.33, 6.2, 0.44, fill=BOXF, line=LINE)
        text(s, 1.30, y + 0.33, 6.0, 0.44, [[(f, 12.5, NAVY, False, EN)]], anchor='m')
    text(s, 1.16, 5.35, 6.2, 1.1, [[('J', 10.5, NAVY, True), (': 구좌표 (r, az, el) -> 안테나 직교 (x, y, z) 변환의 미분 (야코비안)', 10.5, SUB)],
                                   [('M', 10.5, NAVY, True), (': 안테나 -> 동체 -> VC-NED 회전 (플랫폼 Roll · Pitch · Yaw)', 10.5, SUB)],
                                   [('VC-NED', 10.5, NAVY, True), (': 플랫폼 무게중심이 원점, 축은 북 · 동 · 아래로 고정. 플랫폼이 기울어도 축이 같이 돌지 않는다', 10.5, SUB)]], lnspc=115)
    picture(s, F('fig_rcov.png'), 7.55, 2.42, w=5.08)

    # 10 간지 02
    d.divider('02. 트랙 관리', '만든 트랙이 진짜 표적인지 가려 확정하거나 지우는 방법입니다. 트랙의 네 가지 상태와 연관을 보고, M/N 방식과 이번에 쓴 트랙 스코어 방식을 비교합니다.',
              [('2.1', '트랙 상태와 연관', PG['s21']), ('2.2', 'M/N 방식', PG['s22']), ('2.3', '트랙 스코어 방식', PG['s23']), ('2.4', '점수가 바뀌는 과정 (예)', PG['s24'])])

    # 11  2.1
    s = d.content(sec2, '2.1 트랙 상태와 연관', '트랙은 임시로 시작해 확정, 메모리, 삭제 상태를 오갑니다. 상태를 바꾸는 근거는 스캔마다 플롯이 트랙에 붙었는지 (연관 결과, Assigned_Flag) 입니다.',
                  tag='2 · 3-2 · 3-3')
    label(s, 0.70, 2.44, 6.2, '연관이란')
    picture(s, F('fig_assoc.png'), 0.70, 2.74, w=6.2)
    rect(s, 0.70, 5.45, 6.2, 1.08, fill=BOXF, line=LINE)
    text(s, 0.88, 5.45, 5.9, 1.08, anchor='m', paras=[[('연관  ', 10.5, NAVY, True), ('이번 스캔 플롯 가운데 하나를 기존 트랙에 이어 붙이는 것', 10.5, NAVY)],
                                    [('Assigned_Flag  ', 10.5, NAVY, True), ('트랙마다 스캔마다 남기는 연관 결과. 1 성공, 0 실패', 10.5, NAVY)],
                                    [('이번 과제  ', 10.5, '2563EB', True), ('연관 기능은 다음 과제라 매 스캔 0 으로 둔다 (과제 조건)', 10.5, NAVY)]], lnspc=135)
    label(s, 7.25, 2.44, 5.38, '트랙 상태 네 가지', 'TrackState')
    states = [('1  임시', TENT, 'FFF1E6', '막 만든 트랙. 표적인지 아직 확인하는 중', '클러터 두 개가 우연히 이어진 트랙일 수도 있다'),
              ('2  확정', '16A34A', 'E8F7EE', '연관이 이어져 표적으로 인정한 트랙', '보통 운용자 화면에 보이는 것은 확정 트랙'),
              ('3  메모리', '9333EA', 'F5EBFF', '확정 트랙이 이번 스캔에 플롯을 놓친 상태', '바로 지우지 않고 다시 잡히기를 기다린다'),
              ('4  삭제', '6B7280', 'F1F2F4', '더 쫓지 않기로 한 트랙', '다음 스캔을 시작할 때 트랙 표에서 지운다')]
    for i, (name, col, fc, l1, l2) in enumerate(states):
        y = 2.78 + i * 1.0
        rect(s, 7.25, y, 5.38, 0.86, fill=fc, line=col, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        text(s, 7.42, y, 1.15, 0.86, [[(name, 12, col, True)]], anchor='m')
        text(s, 8.55, y, 4.0, 0.86, [[(l1, 10.5, NAVY, True)], [(l2, 10, SUB)]], anchor='m', lnspc=115)

    # 12  2.2
    s = d.content(sec2, '2.2 M/N 방식', '최근 N 번 가운데 연관에 M 번 성공하면 확정하고, M 번 연속 실패하면 지우는 방식입니다. 횟수만 세면 되어 단순합니다.', tag='2.1')
    picture(s, F('fig_mn.png'), 0.70, 2.45, w=6.85)
    num_step(s, 1, 7.75, 2.60, 4.88, '임시 트랙', ['N 번 중 M 번 트랙 연관에 성공 -> 확정', 'M 번 연속 결합 실패 -> 삭제'])
    num_step(s, 2, 7.75, 3.65, 4.88, '확정 트랙', ['트랙 연관 실패 -> 메모리 트랙으로', '정한 시간 이상 메모리 상태 -> 삭제'])
    num_step(s, 3, 7.75, 4.70, 4.88, '특징', ['성공 · 실패 횟수만 세면 되어 단순하다', '대신 최근 N 번의 성공 · 실패 기록을 트랙마다 들고 있어야 한다'])
    rect(s, 7.75, 5.75, 4.88, 0.62, fill=BOXF, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, 7.90, 5.75, 4.6, 0.62, [[('이번 과제는 과제 조건대로 다음 장의 트랙 스코어 방식을 썼습니다.', 10, NAVY)]], anchor='m')

    # 13  2.3
    s = d.content(sec2, '2.3 트랙 스코어 방식  (이번 과제)', '트랙마다 점수 두 개 (TrackScore, MemoryScore) 를 두고, 연관 성공 · 실패에 따라 점수를 올리고 내리다가 문턱을 넘으면 상태를 바꿉니다.',
                  tag='2.2 · 3-2')
    picture(s, F('fig_state.png'), 0.70, 2.45, w=7.0)
    num_step(s, 1, 7.95, 2.55, 4.68, 'TrackState (과제 값)', ['1 임시,  2 확정,  3 메모리,  4 삭제'])
    num_step(s, 2, 7.95, 3.25, 4.68, '문턱 (과제 값)', ['임시 트랙 확정 5,  임시 트랙 삭제 0', '메모리 트랙 삭제 5'])
    num_step(s, 3, 7.95, 4.20, 4.68, '메모리 트랙이 다시 연관되면', ['바로 확정 트랙으로 되돌리고 MemoryScore = 0'])
    num_step(s, 4, 7.95, 4.90, 4.68, '과제에 없어 정한 값', ['개시 직후 TrackScore = 2 (플롯 두 개가 이미 이어졌으므로)', '한 번에 올리고 내리는 양 = 1'])
    rect(s, 7.95, 5.85, 4.68, 0.58, fill=BOXF, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, 8.10, 5.85, 4.4, 0.58, [[('M/N 과 달리 정수 두 개만 있으면 같은 판단을 할 수 있습니다.', 10, NAVY)]], anchor='m')

    # 14  2.4
    s = d.content(sec2, '2.4 점수가 바뀌는 과정 (예)', '왼쪽은 연관 결과를 직접 넣어 본 예, 오른쪽은 이번 과제 조건 (매 스캔 실패) 입니다. 맨 위가 스캔마다의 Assigned_Flag, 맨 아래가 그때 점수가 바뀐 계산입니다.',
                  tag='3-2 · 3-3')
    picture(s, F('fig_score.png'), 0.70, 2.42, w=11.93)

    # 15 간지 03
    d.divider('03. 구현', '쓰레드 · UDP 구조와 플롯 메시지를 보고, 송신 · 수신 코드와 트랙 개시 · 관리 코드를 차례로 봅니다.',
              [('3.1', '전체 구조: 쓰레드와 UDP', PG['s31']), ('3.2', '플롯 메시지 STRUCT_PlotMsg', PG['s32']), ('3.3', '송신: 플롯 메시지 만들기', PG['s33']),
               ('3.4', '수신: 좌표 변환과 트랙 처리', PG['s34']), ('3.5', '2-Point 개시: 쌍 고르기', PG['s35']),
               ('3.6', '2-Point 초기화', PG['s36']), ('3.7', '트랙 스코어 로직', PG['s37'])])

    # 16  3.1
    s = d.content(sec3, '3.1 전체 구조: 쓰레드와 UDP', '과제 4 의 측정 결과를 신호처리가 보낸 플롯이라고 보고, 송신 쓰레드가 스캔마다 UDP 로 보내면 수신 쓰레드가 받아 트랙을 처리합니다.',
                  tag='3-3')
    picture(s, F('fig_arch.png'), 0.70, 2.42, w=11.93)
    text(s, 0.70, 5.27, 11.93, 0.30, [[('왜 TCP 가 아니라 UDP 인가', 11.5, NAVY, True), ('   과제는 UDP 또는 TCP/IP 중 하나를 고르게 함', 9.5, SUB)]], anchor='m')
    U = lambda t: (t, '1D4ED8', False)
    ut = [('', 'UDP  (이번 과제)', 'TCP'),
          ('메시지 단위', U('보낸 메시지 1 개 = 받는 메시지 1 개 (스캔 1 개)'), '바이트가 이어서 와서, 메시지를 직접 잘라야 함'),
          ('지연', U('연결 · 재전송 없이 바로 전달'), '하나가 늦으면 뒤 스캔도 함께 기다림'),
          ('손실', U('생길 수 있음 -> 스캔 번호로 확인 (1,160 / 1,160, 빠짐 0)'), '없음 (잃으면 다시 보냄)')]
    table(s, 0.70, 5.60, 11.93, [1.55, 5.40, 4.98], ut, head_size=10, size=10.5, row_h=0.29)

    # 17  3.2  플롯 메시지: 그림으로
    s = d.content(sec3, '3.2 플롯 메시지 STRUCT_PlotMsg', '스캔 하나의 플롯을 메시지 하나에 담습니다. 보내는 쪽은 머리와 안테나 구좌표만 채우고, 나머지 좌표 칸은 받는 쪽이 계산해 채웁니다.',
                  'TargetSimCore/include/Idd.h', tag='3-3')
    # 머리
    hx, hy, hw = 0.70, 2.50, 3.55
    rect(s, hx, hy, hw, 3.96, fill=WHITE, line='2563EB', lw=1.25)
    rect(s, hx, hy, hw, 0.36, fill='2563EB')
    text(s, hx, hy, hw, 0.36, [[('메시지 머리  (스캔마다 1 개)', 11, WHITE, True)]], anchor='m', align='c')
    head = [('MsgId / MsgSize', '메시지 종류 · 보낼 바이트 수'), ('ScanNum', '스캔 번호'), ('TimeStamp_s / _ns', '스캔 시각'),
            ('PlotNum', '플롯 수 (표적 + 클러터)'), ('Sigma_Range / Azimuth / Elevation', '측정 표준편차 -> 받는 쪽 R 계산'),
            ('st_Platform', '플랫폼 위치 (LLA, ECEF) 와 자세 -> 좌표 변환')]
    for i, (a, b) in enumerate(head):
        y = hy + 0.46 + i * 0.57
        text(s, hx + 0.15, y, hw - 0.3, 0.26, [[(a, 10, NAVY, True, CODE)]], anchor='m')
        text(s, hx + 0.15, y + 0.25, hw - 0.3, 0.26, [[(b, 9.5, SUB)]], anchor='m')
        if i < len(head) - 1:
            hline(s, hx + 0.12, y + 0.54, hx + hw - 0.12)
    # 플롯 배열
    px, py, pw = 4.50, 2.50, 3.95
    rect(s, px, py, pw, 3.96, fill=WHITE, line='D97706', lw=1.25)
    rect(s, px, py, pw, 0.36, fill='D97706')
    text(s, px, py, pw, 0.36, [[('플롯 st_Plot[ ]  (PlotNum 개, 최대 64)', 11, WHITE, True)]], anchor='m', align='c')
    grp = [('Id, ScanNum, TimeStamp, BeamType', '플롯 ID, 스캔 번호, 탐지 시각, 빔 종류', 'DC2626', 'FDEEEE', '송신이 채움'),
           ('Range / Azimuth / Elevation_Ant', '안테나 기준 구좌표 (과제 조건)', 'DC2626', 'FDEEEE', '송신이 채움'),
           ('Range / Azimuth / Elevation_VcNed', 'VC-NED 기준 구좌표', '2563EB', 'EAF1FD', '수신이 계산'),
           ('X / Y / Z_Ant,  X / Y / Z_Body', '안테나 · 동체 직교좌표', '2563EB', 'EAF1FD', '수신이 계산'),
           ('X / Y / Z_VcNed', 'VC-NED 직교좌표 (북 · 동 · 아래) -> 트랙 개시에 씀', '2563EB', 'EAF1FD', '수신이 계산')]
    for i, (a, b, col, fc, who) in enumerate(grp):
        y = py + 0.48 + i * 0.68
        rect(s, px + 0.12, y, pw - 0.24, 0.60, fill=fc, line=col, lw=0.75)
        text(s, px + 0.22, y + 0.03, pw - 1.2, 0.28, [[(a, 8.8, NAVY, True, CODE)]], anchor='m')
        text(s, px + 0.22, y + 0.30, pw - 1.2, 0.28, [[(b, 9, SUB)]], anchor='m')
        text(s, px + pw - 1.05, y, 0.85, 0.60, [[(who, 9, col, True)]], anchor='m', align='r')
    # 오른쪽 설명
    num_step(s, 1, 8.80, 2.55, 3.83, '보낼 때는 필요한 만큼만', ['머리 + 플롯 PlotNum 개 = MsgSize 바이트', '7 개 플롯이면 919 B (최대 64 개)'], size=10)
    num_step(s, 2, 8.80, 3.45, 3.83, '송신은 빨간 칸만 채움', ['과제 조건: 안테나 구좌표 성분으로 구성', '표적 플롯과 클러터가 섞여 있고 구분 표시 없음'], size=10)
    num_step(s, 3, 8.80, 4.35, 3.83, '수신이 파란 칸을 채움', ['플랫폼 자세로 안테나 -> 동체 -> VC-NED', '측정 표준편차로 오차 공분산 R 계산 (1.5)'], size=10)
    num_step(s, 4, 8.80, 5.25, 3.83, '참고 소스의 Idd.h 를 이어 씀', ['Y_Body 에서 끊겨 있던 STRUCT_PlotInfo 를', 'Z_Body, VC-NED 직교까지 채우고 STRUCT_PlotMsg 추가'], size=10)
    text(s, 8.80, 6.15, 3.83, 0.3, [[('#pragma pack(1): 칸 사이 빈 공간 없이 붙여서 보냄', 9, SUB)]], anchor='m')

    # 18  3.3
    s = d.content(sec3, '3.3 송신: 플롯 메시지 만들기 (f_Lnk_BuildMsg)', '표적 플롯과 클러터를 한 배열에 섞어 방위각 순서로 담고, 플롯마다 ID 를 매겨 안테나 구좌표만 채웁니다.',
                  'TargetSimCore/src/PlotLink.c', tag='3-3')
    rows = excerpt('TargetSimCore/src/PlotLink.c', [
        ('static UINT32 f_Lnk_BuildMsg', 1),
        ('for (nIndex = 0; (nIndex < st_Scan->nTargetNum)', 4),
        ('for (nIndex = 0; (nIndex < st_Scan->nClutterNum)', 4),
        ('// 방위각 오름차순', 1),
        ('st_Msg->MsgId', 'st_Msg->PlotNum'),
        ('for (nIndex = 0; nIndex < nPlotNum; nIndex++)', 4),
        ('st_Info->Range_Ant', 'st_Info->Elevation_Ant'),
        ('// 플롯 ID 는 10000', '*pt_PlotId = ')])
    rows = annotate(rows, [('f_Lnk_BuildMsg(', '스캔 하나의 메시지 만들기'), ('nTargetNum)', '표적 플롯부터'),
                           ('st_Scan->st_Plot[nIndex];', '거리 · 방위각 · 고각'), ('= nIndex + 1;', '출처: 표적 번호 (송신 쪽 기록용)'),
                           ('nClutterNum)', '이어서 클러터를 같은 배열에'), ('st_Scan->st_Clutter[nIndex];', '클러터 플롯'), ('= 0;', '출처: 0 은 클러터'),
                           ('st_Msg->MsgId', '메시지 종류: 플롯'), ('st_Msg->MsgSize', '보낼 바이트 수'), ('st_Msg->ScanNum', '스캔 번호'),
                           ('st_Msg->TimeStamp_s', '스캔 시각 (초)'), ('st_Msg->TimeStamp_ns', '스캔 시각 (나노초)'), ('st_Msg->PlotNum', '표적 + 클러터 수'),
                           ('nIndex < nPlotNum;', '정렬된 순서대로 플롯마다'), ('&st_Msg->st_Plot[nIndex];', '메시지의 플롯 칸'), ('st_Info->Id', '플롯 ID'),
                           ('st_Info->Range_Ant', '안테나 기준 거리'), ('st_Info->Azimuth_Ant', '방위각'), ('st_Info->Elevation_Ant', '고각'),
                           ('*pt_PlotId = ', '다음 플롯 ID')])
    b1 = row_of(rows, 'nTargetNum)')
    b2 = row_of(rows, '방위각 오름차순')
    b3 = row_of(rows, 'st_Msg->MsgId')
    b4 = row_of(rows, 'st_Info->Id')
    b5 = row_of(rows, 'Range_Ant')
    b6 = row_of(rows, '// 플롯 ID 는')
    b7 = row_of(rows, 'nClutterNum)')
    code_box(s, 0.70, 2.45, 7.85, rows, bands=[(b1, 4, 0, 1), (b7, 4, 0, None), (b2, 1, 1, 2), (b3, 6, 2, 3), (b4, 1, 3, 4), (b5, 3, 3, None), (b6, 2, 4, 5)],
             size=7.6, step=9)
    num_step(s, 1, 8.85, 2.50, 3.78, '표적 플롯 + 클러터를 한 배열에', ['표적 플롯 뒤에 클러터를 이어 붙임', '출처 (표적 k / 클러터) 는 송신 쪽에만 기록'], size=10)
    num_step(s, 2, 8.85, 3.45, 3.78, '방위각 순서로 정렬', ['빔이 지나가는 순서처럼 보이게', '배열 위치만으로는 표적을 알 수 없게'], size=10)
    num_step(s, 3, 8.85, 4.40, 3.78, '메시지 머리', ['스캔 번호, 시각, 플롯 수'], size=10)
    num_step(s, 4, 8.85, 5.05, 3.78, '플롯 ID 와 안테나 구좌표', ['거리 · 방위각 · 고각만 채움 (과제 조건)'], size=10)
    num_step(s, 5, 8.85, 5.70, 3.78, '플롯 ID 10000 ~ 1000000', ['하나씩 올리고, 1000000 다음은 10000 으로'], size=10)

    # 19  3.4
    s = d.content(sec3, '3.4 수신: 좌표 변환과 트랙 처리 (f_Lnk_ProcessMsg)', '받은 플롯을 VC-NED 로 바꾸고 측정 공분산 R 을 구한 뒤, 트랙 관리와 2-Point 개시를 한 번에 돌립니다.',
                  'TargetSimCore/src/PlotLink.c', tag='3-3')
    rows = excerpt('TargetSimCore/src/PlotLink.c', [
        ('static VOID f_Lnk_ProcessMsg', 1),
        ('const FLOAT64\t\t\ttime = ', 1),
        ('st_Sigma.r\t= st_Msg->Sigma_Range;', 'st_Att.Yaw'),
        ('for (nPlot = 0; nPlot < nPlotNum; nPlot++)', 4),
        ('st_Plot[nPlot].pos[0]', 'st_Plot[nPlot].pos[2]'),
        ('f_Trk_MeasCov(st_Plot[nPlot].cov', 1),
        ('f_Trk_MoveOrigin(&s_Link.st_Table', 1),
        ('f_Trk_ProcessScan(&s_Link.st_Table', 1)])
    rows = annotate(rows, [('f_Lnk_ProcessMsg(', '받은 메시지 하나 처리'), ('time = ', '스캔 시각 [s]'), ('st_Sigma.r', '측정 표준편차: 거리'),
                           ('st_Sigma.az', '방위각'), ('st_Sigma.el', '고각'), ('st_Att.Roll', '플랫폼 자세'),
                           ('nPlot < nPlotNum;', '플롯마다'), ('&st_Msg->st_Plot[nPlot];', '받은 플롯'), ('f_Lnk_FillPlotCoord(', '좌표 변환 (아래 세 줄)'),
                           ('pos[0]', '북'), ('pos[1]', '동'), ('pos[2]', '아래'), ('f_Trk_MeasCov(', '측정 공분산 R (1.5)'),
                           ('f_Trk_MoveOrigin(', '앞 스캔 값을 이번 원점 기준으로'), ('f_Trk_ProcessScan(', '트랙 관리 + 개시')])
    i1 = row_of(rows, 'st_Sigma.r')
    i2 = row_of(rows, 'f_Lnk_FillPlotCoord(')
    i3 = row_of(rows, 'pos[0]')
    i4 = row_of(rows, 'f_Trk_MeasCov(')
    i5 = row_of(rows, 'f_Trk_MoveOrigin(')
    i6 = row_of(rows, 'f_Trk_ProcessScan(')
    hm = code_box(s, 0.70, 2.45, 7.85, rows, bands=[(i1, 6, 0, 1), (i2, 1, 1, 2), (i3, 3, 2, 3), (i4, 1, 2, None), (i5, 1, 4, 4), (i6, 1, 3, 5)],
                  size=7.9, step=9)
    rowsc = excerpt('TargetSimCore/src/PlotLink.c', [('st_Ant\t= f_Trans_Ant_Sph_To_XYZ(st_Info->Range_Ant', 'st_Ned\t= f_Trans_Body_To_Ned(st_Body.x')], dedent=2)
    rowsc = annotate(rowsc, [('st_Ant', '구좌표 -> 안테나 직교'), ('st_Body', '안테나 -> 동체'), ('st_Ned', '동체 -> VC-NED (플랫폼 자세)')])
    label(s, 0.70, 2.45 + hm + 0.08, 7.85, '2  f_Lnk_FillPlotCoord 의 변환 세 줄', 'PlotLink.c')
    code_box(s, 0.70, 2.45 + hm + 0.36, 11.93, rowsc, size=7.4, step=9)
    num_step(s, 1, 8.85, 2.50, 3.78, '머리에서 꺼냄', ['시각, 측정 표준편차, 플랫폼 자세'], size=10)
    num_step(s, 2, 8.85, 3.15, 3.78, '안테나 -> 동체 -> VC-NED', ['아래 세 줄. 지난 과제의 좌표변환 함수 그대로'], size=10)
    num_step(s, 3, 8.85, 3.80, 3.78, '트랙 개시용 플롯', ['위치는 VC-NED, 공분산은 f_Trk_MeasCov (1.5)'], size=10)
    num_step(s, 4, 8.85, 4.45, 3.78, '원점 옮기기', ['플랫폼이 움직이면 앞 스캔 플롯 · 트랙을', '이번 스캔 VC-NED 로 (이번 시나리오는 정지, 0)'], size=10)
    num_step(s, 5, 8.85, 5.35, 3.78, '트랙 관리 + 2-Point 개시', ['결과는 스캔 단위로 쌓아 두고 UI 가 가져감'], size=10)

    # 20  3.5
    s = d.content(sec3, '3.5 2-Point 개시: 쌍 고르기 (f_Trk_ProcessScan)', '먼저 있던 트랙을 Assigned_Flag 0 으로 판정하고, 앞 스캔 플롯과 이번 스캔 플롯의 모든 쌍 가운데 게이트를 통과한 가장 느린 쌍부터 트랙으로 만듭니다.',
                  'TargetSimCore/src/TrackMgr.c', tag='3-1 · 3-3')
    rows = excerpt('TargetSimCore/src/TrackMgr.c', [
        ('st_Track->Assigned_Flag\t= 0;', 2),
        ('while (isSearching != 0)', 3),
        ('for (nPrev = 0; nPrev < st_Table->nPrevPlotNum; nPrev++)', 'bestCur\t\t= nCur;'),
        ('st_Table->st_PrevPlot[bestPrev].isUsed\t= 1;', 2),
        ('newId = (st_Table->nTrackNum < TRK_MAX_TRACK)', 1),
        ('f_Trk_TwoPointInit(st_Track, &st_Table->st_PrevPlot[bestPrev]', 1),
        ('st_Track->TrackId\t\t= newId;', 'st_Track->TrackScore\t= TRK_SCORE_INIT;'),
        ('st_Track->PlotId[0]', 'st_Track->PlotId[1]')], dedent=2)
    rows = annotate(rows, [('Assigned_Flag', '연관 기능이 없어 늘 실패 (과제 조건)'), ('f_Trk_UpdateState(', '스코어로 상태 판정 (3.7)'),
                           ('while (isSearching', '고를 쌍이 없을 때까지 반복'), ('TRK_VMAX;', 'V_max 에서 시작 = 속도 게이트'),
                           ('for (nPrev', '앞 스캔 플롯'), ('for (nCur', '이번 스캔 플롯 (모든 쌍)'), ('isUsed == 0) && (', '아직 안 쓴 플롯끼리만'),
                           ('f_Trk_PairSpeed(', '|z_k - z_k-1| / T'), ('if (speed < bestSpeed)', '지금까지 가장 느린 쌍이면'),
                           ('bestSpeed  = speed;', '그 쌍을 기억'), ('[bestPrev].isUsed', '고른 두 플롯은 다시 쓰지 않음'),
                           ('f_Trk_NewId(', '트랙 ID (100 ~ 10000)'), ('f_Trk_TwoPointInit(', '초기 상태변수 · 공분산 (3.6)'),
                           ('TrackState', '임시 트랙 (1)'), ('TrackScore', '개시 스코어 2'), ('PlotId[0]', '개시에 쓴 플롯 ID')])
    a1 = row_of(rows, 'Assigned_Flag')
    a2 = row_of(rows, 'TRK_VMAX;')
    a3 = row_of(rows, 'for (nPrev')
    a4 = row_of(rows, '[bestPrev].isUsed')
    a5 = row_of(rows, 'f_Trk_NewId')
    a6 = row_of(rows, 'f_Trk_TwoPointInit')
    a7 = row_of(rows, 'TrackId')
    code_box(s, 0.70, 2.45, 7.85, rows, bands=[(a1, 2, 0, 1), (a2, 1, 1, 2), (a3, 13, 1, None), (a4, 2, 2, 3), (a5, 1, 3, 4), (a6, 1, 4, 5),
                                               (a7, 3, 4, None)], size=7.6, step=9)
    num_step(s, 1, 8.85, 2.50, 3.78, '트랙 관리 먼저', ['연관이 없어 모든 트랙 Assigned_Flag = 0', 'f_Trk_UpdateState 로 판정 (3.7)'], size=10)
    num_step(s, 2, 8.85, 3.45, 3.78, '속도 게이트', ['bestSpeed 를 V_max 로 시작하면', '800 m/s 이상인 쌍은 처음부터 후보가 안 됨'], size=10)
    num_step(s, 3, 8.85, 4.40, 3.78, '쓴 플롯 표시', ['한 플롯은 한 쌍에만. 가장 느린 쌍부터 차례로'], size=10)
    num_step(s, 4, 8.85, 5.05, 3.78, '트랙 ID', ['100 부터 차례로, 10000 다음은 100', '쓰는 중인 ID 는 건너뜀'], size=10)
    num_step(s, 5, 8.85, 5.95, 3.78, '초기화 (3.6)', ['임시 (1), TrackScore 2, 개시 플롯 ID 2 개 기록'], size=10)

    # 21  3.6
    s = d.content(sec3, '3.6 2-Point 초기화 (f_Trk_TwoPointInit)', '9 x 1 상태변수와 9 x 9 오차 공분산을 0 으로 만든 뒤, 축마다 위치 · 속도 칸만 채웁니다. 가속도 칸은 손대지 않아 0 으로 남습니다 (과제 조건).',
                  'TargetSimCore/src/TrackMgr.c', tag='3-1')
    rows = excerpt('TargetSimCore/src/TrackMgr.c', [('VOID f_Trk_TwoPointInit(', 'st_P->e[(3 * nAxis) + 1][(3 * nAxisB) + 1]')])
    rows.append((rows[-1][0] + 1, '    }'))
    rows.append((rows[-1][0] + 1, '  }'))
    rows.append((rows[-1][0] + 1, '}'))
    rows = annotate(rows, [('f_Trk_TwoPointInit(', '앞 플롯, 이번 플롯'), ('dt = ', '스캔 간격 T [s]'), ('*st_X = ', '상태변수 X (9 x 1)'),
                           ('*st_P = ', '오차 공분산 P (9 x 9)'), ('nAxis;', '축 a'), ('nAxisB;', '축 b'),
                           ('Matrix_initialize(st_X', '9 x 1, 모두 0'), ('Matrix_initialize(st_P', '9 x 9, 모두 0'),
                           ('for (nAxis = 0', '축 a: 0 북, 1 동, 2 아래'), ('st_X->e[(3 * nAxis)][0]', '위치 = 이번 플롯'),
                           ('st_X->e[(3 * nAxis) + 1][0]', '속도 = 위치 차 / T'), ('st_X->e[(3 * nAxis) + 2][0]', '가속도 = 0'),
                           ('for (nAxisB = 0', '축 b'), ('st_P->e[(3 * nAxis)][(3 * nAxisB)]', '위치 · 위치 = R_k'),
                           ('st_P->e[(3 * nAxis)][(3 * nAxisB) + 1]', '위치 · 속도 = R_k / T'), ('st_P->e[(3 * nAxis) + 1][(3 * nAxisB)]', '속도 · 위치 = R_k / T'),
                           ('st_P->e[(3 * nAxis) + 1][(3 * nAxisB) + 1]', '속도 · 속도 = (R_k + R_k-1) / T²', 'above')])
    c1 = row_of(rows, 'Matrix_initialize(st_X')
    c2 = row_of(rows, 'st_X->e[(3 * nAxis)][0]')
    c3 = row_of(rows, 'st_P->e[(3 * nAxis)][(3 * nAxisB)]')
    code_box(s, 0.70, 2.45, 7.85, rows, bands=[(c1, 2, 0, 1), (c2, 3, 1, 2), (c3, 5, 2, 3)], size=7.6, step=10)
    num_step(s, 1, 8.85, 2.50, 3.78, '0 으로 채운 행렬', ['참고 소스 코드의 Matrix_initialize', '가속도 칸은 손대지 않아 0'], size=10)
    num_step(s, 2, 8.85, 3.45, 3.78, '상태변수 (축 a 마다)', ['위치 = 이번 플롯  (3a 번째 칸)', '속도 = 위치 차 / 시간 차  (3a+1)', '가속도 = 0  (3a+2)'], size=10)
    num_step(s, 3, 8.85, 4.60, 3.78, '오차 공분산 (1.3 의 3 x 3 표 그대로)', ['위치 · 위치 = R_k', '위치 · 속도 = R_k / T', '속도 · 속도 = ( R_k + R_k-1 ) / T²'], size=10)
    text(s, 8.85, 5.75, 3.78, 0.6, [[('코드의 3 * nAxis 는 축 a 의 위치 칸, + 1 은 속도 칸, + 2 는 가속도 칸입니다.', 9.5, SUB)]], lnspc=110)

    # 22  3.7
    s = d.content(sec3, '3.7 트랙 스코어 로직 (f_Trk_UpdateState)', '2.3 의 상태 그림을 그대로 코드로 옮긴 것입니다. 상태마다 case 하나씩이고, 임시는 TrackScore 로, 확정 · 메모리는 MemoryScore 로 다음 상태를 정합니다.',
                  'TargetSimCore/src/TrackMgr.c', tag='3-2')
    rows = excerpt('TargetSimCore/src/TrackMgr.c', [('VOID f_Trk_UpdateState(', 'case TRK_STATE_CONFIRMED:')])
    rows = rows[:-1]
    rows = annotate(rows, [('f_Trk_UpdateState(', '스캔마다 트랙 하나씩'), ('isAssigned = ', '연관 결과 1 / 0'),
                           ('switch (', '지금 상태에 따라'), ('case TRK_STATE_TENTATIVE:', '임시 트랙'),
                           ('TrackScore + ', '성공 +1, 실패 -1', 'above'), ('>= TRK_SCORE_CONFIRM', '5 이상이면'), ('= TRK_STATE_CONFIRMED;', '확정'),
                           ('<= TRK_SCORE_DELETE', '0 이하면'), ('= TRK_STATE_DELETED;', '삭제')])
    d1 = row_of(rows, 'TrackScore + ')
    d2 = row_of(rows, '>= TRK_SCORE_CONFIRM')
    d3 = row_of(rows, '<= TRK_SCORE_DELETE')
    label(s, 0.70, 2.44, 6.0, '임시 트랙')
    code_box(s, 0.70, 2.72, 6.05, rows, bands=[(d1 - 1, 2, 0, 1), (d2, 4, 2, 2), (d3, 4, 4, None)], size=7.6, step=9)
    rows2 = excerpt('TargetSimCore/src/TrackMgr.c', [('case TRK_STATE_CONFIRMED:', 'default:')])
    rows2 = rows2[:-1]
    rows2 = annotate(rows2, [('case TRK_STATE_CONFIRMED:', '확정 트랙'), ('isAssigned == 0', '연관 실패면'), ('= TRK_STATE_MEMORY;', '메모리로'),
                             ('MemoryScore + 1;', 'MemoryScore +1'), ('case TRK_STATE_MEMORY:', '메모리 트랙'), ('isAssigned != 0', '연관 성공이면'),
                             ('= TRK_STATE_CONFIRMED;', '바로 확정'), ('MemoryScore  = 0;', 'MemoryScore 0 으로'), ('else', '연관 실패면'),
                             ('MemoryScore + 1;', 'MemoryScore +1'), ('>= TRK_MEMORY_DELETE', '5 이상이면'), ('= TRK_STATE_DELETED;', '삭제')])
    e1 = row_of(rows2, 'TRK_STATE_MEMORY;')
    e3 = row_of(rows2, '= 0;')
    e4 = row_of(rows2, '>= TRK_MEMORY_DELETE')
    label(s, 6.95, 2.44, 5.7, '확정 트랙, 메모리 트랙')
    code_box(s, 6.95, 2.72, 5.68, rows2, bands=[(e1, 2, 4, 3), (e3 - 1, 2, 2, 4), (e4, 4, 3, 5)], size=7.6, step=9)
    y = 6.42
    marks = [(1, '성공 +1, 실패 -1'), (2, '5 이상 확정, 0 이하 삭제'), (3, '확정 -> 메모리, MemoryScore +1'), (4, '메모리 -> 바로 확정, MemoryScore 0'), (5, 'MemoryScore 5 이상 삭제')]
    for k, (n, t) in enumerate(marks):
        x = 0.70 + k * 2.40
        badge(s, n, x, y, d=0.22)
        text(s, x + 0.28, y - 0.03, 2.1, 0.5, [[(t, 9.5, NAVY)]], anchor='t')

    # 23 간지 04
    d.divider('04. 결과', '시뮬레이션 조건과 2D · 3D 결과 화면, 트랙이 생기고 사라지는 과정, 거리에 따라 개시 결과가 어떻게 달라지는지 봅니다.',
              [('4.1', '시뮬레이션 조건과 결과 화면 (True · Plot · Track)', PG['s41']), ('4.2', '트랙 수명과 스코어 확인', PG['s42']),
               ('4.3', '거리에 따른 개시 결과', PG['s43'])])

    # 24  4.1
    s = d.content(sec4, '4.1 시뮬레이션 조건과 결과 화면', '1,160 스캔을 UDP 로 모두 받아 트랙 355 개를 개시했습니다 (같은 표적 플롯 352, 클러터끼리 3). 아래는 스캔 546 의 참값 · 플롯 · 트랙 화면입니다.',
                  tag='3-3')
    conds = [('시나리오', '플랫폼 정지 (32°N, 126°E)', '표적 1 대함: 30 m/s, 약 14 km, 서쪽으로', '표적 2 대공: 200 m/s, 고도 2,000 m, 13.4 -> 4.8 km 로 접근'),
             ('측정 (과제 4)', '116 s, 스캔 간격 0.1 s = 1,160 스캔', 'SNR 20, 스캔당 클러터 5 개', '플롯 7 개 / 스캔, 메시지 919 B'),
             ('UDP 송수신', '보냄 1,160 / 받음 1,160 / 손실 0', '트랙 개시 355, 삭제 상태 353 (끝에 2 개 남음)', '클러터끼리 만든 트랙 3 개는 모두 레이다 350 m 안쪽')]
    for i, (t, a, b, c) in enumerate(conds):
        x = 0.70 + i * 4.02
        rect(s, x, 2.42, 3.88, 0.84, fill=BOXF, line=LINE)
        text(s, x + 0.12, 2.42, 1.0, 0.84, [[(t, 10.5, NAVY, True)]], anchor='m')
        text(s, x + 1.12, 2.42, 2.7, 0.84, [[(a, 9, NAVY)], [(b, 9, NAVY)], [(c, 9, NAVY)]], anchor='m', lnspc=108)
    picture(s, F('res_view.png'), 0.70, 3.38, w=11.93, border=True)
    text(s, 0.70, 6.52, 11.93, 0.3, [[('표시 규칙은 프로그램 화면과 같습니다.  선: 참값 궤적,  마름모: 지금 위치,  점: 플롯 (회색은 클러터),  빈 주황 사각형: 임시 트랙,  회색 X: 삭제 트랙,  사각형에서 뻗은 선: 트랙 속도로 1 초 동안 가는 거리', 9, SUB)]], anchor='m')

    # 25  4.2
    s = d.content(sec4, '4.2 트랙 수명과 스코어 확인', '연관이 늘 실패인 과제 조건에서는 모든 트랙이 TrackScore 2 -> 1 -> 0 으로 내려가 세 스캔 만에 삭제됩니다. 확정 · 메모리는 스코어 확인 창으로 검증했습니다.',
                  tag='3-2 · 3-3')
    label(s, 0.70, 2.44, 6.9, '스캔 545 의 트랙 표', '프로그램 결과 표의 값')
    T2 = lambda t, col: (t, col, True)
    tt = [('트랙 ID', 'TrackState', 'TrackScore', 'MemoryScore', 'Assigned_Flag', '개시 스캔', '개시 플롯 (앞, 이번)', '참값'),
          ('240', T2('4 삭제', '6B7280'), '0', '0', '0', '543', '13796, 13804', '표적 2'),
          ('241', T2('1 임시', TENT), '1', '0', '0', '544', '13805, 13814', '표적 1'),
          ('242', T2('1 임시', TENT), '2', '0', '0', '545', '13813, 13819', '표적 2')]
    table(s, 0.70, 2.74, 6.9, [0.65, 0.85, 0.85, 0.95, 1.0, 0.75, 1.25, 0.6], tt, head_size=8.5, size=9.5, row_h=0.30)
    num_step(s, 1, 0.75, 4.05, 6.85, '한 스캔에 나이가 다른 트랙 세 개', ['242: 이번 스캔에 개시, TrackScore 2  /  241: 한 스캔 전 개시, 1', '240: 두 스캔 전 개시, 0 -> 삭제 (4). 다음 스캔에 표에서 빠짐'], size=10)
    picture(s, F('res_life.png'), 0.70, 4.95, w=6.9)
    label(s, 7.95, 2.44, 4.68, '스코어 확인 창', '연관 결과를 직접 넣어 봄')
    picture(s, F('score.png'), 7.95, 2.74, h=3.05, border=True)
    num_step(s, 2, 8.00, 5.95, 4.6, '1 1 1 0 0 1 1 을 넣으면', ['3 번째 스캔에 확정, 두 번 실패해 메모리, 다시 성공하면', '바로 확정 (MemoryScore 0). 같은 f_Trk_UpdateState 를 부름'], size=10)

    # 26  4.3
    s = d.content(sec4, '4.3 거리에 따른 개시 결과', '멀리 있는 표적일수록 플롯이 크게 흔들려 같은 표적의 두 플롯도 80 m 게이트를 넘기 쉽습니다. 그래서 먼 표적은 트랙이 덜 만들어집니다.',
                  tag='1.2 · 속도 게이팅')
    picture(s, F('fig_rate.png'), 0.70, 2.50, w=6.9)
    tab = [('표적', '거리 [km]', '스캔 쌍', '게이트 통과', '개시 트랙', '개시 비율'),
           ('2', '4.8 ~ 6.5', '305', '45 %', '96', '31 %'),
           ('2', '6.5 ~ 8.0', '87', '32 %', '21', '24 %'),
           ('2', '8.0 ~ 10.5', '414', '23 %', '71', '17 %'),
           ('2', '10.5 ~ 13.4', '348', '15 %', '46', '13 %'),
           ('1', '13.6 ~ 14.2', '1,159', '12 %', '117', '10 %')]
    table(s, 7.80, 2.60, 4.83, [0.50, 1.00, 0.70, 0.90, 0.83, 0.90], tab, head_size=10, size=10.5, row_h=0.36)
    text(s, 7.80, 4.85, 4.83, 1.5, [[('읽는 법', 10.5, NAVY, True)],
                                    [('왼쪽 파란 선은 측정 잡음만으로 계산한 이론 통과율, 주황 점은 시뮬레이션입니다.', 10, SUB)],
                                    [('표는 가까운 순서입니다. 표적 2 는 다가오는 표적이라 거리 구간으로 나눴고, 개시 비율 = 개시 트랙 / 스캔 쌍 입니다.', 10, SUB)],
                                    [('각도 오차 0.48° 가 거리를 곱한 만큼 옆 방향 위치 오차가 됩니다. 5 km 면 ±42 m, 14 km 면 ±117 m.', 10, SUB)]], lnspc=112)

    # 27 정리
    s = d.content('05. 정리', '정리', '이번 과제에서 만든 것과 확인한 것, 그리고 다음 과제로 이어지는 부분입니다.', tag='1 · 2 · 3')
    sums = [('트랙 개시', '2563EB', 'EAF1FD', ['연속 두 스캔의 플롯 쌍 가운데 속도 게이트 (800 m/s) 를 통과한 쌍으로 2-Point 개시', '상태변수 9 개 (가속도 0), 오차 공분산 9 x 9, 트랙 ID 100 ~ 10000']),
            ('트랙 관리', '059669', 'E6F5EE', ['TrackScore · MemoryScore 와 문턱 5 / 0 / 5 로 임시 -> 확정 -> 메모리 -> 삭제', '메모리 트랙이 다시 연관되면 바로 확정 (MemoryScore 0)']),
            ('시뮬레이션', 'D97706', 'FEF3E2', ['STRUCT_PlotMsg 를 UDP 로 스캔마다 송수신, 플롯 ID 10000 ~ 1000000', '1,160 스캔 손실 없이 처리, 트랙 355 개 개시, 2D · 3D 에 참값 · 플롯 · 트랙 표시']),
            ('확인한 것', '7C3AED', 'F2ECFD', ['Assigned_Flag 가 늘 0 이라 트랙은 세 스캔 뒤 삭제. 확정 · 메모리는 스코어 확인 창으로 검증', '먼 표적일수록 게이트 통과율이 낮아 트랙이 덜 생김 (14 km 에서 약 10 %)'])]
    for i, (t, col, fc, lines) in enumerate(sums):
        chip(s, 0.70, 2.45 + i * 0.93, 8.0, 0.82, t, col, fc, lines, title_size=11.5, size=10)
    rect(s, 9.00, 2.45, 3.63, 3.54, fill=BOXF, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, 9.18, 2.55, 3.3, 3.3, [[('다음 과제로 이어지는 것', 11.5, NAVY, True)], [('', 6, NAVY)],
                                    [('트랙 연관', 10.5, NAVY, True)], [('이번 스캔 플롯을 기존 트랙에 붙여 Assigned_Flag 를 실제 값으로', 10, SUB)], [('', 6, NAVY)],
                                    [('추적 필터', 10.5, NAVY, True)], [('붙인 플롯으로 상태변수와 오차 공분산을 갱신', 10, SUB)], [('', 6, NAVY)],
                                    [('그러면 이번에 만든 확정 · 메모리 로직이 실제 시뮬레이션에서 동작합니다.', 10, SUB)]], lnspc=112)

    # 28 끝
    d.end('감사합니다', '질문 주시면 답변드리겠습니다')
    d.save(out)
    print('saved', out, len(d.prs.slides))


NOTES = {}


def write_script_md(path, appendix=''):
    """노트를 발표 대본 md 로 같이 써 둠. appendix 가 있으면 끝에 질의응답 부록으로 붙임"""
    lines = ['# Chapter 5 발표 대본 (2 판)', '', 'Part.5 : 트랙 개시 · 관리 (Track Initiation and Management)', '',
             '- 자료 : `Chapter5_TrackInitMgmt.pptx` (%d장) — 같은 내용이 각 슬라이드의 슬라이드 노트에도 들어 있다.' % PG['end'],
             '- 발표자 : Janghwan Kim (Harley)', '- 코드 : `Part5_Track_Init_Mgmt/src/TargetSim` — TargetSimCore (TrackMgr, PlotLink, Idd.h) / TargetSimUI (MFC)', '', '---', '']
    for n in sorted(NOTES):
        lines += ['## %d. %s' % (n, NOTES[n][0]), '', NOTES[n][1].strip(), '', '---', '']
    if appendix:
        lines += [appendix.strip(), '']
    io.open(path, 'w', encoding='utf-8', newline='\n').write('\n'.join(lines))


if __name__ == '__main__':
    import notes_v2
    NOTES.update(notes_v2.NOTES)
    build(sys.argv[1], sys.argv[2], sys.argv[3])
    write_script_md(os.path.join(os.path.dirname(os.path.abspath(sys.argv[3])), 'Chapter5_발표대본.md'), getattr(notes_v2, 'APPENDIX', ''))
