# -*- coding: utf-8 -*-
"""슬라이드에 박혀 있는 그림 파일 하나를 새로 그린 png 로 갈아 끼운다.

    python swap_figure.py 8 fig17_cartesian.png          # 8쪽의 그림을 바꾼다
    python swap_figure.py 8 fig17_cartesian.png --box    # 새 그림 비율에 맞춰 높이도 고친다

pptx 는 그림을 파일 경로가 아니라 통째로 안에 품고 있다. 그래서 make_figures.py 로 png 를
다시 그려도 발표자료는 옛 그림을 그대로 들고 있다. 이 스크립트가 그 사이를 잇는다.

한 장에 그림이 둘 이상이면 이름을 하나 더 받아 고른다 (세 번째 인자).
--box 는 가로를 그대로 두고 새 png 의 가로세로 비율로 세로만 다시 잡는다.
"""
import io
import os
import sys

from PIL import Image
from pptx import Presentation
from pptx.util import Emu

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.normpath(os.path.join(HERE, ".."))
FIG_DIR = os.path.normpath(os.path.join(HERE, "..", "..", "참고 자료", "figures"))
DECK = os.path.join(OUT_DIR, "Chapter2_coordinate_systems.pptx")

PICTURE = 13                                     # MSO_SHAPE_TYPE.PICTURE


def pictures(slide):
    return [sh for sh in slide.shapes if sh.shape_type == PICTURE]


def run(page, png, name=None, fit=False):
    path = png if os.path.isabs(png) else os.path.join(FIG_DIR, png)
    if not os.path.exists(path):
        sys.exit("그림을 못 찾음: %s" % path)

    prs = Presentation(DECK)
    slides = list(prs.slides)
    if not 1 <= page <= len(slides):
        sys.exit("쪽번호는 1 ~ %d" % len(slides))
    slide = slides[page - 1]

    found = [sh for sh in pictures(slide) if name is None or sh.name == name]
    if len(found) != 1:
        sys.exit("%d쪽에서 그림을 하나로 못 좁힘: %s"
                 % (page, [sh.name for sh in pictures(slide)]))
    shape = found[0]

    blob = io.open(path, "rb").read()
    old = shape.image.size
    part = shape.part.related_part(shape._element.blip_rId)
    part._blob = blob                            # 같은 파트를 그대로 쓰고 알맹이만 바꾼다
    part.__dict__.pop("image", None)             # lazyproperty 로 물고 있던 옛 그림을 버린다

    new = Image.open(path).size
    if fit:
        shape.height = Emu(int(round(shape.width * new[1] / float(new[0]))))

    prs.save(DECK)
    print("%d쪽 %s : %dx%d -> %dx%d%s"
          % (page, shape.name, old[0], old[1], new[0], new[1],
             "  (높이 %.2f in 로 맞춤)" % (shape.height / 914400.0) if fit else ""))


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if a != "--box"]
    if len(args) not in (2, 3):
        sys.exit("사용법: python swap_figure.py <쪽번호> <그림.png> [도형이름] [--box]")
    run(int(args[0]), args[1], args[2] if len(args) == 3 else None,
        fit="--box" in sys.argv[1:])
